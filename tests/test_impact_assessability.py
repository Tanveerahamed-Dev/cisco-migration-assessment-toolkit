"""W33 (follow-up F5): the ENGINE owns row-level failure-impact assessability.

``cisco_toolkit/impact_assessability.py`` is the one owner of the rules that decide whether a stored
``failure_impact`` row is a measurement: the producer's INDETERMINATE detail, a row older than its
``off_scan_gw_vlans`` marker, an unreadable off-scan count or host, a device whose scoped interface running-config
was not captured (``run_config_observed``), a positive off-scan count with nothing simulated, a partial simulation,
W32's per-row ``blind_links`` count (unreadable, alone, positive, or absent on a row older than it; section 4), an
uncollected neighbour that can carry endpoints (read from the stored cable map), and two rows naming one host.
It never re-simulates.

Three things are pinned here:

1. **The owner's predicates**, each driven by the REAL producers (``analyze.compute_failure_impact`` and
   ``analyze.compute_cable_map`` over constructed fleets), never by a hand-written row in the shape the owner expects.
2. **Projection parity.** ui_projection's shared row builder consumes the owner and must keep every pre-refactor
   state, reason and witness byte for byte: the full reason texts below are copied from ui_projection at origin/main
   d0e10888 (before W33), and the owner's verdict on each cell must equal the projection's state.
3. **The engine deliverables** (the workbook's Failure Impact and Executive Summary sheets, the design document, the
   executive deck, the architecture review, the device dossier and the explorer embed) read the same verdict: an
   INDETERMINATE, held or neighbour-bounded row never renders as Info / no impact / 0, and a held row never ranks as
   a keystone. Rankings follow the owner's per-cell decision: a lower-bound row whose positive stranded floor the
   owner publishes ranks by that floor and is written as one. A withheld impact never lowers a dossier's risk.

No test here runs a pipeline.
"""
from __future__ import annotations

import ast
import copy
import json
import math
import pathlib
import re
import shutil
from dataclasses import asdict

import pytest

from cisco_toolkit import analyze
from cisco_toolkit import impact_assessability as ia
from cisco_toolkit import ui_projection as ui
from cisco_toolkit.model import InterfaceData

ROOT = pathlib.Path(__file__).resolve().parent.parent
OWNER = ROOT / "cisco_toolkit" / "impact_assessability.py"
PUB, CBE, NC, AU, UV = "published", "collected_but_empty", "not_collected", "analysis_unavailable", "unverified"
MEASURES = ("severity", "vlans_impacted", "stranded", "hard", "backup", "fhrp")

# --------------------------------------------------------------------------------------------------
# the pre-refactor reasons, copied verbatim from cisco_toolkit/ui_projection.py at origin/main d0e10888
# --------------------------------------------------------------------------------------------------
PRE_INDETERMINATE = ("not collected: analyze.compute_failure_impact could not simulate this switch's blast "
                     "radius (its detail says why), so its severity and counts are not measurements")
PRE_LEGACY = ("not collected: this stored row carries no off_scan_gw_vlans, so it predates the producer's "
              "assessability marker (analyze.compute_failure_impact). An engine that old wrote 'No reachability "
              "impact' with Info and zero counts for a switch it could not simulate, so this row's severity and "
              "counts are not measurements")
PRE_OFF_SCAN_ONLY = ("not collected: every VLAN analyze.compute_failure_impact found on this switch has an "
                     "off-scan gateway (1 counted in off_scan_gw_vlans), so it simulated none of them and its "
                     "severity and counts are not measurements")
PRE_OFF_SCAN_UNREAD = ("unverified: off_scan_gw_vlans is not a count, so whether the simulation covered this "
                       "switch's whole blast radius cannot be read")
PRE_NO_RUN_CONFIG = ("not collected: no interface of this device carries run_config_observed: true. build.py "
                     "marks every interface its scoped interface running-config capture ('show running-config "
                     "interface' or '| section ^interface') parsed, and takes SVI gateway addresses (svi_ip) "
                     "only from that capture, so this device's own gateways never reached the simulation and "
                     "its severity and counts are not measurements. A snapshot that predates the marker, or "
                     "drops it as false (html.sparsify_interfaces), reads the same: not captured")
PRE_UNDERSTATED = ("not collected: this severity may understate the blast radius: 1 VLAN(s) on this switch "
                   "have an off-scan gateway the simulation could not assess (off_scan_gw_vlans), and only the "
                   "worst band (High) cannot be understated")
PRE_ZERO_BOUND = ("not collected: this 0 is only a lower bound: 1 VLAN(s) on this switch have an off-scan "
                  "gateway the simulation could not assess (off_scan_gw_vlans), so it is not a measurement of "
                  "none")
PRE_PEER_CLAUSE = ("this row cannot account for endpoints behind 1 uncollected neighbour(s): the stored cable map "
                   "cables this switch to 1 peer(s) it does not show as collected that can carry endpoints or "
                   "transit (a cable_map.nodes row with collected: false and a kind other than ap, phone or endpoint, "
                   "or a cable end that does not join exactly one node), and analyze.compute_failure_impact "
                   "counts only endpoints on scanned switches")
PRE_PEER_SEVERITY = ("not collected: " + PRE_PEER_CLAUSE + ", so this severity may understate the blast radius, "
                     "and only the worst band (High) cannot be understated")
PRE_PEER_ZERO = "not collected: " + PRE_PEER_CLAUSE + ", so this 0 is only a lower bound, not a measurement of none"
PRE_PEER_DETAIL = ("not collected: " + PRE_PEER_CLAUSE + ", so this detail, which names no simulated VLAN, is not a "
                   "clean bill: it was never checked against what lies behind them")
PRE_DUP = ("unverified: 2 rows in failure_impact name this exact host, but analyze.compute_failure_impact "
           "writes one row per host, so no single row can be chosen")


# --------------------------------------------------------------------------------------------------
# fleets for the REAL producers (copied from tests/test_ui_projection_device_impact.py, not imported)
# --------------------------------------------------------------------------------------------------
def _trunk(port, peer, peer_port, vlans="", blocked=""):
    return InterfaceData(port=port, status="connected", switchport_mode="Trunk", cdp_neighbor=peer,
                         neighbor_port=peer_port, endpoint_type="Switch", trunk_allowed_vlans=vlans,
                         stp_fwd_vlans="" if blocked else vlans, stp_blk_vlans=blocked, run_config_observed=True)


def _access(port, vid, mac):
    return InterfaceData(port=port, status="connected", switchport_mode="Access", vlan=str(vid), end_host_mac=mac,
                         run_config_observed=True)


def _svi(vid, address, fhrp=""):
    return InterfaceData(port=f"Vlan{vid}", svi_ip=address, hsrp_behavior=fhrp, run_config_observed=True)


def _edge(port, peer, endpoint_type, platform=""):
    return InterfaceData(port=port, status="connected", switchport_mode="Trunk", cdp_neighbor=peer,
                         neighbor_port="Gi0/1", endpoint_type=endpoint_type, neighbor_platform=platform,
                         trunk_allowed_vlans="10", stp_fwd_vlans="10", run_config_observed=True)


def _mixed_fleet(peer="dsw", endpoint_type="Switch", platform=""):
    """`gw` gateways VLAN 10 for `acc` (removing `gw` is a measured hard partition: High, 1 stranded) and trunks up to
    the router `wan`; `acc` trunks down to `peer`; the collection reached neither. `x1` and `x2` share a trunk with
    no VLAN-carriage evidence on either end, so the producer writes its INDETERMINATE disclosure for both."""
    return {"gw": {"Gi1": _trunk("Gi1", "acc", "Gi1", "10"),
                   "Gi47": _edge("Gi47", "wan", "Router", "cisco ISR4331/K9"), "Vlan10": _svi(10, "10.10.0.1/24")},
            "acc": {"Gi1": _trunk("Gi1", "gw", "Gi1", "10"), "Gi10": _access("Gi10", 10, "0000.0000.000a"),
                    "Gi48": _edge("Gi48", peer, endpoint_type, platform)},
            "x1": {"Gi1": _trunk("Gi1", "x2", "Gi1")}, "x2": {"Gi1": _trunk("Gi1", "x1", "Gi1")}}


def _offscan_fleet():
    """`g1` and `g2` are FHRP peers gatewaying VLAN 10 for `acc` (removing either is FHRP-covered: Low); both also
    carry VLAN 20 to `acc`, whose gateway was never scanned. `acc` carries only what it cannot simulate."""
    return {"g1": {"Gi1": _trunk("Gi1", "acc", "Gi1", "10,20"), "Vlan10": _svi(10, "10.10.0.2/24", "Active")},
            "g2": {"Gi1": _trunk("Gi1", "acc", "Gi2", "10,20"), "Vlan10": _svi(10, "10.10.0.3/24", "Standby")},
            "acc": {"Gi1": _trunk("Gi1", "g1", "Gi1", "10,20"), "Gi2": _trunk("Gi2", "g2", "Gi1", "10,20"),
                    "Gi10": _access("Gi10", 10, "0000.0000.000a"), "Gi20": _access("Gi20", 20, "0000.0000.0014")}}


def _keystone_fleet():
    """The independent review's counterexample (P2-1), from the real producers: `core1` is the only gateway of VLAN 10
    for `acc1`'s 300 endpoints, so removing it strands all of them (High, 300), and it is cabled to the uncollected
    router `wan1`, so that count is only a lower bound. `acc1` strands nobody (Info, 0) and faces only an uncollected
    access point (edge gear), so its row is a published measurement."""
    acc1 = {"Gi1": _trunk("Gi1", "core1", "Gi1", "10"),
            "Gi48": _edge("Gi48", "ap1", "Switch", "cisco AIR-AP2802I-E-K9")}
    acc1.update({f"Gi1/0/{i}": _access(f"Gi1/0/{i}", 10, f"0000.0000.{i:04x}") for i in range(1, 301)})
    return {"core1": {"Gi1": _trunk("Gi1", "acc1", "Gi1", "10"),
                      "Gi47": _edge("Gi47", "wan1", "Router", "cisco ISR4331/K9"),
                      "Vlan10": _svi(10, "10.10.0.1/24")},
            "acc1": acc1}


def _snapshot(interfaces):
    """The real producers' rows over `interfaces`, in the snapshot's own shape."""
    return {"schema": "collect_parse_snapshot/1", "devices": {host: {"hostname": host} for host in interfaces},
            "interfaces": {host: {port: asdict(row) for port, row in ports.items()}
                           for host, ports in interfaces.items()},
            "failure_impact": analyze.compute_failure_impact(interfaces),
            "cable_map": analyze.compute_cable_map(interfaces), "routes": {}}


def _k(snap, host):
    return next(i for i, row in enumerate(snap["failure_impact"]) if row.get("host") == host)


def _verdicts(snap):
    return {v.index: v for v in ia.assess_failure_impact(snap)}


def _node(snap, host):
    same = [n for n in snap["cable_map"]["nodes"] if n["host"] == host]
    assert len(same) == 1, (host, same)
    return same[0]


@pytest.fixture()
def mixed():
    return _snapshot(_mixed_fleet())


@pytest.fixture()
def keystone():
    """The P2-1 counterexample, with its preconditions read back from the real producers."""
    snap = _snapshot(_keystone_fleet())
    v = _verdicts(snap)
    core1, acc1 = v[_k(snap, "core1")], v[_k(snap, "acc1")]
    assert _node(snap, "wan1")["collected"] is False and _node(snap, "ap1")["kind"] == "ap"
    assert core1.assessable == ia.LOWER_BOUND and core1.codes == ["uncollected_neighbours"], core1.as_dict()
    assert (core1.raw["severity"], core1.raw["stranded"]) == ("High", 300), core1.raw
    assert acc1.published and (acc1.raw["severity"], acc1.raw["stranded"]) == ("Info", 0), acc1.as_dict()
    return snap


# --------------------------------------------------------------------------------------------------
# 1. the owner's predicates over the REAL producers
# --------------------------------------------------------------------------------------------------
def test_the_indeterminate_marker_is_the_real_producers_opening(mixed):
    """Both INDETERMINATE branches of analyze.compute_failure_impact open with the owner's marker, and the
    projection re-exports the owner's constant rather than keeping its own."""
    assert ui.IMPACT_INDETERMINATE_PREFIX is ia.IMPACT_INDETERMINATE_PREFIX
    assert ui.IMPACT_SEVERITIES is ia.IMPACT_SEVERITIES
    source = pathlib.Path(analyze.__file__).read_text(encoding="utf-8")
    assert source.count(f'f"{ia.IMPACT_INDETERMINATE_PREFIX} —') == 2
    for host in ("x1", "x2"):
        row = mixed["failure_impact"][_k(mixed, host)]
        assert row["detail"].startswith(ia.IMPACT_INDETERMINATE_PREFIX) and row["severity"] == "Info", row
        assert all(row[f] == 0 for f in MEASURES[1:]), row                    # zeros that look like measurements
    offscan = _snapshot(_offscan_fleet())
    row = offscan["failure_impact"][_k(offscan, "acc")]
    assert row["detail"].startswith(ia.IMPACT_INDETERMINATE_PREFIX) and row["off_scan_gw_vlans"] == 1, row


def test_each_verdict_from_the_real_producers(mixed):
    v = _verdicts(mixed)
    gw, acc, x1, x2 = (v[_k(mixed, h)] for h in ("gw", "acc", "x1", "x2"))
    for held in (x1, x2):                                    # the producer's own INDETERMINATE disclosure
        assert held.assessable == ia.NOT_ASSESSED and held.codes == ["indeterminate"], held.as_dict()
        assert all(held.withholds(f) for f in MEASURES)
        assert not held.withholds("detail")                  # its disclosure stays readable
        assert held.facts.hold.reason == PRE_INDETERMINATE
    # acc: the producer's clean bill, but it is cabled to an uncollected switch
    assert _node(mixed, "dsw")["collected"] is False and _node(mixed, "dsw")["kind"] == "switch"
    assert acc.assessable == ia.LOWER_BOUND and acc.codes == ["uncollected_neighbours"], acc.as_dict()
    assert all(acc.withholds(f) for f in MEASURES) and acc.withholds("detail")   # Info, zeros, clean-bill detail
    # gw: High, one stranded, cabled to an uncollected router: High and the positive counts are lower bounds
    src = mixed["failure_impact"][_k(mixed, "gw")]
    assert src["severity"] == "High" and src["stranded"] == src["hard"] == src["vlans_impacted"] == 1, src
    assert gw.assessable == ia.LOWER_BOUND, gw.as_dict()
    assert [f for f in MEASURES if gw.withholds(f)] == ["backup", "fhrp"]
    assert not gw.withholds("detail")                        # a per-VLAN detail lists what was simulated
    assert ia.table_value(gw, "severity") == "High (lower bound)" and ia.table_value(gw, "stranded") == "≥ 1"
    # the control: once the stored cable map shows the router collected, gw is the producer's measurement
    _node(mixed, "wan")["collected"] = True
    gw = _verdicts(mixed)[_k(mixed, "gw")]
    assert gw.assessable == ia.PUBLISHED and gw.summary == "published" and gw.codes == []
    assert not any(gw.withholds(f) for f in MEASURES + ("detail",))
    assert ia.table_value(gw, "stranded") == 1 and ia.table_detail(gw) == src["detail"]


def test_a_partial_simulation_is_a_lower_bound_and_off_scan_only_is_held():
    snap = _snapshot(_offscan_fleet())
    v = _verdicts(snap)
    g1 = v[_k(snap, "g1")]
    src = snap["failure_impact"][_k(snap, "g1")]
    assert src["severity"] == "Low" and src["off_scan_gw_vlans"] == 1 and src["vlans_impacted"] == 1, src
    assert g1.assessable == ia.LOWER_BOUND and g1.codes == ["off_scan_partial"], g1.as_dict()
    assert g1.withholds("severity") and not g1.withholds("vlans_impacted") and not g1.withholds("detail")
    acc = v[_k(snap, "acc")]
    assert acc.assessable == ia.NOT_ASSESSED and acc.codes == ["indeterminate"]
    # the same row without its prose marker: the count still says nothing was simulated
    snap["failure_impact"][_k(snap, "acc")]["detail"] = "No reachability impact from removing this switch."
    acc = _verdicts(snap)[_k(snap, "acc")]
    assert acc.codes == ["off_scan_only"] and acc.facts.hold.reason == PRE_OFF_SCAN_ONLY
    assert acc.withholds("detail")                             # a clean bill beside a hold is held with it


def test_holds_in_their_first_match_order(mixed):
    k = _k(mixed, "gw")
    _node(mixed, "wan")["collected"] = True                   # gw published: each mutation alone holds it
    row = mixed["failure_impact"][k]
    cases = []
    legacy = copy.deepcopy(mixed); del legacy["failure_impact"][k]["off_scan_gw_vlans"]
    cases.append((legacy, "legacy_row", NC, PRE_LEGACY, f"/failure_impact/{k}"))
    unread = copy.deepcopy(mixed); unread["failure_impact"][k]["off_scan_gw_vlans"] = "1"
    cases.append((unread, "off_scan_unreadable", UV, PRE_OFF_SCAN_UNREAD, f"/failure_impact/{k}/off_scan_gw_vlans"))
    nohost = copy.deepcopy(mixed); nohost["failure_impact"][k]["host"] = 7
    cases.append((nohost, "no_host", UV, ia.R_NO_HOST, f"/failure_impact/{k}"))
    norc = copy.deepcopy(mixed)
    for port in norc["interfaces"]["gw"].values():
        port.pop("run_config_observed", None)                 # what html.sparsify_interfaces writes for false
    cases.append((norc, "no_run_config", NC, PRE_NO_RUN_CONFIG, "/interfaces/gw"))
    marker_text = copy.deepcopy(mixed)
    for port in marker_text["interfaces"]["gw"].values():
        port["run_config_observed"] = "true"                  # present but not true reads as not captured
    cases.append((marker_text, "no_run_config", NC, PRE_NO_RUN_CONFIG, "/interfaces/gw"))
    gone = copy.deepcopy(mixed); del gone["interfaces"]["gw"]
    cases.append((gone, "no_run_config", NC, PRE_NO_RUN_CONFIG, "/interfaces"))
    for snap, code, state, reason, pointer in cases:
        v = _verdicts(snap)[k]
        assert v.assessable == ia.NOT_ASSESSED and v.codes == [code], (code, v.as_dict())
        assert v.facts.hold.state == state and v.facts.hold.reason == reason, code
        assert pointer in v.pointers, (code, v.pointers)
        assert all(v.withholds(f) for f in MEASURES + ("detail",)), code
    assert row["detail"] and not row["detail"].startswith(ia.IMPACT_INDETERMINATE_PREFIX)


@pytest.mark.parametrize("peer, platform, kind", [("ap1", "cisco AIR-AP2802I-E-K9", "ap"),
                                                   ("sep1", "Cisco IP Phone 8865", "phone")])
def test_uncollected_edge_gear_bounds_nothing_and_every_other_kind_does(peer, platform, kind):
    snap = _snapshot(_mixed_fleet(peer, "Switch", platform))
    k = _k(snap, "acc")
    assert _node(snap, peer)["collected"] is False and _node(snap, peer)["kind"] == kind
    assert _verdicts(snap)[k].assessable == ia.PUBLISHED
    for mutate in (lambda n: n.update(kind="unknown"), lambda n: n.pop("kind"), lambda n: n.update(kind="AP"),
                   lambda n: n.update(kind="device"), lambda n: n.pop("collected"),
                   lambda n: n.update(collected="false")):
        held = copy.deepcopy(snap)
        mutate(_node(held, peer))
        v = _verdicts(held)[k]
        assert v.assessable == ia.LOWER_BOUND and v.codes == ["uncollected_neighbours"], v.as_dict()
        assert v.facts.bounds[-1].severity_reason == PRE_PEER_SEVERITY


@pytest.mark.parametrize("mode, state", [("duplicate_node", NC), ("missing_node", NC), ("unreadable_cable", NC),
                                         ("malformed_cables", UV), ("absent_cables", NC), ("failed_cable_map", AU)])
def test_a_neighbour_the_join_cannot_resolve_fails_closed(mode, state):
    snap = _snapshot(_mixed_fleet("ap1", "Switch", "cisco AIR-AP2802I-E-K9"))
    k = _k(snap, "acc")
    assert _verdicts(snap)[k].assessable == ia.PUBLISHED
    nodes, cables = snap["cable_map"]["nodes"], snap["cable_map"]["cables"]
    if mode == "duplicate_node":
        nodes.append(copy.deepcopy(_node(snap, "ap1")))
    elif mode == "missing_node":
        nodes.remove(_node(snap, "ap1"))
    elif mode == "unreadable_cable":
        cables.append(None)
    elif mode == "malformed_cables":
        snap["cable_map"]["cables"] = 7
    elif mode == "absent_cables":
        del snap["cable_map"]["cables"]
    else:
        snap["assessment_integrity"] = {"cable_map": "failed"}
    v = _verdicts(snap)[k]
    assert v.assessable == ia.LOWER_BOUND, (mode, v.as_dict())
    assert v.facts.bounds[-1].state == state, mode
    code = "uncollected_neighbours" if mode in ("duplicate_node", "missing_node", "unreadable_cable") \
        else "neighbours_unreadable"
    assert v.codes == [code], (mode, v.codes)
    assert all(v.withholds(f) for f in MEASURES) and v.withholds("detail")


def test_two_rows_naming_one_host_are_ambiguous_never_picked(mixed):
    _node(mixed, "wan")["collected"] = True
    k = _k(mixed, "gw")
    mixed["failure_impact"].append(dict(mixed["failure_impact"][k], severity="Info", stranded=0))
    v = _verdicts(mixed)
    for i in (k, len(mixed["failure_impact"]) - 1):
        assert v[i].assessable == ia.AMBIGUOUS and v[i].codes == ["duplicate_host"], v[i].as_dict()
        assert v[i].facts.doubt.reason == PRE_DUP
        assert {f"/failure_impact/{k}", f"/failure_impact/{len(mixed['failure_impact']) - 1}"} <= set(v[i].pointers)
        assert all(v[i].withholds(f) for f in MEASURES + ("detail",)) and not v[i].withholds("host")
    # a case variant names another host: the key is exact text
    mixed["failure_impact"][-1]["host"] = "GW"
    assert _verdicts(mixed)[k].assessable == ia.PUBLISHED


def test_rankings_follow_the_per_cell_decision_not_the_row_verdict(mixed):
    """Review P2-1: a lower-bound row whose positive stranded count the owner publishes ranks by that floor and is
    written as one; a bounded zero, a held row and an ambiguous row never rank; a published row ranks by its own
    measurement, written as stored."""
    v = _verdicts(mixed)
    gw, acc, x1 = (v[_k(mixed, h)] for h in ("gw", "acc", "x1"))
    assert not gw.withholds("stranded") and ia.ranking_floor(gw) == 1 and ia.ranks(gw)
    assert ia.ranked_value(gw, "stranded") == "≥ 1 (lower bound)"
    assert ia.ranked_value(gw, "severity") == "High (lower bound)"
    assert ia.ranked_value(gw, "backup") == ia.NOT_ASSESSED_CELL           # a bounded zero stays withheld
    assert acc.withholds("stranded") and ia.ranking_floor(acc) is None and not ia.ranks(acc)
    assert ia.ranking_floor(x1) is None and not ia.ranks(x1)
    _node(mixed, "wan")["collected"] = True
    k = _k(mixed, "gw")
    gw = _verdicts(mixed)[k]
    assert gw.published and ia.ranking_floor(gw) is None and ia.ranks(gw)
    assert ia.ranked_value(gw, "stranded") == 1                             # a measurement is written as stored
    mixed["failure_impact"].append(copy.deepcopy(mixed["failure_impact"][k]))
    doubted = [d for d in _verdicts(mixed).values() if d.host == "gw"]
    assert len(doubted) == 2 and not any(ia.ranks(d) for d in doubted)      # ambiguous never ranks
    unavailable = ia.unavailable_document()
    assert unavailable["schema"] == ia.SCHEMA and unavailable["rows"] == [] and unavailable["unavailable"] is True


def test_a_failed_section_reads_every_row_not_assessed(mixed):
    mixed["assessment_integrity"] = {"failed_phases": ["Failure Impact"]}
    verdicts = ia.assess_failure_impact(mixed)
    assert verdicts and all(v.assessable == ia.NOT_ASSESSED and v.codes == ["section_unavailable"] for v in verdicts)
    assert all(v.withholds(f) for v in verdicts for f in ia.IMPACT_FIELDS)


def test_verdict_pointers_resolve_and_the_document_is_aligned(mixed):
    doc = ia.assessment_document(mixed)
    assert doc["schema"] == ia.SCHEMA == "failure_impact_assessability/1"
    assert [r["index"] for r in doc["rows"]] == list(range(len(mixed["failure_impact"])))
    assert sum(doc["counts"].values()) == len(doc["rows"]) and set(doc["counts"]) == set(ia.VERDICTS)
    for row in doc["rows"]:
        assert row["assessable"] in ia.VERDICTS and set(row["codes"]) <= set(ia.CODE_PHRASES), row
        for pointer in row["pointers"]:
            node = mixed
            for tok in pointer.split("/")[1:]:
                tok = tok.replace("~1", "/").replace("~0", "~")
                node = node[int(tok)] if isinstance(node, list) else node[tok]
        assert set(row["withheld"]) <= set(ia.IMPACT_FIELDS)
        json.dumps(row, allow_nan=False)
    # the reader-facing reasons name no engine internals
    for phrase in ia.CODE_PHRASES.values():
        assert not re.search(r"analyze|compute_|build\.py|_gw_vlans|snapshot", phrase), phrase


@pytest.mark.parametrize("poison", [None, 7, "x", {"a": 1}, [None, 7, "x", [1], {"host": None}],
                                    [{"host": "h", "off_scan_gw_vlans": math.inf, "detail": 5}]])
def test_the_owner_is_total(poison):
    for snap in ({"failure_impact": poison}, {"failure_impact": poison, "interfaces": 5, "cable_map": [1]},
                 {"failure_impact": poison, "cable_map": {"cables": [None, {"a": 1}], "nodes": "x"}}):
        verdicts = ia.assess_failure_impact(snap)
        assert len(verdicts) == (len(poison) if isinstance(poison, list) else 0)
        assert not any(v.published for v in verdicts if not isinstance(v.raw, dict))
        json.dumps(ia.assessment_document(snap), allow_nan=False)


# --------------------------------------------------------------------------------------------------
# the owner never re-simulates and imports nothing but ssot
# --------------------------------------------------------------------------------------------------
SIMULATION_NAMES = frozenset({
    "compute_failure_impact", "compute_link_centrality", "build_network_model", "_topology_adjacency",
    "compute_topology_links", "compute_cable_map", "compute_causality_chains", "_vlan_components", "_link_carries",
    "_carry", "_link_has_vlan_evidence"})


def test_the_owner_never_names_a_topology_producer_and_imports_only_ssot():
    tree = ast.parse(OWNER.read_text(encoding="utf-8"))
    names = set()
    imported = set()
    for node in ast.walk(tree):
        if isinstance(node, ast.Name):
            names.add(node.id)
        elif isinstance(node, ast.Attribute):
            names.add(node.attr)
        elif isinstance(node, ast.Import):
            imported.update(a.name for a in node.names)
        elif isinstance(node, ast.ImportFrom):
            imported.update(f"{node.module}:{a.name}" for a in node.names)
    assert not names & SIMULATION_NAMES
    assert {"getattr", "__import__", "eval", "exec"}.isdisjoint(names)
    extra = {n for n in imported if not n.startswith(("typing:", "__future__:", "types:")) and n != "math"}
    assert extra == {"cisco_toolkit:ssot"}, extra
    # and the owner names no state of the protocol receipt (it is not one of its readers)
    assert "protocol_assessability" not in OWNER.read_text(encoding="utf-8")


# --------------------------------------------------------------------------------------------------
# 2. projection parity: the shared row builder consumes the owner, byte for byte as before W33
# --------------------------------------------------------------------------------------------------
@pytest.mark.parametrize("value", [0, 1, 2 ** 53 - 1, 2 ** 53, -1, 1.0, 2.5, math.inf, math.nan, True, False, None,
                                   "1", "", "core1", "\ud800", "café", [], {}])
def test_the_owner_value_checks_equal_the_projections(value):
    assert ia._count(value) == ui._count(value) or (value != value)      # NaN never equals itself
    assert ia._count(value)[0] == ui._count(value)[0]
    assert ia._is_text(value) == ui._is_text(value)
    assert ia.JS_MAX_SAFE_INT == ui.JS_MAX_SAFE_INT


def _cable_states(snap):
    ctx = ui._Ctx(snap)
    ui_state, _reason, _raw = ui._topology_source(ctx, ("cable_map", "cables"))
    src = ia.read_cable_source(snap)
    return (None if ui_state in (PUB, CBE) else ui_state if ui_state in (AU, UV) else NC), src.state


@pytest.mark.parametrize("cable_map", [
    {"nodes": [], "cables": []}, {"nodes": [], "cables": [{"a": "x", "b": "y"}]}, {"nodes": []}, None, 7, [],
    {"nodes": [], "cables": 7}, {"nodes": [], "cables": [None, {}]}, {"cables": [{"a": "x"}], "nodes": "n"}])
@pytest.mark.parametrize("integrity", [None, {"cable_map": "failed"}, {"failed_phases": ["Cable map"]},
                                       {"failed_phases": ["nobody classified this"]}])
def test_the_owners_cable_map_reading_follows_the_projections_list_states(cable_map, integrity):
    snap = {"failure_impact": [{"host": "x"}]}
    if cable_map is not None:
        snap["cable_map"] = cable_map
    if integrity is not None:
        snap["assessment_integrity"] = integrity
    projected, owned = _cable_states(snap)
    assert projected == owned, (cable_map, integrity, projected, owned)


def test_the_owners_joins_equal_the_projections(mixed):
    ctx = ui._Ctx(mixed)
    cables = mixed["cable_map"]["cables"] + [None, {"a": "gw"}, {"a": "gw", "b": "gw"}]
    mixed["cable_map"]["cables"] = cables
    assert ia.index_rows(cables, ("a", "b")) == ctx.index(("cable_map", "cables"), ("a", "b"))
    assert ia.unjoinable_rows(cables, ("a", "b")) == ui._unjoinable_rows(cables, ("a", "b"))
    assert ia.index_rows(mixed["failure_impact"], ("host",)) == ctx.index(("failure_impact",), ("host",))


def _reference_neighbour_join(snap, host):
    """The pre-W33 join loop of ui_projection._impact_peers (origin/main d0e10888), copied for parity: the bounding
    cable indices in iteration order, the far-end fail-closed map and the unreadable count."""
    cables = snap["cable_map"]["cables"]
    nodes = snap["cable_map"]["nodes"]
    index = ia.index_rows(nodes, ("host",))
    hits, peers, unreadable = [], {}, 0
    bad = set(ia.unjoinable_rows(cables, ("a", "b")))
    for j in sorted(set(ia.index_rows(cables, ("a", "b")).get(host, [])) | bad):
        if j in bad:
            hits.append(j)
            unreadable += 1
            continue
        ends = (cables[j]["a"], cables[j]["b"])
        far = ends[1] if ends[0] == host else ends[0]
        found = index.get(far, []) if far else []
        if len(found) == 1:
            node = nodes[found[0]]
            kind = node.get("kind")
            if node.get("collected") is True or (
                    node.get("collected") is False and ia._is_text(kind) and kind in ia.IMPACT_EDGE_KINDS):
                continue
        hits.append(j)
        peers[far] = peers.get(far, False) or len(found) != 1
    return hits, peers, unreadable


@pytest.mark.parametrize("extra", [
    [], [None], [{"a": "acc"}], [{"a": "acc", "b": 7}], [{"b": "gw"}, None, {"a": "x1", "b": "x1"}],
    [{"a": "acc", "b": "ghost"}, {"a": "ghost", "b": "acc"}], [{"a": "gw", "b": ""}]])
def test_the_neighbour_join_equals_the_pre_refactor_loop(mixed, extra):
    """The owner's join is restructured (readable rows first, the unreadable ones once per list), never re-ruled: for
    every host its witnesses, peer count and fail-closed count equal the pre-W33 loop, half-readable rows included."""
    mixed["cable_map"]["cables"] = mixed["cable_map"]["cables"] + extra
    src = ia.read_cable_source(mixed)
    for host in ("gw", "acc", "x1", "x2", "ghost"):
        hits, peers, unreadable = _reference_neighbour_join(mixed, host)
        bound = ia.neighbour_bound(host, src)
        if not hits:
            assert bound is None, (host, extra)
            continue
        closed = unreadable + sum(peers.values())
        clause = ia.R_PEERS.format(n=len(peers) + unreadable,
                                   closed=ia.R_PEERS_CLOSED.format(k=closed) if closed else "")
        assert bound.witnesses == [("witness", ("cable_map", "cables", j)) for j in hits], (host, extra)
        assert bound.reason == "not collected: " + clause and bound.n == len(peers) + unreadable, (host, extra)


def test_a_hostile_cable_list_keeps_verdicts_exact_and_witnesses_bounded(mixed):
    """A cable row the join cannot read bounds EVERY row. The deliverable verdicts keep their exact reasons and codes
    and cap only the witness pointers; the projection's uncapped reading keeps every witness."""
    mixed["cable_map"]["cables"] = mixed["cable_map"]["cables"] + [None] * 500
    exact = ia.ImpactSnapshot(mixed)                                   # no cap: the projection's reading
    for v in ia.assess_failure_impact(mixed):
        full = ia.RowVerdict(mixed, v.index, v.raw, exact.row(v.index, v.raw), None)
        assert (v.assessable, v.codes, v.reasons) == (full.assessable, full.codes, full.reasons), v.host
        assert len(v.pointers) <= 2 * ia.DELIVERABLE_WITNESS_CAP + 4, (v.host, len(v.pointers))
        assert [f for f in ia.IMPACT_FIELDS if v.withholds(f)] == [f for f in ia.IMPACT_FIELDS if full.withholds(f)]
    k = _k(mixed, "acc")
    assert len(exact.row(k, mixed["failure_impact"][k]).bounds[-1].witnesses) == 501   # dsw + every None row
    assert "501 uncollected neighbour(s)" in ia.assess_failure_impact(mixed)[k].reasons[0]


def _many_rows(n_rows, n_cables):
    """`n_rows` stored rows naming the one host `h` (each ambiguous, none held) plus one row for `g`. `h` has `n_cables`
    interfaces, only the last carrying the running-config mark, and is cabled to `n_cables` uncollected switches.
    Hand-built on purpose: the subject is the owner's cost per host, not a producer's row. Each row carries both of a
    current producer row's markers (off_scan_gw_vlans and the W32 blind_links count, 0 each), so the only bounds are
    the ones under test."""
    row = {"host": "h", "severity": "High", "vlans_impacted": 1, "stranded": 5, "hard": 1, "backup": 0, "fhrp": 0,
           "off_scan_gw_vlans": 0, "detail": "VLAN 10: Hard partition (5 ep)", "blind_links": 0}
    return {"failure_impact": [dict(row) for _ in range(n_rows)] + [dict(row, host="g")],
            "interfaces": {"h": {f"Gi{i}": {"run_config_observed": i == n_cables - 1} for i in range(n_cables)},
                           "g": {"Gi0": {"run_config_observed": True}}},
            "cable_map": {"nodes": [{"host": "h", "collected": True, "kind": "switch"},
                                    {"host": "g", "collected": True, "kind": "switch"}]
                                   + [{"host": f"peer{i}", "collected": False, "kind": "switch"}
                                      for i in range(n_cables)],
                          "cables": [{"a": "h", "b": f"peer{i}"} for i in range(n_cables)]}}


def test_the_per_host_scans_run_once_however_many_rows_name_the_host(monkeypatch):
    """Review P2-3: the interface running-config scan and the uncollected-neighbour join are pure functions of the host
    and the cached sources, so the owner runs each once per host. Recomputing them per row made R rows on one host cost
    R x (its interfaces + its cables): quadratic on a hostile upload."""
    calls = {"captured": [], "peers": []}
    real_captured, real_peers = ia.run_config_captured, ia.neighbour_bound

    def captured(interfaces, host):
        calls["captured"].append(host)
        return real_captured(interfaces, host)

    def peers(host, src, **kwargs):
        calls["peers"].append(host)
        return real_peers(host, src, **kwargs)

    monkeypatch.setattr(ia, "run_config_captured", captured)
    monkeypatch.setattr(ia, "neighbour_bound", peers)
    snap = _many_rows(40, 30)
    verdicts = ia.assess_failure_impact(snap)
    assert sorted(calls["captured"]) == ["g", "h"] and sorted(calls["peers"]) == ["g", "h"], calls
    # the memo changes no verdict: every `h` row is ambiguous and neighbour-bounded, `g` is a measurement
    assert all(v.codes == ["duplicate_host", "uncollected_neighbours"] for v in verdicts[:40]), verdicts[0].as_dict()
    assert "30 uncollected neighbour(s)" in verdicts[0].reasons[1] and verdicts[-1].published


class _CountingRows(list):
    """A cable list that counts the owner's indexed reads of its rows (list iteration does not call __getitem__)."""

    def __init__(self, rows):
        super().__init__(rows)
        self.reads = 0

    def __getitem__(self, index):
        self.reads += 1
        return super().__getitem__(index)


def test_the_neighbour_join_reads_one_hosts_cables_once_and_hands_out_copies():
    """Review P2-3, sized without a clock: the cable rows the join reads stay one host's cables whether 10 or 200 rows
    name the host (per-row joins read rows x cables), and the memo hands each row its own witness list."""
    reads = {}
    for n_rows in (10, 200):
        snap = _many_rows(n_rows, 50)
        cables = _CountingRows(snap["cable_map"]["cables"])
        src = ia.readable_cables(cables, snap["cable_map"]["nodes"])
        owner = ia.ImpactSnapshot(snap, cables=lambda src=src: src)        # uncapped: every witness kept
        facts = [owner.row(i, raw) for i, raw in enumerate(snap["failure_impact"])]
        reads[n_rows] = cables.reads
        assert all(len(f.bounds[-1].witnesses) == 50 for f in facts[:n_rows]), n_rows
        assert not facts[-1].bounds, facts[-1]                             # `g` faces no neighbour
        facts[0].bounds[-1].witnesses.append(("witness", ("tampered",)))  # a caller mutating its own copy ...
        again = owner.row(1, snap["failure_impact"][1]).bounds[-1].witnesses
        assert len(again) == 50 and ("witness", ("tampered",)) not in again  # ... never reaches the cache
    assert reads[10] == reads[200] and 0 < reads[200] <= 2 * 50, reads


def _topology(snap):
    return ui.project_topology(snap)["failure_impact"]["items"]


def _refs(fact):
    return {(r["pointer"], r["role"]) for r in fact["refs"]}


def _assert_cells_follow_the_owner(snap):
    """Every measure and detail cell is withheld exactly when the owner's verdict withholds it (a well-typed row)."""
    rows = _topology(snap)
    for v in ia.assess_failure_impact(snap):
        for field in MEASURES + ("detail",):
            assert (rows[v.index][field]["state"] != PUB) == v.withholds(field), (v.host, field, v.as_dict())


def test_projection_reasons_and_witnesses_are_the_pre_refactor_ones(mixed):
    rows = _topology(mixed)
    _assert_cells_follow_the_owner(mixed)
    # INDETERMINATE: every measure held with the producer's detail as witness; the disclosure stays published
    for host in ("x1", "x2"):
        k = _k(mixed, host)
        for field in MEASURES:
            fact = rows[k][field]
            assert (fact["state"], fact["reason"]) == (NC, PRE_INDETERMINATE), (host, field)
            assert (f"/failure_impact/{k}/detail", "witness") in _refs(fact)
        assert rows[k]["detail"]["state"] == PUB
    # an uncollected neighbour: the exact severity / zero / clean-bill reasons and the cable witness
    k = _k(mixed, "acc")
    j = next(i for i, c in enumerate(mixed["cable_map"]["cables"]) if "dsw" in (c["a"], c["b"]))
    assert rows[k]["severity"]["reason"] == PRE_PEER_SEVERITY
    assert all(rows[k][f]["reason"] == PRE_PEER_ZERO for f in MEASURES[1:])
    assert rows[k]["detail"]["reason"] == PRE_PEER_DETAIL
    assert all((f"/cable_map/cables/{j}", "witness") in _refs(rows[k][f]) for f in MEASURES + ("detail",))
    # a partial simulation: the off-scan bound's understated / zero reasons
    snap = _snapshot(_offscan_fleet())
    rows = _topology(snap)
    _assert_cells_follow_the_owner(snap)
    k = _k(snap, "g1")
    assert rows[k]["severity"]["reason"] == PRE_UNDERSTATED
    assert rows[k]["stranded"]["reason"] == PRE_ZERO_BOUND
    assert (f"/failure_impact/{k}/off_scan_gw_vlans", "witness") in _refs(rows[k]["vlans_impacted"])


def test_projection_holds_and_doubts_are_the_pre_refactor_ones(mixed):
    _node(mixed, "wan")["collected"] = True
    k = _k(mixed, "gw")
    for mutate, reason, state in (
            (lambda s: s["failure_impact"][k].pop("off_scan_gw_vlans"), PRE_LEGACY, NC),
            (lambda s: s["failure_impact"][k].update(off_scan_gw_vlans="1"), PRE_OFF_SCAN_UNREAD, UV),
            (lambda s: [p.pop("run_config_observed", None) for p in s["interfaces"]["gw"].values()],
             PRE_NO_RUN_CONFIG, NC)):
        snap = copy.deepcopy(mixed)
        mutate(snap)
        row = _topology(snap)[k]
        _assert_cells_follow_the_owner(snap)
        for field in MEASURES + ("detail",):
            assert (row[field]["state"], row[field]["reason"]) == (state, reason), (reason[:30], field)
    snap = copy.deepcopy(mixed)
    snap["failure_impact"].append(copy.deepcopy(snap["failure_impact"][k]))
    rows = _topology(snap)
    for i in (k, len(snap["failure_impact"]) - 1):
        for field in ("host",) + MEASURES + ("detail",):
            assert (rows[i][field]["state"], rows[i][field]["reason"]) == (UV, PRE_DUP), (i, field)
    _assert_cells_follow_the_owner(snap)


def test_the_projection_reads_the_cable_list_through_its_own_envelope(mixed):
    """The projection injects its envelope reading of the cable list (reason text and failure records included):
    a failed cable map keeps the projection's own reason and its failure-record ref, as before W33."""
    _node(mixed, "wan")["collected"] = True
    mixed["assessment_integrity"] = {"cable_map": "failed"}
    k = _k(mixed, "acc")
    row = _topology(mixed)[k]
    fact = row["severity"]
    assert fact["state"] == AU and fact["reason"].startswith("analysis unavailable: whether this switch faces")
    assert "(analysis unavailable:" in fact["reason"]                          # the projection's own list reason
    assert ("/assessment_integrity/cable_map", "failure_record") in _refs(fact)


# --------------------------------------------------------------------------------------------------
# 3. the engine deliverables read the owner's verdict
# --------------------------------------------------------------------------------------------------
openpyxl = pytest.importorskip("openpyxl")


def _sheet_rows(ws):
    return [[c.value for c in row] for row in ws.iter_rows(min_row=2)]


def test_failure_impact_sheet_never_writes_info_or_zero_for_a_withheld_value(mixed):
    from openpyxl import Workbook
    from cisco_toolkit.excel import FAILURE_SHEET_NAME, write_failure_impact_sheet
    wb = Workbook()
    wb.active.title = "First"
    write_failure_impact_sheet(wb, mixed["failure_impact"])           # no evidence: every row is not assessed
    wb.create_sheet("Later")
    assert all(r[0] == ia.NOT_ASSESSED_CELL for r in _sheet_rows(wb[FAILURE_SHEET_NAME]))
    write_failure_impact_sheet(wb, mixed["failure_impact"], mixed)    # rewritten with the evidence, in place
    assert wb.sheetnames == ["First", FAILURE_SHEET_NAME, "Later"]
    rows = {r[1]: r for r in _sheet_rows(wb[FAILURE_SHEET_NAME])}
    for host in ("x1", "x2"):
        severity, _h, *counts, detail = rows[host]
        assert severity == ia.NOT_ASSESSED_CELL and counts == [ia.NOT_ASSESSED_CELL] * 5, rows[host]
        assert detail.startswith("Not assessed") and "Producer detail: Blast radius INDETERMINATE" in detail
    severity, _h, *counts, detail = rows["acc"]
    assert severity == ia.NOT_ASSESSED_CELL and counts == [ia.NOT_ASSESSED_CELL] * 5
    assert detail.startswith("Lower bound") and "No reachability impact" not in detail
    severity, _h, vlans, stranded, hard, backup, fhrp, detail = rows["gw"]
    assert severity == "High (lower bound)" and (vlans, stranded, hard) == ("≥ 1",) * 3
    assert backup == fhrp == ia.NOT_ASSESSED_CELL and detail.startswith("Lower bound")
    _node(mixed, "wan")["collected"] = True                            # the control: a measurement is written as is
    write_failure_impact_sheet(wb, mixed["failure_impact"], mixed)
    src = mixed["failure_impact"][_k(mixed, "gw")]
    assert {r[1]: r for r in _sheet_rows(wb[FAILURE_SHEET_NAME])}["gw"] == [
        src["severity"], "gw", src["vlans_impacted"], src["stranded"], src["hard"], src["backup"], src["fhrp"],
        src["detail"]]


def _summary_text(wb):
    from cisco_toolkit.excel import EXEC_SUMMARY_SHEET_NAME
    return "\n".join(str(c.value) for row in wb[EXEC_SUMMARY_SHEET_NAME].iter_rows() for c in row
                     if c.value is not None)


def test_executive_summary_ranks_published_rows_and_lower_bound_floors(mixed):
    from openpyxl import Workbook
    from cisco_toolkit.excel import write_executive_summary_sheet
    gw = _verdicts(mixed)[_k(mixed, "gw")]
    wb = Workbook()
    write_executive_summary_sheet(wb, [], [], [], mixed["failure_impact"], impact_evidence=mixed)
    text = _summary_text(wb)
    # gw strands at least 1 (a lower bound): ranked by that floor and written as one; acc (a bounded zero) and the
    # INDETERMINATE x1 / x2 are not ranked, and are named
    assert f"gw is the top keystone — its loss strands ≥ 1 endpoint(s) ({gw.summary})." in text, text
    assert "Not ranked — 3 switch(es)" in text and "x1 (" in text and "acc (" in text
    assert "No switch's blast radius is a measurement" not in text
    # nothing ranks at all: every zero below is withheld (bounded) or held, so no switch is named a keystone
    zero = copy.deepcopy(mixed)
    for row in zero["failure_impact"]:
        row["stranded"] = 0
    wb = Workbook()
    write_executive_summary_sheet(wb, [], [], [], zero["failure_impact"], impact_evidence=zero)
    text = _summary_text(wb)
    assert "Not ranked — 4 switch(es)" in text and "top keystone" not in text.lower()
    assert "No switch's blast radius is a measurement on this evidence (4 not ranked)" in text
    _node(mixed, "wan")["collected"] = True                           # the control: gw is a measurement
    wb = Workbook()
    write_executive_summary_sheet(wb, [], [], [], mixed["failure_impact"], impact_evidence=mixed)
    text = _summary_text(wb)
    assert "gw is the top keystone — its loss strands 1 endpoint(s). Harden" in text and "Not ranked — 3" in text


def test_executive_summary_never_names_a_measured_zero_over_a_lower_bound_keystone(keystone):
    """The review's P2-1 counterexample: ranking on the row verdict dropped core1 (High, at least 300 stranded) and
    named acc1 (Info, 0) the top keystone."""
    from openpyxl import Workbook
    from cisco_toolkit.excel import EXEC_SUMMARY_SHEET_NAME, write_executive_summary_sheet
    core1 = _verdicts(keystone)[_k(keystone, "core1")]
    wb = Workbook()
    write_executive_summary_sheet(wb, [], [], [], keystone["failure_impact"], impact_evidence=keystone)
    text = _summary_text(wb)
    assert "acc1 is the top keystone" not in text
    assert f"core1 is the top keystone — its loss strands ≥ 300 endpoint(s) ({core1.summary})." in text, text
    table = [[c.value for c in row][:5] for row in wb[EXEC_SUMMARY_SHEET_NAME].iter_rows()]
    table = [row for row in table if isinstance(row[0], int) and row[1] in ("core1", "acc1")]
    assert table == [[1, "core1", "High (lower bound)", "≥ 300 (lower bound)", "≥ 1 (lower bound)"],
                     [2, "acc1", "Info", 0, 0]], table
    assert "Not ranked" not in text                                   # nothing here is withheld


def test_architecture_review_ranks_a_lower_bound_by_its_floor_and_never_conforms_over_a_withheld_zero(mixed):
    from cisco_toolkit.archreview import compute_architecture_review
    from impact_fixtures import assessable

    def res4(snap):
        return next(c for c in compute_architecture_review(snap)["checks"] if c["id"] == "RES-4")
    zero = copy.deepcopy(mixed)
    for row in zero["failure_impact"]:
        row["stranded"] = 0
    c = res4(zero)
    assert c["verdict"] == "not-assessable", c                       # INDETERMINATE + bounded zeros: never conforms
    assert "not graded" in c["observed"] and "x1" in c["observed"]
    gw = _verdicts(mixed)[_k(mixed, "gw")]
    c = res4(mixed)                                                  # gw strands at least 1: a keystone, as a floor
    assert c["verdict"] == "advisory" and c["evidence"] == ["gw"], c
    assert c["observed"].startswith(f"Losing gw strands ≥ 1 endpoint(s) ({gw.summary}).")
    assert "3 simulated device(s) are not graded" in c["observed"] and "x1 (" in c["observed"]
    _node(mixed, "wan")["collected"] = True
    c = res4(mixed)
    assert c["verdict"] == "advisory" and c["evidence"] == ["gw"], c
    assert c["observed"].startswith("Losing gw strands 1 endpoint(s).") and "3 simulated device(s)" in c["observed"]
    # a floor in the hidden tail is disclosed as one, and the band reads its floor
    fi = [{"host": f"h{i}", "severity": "High", "stranded": 70 - i, "vlans_impacted": 1,
           "detail": f"VLAN {i}: Hard partition"} for i in range(7)]
    snap = assessable({"failure_impact": fi})
    snap["cable_map"] = {"nodes": [{"host": "rtr", "collected": False, "kind": "router"}],
                         "cables": [{"a": "h6", "b": "rtr"}]}
    c = res4(snap)
    assert c["observed"].startswith("Losing h0 strands 70 endpoint(s); h1 strands 69 endpoint(s);"), c
    assert c["observed"].endswith("and 2 further device(s) strand 64-65 endpoint(s) each "
                                  "(1 of them only as a lower bound)."), c
    assert len(c["evidence"]) == 7


def test_architecture_review_ranks_the_lower_bound_keystone_over_a_measured_zero(keystone):
    from cisco_toolkit.archreview import compute_architecture_review
    core1 = _verdicts(keystone)[_k(keystone, "core1")]
    c = next(c for c in compute_architecture_review(keystone)["checks"] if c["id"] == "RES-4")
    assert c["verdict"] == "advisory" and c["evidence"] == ["core1"], c
    assert c["observed"] == f"Losing core1 strands ≥ 300 endpoint(s) ({core1.summary}).", c


def _design_text(snap, path):
    from docx import Document
    from cisco_toolkit.design import write_design_doc_docx
    write_design_doc_docx(path, snap, "W33")
    doc = Document(path)
    return [p.text for p in doc.paragraphs] + [" | ".join(c.text for c in r.cells) for t in doc.tables for r in t.rows]


def _deck_text(snap, path):
    from pptx import Presentation
    from cisco_toolkit.deck import write_executive_deck_pptx
    write_executive_deck_pptx(path, snap, "W33")
    return "\n".join(sh.text_frame.text for sl in Presentation(path).slides for sh in sl.shapes if sh.has_text_frame)


def test_design_and_deck_rank_published_rows_and_lower_bound_floors(mixed, tmp_path):
    pytest.importorskip("docx")
    pytest.importorskip("pptx")
    gw = _verdicts(mixed)[_k(mixed, "gw")]
    blocks = _design_text(mixed, str(tmp_path / "held.docx"))
    assert [b for b in blocks if b.startswith("Concentrated dependency:") and f"gw (≥ 1 endpoints; {gw.summary})" in b]
    assert [b for b in blocks if b.startswith("Blast radius not a measurement:") and "3 switch(es)" in b
            and "x1 (" in b]
    assert "Keystone devices (strand endpoints if lost) | 1" in blocks
    assert "…blast radius not a measurement (not ranked; §2.1) | 3" in blocks
    deck = _deck_text(mixed, str(tmp_path / "held.pptx"))
    assert "≥ 1" in deck and "stranded (lower bound)" in deck and "3 switch(es) not ranked" in deck
    assert "well distributed" not in deck and "Blast radius INDETERMINATE — " not in deck
    _node(mixed, "wan")["collected"] = True
    blocks = _design_text(mixed, str(tmp_path / "ranked.docx"))
    assert [b for b in blocks if b.startswith("Concentrated dependency:") and "gw (1 endpoints)" in b]
    assert "Keystone devices (strand endpoints if lost) | 1" in blocks
    assert "…blast radius not a measurement (not ranked; §2.1) | 3" in blocks
    deck = _deck_text(mixed, str(tmp_path / "ranked.pptx"))
    assert "3 switch(es) not ranked" in deck and "well distributed" not in deck
    assert "stranded (lower bound)" not in deck
    # nothing ranks: every row is held or a bounded zero, so the slide reads as a coverage gap
    zero = copy.deepcopy(mixed)
    _node(zero, "wan")["collected"] = False
    for row in zero["failure_impact"]:
        row["stranded"] = 0
    blocks = _design_text(zero, str(tmp_path / "zero.docx"))
    assert not [b for b in blocks if b.startswith("Concentrated dependency:")]
    assert "Keystone devices (strand endpoints if lost) | 0" in blocks
    deck = _deck_text(zero, str(tmp_path / "zero.pptx"))
    assert "well distributed" not in deck and "Blast radius INDETERMINATE — 4 switch(es)" in deck


def test_design_and_deck_rank_the_lower_bound_keystone(keystone, tmp_path):
    """The P2-1 counterexample on the two documents: core1 is the keystone, at least 300, never dropped."""
    pytest.importorskip("docx")
    pytest.importorskip("pptx")
    core1 = _verdicts(keystone)[_k(keystone, "core1")]
    blocks = _design_text(keystone, str(tmp_path / "k.docx"))
    assert [b for b in blocks if b.startswith("Concentrated dependency:")
            and f"core1 (≥ 300 endpoints; {core1.summary})" in b and "acc1" not in b]
    assert "Keystone devices (strand endpoints if lost) | 1" in blocks
    assert not [b for b in blocks if b.startswith("Blast radius not a measurement:")
                or b.startswith("…blast radius not a measurement")]
    deck = _deck_text(keystone, str(tmp_path / "k.pptx"))
    assert "≥ 300" in deck and "stranded (lower bound)" in deck and core1.raw["detail"] in deck
    assert "well distributed" not in deck and "not ranked" not in deck and "Blast radius INDETERMINATE" not in deck


def _dossiers(snap, lifecycle=None, doc=None):
    kwargs = {"failure_impact": snap["failure_impact"], "lifecycle_risk": lifecycle}
    if doc is not None:
        kwargs["failure_impact_assessability"] = doc
    return {d["host"]: d for d in analyze.compute_device_dossiers(**kwargs)["per_device"]}


#: The dossier fields the impact term scores or drives (the multiplicand, the band and the compound patterns).
_SCORED = ("impact_score", "impact_severity", "stranded", "vlans_impacted", "exposure_score", "n_na", "risk_index",
           "risk_band")


def test_the_dossier_impact_term_never_reads_a_withheld_row_as_clean(mixed):
    plain = _dossiers(mixed)
    assert plain["x1"]["impact_severity"] == "Info" and "impact_assessability" not in plain["x1"]   # unchanged
    dd = _dossiers(mixed, doc=ia.assessment_document(mixed))
    for host in ("x1", "x2"):
        d = dd[host]
        assert d["impact_assessability"]["assessable"] == ia.NOT_ASSESSED, d
        assert d["impact_assessability"]["pointer"] == f"/failure_impact/{_k(mixed, host)}"
        assert all(d[key] == plain[host][key] for key in _SCORED), host      # the score is never moved
        assert "no modeled reachability impact" not in d["verdict"]
        assert d["verdict"] != "No stacked risk — routine migration handling."      # never routine by default
        assert "the blast radius is not assessed" in d["verdict"], d["verdict"]
        assert plain[host]["verdict"] == "No stacked risk — routine migration handling."  # the pre-W33 clean bill
    acc = dd["acc"]
    assert acc["impact_assessability"]["assessable"] == ia.LOWER_BOUND
    assert "the blast radius is only a lower bound" in acc["verdict"], acc["verdict"]
    gw = dd["gw"]
    assert gw["impact_assessability"]["assessable"] == ia.LOWER_BOUND and gw["impact_severity"] == "High"
    assert gw["impact_score"] == plain["gw"]["impact_score"]          # a lower bound's High still floors the term


_PAST_LDOS = {"per_device": [{"host": "core1", "model": "WS-C3850-48P", "band": "Past-LDoS"}]}


def test_a_withheld_impact_never_lowers_dossier_risk(keystone):
    """Review P2-2 counterexample: a past-LDoS switch whose stored row is High with hundreds stranded, but whose
    scoped interface running-config was never captured, so the owner holds the row. Before the fix the hold dropped
    the term to the absent-row floor: impact 1, severity '—', CR-01 gone and the band no longer Severe. The score is
    the pre-W33 one; the hold is disclosed in the CR-01 basis, the verdict and `impact_assessability`."""
    k = _k(keystone, "core1")
    for port in keystone["interfaces"]["core1"].values():
        port.pop("run_config_observed", None)                    # what html.sparsify_interfaces writes for false
    doc = ia.assessment_document(keystone)
    assert doc["rows"][k]["assessable"] == ia.NOT_ASSESSED and doc["rows"][k]["codes"] == ["no_run_config"]
    plain = _dossiers(keystone, _PAST_LDOS)["core1"]
    assert [c["code"] for c in plain["compound"]] == ["CR-01"] and plain["compound"][0]["severity"] == "Critical"
    assert plain["risk_band"] == "Severe" and plain["impact_score"] == 10, plain
    for supplied, why in ((doc, doc["rows"][k]["why"]),
                          (ia.unavailable_document(), "no assessability verdict names this device's row")):
        held = _dossiers(keystone, _PAST_LDOS, supplied)["core1"]
        assert all(held[key] == plain[key] for key in _SCORED), {key: (held[key], plain[key]) for key in _SCORED}
        assert [(c["code"], c["severity"]) for c in held["compound"]] == [("CR-01", "Critical")], held["compound"]
        basis = held["compound"][0]["basis"]
        assert "unverified" in basis and "300 endpoint(s)" in basis and "the blast radius is not assessed" in basis
        assert "the blast radius is not assessed" in held["verdict"] and held["verdict"] != plain["verdict"]
        assert held["impact_assessability"] == {"assessable": ia.NOT_ASSESSED, "why": why,
                                                "pointer": f"/failure_impact/{k}"}, held["impact_assessability"]


def test_a_published_row_scores_and_reads_exactly_as_before(keystone):
    _node(keystone, "wan1")["collected"] = True                      # core1 is now the producer's measurement
    plain = _dossiers(keystone, _PAST_LDOS)
    dd = _dossiers(keystone, _PAST_LDOS, ia.assessment_document(keystone))
    for host in ("core1", "acc1"):
        assert dd[host]["impact_assessability"] == {"assessable": ia.PUBLISHED, "why": "",
                                                    "pointer": f"/failure_impact/{_k(keystone, host)}"}
        assert {key: value for key, value in dd[host].items() if key != "impact_assessability"} == plain[host]


def test_a_lower_bound_row_scores_by_its_floor(keystone):
    plain = _dossiers(keystone, _PAST_LDOS)["core1"]
    d = _dossiers(keystone, _PAST_LDOS, ia.assessment_document(keystone))["core1"]
    assert d["impact_assessability"]["assessable"] == ia.LOWER_BOUND
    assert all(d[key] == plain[key] for key in _SCORED)
    assert [c["code"] for c in d["compound"]] == ["CR-01"]
    assert ("removal strands at least 300 endpoint(s) across at least 1 VLAN(s); the blast radius is only a lower "
            "bound") in d["compound"][0]["basis"], d["compound"][0]["basis"]


def test_the_explorer_embed_carries_the_owners_verdicts(mixed):
    from cisco_toolkit.html import _slim_for_embed
    out = _slim_for_embed(mixed)
    verdicts = ia.assess_failure_impact(mixed)
    assert [v["assessable"] for v in out["failure_impact_assessability"]] == [v.assessable for v in verdicts]
    assert [v.get("floor") for v in out["failure_impact_assessability"]] == [ia.ranking_floor(v) for v in verdicts]
    assert out["failure_impact_assessability"][_k(mixed, "gw")]["floor"] == 1
    assert "failure_impact_assessability" not in _slim_for_embed({"interfaces": {}})
    json.dumps(out, allow_nan=False)


NODE = shutil.which("node")


@pytest.mark.skipif(not NODE, reason="node is not available")
def test_the_explorer_keystone_card_ranks_published_rows_and_lower_bound_floors(mixed, keystone, tmp_path):
    """Executed, not grepped: the embedded explorer script ranks a stored row when the engine published it, or by the
    positive stranded floor the engine embedded for a lower bound (marked as a floor), and names every other switch
    (a held row, a row without a verdict, a scanned switch without a row) with the reason."""
    import test_explorer_render_safety as rs                    # the shared DOM stub and script runner
    from cisco_toolkit.html import _slim_for_embed
    driver = """
      const P=JSON.parse(require('fs').readFileSync(process.argv[2],'utf-8'));
      const out={};
      for(const k of Object.keys(P)){__EV('load')(P[k],'T',false);
        out[k]={ranked:__EV('cockpitKeystones')().map(r=>[r.host,r.blast]),card:__EV('keystoneCard')()};}
      process.stdout.write(JSON.stringify(out));
    """
    gw_summary = _verdicts(mixed)[_k(mixed, "gw")].summary
    core1_summary = _verdicts(keystone)[_k(keystone, "core1")].summary
    held = _slim_for_embed(mixed)
    published = copy.deepcopy(mixed)
    _node(published, "wan")["collected"] = True
    published = _slim_for_embed(published)
    no_verdicts = copy.deepcopy(published)
    del no_verdicts["failure_impact_assessability"]
    out = rs._run(driver, tmp_path, payload={"held": held, "published": published, "no_verdicts": no_verdicts,
                                             "keystone": _slim_for_embed(keystone)})
    assert out["held"]["ranked"] == [["gw", 1]], out["held"]           # the floor ranks; nothing else is published
    assert '≥ <span class="num" data-to="1">' in out["held"]["card"]
    assert f"endpoints stranded if it fails ({gw_summary})" in out["held"]["card"]
    assert "3 switches not ranked" in out["held"]["card"] and "x1" in out["held"]["card"]
    assert out["published"]["ranked"][0] == ["gw", 1], out["published"]
    assert "3 switches not ranked" in out["published"]["card"] and "≥ " not in out["published"]["card"]
    assert out["no_verdicts"]["ranked"] == [] and "no engine assessability verdict" in out["no_verdicts"]["card"]
    assert out["keystone"]["ranked"][0] == ["core1", 300], out["keystone"]
    assert '≥ <span class="num" data-to="300">' in out["keystone"]["card"]
    assert f"endpoints stranded if it fails ({core1_summary})" in out["keystone"]["card"]
    assert "not ranked" not in out["keystone"]["card"]


# --------------------------------------------------------------------------------------------------
# 4. W32's per-row blind-link count is one more owner rule (W45 integration): every surface reads it from here
# --------------------------------------------------------------------------------------------------
def _blind_fleet():
    """W32's fleet (tests/test_ui_projection_device_impact.py::_blind_fleet, copied, not imported). `g1` and `g2` are
    FHRP peers gatewaying VLAN 10 for `acc` (removing `g1` is FHRP-covered: Low); `g2` is also the sole gateway of
    VLAN 30 for `acc` (removing it is a hard partition: High). Every trunk among those three carries VLAN evidence.
    `g1` also trunks to `x` over a link with NO trunk/STP evidence on either end, and `x` has nothing else."""
    return {"g1": {"Gi1": _trunk("Gi1", "acc", "Gi1", "10"), "Gi9": _trunk("Gi9", "x", "Gi1"),
                   "Vlan10": _svi(10, "10.10.0.2/24", "Active")},
            "g2": {"Gi1": _trunk("Gi1", "acc", "Gi2", "10,30"), "Vlan10": _svi(10, "10.10.0.3/24", "Standby"),
                   "Vlan30": _svi(30, "10.30.0.1/24")},
            "acc": {"Gi1": _trunk("Gi1", "g1", "Gi1", "10"), "Gi2": _trunk("Gi2", "g2", "Gi1", "10,30"),
                    "Gi10": _access("Gi10", 10, "0000.0000.000a"), "Gi30": _access("Gi30", 30, "0000.0000.001e")},
            "x": {"Gi1": _trunk("Gi1", "g1", "Gi9")}}


@pytest.fixture()
def blind():
    """The REAL producers over :func:`_blind_fleet`, with the counts the producer writes read back."""
    snap = _snapshot(_blind_fleet())
    assert {row["host"]: row["blind_links"] for row in snap["failure_impact"]} == {"g1": 1, "g2": 0, "acc": 0, "x": 1}
    return snap


def test_a_positive_blind_link_count_is_an_owner_bound_and_nothing_simulated_is_an_owner_hold(blind):
    v = _verdicts(blind)
    g1, g2, acc, x = (v[_k(blind, h)] for h in ("g1", "g2", "acc", "x"))
    k = _k(blind, "g1")
    src = blind["failure_impact"][k]
    assert (src["severity"], src["vlans_impacted"], src["fhrp"]) == ("Low", 1, 1), src
    assert src["detail"] == "VLAN 10: FHRP-covered", src
    # g1 simulated in part beside an evidence-less link: a lower bound, decided by the owner, citing the count
    assert g1.assessable == ia.LOWER_BOUND and g1.codes == ["blind_links"], g1.as_dict()
    assert g1.why == ia.CODE_PHRASES["blind_links"].format(n=1)
    assert f"/failure_impact/{k}/blind_links" in g1.pointers
    assert [f for f in MEASURES if g1.withholds(f)] == ["severity", "stranded", "hard", "backup"]
    assert not g1.withholds("detail")                          # a per-VLAN detail lists what was simulated
    bound = g1.facts.bounds[0]
    assert bound.reason == "not collected: " + ia.R_BLIND.format(n=1) and bound.detail_reason.endswith(ia.R_BLIND_TAIL)
    assert ia.ranking_floor(g1) is None and not ia.ranks(g1)  # a bounded zero is never a keystone
    # the fully evidenced controls are measurements
    assert g2.published and acc.published and g2.codes == acc.codes == []
    # x simulated nothing: the producer's INDETERMINATE hold wins, and the count is still cited beside it
    kx = _k(blind, "x")
    assert x.assessable == ia.NOT_ASSESSED and x.codes == ["indeterminate"], x.as_dict()
    assert f"/failure_impact/{kx}/blind_links" in x.pointers
    # the same row without its prose marker: the count alone says nothing was simulated
    blind["failure_impact"][kx]["detail"] = "No reachability impact from removing this switch (within the scan)."
    x = _verdicts(blind)[kx]
    assert x.codes == ["blind_links_only"] and x.facts.hold.reason == ia.R_BLIND_ONLY.format(n=1), x.as_dict()
    assert all(x.withholds(f) for f in MEASURES + ("detail",))
    # a count of 0 is the producer's measurement of none: the clean bill is published
    blind["failure_impact"][kx]["blind_links"] = 0
    assert _verdicts(blind)[kx].published


@pytest.mark.parametrize("bad", ["1", None, -1, True, 1.5, [1]])
def test_a_blind_link_count_that_is_not_a_count_is_an_owner_hold(blind, bad):
    k = _k(blind, "g2")
    blind["failure_impact"][k]["blind_links"] = bad
    v = _verdicts(blind)[k]
    assert v.assessable == ia.NOT_ASSESSED and v.codes == ["blind_links_unreadable"], v.as_dict()
    assert v.facts.hold.state == UV and v.facts.hold.reason == ia.R_BLIND_UNREAD
    assert f"/failure_impact/{k}/blind_links" in v.pointers
    assert all(v.withholds(f) for f in MEASURES + ("detail",)) and not v.withholds("host")
    # an unreadable off-scan count is read first (first match wins)
    blind["failure_impact"][k]["off_scan_gw_vlans"] = "1"
    assert _verdicts(blind)[k].codes == ["off_scan_unreadable"]


def test_a_row_older_than_the_blind_link_count_is_an_owner_lower_bound(blind):
    for row in blind["failure_impact"]:
        del row["blind_links"]
    v = _verdicts(blind)
    g2, acc, x = (v[_k(blind, h)] for h in ("g2", "acc", "x"))
    # High and positive counts stay as lower bounds; the zero and the clean bill are withheld, citing the row itself
    assert g2.assessable == ia.LOWER_BOUND and g2.codes == ["blind_links_legacy"], g2.as_dict()
    assert [f for f in MEASURES if g2.withholds(f)] == ["backup"] and ia.ranking_floor(g2) == 1
    assert g2.facts.bounds[0].witnesses == [("witness", ("failure_impact", _k(blind, "g2")))]
    assert acc.assessable == ia.LOWER_BOUND and all(acc.withholds(f) for f in MEASURES + ("detail",))
    assert x.codes == ["indeterminate"]                       # the producer's own disclosure still wins
    # a row older than the off-scan marker as well keeps that hold, which wins over the bound
    del blind["failure_impact"][_k(blind, "acc")]["off_scan_gw_vlans"]
    assert _verdicts(blind)[_k(blind, "acc")].codes == ["legacy_row"]


def test_every_owner_consumer_reads_the_blind_link_rule_from_the_owner(blind):
    """One structural rule, no parallel path: the projection's cells, the workbook sheet, RES-4, the dossier and the
    explorer embed all follow the owner's verdict on the blind-link-bounded row."""
    from openpyxl import Workbook
    from cisco_toolkit.archreview import compute_architecture_review
    from cisco_toolkit.excel import FAILURE_SHEET_NAME, write_failure_impact_sheet
    from cisco_toolkit.html import _slim_for_embed
    k = _k(blind, "g1")
    g1 = _verdicts(blind)[k]
    # the projection: each cell withheld exactly when the owner withholds it, with the owner's own bound reasons
    _assert_cells_follow_the_owner(blind)
    row = _topology(blind)[k]
    assert row["severity"]["reason"] == g1.facts.bounds[0].severity_reason
    assert (f"/failure_impact/{k}/blind_links", "witness") in _refs(row["fhrp"])
    # the workbook's Failure Impact sheet
    wb = Workbook()
    write_failure_impact_sheet(wb, blind["failure_impact"], blind)
    severity, _h, vlans, stranded, hard, backup, fhrp, detail = {
        r[1]: r for r in _sheet_rows(wb[FAILURE_SHEET_NAME])}["g1"]
    assert severity == stranded == hard == backup == ia.NOT_ASSESSED_CELL and vlans == fhrp == "≥ 1"
    assert detail == ia.table_detail(g1) and detail.startswith("Lower bound") and "VLAN 10: FHRP-covered" in detail
    # RES-4: g1's bounded zero is disclosed, never graded or ranked
    c = next(c for c in compute_architecture_review(blind)["checks"] if c["id"] == "RES-4")
    assert c["verdict"] == "advisory" and c["evidence"] == ["g2"], c
    assert f"g1 ({g1.summary})" in c["observed"] and "2 simulated device(s) are not graded" in c["observed"], c
    # the dossier: disclosed as a lower bound, never scored lower
    plain = _dossiers(blind)["g1"]
    d = _dossiers(blind, doc=ia.assessment_document(blind))["g1"]
    assert d["impact_assessability"] == {"assessable": ia.LOWER_BOUND, "why": g1.why, "pointer": f"/failure_impact/{k}"}
    assert all(d[key] == plain[key] for key in _SCORED)
    assert "the blast radius is only a lower bound" in d["verdict"], d["verdict"]
    # the explorer embed
    embed = _slim_for_embed(blind)["failure_impact_assessability"][k]
    assert embed == {"assessable": ia.LOWER_BOUND, "summary": g1.summary}, embed

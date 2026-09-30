"""G15: a fleet in which NO device was health-scored must never publish a measured-looking headline.

`compute_executive_brief` averaged only the genuinely-scored rows (an 'Insufficient Data' row keeps
its near-perfect, deduction-free score) but fell back to a hard ``0`` when that set was empty. The
fleet posture then published ``avg_health: 0`` and a "Fleet health" axis of "0/100 avg" at severity
``Low`` -- absence rendered as a measurement -- and the posture statement concluded "no top-tier
blockers ... proceed with the standard wave plan" about a fleet nothing had looked at.

The fix is a CONTRACT, not a one-line default, because three consumers recomputed a fallback mean
over EVERY health row (the 'Insufficient Data' rows included) whenever the canonical value was not a
number: turning the producer's 0 into ``None`` alone would have rendered a fabricated "100 / 100".
These tests pin both halves: the producer abstains, and no consumer re-derives a number from the
unscored rows.

The scored path is pinned too -- it must stay byte-for-byte what it was (the golden snapshot and the
sample fleet are regenerated from it).
"""
import json
import pathlib
import re
import shutil
import subprocess

import pytest

from cisco_toolkit import ssot
from cisco_toolkit.analyze import compute_executive_brief

ROOT = pathlib.Path(__file__).resolve().parent.parent
EXPLORER = ROOT / "cisco_toolkit" / "blast_radius_explorer.html"
SAMPLE = ROOT / "webapp" / "sample_data" / "sample_fleet.snapshot.json"
NODE = shutil.which("node")


def _insufficient(n=5, score=100):
    """`n` never-collected rows exactly as compute_health_scores emits them: band 'Insufficient Data'
    with the deduction-free score the scorer still assigns (no evidence -> no deductions)."""
    return [{"switch": f"sw{i}", "band": "Insufficient Data", "score": score} for i in range(n)]


def _axis(brief, name):
    return next(a for a in brief["axes"] if a["axis"] == name)


# --------------------------------------------------------------------------------------------------
# producer: posture + Fleet-health axis
# --------------------------------------------------------------------------------------------------
def test_all_insufficient_fleet_publishes_no_average():
    b = compute_executive_brief(health_scores=_insufficient())
    p = b["posture"]
    assert p["avg_health"] is None, "no device was scored: a 0 (or any number) is a fabricated measurement"
    assert p["worst_band"] is None, "no band was observed: '' is not a published fact"
    # A zero band count over a fleet in which nothing was scored is the same absence as the average: a
    # user assertion `executive_brief.posture.n_critical == 0` PASSED over it. None abstains.
    assert p["n_critical"] is None and p["n_poor"] is None
    assert p["n_scored"] == 0
    assert p["not_assessed"] == "all_insufficient_data"
    assert b["scale"]["n_devices"] == 5                        # the inventory is still counted


def test_empty_health_input_publishes_no_average():
    b = compute_executive_brief()
    p = b["posture"]
    assert p["avg_health"] is None and p["worst_band"] is None
    assert p["n_critical"] is None and p["n_poor"] is None
    assert p["n_scored"] == 0 and p["not_assessed"] == "no_health_rows"
    # len([]) is not an inventory: a crashed Health Scores phase published `n_devices: 0` beside a
    # collection inventory of N, and every docx printed "(read-only evidence; -N not reached)".
    assert b["scale"]["n_devices"] is None


def test_fleet_health_axis_is_not_assessed_not_low():
    b = compute_executive_brief(health_scores=_insufficient())
    ax = _axis(b, "Fleet health")
    assert ax["severity"] == "Info", "Low is the clean-fleet value; nothing was measured"
    assert "NOT ASSESSED" in ax["headline"] and "0 of 5" in ax["headline"]
    assert "/100" not in ax["headline"], ax["headline"]
    assert "0 switch(es) assessed" in ax["detail"] and "5 not collected" in ax["detail"]
    empty = _axis(compute_executive_brief(), "Fleet health")
    assert empty["severity"] == "Info" and "NOT ASSESSED" in empty["headline"]
    assert "/100" not in empty["headline"]


@pytest.mark.parametrize("kwargs", [dict(health_scores=_insufficient()), dict()],
                         ids=["all-insufficient", "no-health-rows"])
def test_posture_statement_never_clears_an_unscored_fleet(kwargs):
    ps = compute_executive_brief(**kwargs)["posture_statement"]
    assert "no top-tier blockers" not in ps and "proceed with the standard wave plan" not in ps, ps
    assert "fleet health is NOT ASSESSED" in ps and "not assessed, not clear" in ps, ps


def test_not_assessed_flag_survives_a_crowded_flag_list():
    """The flag must sit inside the rendered flags[:4] even when every other flag fires."""
    b = compute_executive_brief(
        health_scores=_insufficient(3),
        migration_readiness=[{"readiness": "NOT READY"}],
        lifecycle_risk={"summary": {"n_devices": 3, "n_past_ldos": 1, "n_near": 0, "n_unknown": 1}},
        segmentation={"summary": {"n_gateways": 2, "flat": True}},
        multicast_intelligence={"summary": {"n_groups": 4, "n_mac_clashes": 1, "n_ptp_clocks": 2,
                                            "n_ptp_dormant": 2}})
    shown = b["posture_statement"].split("; +")[0]
    assert "fleet health is NOT ASSESSED" in shown, b["posture_statement"]
    assert "wave(s)" not in b["posture_statement"]


def test_no_health_call_keeps_the_not_ready_flag_visible():
    ps = compute_executive_brief(
        migration_readiness=[{"readiness": "NOT READY"}, {"readiness": "READY"}])["posture_statement"]
    assert "move-group(s) are NOT READY" in ps.split("; +")[0]
    assert "wave(s)" not in ps


def test_a_critical_band_without_a_numeric_score_still_leads():
    """An adverse OBSERVED band is still a finding even when its row carried no usable score."""
    b = compute_executive_brief(health_scores=[{"switch": "a", "band": "Critical", "score": None}])
    assert b["posture"]["avg_health"] is None and b["posture"]["worst_band"] == "Critical"
    assert b["posture"]["n_critical"] == 1, "an OBSERVED adverse band is a count, not an abstention"
    assert b["posture"]["n_poor"] is None
    assert _axis(b, "Fleet health")["severity"] == "Critical"


# --------------------------------------------------------------------------------------------------
# producer: the fully-scored path is unchanged (the sample fleet regenerates from it)
# --------------------------------------------------------------------------------------------------
def test_fully_scored_path_shape_and_text_are_unchanged():
    b = compute_executive_brief(health_scores=[
        {"switch": "a", "band": "Critical", "score": 20},
        {"switch": "b", "band": "Good", "score": 80}])
    assert b["posture"] == {"avg_health": 50, "n_critical": 1, "n_poor": 0, "worst_band": "Critical"}
    ax = _axis(b, "Fleet health")
    assert ax == {"axis": "Fleet health", "severity": "Critical",
                  "headline": "50/100 avg · 1 Critical, 0 Poor band",
                  "detail": "2 switch(es) assessed."}
    assert set(b) == {"scale", "posture", "axes", "top_gating", "posture_statement"}
    assert "NOT ASSESSED" not in b["posture_statement"]


def test_sample_fleet_brief_recomputes_byte_equal():
    """The sample fleet's embedded brief, recomputed from the sample's OWN inputs by the current
    producer, must equal what is stored: a change here that moved it would require regenerating the
    committed sample, which this branch must not do."""
    s = json.loads(SAMPLE.read_text(encoding="utf-8"))
    eb = s["executive_brief"]
    assert set(eb["posture"]) == {"avg_health", "n_critical", "n_poor", "worst_band"}   # no abstention keys
    b = compute_executive_brief(
        health_scores=s.get("health_scores"), punchlist=s.get("punchlist"),
        migration_readiness=s.get("migration_readiness"),
        application_intelligence=s.get("application_intelligence"), lifecycle_risk=s.get("lifecycle_risk"),
        segmentation=s.get("segmentation"), multicast_intelligence=s.get("multicast_intelligence"),
        remediation_plan=s.get("remediation_plan"), syslog_intelligence=s.get("syslog_intelligence"),
        qos_audit=s.get("qos_audit"), software_risk=s.get("software_risk"),
        platform_health=s.get("platform_health"), device_dossiers=s.get("device_dossiers"),
        endpoint_identity=s.get("endpoint_identity"))
    for key in ("posture", "axes", "top_gating", "posture_statement"):
        assert b[key] == eb[key], key
    for key in ("n_devices", "n_domains", "n_endpoints"):     # n_vlans / n_collected are injected by main()
        assert b["scale"][key] == eb["scale"][key], key
    assert ssot.summary(s) == eb["ssot"]                      # n_facts / n_checked unchanged
    assert ssot.reconcile(s) == []
    assert ssot.fleet_avg_health(s)["state"] == "measured"


# --------------------------------------------------------------------------------------------------
# producer: PARTIAL coverage is not the clean-fleet Low either (the EoL axis's own rule)
# --------------------------------------------------------------------------------------------------
def test_partially_scored_fleet_is_not_low_and_never_proceeds():
    """1 Excellent row + 99 'Insufficient Data' rows read 'Low · 97/100 avg' and the posture said
    'proceed with the standard wave plan' -- a clean verdict on 1% coverage."""
    b = compute_executive_brief(health_scores=[{"switch": "x", "band": "Excellent", "score": 97}]
                                + _insufficient(99))
    ax = _axis(b, "Fleet health")
    assert ax["severity"] == "Info", ax
    assert ax["headline"].startswith("97/100 avg") and "99 of 100 NOT ASSESSED" in ax["headline"], ax
    ps = b["posture_statement"]
    assert "proceed with the standard wave plan" not in ps, ps
    assert "fleet health is NOT ASSESSED on 99 of 100 switch(es)" in ps.split("; +")[0], ps
    assert b["posture"]["avg_health"] == 97                    # the measured mean is still published


def test_partial_coverage_keeps_an_adverse_band_leading():
    b = compute_executive_brief(health_scores=[{"switch": "a", "band": "Critical", "score": 20},
                                               {"switch": "b", "band": "Good", "score": 80}]
                                + _insufficient(1))
    ax = _axis(b, "Fleet health")
    assert ax["severity"] == "Critical"
    assert ax["headline"] == "50/100 avg · 1 Critical, 0 Poor band · 1 of 3 NOT ASSESSED", ax
    assert "fleet health is NOT ASSESSED on 1 of 3 switch(es)" in b["posture_statement"]


def test_partially_unassessed_asset_register_is_not_low():
    b = compute_executive_brief(
        health_scores=[{"switch": "a", "band": "Good", "score": 90}],
        device_dossiers={"summary": {"n_devices": 5, "n_compound": 0, "worst": [],
                                     "bands": {"Severe": 0, "Elevated": 0, "Guarded": 0, "Low": 1,
                                               "Unassessed": 4}}})
    ax = _axis(b, "Asset risk register")
    assert ax["severity"] == "Info", ax
    assert "4 of 5 asset(s) Unassessed (no evidence)" in ax["headline"], ax["headline"]


def test_cutover_sequence_over_no_move_groups_is_not_low():
    b = compute_executive_brief(
        health_scores=[{"switch": "a", "band": "Good", "score": 90}],
        application_intelligence={"summary": {"n_domains": 2, "pilot_domain": "A", "last_domain": "B"}})
    ax = _axis(b, "Cutover sequence")
    assert ax["severity"] == "Info", ax
    assert "move-group readiness not assessed" in ax["headline"], ax["headline"]


@pytest.mark.parametrize("score", [True, float("nan"), float("inf"), 10 ** 400],
                         ids=["bool", "nan", "inf", "huge-int"])
def test_a_non_finite_or_boolean_score_is_not_a_score(score):
    """The producer counted `isinstance(score, (int, float))` -- a bool scored, NaN / inf crashed the
    round() -- while ssot.reconcile uses is_finite_num; the two-way avg check then false-fired."""
    hs = [{"switch": "a", "band": "Good", "score": score}]
    b = compute_executive_brief(health_scores=hs)              # must not raise
    assert b["posture"]["avg_health"] is None
    assert ssot.reconcile({"health_scores": hs, "executive_brief": b}) == []


def test_lifecycle_rollup_without_a_coverage_count_is_not_complete():
    """A lifecycle summary that omits `n_unknown` while its bands do NOT partition the fleet read as
    complete coverage (`lc.get("n_unknown", 0)`) -> Low + 'proceed'."""
    b = compute_executive_brief(health_scores=[{"switch": "a", "band": "Good", "score": 90}],
                                lifecycle_risk={"summary": {"n_devices": 4, "n_past_ldos": 0, "n_near": 0}})
    ax = _axis(b, "Hardware lifecycle (EoL)")
    assert ax["severity"] == "Info", ax
    assert "4 NOT ASSESSED" in ax["headline"], ax["headline"]
    assert "proceed with the standard wave plan" not in b["posture_statement"]
    assert "UNDETERMINED on 4 device(s)" in b["posture_statement"], b["posture_statement"]
    # NON-VACUITY: bands that DO partition the fleet are complete coverage even without the key.
    full = compute_executive_brief(health_scores=[{"switch": "a", "band": "Good", "score": 90}],
                                   lifecycle_risk={"summary": {"n_devices": 4, "n_active": 4}})
    assert _axis(full, "Hardware lifecycle (EoL)")["severity"] == "Low"
    assert "UNDETERMINED" not in full["posture_statement"]


def test_scored_rows_with_no_recognised_band_publish_no_worst_band():
    b = compute_executive_brief(health_scores=[{"switch": "a", "band": "Mystery", "score": 90}])
    assert b["posture"]["avg_health"] == 90 and b["posture"]["worst_band"] is None
    assert ssot.reconcile({"health_scores": [{"switch": "a", "band": "Mystery", "score": 90}],
                           "executive_brief": b}) == []


def test_segmentation_share_over_zero_gateways_is_not_a_percentage():
    f = ssot.segmentation_facts({"segmentation": {"summary": {"n_gateways": 0},
                                                  "gateway_acl": {"n_gateways": 0, "n_with_acl": 0,
                                                                  "coverage_pct": 0.0}}})
    assert f["n_gateways"] == 0 and f["coverage_pct"] is None


# --------------------------------------------------------------------------------------------------
# producer: the other zero-for-nothing members of the class inside compute_executive_brief
# --------------------------------------------------------------------------------------------------
def test_lifecycle_share_over_zero_assessable_is_not_a_percentage():
    b = compute_executive_brief(
        health_scores=[{"switch": "a", "band": "Good", "score": 90}],
        lifecycle_risk={"summary": {"n_devices": 4, "n_unknown": 4, "n_past_ldos": 0, "n_near": 0}})
    head = _axis(b, "Hardware lifecycle (EoL)")["headline"]
    assert "0% of 0" not in head and "share not computable" in head, head


def test_lifecycle_flag_never_prints_zero_percent_beside_a_past_ldos_device():
    """A (malformed) rollup whose assessable denominator is 0 while a Past-LDoS count is published:
    the flag stated '(0% past/near)' beside a real finding."""
    b = compute_executive_brief(
        health_scores=[{"switch": "a", "band": "Good", "score": 90}],
        lifecycle_risk={"summary": {"n_devices": 4, "n_unknown": 4, "n_past_ldos": 1, "n_near": 0}})
    ps = b["posture_statement"]
    assert "hardware end-of-support is a primary driver" in ps
    assert "(0% past/near)" not in ps and "share not computable" in ps, ps


def test_lifecycle_share_with_assessable_devices_is_unchanged():
    b = compute_executive_brief(
        health_scores=[{"switch": "a", "band": "Good", "score": 90}],
        lifecycle_risk={"summary": {"n_devices": 10, "n_unknown": 0, "n_past_ldos": 5, "n_near": 0}})
    assert "(50% of 10 assessable)" in _axis(b, "Hardware lifecycle (EoL)")["headline"]
    assert "(50% past/near)" in b["posture_statement"]


def test_asset_register_of_only_unassessed_assets_is_not_low():
    b = compute_executive_brief(
        health_scores=_insufficient(),
        device_dossiers={"summary": {"n_devices": 5, "n_compound": 0, "worst": [],
                                     "bands": {"Severe": 0, "Elevated": 0, "Guarded": 0, "Low": 0,
                                               "Unassessed": 5}}})
    ax = _axis(b, "Asset risk register")
    assert ax["severity"] == "Info", ax
    assert "all 5 asset(s) Unassessed" in ax["headline"], ax["headline"]


def test_asset_register_with_assessed_assets_is_unchanged():
    b = compute_executive_brief(
        health_scores=[{"switch": "a", "band": "Good", "score": 90}],
        device_dossiers={"summary": {"n_devices": 2, "n_compound": 0, "worst": [],
                                     "bands": {"Severe": 0, "Elevated": 0, "Guarded": 0, "Low": 2,
                                               "Unassessed": 0}}})
    ax = _axis(b, "Asset risk register")
    assert ax["severity"] == "Low" and ax["headline"] == "0 Severe, 0 Elevated of 2 asset(s)"


# --------------------------------------------------------------------------------------------------
# ssot.reconcile: the abstention is verified in BOTH directions
# --------------------------------------------------------------------------------------------------
def _snap(health, posture):
    return {"health_scores": health, "executive_brief": {"posture": posture}}


def test_reconcile_flags_a_number_published_for_an_unscored_fleet():
    """The pre-fix producer's output (and every stored snapshot it wrote)."""
    v = ssot.reconcile(_snap(_insufficient(), {"avg_health": 0, "n_critical": 0, "n_poor": 0,
                                               "worst_band": ""}))
    assert any("avg_health" in x for x in v), v


def test_reconcile_flags_an_abstention_published_for_a_scored_fleet():
    v = ssot.reconcile(_snap([{"switch": "a", "band": "Good", "score": 80}],
                             {"avg_health": None, "n_critical": 0, "n_poor": 0, "worst_band": "Good"}))
    assert any("avg_health" in x for x in v), v


def test_reconcile_flags_a_withheld_worst_band_while_a_band_is_observed():
    v = ssot.reconcile(_snap([{"switch": "a", "band": "Poor", "score": 50}],
                             {"avg_health": 50, "n_critical": 0, "n_poor": 1, "worst_band": None}))
    assert any("worst_band" in x for x in v), v


def test_reconcile_accepts_and_counts_the_honest_abstention():
    b = compute_executive_brief(health_scores=_insufficient())
    snap = {"health_scores": _insufficient(), "executive_brief": b}
    ran = []
    assert ssot.reconcile(snap, _ran=ran) == []
    for path in ("avg_health", "n_critical", "n_poor"):       # verified, not skipped
        assert f"executive_brief.posture.{path}" in ran, (path, ran)
    facts = ssot.canonical_facts(snap)
    assert facts["avg_health"] is None and facts["worst_band"] is None
    assert facts["n_critical"] is None and facts["n_poor"] is None
    assert ssot.summary(snap)["verified"] is True


@pytest.mark.parametrize("rows, why", [
    ([{"switch": "a", "band": "Critical", "score": None}], "carry the Critical band"),
    ([{"switch": "a", "band": "Good", "score": 80}], "are scored"),
], ids=["band-observed", "rows-scored"])
def test_reconcile_flags_a_band_count_withheld_while_measurable(rows, why):
    v = ssot.reconcile(_snap(rows, {"avg_health": 80 if rows[0]["score"] else None,
                                    "n_critical": None, "n_poor": 0,
                                    "worst_band": rows[0]["band"]}))
    assert any("n_critical=None" in x and why in x for x in v), v


def test_summary_does_not_certify_a_check_over_a_crashed_health_phase():
    """With 'Health Scores' failed, `health_scores == []` is the phase's fallback: verifying the
    abstention against it certified `verified: True` over nothing (origin reported verified False)."""
    snap = {"assessment_integrity": {"failed_phases": ["Health Scores"]}, "health_scores": [],
            "executive_brief": compute_executive_brief()}
    ran = []
    assert ssot.reconcile(snap, _ran=ran) == []
    assert not [r for r in ran if r.startswith("executive_brief.posture.")], ran
    assert ssot.summary(snap)["verified"] is False
    # an unattributed failure beside an EMPTY list is the same fallback
    snap["assessment_integrity"] = {"failed_phases": ["dependency map"]}
    ran = []
    ssot.reconcile(snap, _ran=ran)
    assert not [r for r in ran if r.startswith("executive_brief.posture.")], ran


# --------------------------------------------------------------------------------------------------
# ssot.fleet_avg_health: the one Python reader every consumer renders from
# --------------------------------------------------------------------------------------------------
_LEGACY_ZERO = {"health_scores": _insufficient(3), "devices": {f"sw{i}": {} for i in range(3)},
                "executive_brief": {"posture": {"avg_health": 0, "n_critical": 0, "n_poor": 0,
                                                "worst_band": ""},
                                    "scale": {"n_devices": 3}}}


@pytest.mark.parametrize("snap, state, value", [
    ({"health_scores": [{"band": "Good", "score": 80}],
      "executive_brief": {"posture": {"avg_health": 80}}}, "measured", 80),
    ({"health_scores": _insufficient(2),
      "executive_brief": {"posture": {"avg_health": None, "n_scored": 0}}}, "not_assessed", None),
    (_LEGACY_ZERO, "not_assessed", None),                     # pre-G15 hard 0 over zero scored rows
    # no scored-row basis -> not a measurement (tests/test_ssot_owner_robustness.py, defect 4)
    ({"executive_brief": {"posture": {"avg_health": 0}}}, "unverified", None),
    ({"health_scores": [{"band": "Good", "score": 80}],
      "executive_brief": {"posture": {"avg_health": "abc"}}}, "unverified", None),
    ({"health_scores": [{"band": "Good", "score": 80}]}, "unpublished", None),
    ({"health_scores": [{"band": "Good", "score": 80}], "executive_brief": {"_unavailable": True}},
     "unpublished", None),
    (None, "unpublished", None),
], ids=["measured", "abstention", "legacy-zero", "no-basis", "malformed", "no-brief", "failed-brief",
        "garbage"])
def test_fleet_avg_health_states(snap, state, value):
    fh = ssot.fleet_avg_health(snap)
    assert fh["state"] == state and fh["value"] == value, fh


def test_fleet_avg_health_counts_scored_rows():
    fh = ssot.fleet_avg_health({"health_scores": [{"band": "Good", "score": 80}] + _insufficient(2)})
    assert fh["n_scored"] == 1 and fh["n_rows"] == 3
    assert ssot.fleet_avg_health({})["n_rows"] is None


# --------------------------------------------------------------------------------------------------
# consumers: none may re-derive a number from the unscored rows
# --------------------------------------------------------------------------------------------------
def _unscored_snapshot():
    hs = _insufficient()
    return {"generated_at": "2026-01-01T00:00:00", "script_version": "V3.23.0",
            "devices": {r["switch"]: {} for r in hs}, "health_scores": hs,
            "executive_brief": compute_executive_brief(health_scores=hs)}


def test_trend_point_abstains_on_a_published_abstention():
    from cisco_toolkit.html import _trend_point
    assert _trend_point(_unscored_snapshot())["avg_health"] == ""
    # canonical-first on the KEY: a published abstention is never replaced by a recompute, even when the
    # rows beside it would yield a number (that disagreement is ssot.reconcile's to report, not to paper over)
    mixed = {"health_scores": [{"switch": "a", "band": "Good", "score": 80}],
             "executive_brief": {"posture": {"avg_health": None, "n_scored": 0}}}
    assert _trend_point(mixed)["avg_health"] == ""


def test_trend_point_legacy_fallback_excludes_insufficient_data():
    from cisco_toolkit.html import _trend_point
    snap = {"health_scores": [{"switch": "a", "band": "Good", "score": 80},
                              {"switch": "b", "band": "Insufficient Data", "score": 100}]}
    assert _trend_point(snap)["avg_health"] == 80.0
    only_unscored = {"health_scores": _insufficient()}
    assert _trend_point(only_unscored)["avg_health"] == ""
    assert _trend_point(only_unscored)["n_critical"] == ""


def test_trend_point_does_not_trust_a_legacy_zero_over_nothing():
    from cisco_toolkit.html import _trend_point
    pt = _trend_point(_LEGACY_ZERO)
    assert pt["avg_health"] == "" and pt["n_critical"] == "", pt


def test_campaign_never_reads_an_unscored_collection_as_fewer_critical_switches():
    """C1: 6 Critical rows; C2: nothing scored. 'Critical-band switches 6 -> 0, improving' was a
    measurement of nothing."""
    from cisco_toolkit.html import compute_campaign_trend
    c1_rows = [{"switch": f"sw{i}", "band": "Critical", "score": 10} for i in range(6)]
    c1 = {"generated_at": "2026-01-01T00:00:00", "health_scores": c1_rows,
          "executive_brief": compute_executive_brief(health_scores=c1_rows)}
    c2_rows = _insufficient(6)
    c2 = {"generated_at": "2026-02-01T00:00:00", "health_scores": c2_rows,
          "executive_brief": compute_executive_brief(health_scores=c2_rows)}
    trend = compute_campaign_trend([c1, c2])
    assert not [r for r in trend["trajectory"] if r["metric"] == "Critical-band switches"], trend["trajectory"]
    assert trend["timeline"][1]["n_critical"] == ""


def _exec_sheet_value(brief, hs, label="Average health score"):
    from openpyxl import Workbook
    import cisco_toolkit.excel as X
    wb = Workbook(); X.harden_workbook(wb)
    X.write_executive_summary_sheet(wb, hs, [], [], [], brief=brief)
    for row in wb[X.EXEC_SUMMARY_SHEET_NAME].iter_rows(values_only=True):
        cells = [c for c in row if c not in (None, "")]
        if cells and str(cells[0]).strip() == label:
            return " ".join(str(c) for c in cells[1:])
    return None


def test_workbook_exec_summary_says_not_assessed():
    snap = _unscored_snapshot()
    val = _exec_sheet_value(snap["executive_brief"], snap["health_scores"])
    assert val is not None and "NOT ASSESSED" in val, val
    assert "/ 100" not in val, val
    mixed = _exec_sheet_value({"posture": {"avg_health": None, "n_scored": 0}},
                              [{"switch": "a", "band": "Good", "score": 80}])
    assert mixed is not None and "NOT ASSESSED" in mixed and "80" not in mixed, mixed


def test_workbook_legacy_fallback_excludes_insufficient_data():
    hs = [{"switch": "a", "band": "Good", "score": 80},
          {"switch": "b", "band": "Insufficient Data", "score": 100}]
    assert _exec_sheet_value(None, hs) == "80 / 100"
    val = _exec_sheet_value(None, _insufficient())
    assert val is not None and "NOT ASSESSED" in val and "100" not in val, val


def test_workbook_does_not_trust_a_legacy_zero_over_nothing():
    val = _exec_sheet_value(_LEGACY_ZERO["executive_brief"], _LEGACY_ZERO["health_scores"])
    assert val is not None and "NOT ASSESSED" in val and "0 / 100" not in val, val


def test_workbook_band_counts_abstain_when_nothing_was_scored():
    snap = _unscored_snapshot()
    for label in ("Critical band", "Poor / Fair", "Good / Excellent"):
        val = _exec_sheet_value(snap["executive_brief"], snap["health_scores"], label=label)
        assert val is not None and val != "0" and "not assessed" in val.lower(), (label, val)
    # NON-VACUITY: a scored fleet keeps its integer counts
    hs = [{"switch": "a", "band": "Critical", "score": 20}]
    assert _exec_sheet_value(compute_executive_brief(health_scores=hs), hs, label="Critical band") == "1"


def test_no_health_rows_never_renders_a_collected_count_above_a_zero_inventory():
    """A crashed Health Scores phase published `n_devices: 0`: the workbook read 'Switches collected /
    inventoried 3 / 0' and every docx '(read-only evidence; -3 not reached)'."""
    from cisco_toolkit import docmeta
    brief = compute_executive_brief()
    brief["scale"]["n_collected"] = 3                      # injected by main() from collection_completeness
    val = _exec_sheet_value(brief, [], label="Switches collected / inventoried")
    assert val is not None and "/ 0" not in val and "not published" in val, val
    scope = dict(docmeta._glance_rows({"health_scores": [], "executive_brief": brief}))["What was assessed?"]
    assert "-3" not in scope and "0 inventoried" not in scope, scope
    # NON-VACUITY: a normal fleet keeps "collected / inventoried"
    hs = [{"switch": "a", "band": "Good", "score": 80}, {"switch": "b", "band": "Good", "score": 70}]
    b2 = compute_executive_brief(health_scores=hs)
    b2["scale"]["n_collected"] = 1
    assert _exec_sheet_value(b2, hs, label="Switches collected / inventoried") == "1 / 2"


def test_workbook_malformed_average_is_unverified_not_zero_of_n():
    """A non-numeric published average beside SCORED rows is malformed, not 'NOT ASSESSED — 0 of N'."""
    hs = [{"switch": "a", "band": "Good", "score": 80}]
    val = _exec_sheet_value({"posture": {"avg_health": "abc"}}, hs)
    assert val is not None and "UNVERIFIED" in val and "0 of" not in val, val


def _deck_text(tmp_path, snap):
    pptx = pytest.importorskip("pptx")
    from cisco_toolkit.deck import write_executive_deck_pptx
    out = tmp_path / "deck.pptx"
    write_executive_deck_pptx(str(out), snap, "Fleet")
    prs = pptx.Presentation(str(out))
    return [sh.text_frame.text for sl in prs.slides for sh in sl.shapes if sh.has_text_frame]


def test_deck_never_renders_a_literal_none(tmp_path):
    txt = "\n".join(_deck_text(tmp_path, _unscored_snapshot()))
    assert "None" not in txt, txt
    assert "NOT ASSESSED" in txt


def test_deck_does_not_trust_a_legacy_zero_and_withholds_the_band_count(tmp_path):
    shapes = _deck_text(tmp_path, _LEGACY_ZERO)
    txt = "\n".join(shapes)
    assert "NOT ASSESSED" in txt, txt
    # the stat helper writes the number and its label as two consecutive text shapes
    idx = next(i for i, t in enumerate(shapes) if t.startswith("Critical-band switches"))
    assert shapes[idx - 1].strip() == "—", shapes[idx - 1]
    # NON-VACUITY: a scored fleet keeps the count
    hs = [{"switch": "a", "band": "Critical", "score": 20}]
    scored = _deck_text(tmp_path, {"health_scores": hs,
                                   "executive_brief": compute_executive_brief(health_scores=hs)})
    idx = next(i for i, t in enumerate(scored) if t.startswith("Critical-band switches"))
    assert scored[idx - 1].strip() == "1", scored[idx - 1]


def test_deck_breadcrumb_discloses_not_assessed_axes_it_cut(tmp_path):
    """A not-assessed axis is Info, which sorts after every Low axis: the 4-row cap dropped it and the
    breadcrumb counted only High/Critical."""
    snap = _unscored_snapshot()
    snap["executive_brief"] = compute_executive_brief(
        health_scores=snap["health_scores"],
        punchlist=[{"severity": "Low"}],
        application_intelligence={"summary": {"n_domains": 1, "pilot_domain": "A", "last_domain": "A"}},
        migration_readiness=[{"readiness": "READY"}],
        segmentation={"summary": {"n_gateways": 2, "flat": False, "gateway_acl_coverage": 100}},
        multicast_intelligence={"summary": {"n_groups": 1, "n_mcast_vlans": 1}})
    axes = snap["executive_brief"]["axes"]
    assert len(axes) > 4 and any(a["severity"] == "Info" for a in axes[4:]), axes
    crumb = [t for t in _deck_text(tmp_path, snap) if "more axis headline(s)" in t]
    assert crumb and "Info" in crumb[0], crumb


def test_at_a_glance_row_says_not_assessed():
    from cisco_toolkit import docmeta
    rows = dict(docmeta._glance_rows(_unscored_snapshot()))
    health = rows["How healthy is the fleet?"]
    assert "NOT ASSESSED" in health and "/100" not in health, health
    legacy = dict(docmeta._glance_rows(_LEGACY_ZERO))["How healthy is the fleet?"]
    assert "NOT ASSESSED" in legacy and "0/100" not in legacy, legacy
    malformed = dict(docmeta._glance_rows({"health_scores": [{"band": "Good", "score": 80}],
                                           "executive_brief": {"posture": {"avg_health": "abc"}}}))
    assert "UNVERIFIED" in malformed["How healthy is the fleet?"], malformed


# --------------------------------------------------------------------------------------------------
# the offline explorer: one reader, decided on key PRESENCE, executed under node
# --------------------------------------------------------------------------------------------------
def _explorer_helper():
    m = re.search(r"/\* FLEET-AVG-HEALTH START \*/(.*?)/\* FLEET-AVG-HEALTH END \*/",
                  EXPLORER.read_text(encoding="utf-8"), re.S)
    assert m, "the explorer's single fleet-average reader (FLEET-AVG-HEALTH block) is missing"
    return m.group(1)


def test_explorer_has_no_all_rows_fallback_left():
    """Every site that showed the fleet average used `posture.avg_health ?? mean(all rows)`: a null
    abstention fell through to a mean that counts the 'Insufficient Data' rows (the fabricated 100)."""
    src = EXPLORER.read_text(encoding="utf-8")
    assert "avg_health??" not in src.replace(" ", "")
    assert src.count("fleetAvgHealth(") >= 4, "modeSummary / narrativeCard / drawer head / Ask-bar"
    _explorer_helper()


@pytest.mark.skipif(not NODE, reason="node is not installed -- executed explorer check skipped")
def test_explorer_fleet_average_reader_executes(tmp_path):
    driver = tmp_path / "drv.js"
    driver.write_text(
        "let SNAP=null;\n" + _explorer_helper() + "\n"
        "const out=[], st=[];\n"
        "const ins=[{switch:'a',band:'Insufficient Data',score:100},{switch:'b',band:'Insufficient Data',score:100}];\n"
        "const good=[{band:'Good',score:80}];\n"
        "SNAP={executive_brief:{posture:{avg_health:null,n_scored:0}}}; out.push(fleetAvgHealth(ins));\n"
        # a number published over rows of which NONE is scored is the pre-G15 hard 0 (or worse): not assessed
        "SNAP={executive_brief:{posture:{avg_health:41}}}; out.push(fleetAvgHealth(ins));\n"
        "SNAP={executive_brief:{posture:{avg_health:41}}}; out.push(fleetAvgHealth(good.concat(ins)));\n"
        "SNAP={}; out.push(fleetAvgHealth(good.concat(ins)));\n"
        "SNAP={}; out.push(fleetAvgHealth(ins));\n"
        "SNAP={executive_brief:{_unavailable:true}}; out.push(fleetAvgHealth([{band:'Poor',score:40}]));\n"
        "SNAP={executive_brief:{posture:{avg_health:null}}}; out.push(fleetAvgHealth(good));\n"
        "SNAP={health_scores:ins,executive_brief:{posture:{avg_health:0}}}; out.push(fleetAvgHealth(ins));\n"
        "SNAP={executive_brief:{posture:{avg_health:'abc'}}}; out.push(fleetAvgHealth(good));\n"
        "for(const s of [{executive_brief:{posture:{avg_health:41}}},{health_scores:ins,executive_brief:{posture:{avg_health:0}}},"
        "{executive_brief:{posture:{avg_health:'abc'}}},{},{executive_brief:{_unavailable:true}}]){SNAP=s;"
        "st.push(fleetHealthState(s.health_scores||good).state);}\n"
        "process.stdout.write(JSON.stringify({out,st}));\n", encoding="utf-8")
    proc = subprocess.run([NODE, str(driver)], capture_output=True, text=True, timeout=60)
    assert proc.returncode == 0, proc.stderr
    got = json.loads(proc.stdout)
    assert got["out"] == [None, None, 41, 80, None, 40, None, None, None], got
    assert got["st"] == ["measured", "not_assessed", "unverified", "unpublished", "unpublished"], got


# --------------------------------------------------------------------------------------------------
# the explorer RENDER sites, executed in the real embedded script (full-page node harness)
# --------------------------------------------------------------------------------------------------
def _explorer_run(driver, tmp_path):
    import sys
    if str(ROOT / "tests") not in sys.path:
        sys.path.insert(0, str(ROOT / "tests"))
    import test_explorer_render_safety as harness      # the repo's DOM-stubbed full-script runner
    return harness._run(driver, tmp_path)


_EXEC_CASES = r"""
  const ins=[{switch:"a",band:"Insufficient Data",score:100},{switch:"b",band:"Insufficient Data",score:100}];
  const sc=[{switch:"a",band:"Fair",score:70},{switch:"b",band:"Critical",score:30}];
  const CASES={
    absent:{health_scores:sc},
    failed:{health_scores:sc,executive_brief:{_unavailable:true}},
    abstain:{health_scores:ins,executive_brief:{posture:{avg_health:null,n_critical:null,n_poor:null,
      worst_band:null,n_scored:0,not_assessed:"all_insufficient_data"},scale:{n_devices:2}}},
    legacy0:{health_scores:ins,executive_brief:{posture:{avg_health:0,n_critical:0,n_poor:0,worst_band:""},
      scale:{n_devices:2}}},
    malformed:{health_scores:sc,executive_brief:{posture:{avg_health:"abc",n_critical:1,n_poor:0,
      worst_band:"Critical"},scale:{n_devices:2}}},
    measured:{health_scores:sc,executive_brief:{posture:{avg_health:50,n_critical:1,n_poor:0,
      worst_band:"Critical"},scale:{n_devices:2}}},
  };
  const strip=h=>String(h||"").replace(/<[^>]+>/g,"").replace(/\s+/g," ").trim();
"""


@pytest.mark.skipif(not NODE, reason="node is not installed -- executed explorer check skipped")
def test_explorer_ask_atlas_exec_summary_never_invents_a_posture(tmp_path):
    """abH_exec rendered `worst band · 0 critical / 0 poor` from a MISSING posture beside a recomputed
    average (brief absent / failed), and trusted a legacy 0 as '0%'."""
    out = _explorer_run(_EXEC_CASES + r"""
      const res={};
      for(const [k,s] of Object.entries(CASES)){globalThis.__S=s;__EV("SNAP=globalThis.__S");
        const h=__EV("abH_exec")().html;
        const m=h.match(/Average health<\/span><span class="v">([\s\S]*?)<\/span><\/div>/);
        res[k]=m?strip(m[1]):null;}
      console.log(JSON.stringify(res));
    """, tmp_path)
    for k in ("absent", "failed"):
        assert out[k] and "unavailable" in out[k] and "critical" not in out[k], (k, out[k])
    assert "not assessed" in out["abstain"] and "0 of 2" in out["abstain"], out["abstain"]
    assert "not assessed" in out["legacy0"] and "0%" not in out["legacy0"], out["legacy0"]
    assert "unverified" in out["malformed"] and "0 of" not in out["malformed"], out["malformed"]
    assert "50%" in out["measured"] and "1 critical / 0 poor" in out["measured"], out["measured"]


@pytest.mark.skipif(not NODE, reason="node is not installed -- executed explorer check skipped")
def test_explorer_health_views_name_the_state_they_render(tmp_path):
    """The drawer head, the mode summary and the narrative printed '0 of N switches health-scored'
    for ANY non-numeric average -- false when rows ARE scored and the published value is malformed."""
    out = _explorer_run(_EXEC_CASES + r"""
      const res={};
      // the narrative keeps its 5 most severe lines: silence the structural detectors so the health line is reached
      __EV("structChains=()=>[]"); __EV("linkPhyConsistency=()=>[]"); __EV("trunkConsistency=()=>[]");
      __EV("addressingConflicts=()=>({total:0})"); __EV('MODEL={hosts:["a","b"],links:[]}');
      for(const k of ["abstain","legacy0","malformed","measured"]){globalThis.__S=CASES[k];__EV("SNAP=globalThis.__S");
        let head="", mode="", narr="";
        try{const h=__EV("drawHealth")(); const m=h.match(/<div class="title">([\s\S]*?)<\/div><\/div>/); head=strip(m&&m[1]);}catch(e){head="ERR "+e.message;}
        try{mode=strip(__EV("modeSummary")("health"));}catch(e){mode="ERR "+e.message;}
        try{narr=strip((__EV("buildNarrative")(__EV("MODEL"))||[]).map(x=>x.text).join(" | "));}catch(e){narr="ERR "+e.message;}
        res[k]={head,mode,narr};}
      console.log(JSON.stringify(res));
    """, tmp_path)
    for k in ("abstain", "legacy0"):
        for view in ("head", "mode", "narr"):
            assert "not assessed" in out[k][view].lower(), (k, view, out[k][view])
    for view in ("head", "mode", "narr"):
        assert "unverified" in out["malformed"][view].lower(), (view, out["malformed"][view])
        assert "0 of 2" not in out["malformed"][view], (view, out["malformed"][view])
    assert "/ 100 avg" in out["measured"]["head"], out["measured"]           # the countUp number
    assert "average 50" in out["measured"]["mode"], out["measured"]["mode"]


@pytest.mark.skipif(not NODE, reason="node is not installed -- executed explorer check skipped")
def test_explorer_brief_card_does_not_print_zero_devices_for_an_unpublished_count(tmp_path):
    out = _explorer_run(r"""
      __EV('MODEL={hosts:["a","b","c"],links:[]}');
      globalThis.__S={health_scores:[],executive_brief:{axes:[{axis:"Fleet health",severity:"Info",
        headline:"NOT ASSESSED — no health scores produced"}],scale:{n_devices:null,n_domains:0,n_endpoints:0},
        posture:{avg_health:null,n_scored:0}}};
      __EV("SNAP=globalThis.__S");
      const h=__EV("briefCard")();
      console.log(JSON.stringify({t:h.replace(/<[^>]+>/g," ").replace(/\s+/g," ")}));
    """, tmp_path)
    assert "3 devices" in out["t"] and " 0 devices" not in out["t"], out["t"]

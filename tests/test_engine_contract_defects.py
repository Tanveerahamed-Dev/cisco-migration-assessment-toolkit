"""Engine contract defects found by the Atlas Scope contract-gap register (G13 / G15 / G49).

Every case here drives the REAL producer (compute_move_groups, parse_spanning_tree_root, the real
config parsers) into the consumer, never a hand-written record in the shape the consumer happens to
expect -- that fabricated shape is exactly how each defect stayed green.

* G13 -- move-group identity. ``compute_move_groups`` published anonymous rows while seven consumers
  read ``g.get("group")`` (always ``""``) and four producers re-derived ``"Group N"`` from list
  position. One owner now writes the label; every consumer reads it through one helper, and a
  multi-label join is ordered by the owner's ordinal (never ``"Group 10" < "Group 2"``).
* G15 -- STP root election. "First sorted host claiming root wins" silently published one root when
  several bridges claimed it (the sample's 17 access switches presenting one bridge identity). One
  owner (``stp_topology.classify_stp_root_election``) now classifies each VLAN as published /
  ambiguous / not_observed; every consumer abstains on an ambiguous VLAN.
* G49 -- device-dossier exposure axes separate "input absent" from "input present, no findings";
  each exposure carries an ``input_state`` from ``ssot.ABSTENTION_STATES``.
"""
from __future__ import annotations

import pytest

from cisco_toolkit import analyze
from cisco_toolkit.analyze import (compute_application_intelligence, compute_device_dossiers,
                                   compute_endpoint_dependencies, compute_migration_punchlist,
                                   compute_migration_readiness, compute_move_groups,
                                   compute_remediation_plan, compute_subnet_intelligence,
                                   compute_validation_plan, compute_vlan_cutover_matrix,
                                   compute_wave_sequencing)
from cisco_toolkit.model import InterfaceData


# =====================================================================================================
# G13 -- move-group identity
# =====================================================================================================
def _acc(port: str, vlan: str, mac: str = "") -> InterfaceData:
    return InterfaceData(port=port, switchport_mode="Access", vlan=vlan, end_host_mac=mac)


def _fleet(reverse: bool = False) -> dict:
    """acc1..acc3 share VLAN 10 (Group 1), dist1/dist2 share VLAN 20 (Group 2), solo alone (Group 3)."""
    rows = [
        ("acc1", {"Gi1/0/1": _acc("Gi1/0/1", "10", "0000.0000.0a01"),
                  "Gi1/0/2": _acc("Gi1/0/2", "10", "0000.0000.0a02")}),
        ("acc2", {"Gi1/0/1": _acc("Gi1/0/1", "10", "0000.0000.0a03")}),
        ("acc3", {"Gi1/0/1": _acc("Gi1/0/1", "10", "0000.0000.0a04")}),
        ("dist1", {"Gi1/0/1": _acc("Gi1/0/1", "20", "0000.0000.0b01")}),
        ("dist2", {"Gi1/0/1": _acc("Gi1/0/1", "20", "0000.0000.0b02")}),
        ("solo", {"Gi1/0/1": _acc("Gi1/0/1", "30", "0000.0000.0c01")}),
    ]
    if reverse:
        rows.reverse()
    return dict(rows)


def _mg():
    return compute_move_groups(_fleet())


def test_g13_owner_writes_the_group_label():
    mg = _mg()
    assert [g["group"] for g in mg] == ["Group 1", "Group 2", "Group 3"]
    assert [g["switches"] for g in mg] == [["acc1", "acc2", "acc3"], ["dist1", "dist2"], ["solo"]]


def test_g13_label_is_independent_of_input_order():
    fwd = {g["group"]: g["switches"] for g in compute_move_groups(_fleet())}
    rev = {g["group"]: g["switches"] for g in compute_move_groups(_fleet(reverse=True))}
    assert fwd == rev


def test_g13_written_label_wins_over_list_position():
    """A consumer handed the owner's rows in another order must keep the owner's labels, not re-derive
    'Group N' from its own enumerate()."""
    ifaces = _fleet()
    mg_rev = list(reversed(_mg()))
    seq = compute_wave_sequencing(ifaces, mg_rev)
    assert [r["group"] for r in seq] == ["Group 3", "Group 2", "Group 1"]

    vp = compute_validation_plan(ifaces, move_groups=mg_rev)
    solo_waves = {it["wave"] for it in vp["items"] if it["device"] == "solo"}
    assert solo_waves == {"Group 3"}, solo_waves

    rows = {r["vlan"]: r for r in compute_vlan_cutover_matrix(ifaces, move_groups=mg_rev)}
    assert rows[30]["wave"] == "Group 3"
    assert rows[10]["wave"] == "Group 1"

    dep = analyze.build_dependency_map(ifaces, [], [])
    mr = compute_migration_readiness(ifaces, mg_rev, [{"switch": h, "band": "Good"} for h in ifaces],
                                     [], [], [], [], dep)
    by_switches = {tuple(r["switches"]): r["group"] for r in mr}
    assert by_switches[("solo",)] == "Group 3"
    assert by_switches[("acc1", "acc2", "acc3")] == "Group 1"


def test_g13_endpoint_dependencies_see_real_group_labels():
    mg = _mg()
    split = [{"host": "acc1", "mac": "aaaa.0000.0001", "vlan": "10", "port": "Gi1/0/1"},
             {"host": "dist1", "mac": "aaaa.0000.0001", "vlan": "20", "port": "Gi1/0/1"}]
    same = [{"host": "acc1", "mac": "aaaa.0000.0002", "vlan": "10", "port": "Gi1/0/1"},
            {"host": "acc2", "mac": "aaaa.0000.0002", "vlan": "10", "port": "Gi1/0/1"}]
    d_split = compute_endpoint_dependencies(split, mg)["dual_homed"][0]
    assert d_split["move_groups"] == ["Group 1", "Group 2"]
    assert d_split["split_across_groups"] is True
    d_same = compute_endpoint_dependencies(same, mg)["dual_homed"][0]
    assert d_same["move_groups"] == ["Group 1"]
    assert d_same["split_across_groups"] is False


def test_g13_subnet_punchlist_remediation_dossier_carry_the_label():
    ifaces, mg = _fleet(), _mg()
    si = compute_subnet_intelligence(ifaces, {}, mg)
    assert [g["group"] for g in si["move_groups"]] == ["Group 1", "Group 2", "Group 3"]

    pl = compute_migration_punchlist([], {}, {}, [{"switch": "dist1", "port": "Gi1/0/1", "risk": "err-disabled"}],
                                     [], [], {}, [], mg)
    row = next(r for r in pl if "dist1" in r["devices"])
    assert row["wave"] == "Group 2"

    rem = compute_remediation_plan(stp_findings={"accidental": [{"host": "acc1", "vlan": "10", "priority": 32778}]},
                                   move_groups=mg)
    assert {it["wave"] for it in rem["items"] if it["device"] == "acc1"} == {"Group 1"}

    dos = compute_device_dossiers(health_scores=[{"switch": "solo", "band": "Good", "score": 90}], move_groups=mg)
    assert dos["per_device"][0]["wave"] == "Group 3"


def test_g13_application_domain_split_across_waves_fires_through_the_real_producer():
    ai = {"SW01-BC-DANTE-A": {"Gi1/0/1": _acc("Gi1/0/1", "10")},
          "SW02-BC-DANTE-B": {"Gi1/0/1": _acc("Gi1/0/1", "11")}}
    mg = compute_move_groups(ai)
    assert sorted(g["group"] for g in mg) == ["Group 1", "Group 2"]
    sm = {"multicast": {"ptp": {}, "classified_groups": [], "igmp_queriers": []}}
    out = compute_application_intelligence(ai, [], {}, sm, [], mg, [])
    audio = next(d for d in out["domains"] if d["id"] == "audio")
    assert audio["waves"] == ["Group 1", "Group 2"]
    assert audio["spans_waves"] is True
    assert any("split across migration waves" in r["title"] for r in audio["risks"])


def _eleven():
    return {f"sw{i:02d}": {"Gi1/0/1": _acc("Gi1/0/1", str(100 + i), f"0000.0000.{i:04x}")} for i in range(1, 12)}


def test_g13_multi_label_joins_follow_the_owner_ordinal_not_string_order():
    ifaces = _eleven()
    mg = compute_move_groups(ifaces)
    assert [g["group"] for g in mg] == [f"Group {i}" for i in range(1, 12)]
    assert mg[1]["switches"] == ["sw02"] and mg[9]["switches"] == ["sw10"]
    l2 = {"addressing": {"dup_ip": [{"ip": "10.0.0.1", "where": [("sw02", "Vlan10", 10), ("sw10", "Vlan20", 20)]}],
                         "dup_subnet": []}}
    pl = compute_migration_punchlist([], {}, {}, [], [], [], {}, [], mg, l2=l2)
    row = next(r for r in pl if r["category"] == "Addressing")
    assert row["wave"] == "Group 2, Group 10"

    dh = compute_endpoint_dependencies([{"host": "sw10", "mac": "aaaa.0000.0009", "vlan": "110"},
                                        {"host": "sw02", "mac": "aaaa.0000.0009", "vlan": "102"}], mg)
    assert dh["dual_homed"][0]["move_groups"] == ["Group 2", "Group 10"]


def test_g13_label_helpers_fall_back_to_the_owner_formula_for_unlabelled_or_malformed_rows():
    assert analyze.move_group_label(3) == "Group 3"
    assert analyze.move_group_labels([{"switches": ["a"]}, {"switches": ["b"]}]) == ["Group 1", "Group 2"]
    assert analyze.move_group_labels([{"group": "Group 2", "switches": ["a"]},
                                      {"group": "Group 1", "switches": ["b"]}]) == ["Group 2", "Group 1"]
    # mixed / duplicated / non-string labels -> the positional owner formula for EVERY row, never a raise
    for bad in ([{"group": "Group 1"}, {"switches": ["b"]}],
                [{"group": "Group 1"}, {"group": "Group 1"}],
                [{"group": 7}, {"group": "Group 2"}],
                [{"group": ["x"]}, 5]):
        assert analyze.move_group_labels(bad) == ["Group 1", "Group 2"]
    host_label, ordinal = analyze.move_group_host_index([{"group": "Group 2", "switches": ["a", 5]},
                                                         {"group": "Group 1", "switches": ["a", "b"]}])
    assert host_label == {"a": "Group 2", "5": "Group 2", "b": "Group 1"}
    assert ordinal == {"Group 2": 1, "Group 1": 2}
    assert analyze.MOVE_GROUP_UNSCHEDULED == "(unscheduled)"


def test_g13_snapshot_readers_name_groups_by_the_written_label():
    from cisco_toolkit import mcp_server, mop, nrfu_export

    snap = {"move_groups": [{"group": "Group 2", "switches": ["b"]}, {"group": "Group 1", "switches": ["a"]}],
            "migration_readiness": [{"group": "Group 1", "switches": ["a"], "readiness": "READY"},
                                    {"group": "Group 2", "switches": ["b"], "readiness": "NOT READY"}]}
    mg = mcp_server.get_move_groups(snap)
    assert [(g["group"], g["switches"], g["readiness"]) for g in mg["groups"]] == [
        ("Group 2", ["b"], "NOT READY"), ("Group 1", ["a"], "READY")]
    assert mop._waves({"move_groups": snap["move_groups"]}) == [("Group 2", ["b"]), ("Group 1", ["a"])]
    assert nrfu_export.UNSCHEDULED_WAVE == analyze.MOVE_GROUP_UNSCHEDULED
    nrfu = nrfu_export.compute_nrfu_commands({"move_groups": snap["move_groups"],
                                              "devices": {"a": {"platform": "ios"}, "b": {"platform": "ios"}}})
    order = [w["wave_id"] for w in nrfu["waves"]]
    assert order == ["Group 2", "Group 1"], order


def test_g13_design_wave_plan_names_its_source_move_groups():
    from cisco_toolkit import design_advisor

    snap = {"move_groups": [{"group": "Group 1", "switches": ["a", "b"]}, {"group": "Group 2", "switches": ["c"]}]}
    wp = design_advisor._wave_plan(snap)
    assert wp["waves"][0]["source_move_groups"] == ["Group 1", "Group 2"]


def test_g13_published_labels_reconcile_on_the_real_pipeline(tmp_path):
    """SSOT reconcile over a fresh real pipeline run: the owner's labels are unique and every derived
    wave label is one of them (or the documented no-group sentinels); a punch row's wave is exactly the
    owner-ordered labels of its devices."""
    import test_pipeline_golden as tg

    snap, _ = tg._run_pipeline(tmp_path, **tg._GOLDEN_CLOCKS)
    owner = [g["group"] for g in snap["move_groups"]]
    assert owner and len(set(owner)) == len(owner)
    allowed = set(owner) | {"", analyze.MOVE_GROUP_UNSCHEDULED}
    host_label, ordinal = analyze.move_group_host_index(snap["move_groups"])
    for r in snap["punchlist"]:
        want = sorted({host_label[d] for d in r["devices"] if d in host_label}, key=ordinal.get)
        assert r["wave"] == ", ".join(want), r
    derived = ([r["group"] for r in snap["wave_sequencing"]] + [r["group"] for r in snap["migration_readiness"]]
               + [it["wave"] for it in snap["remediation_plan"]["items"]]
               + [d["wave"] for d in snap["device_dossiers"]["per_device"]]
               + [g["group"] for g in snap["subnet_intelligence"]["move_groups"]]
               + [it["wave"] for it in snap["validation_plan"]["items"]])
    assert set(derived) <= allowed, set(derived) - allowed
    for row in snap["vlan_cutover"]:
        assert {w for w in row["wave"].split(", ") if w} <= set(owner)


def test_g13_explorer_label_mirror_matches_the_python_owner(tmp_path):
    """The explorer's moveGroupLabels() is a JS mirror of analyze.move_group_labels -- held equal on the
    labelled, unlabelled and malformed shapes so the wave cards can never name a group differently."""
    import json
    import shutil
    import subprocess
    from pathlib import Path

    node = shutil.which("node")
    if not node:
        pytest.skip("node is not installed")
    html = (Path(analyze.__file__).parent / "blast_radius_explorer.html").read_text(encoding="utf-8")
    start = html.index("function moveGroupLabels(mg)")
    fn = html[start:html.index("function wavesData()", start)]
    cases = [[{"group": "Group 2", "switches": ["a"]}, {"group": "Group 1", "switches": ["b"]}],
             [{"switches": ["a"]}, {"switches": ["b"]}],
             [{"group": "Group 1"}, {"group": "Group 1"}],
             [{"group": 7}, {"group": "Group 2"}],
             [{"group": "  "}, 5], []]
    driver = tmp_path / "labels.js"
    driver.write_text(fn + "\nprocess.stdout.write(JSON.stringify(" + json.dumps(cases)
                      + ".map(moveGroupLabels)));\n", encoding="utf-8")
    out = subprocess.run([node, str(driver)], capture_output=True, text=True, encoding="utf-8", timeout=60)
    assert out.returncode == 0, out.stderr
    assert json.loads(out.stdout) == [analyze.move_group_labels(c) for c in cases]


# =====================================================================================================
# G15 -- STP root election: one owner, and an ambiguous VLAN is never published as a bare root
# =====================================================================================================
def _stp_fixture():
    """F: REAL parser output. access1's capture (with its 'This bridge is the root' for VLAN 30) cloned
    onto access2 -- exactly how the sample fleet's access clones inherit one bridge identity -- plus core1,
    which roots VLAN 10 and observes VLAN 30's root identity."""
    import synthetic_fixtures as fx
    from cisco_toolkit.parse import parse_spanning_tree_root

    return {"access1": parse_spanning_tree_root(fx._ACCESS1["show spanning-tree"]),
            "access2": parse_spanning_tree_root(fx._ACCESS1["show spanning-tree"]),
            "core1": parse_spanning_tree_root(fx._CORE1["show spanning-tree"])}


def _row(prio, addr, is_root, **kw):
    return {"root_priority": prio, "root_address": addr, "is_root": is_root, "is_mst": False, **kw}


def _election(stp):
    from cisco_toolkit.stp_topology import classify_stp_root_election
    return classify_stp_root_election(stp)


def _nrfu_cases(nrfu):
    return [c for w in nrfu["waves"] for d in w["devices"] for c in d["cases"]]


def test_g15_owner_classifies_the_real_parser_fixture():
    el = _election(_stp_fixture())
    v10, v30 = el["pvst_vlan"]["10"], el["pvst_vlan"]["30"]
    assert (v10["state"], v10["reason"], v10["root"]) == ("published", "single_claimant", "core1")
    assert v10["identities"][0]["observers"] == ["access1", "access2"]
    assert v10["default_election"] is False
    assert (v30["state"], v30["reason"], v30["root"]) == ("ambiguous", "duplicate_bridge_identity", None)
    assert v30["claimants"] == ["access1", "access2"]
    assert v30["identities"] == [{"root_address": "cccc.0003.0003", "root_priority": 32798,
                                  "root_priorities": [32798], "claimants": ["access1", "access2"],
                                  "observers": ["core1"]}]
    assert v30["default_election"] is None
    assert el["mst_instance"] == {}


def test_g15_owner_states_cover_every_evidence_shape():
    el = _election({"a": {"30": _row(32798, "aaaa.0000.0001", True)},
                    "b": {"30": _row(32798, "bbbb.0000.0002", True)}})["pvst_vlan"]["30"]
    assert (el["state"], el["reason"], len(el["identities"])) == ("ambiguous", "multiple_root_identities", 2)
    # one claimant, but an observer disagrees about who the root is
    el = _election({"a": {"30": _row(24606, "aaaa.0000.0001", True)},
                    "b": {"30": _row(24606, "cccc.0000.0003", False)}})["pvst_vlan"]["30"]
    assert (el["state"], el["reason"]) == ("ambiguous", "multiple_root_identities")
    # observers only: the root is off-scan
    el = _election({"a": {"30": _row(24606, "aaaa.0000.0001", False)},
                    "b": {"30": _row(24606, "aaaa.0000.0001", False)}})["pvst_vlan"]["30"]
    assert (el["state"], el["reason"], el["root"], el["default_election"]) == (
        "not_observed", "root_not_collected", None, None)
    el = _election({"c": {"20": _row(None, "", False)}})["pvst_vlan"]["20"]
    assert (el["state"], el["reason"]) == ("not_observed", "no_root_evidence")
    # a malformed claim next to a real claimant is never published
    el = _election({"a": {"30": _row(24606, "aaaa.0000.0001", True)},
                    "b": {"30": _row(24606, "aaaa.0000.0001", "true")}})["pvst_vlan"]["30"]
    assert (el["state"], el["reason"]) == ("ambiguous", "malformed_root_rows")
    # one identity, conflicting priorities: published, but no single priority and no default verdict
    el = _election({"a": {"30": _row(24586, "aaaa.0000.0001", True)},
                    "b": {"30": _row(24606, "aaaa.0000.0001", False)}})["pvst_vlan"]["30"]
    assert (el["state"], el["root_priority"], el["default_election"]) == ("published", None, None)
    assert el["identities"][0]["root_priorities"] == [24586, 24606]


def test_g15_priority_rules_are_one_owner_and_total():
    from cisco_toolkit.stp_topology import stp_default_priority

    assert stp_default_priority("30", 32768 + 30) is True
    assert stp_default_priority("30", 32768) is True           # legacy extend-system-id off
    assert stp_default_priority("30", 24606) is False
    for bad in (float("inf"), float("nan"), True, {"x": 1}, None, "abc"):
        assert stp_default_priority("30", bad) is None
    el = _election({"a": {"30": _row("24576", "aaaa.0000.0001", True)}})["pvst_vlan"]["30"]
    assert el["root_priority"] == 24576 and el["default_election"] is False
    for bad in (float("inf"), True, {"x": 1}):
        el = _election({"a": {"30": _row(bad, "aaaa.0000.0001", True)}})["pvst_vlan"]["30"]
        assert el["state"] == "published" and el["root_priority"] is None and el["default_election"] is None


def test_g15_mst_instances_never_merge_with_pvst_vlans():
    el = _election({"a": {"1": {**_row(32769, "aaaa.0000.0001", True), "is_mst": True}},
                    "b": {"1": _row(24577, "bbbb.0000.0002", True)}})
    assert el["mst_instance"]["1"]["root"] == "a"
    assert el["pvst_vlan"]["1"]["root"] == "b"
    assert el["pvst_vlan"]["1"]["state"] == "published"


def test_g15_every_consumer_abstains_on_an_ambiguous_vlan():
    """The whole consumer class over F plus a VLAN 30 gateway SVI on core1: no surface may name
    access1 (the first sorted claimant) as VLAN 30's root."""
    import openpyxl

    from cisco_toolkit import archreview, failover, nrfu_export
    from cisco_toolkit.excel import write_stp_roots_sheet, write_vlan_cutover_sheet

    stp = _stp_fixture()
    ifaces = {"access1": {"Gi1/0/1": _acc("Gi1/0/1", "30")}, "access2": {"Gi1/0/1": _acc("Gi1/0/1", "30")},
              "core1": {"Vlan30": InterfaceData(port="Vlan30", svi_ip="10.30.0.1/24"),
                        "Vlan10": InterfaceData(port="Vlan10", svi_ip="10.10.0.1/24")}}
    rows = {r["vlan"]: r for r in compute_vlan_cutover_matrix(ifaces, stp)}
    assert rows[30]["stp_root"] == analyze.VLAN_CUTOVER_AMBIGUOUS == "[AMBIGUOUS]"
    assert rows[30]["stp_root_state"] == "ambiguous"
    assert rows[30]["stp_root_claimants"] == ["access1", "access2"]
    assert rows[30]["stp_root_default_election"] is None
    assert rows[10]["stp_root"] == "core1" and rows[10]["stp_root_state"] == "published"

    sf = analyze.stp_root_findings(stp, ifaces)
    assert not any(x["vlan"] == "30" for x in sf["accidental"] + sf["misaligned"])
    assert [(a["vlan"], a["reason"], a["claimants"]) for a in sf["ambiguous"]] == [
        ("30", "duplicate_bridge_identity", ["access1", "access2"])]

    pl = compute_migration_punchlist([], {}, {}, [], [], [], sf, [], [])
    stp_rows = [r for r in pl if r["category"] == "STP"]
    assert [r["title"] for r in stp_rows] == ["STP root ambiguous (VLAN 30)"]
    assert stp_rows[0]["devices"] == ["access1", "access2"]
    assert {ref["ref"] for ref in stp_rows[0]["evidence_refs"]} >= {"/stp_roots/access1/30", "/stp_roots/access2/30"}

    rem = compute_remediation_plan(stp_findings=sf)
    assert not [it for it in rem["items"] if it["category"] == "STP"]

    vp = compute_validation_plan(ifaces, stp_roots=stp)
    v30_checks = [it for it in vp["items"] if it["category"] == "STP" and "VLAN 30" in it["check"]]
    assert v30_checks and all("ambiguous" in it["check"] and it["severity"] == "High"
                              and it.get("evidence_state") == "review" for it in v30_checks)
    assert {it["device"] for it in v30_checks} == {"access1", "access2"}
    assert not any("VLAN 30 unchanged" in it["check"] for it in vp["items"])
    assert any("VLAN 10 unchanged" in it["check"] and it["device"] == "core1" for it in vp["items"])
    assert analyze.compute_current_baseline_gate(vp)["integrity"]["valid"] is True

    nrfu = nrfu_export.compute_nrfu_commands({"stp_roots": stp,
                                              "devices": {h: {"platform": "ios"} for h in stp}})
    v30 = [c for c in _nrfu_cases(nrfu) if c["command"] == "show spanning-tree vlan 30"]
    assert v30 and not any("This bridge is the root for VLAN 30" in c["expected"] for c in v30)
    amb = [c for c in v30 if c["expected"].startswith("AMBIGUOUS")]
    assert len(amb) == 2 and {c.get("evidence_state") for c in amb} == {"review"}

    dos = {d["host"]: d for d in compute_device_dossiers(
        health_scores=[{"switch": h, "band": "Good", "score": 90} for h in stp], stp_roots=stp)["per_device"]}
    for h in ("access1", "access2"):
        assert dos[h]["stp_root_vlans"] == 0 and dos[h]["stp_root_ambiguous_vlans"] == 1
    assert dos["core1"]["stp_root_vlans"] == 1 and dos["core1"]["stp_root_ambiguous_vlans"] == 0

    snap = {"devices": {h: {} for h in stp}, "stp_roots": stp,
            "l3_forwarding": [{"switch": "core1", "vlan": 30, "svi_ip": "10.30.0.1/24"}]}
    l21 = next(c for c in archreview.compute_architecture_review(snap)["checks"] if c["id"] == "L2-1")
    assert "access1 roots" not in l21["observed"] and "ambiguous" in l21["observed"]
    assert l21["verdict"] != "conforms"

    ready = failover.compute_failover_readiness({"stp_roots": stp})
    assert any(r.get("kind") == "stp" and r.get("vlan") == "30" and r.get("host") is None
               for r in ready["at_risk"])
    assert ready["n_stp_roots"] == (ready["n_stp_roots_with_backup"] + ready["n_stp_default_election"]
                                    + ready["n_stp_indeterminate"])

    wb = openpyxl.Workbook()
    write_stp_roots_sheet(wb, stp, ifaces)
    write_vlan_cutover_sheet(wb, list(rows.values()))
    for ws in wb.worksheets[1:]:
        cells = [c.value for row in ws.iter_rows() for c in row]
        assert "access1" not in cells, ws.title
    cut = wb.worksheets[-1]
    by_vlan = {row[0].value: [c.value for c in row] for row in cut.iter_rows(min_row=2)}
    assert by_vlan[30][3] == "undetermined (root ambiguous)"
    assert "access1" in by_vlan[30][2] and "[AMBIGUOUS]" in by_vlan[30][2]


def test_g15_vlan_cutover_sheet_never_renders_no_for_an_undetermined_election():
    import openpyxl

    from cisco_toolkit.excel import write_vlan_cutover_sheet

    wb = openpyxl.Workbook()
    write_vlan_cutover_sheet(wb, [{"vlan": 5, "stp_root": "sw1", "stp_root_default_election": None},
                                  {"vlan": 6, "stp_root": "sw1", "stp_root_state": "published",
                                   "stp_root_default_election": False}])
    ws = wb.worksheets[-1]
    got = {row[0].value: row[3].value for row in ws.iter_rows(min_row=2)}
    assert got[5] != "no"
    assert got[6] == "no"


def _sample():
    import json
    from pathlib import Path

    return json.loads((Path(__file__).resolve().parents[1] / "webapp" / "sample_data"
                       / "sample_fleet.snapshot.json").read_text(encoding="utf-8"))


def test_g15_sample_fleet_vlan30_is_ambiguous_not_access1():
    rows = {r["vlan"]: r for r in compute_vlan_cutover_matrix({}, _sample()["stp_roots"])}
    r30 = rows[30]
    assert r30["stp_root"] == "[AMBIGUOUS]" and r30["stp_root_reason"] == "duplicate_bridge_identity"
    assert len(r30["stp_root_claimants"]) == 17
    assert r30["stp_root_identities"][0]["observers"] == ["core1"]
    assert rows[10]["stp_root"] == "core1"
    for r in rows.values():
        if r["stp_root_state"] != "published":
            assert r["stp_root_default_election"] is None, r["vlan"]


def test_g15_failover_current_root_abstains_when_an_observer_disagrees():
    from cisco_toolkit import failover

    snap = {"stp_roots": {"a": {"30": _row(24606, "aaaa.0000.0001", True, bridge_priority=24606)},
                          "b": {"30": _row(24606, "cccc.0000.0003", False, bridge_priority=32798)}}}
    rows = failover.compute_stp_failover(snap, ["a"])
    assert rows and rows[0]["indeterminate"] is True


def test_g15_explorer_election_mirror_matches_the_python_owner(tmp_path):
    import json
    import shutil
    import subprocess
    from pathlib import Path

    node = shutil.which("node")
    if not node:
        pytest.skip("node is not installed")
    html = (Path(analyze.__file__).parent / "blast_radius_explorer.html").read_text(encoding="utf-8")
    start = html.index("function stpElectionPriority(")
    fn = html[start:html.index("function stpRootFor(", start)]
    cases = [_sample()["stp_roots"], _stp_fixture(),
             {"a": {"30": _row(24606, "aaaa.0000.0001", True)}, "b": {"30": _row(24606, "cccc.0000.0003", False)}},
             {"a": {"30": _row(24606, "aaaa.0000.0001", True)}, "b": {"30": _row(24606, "aaaa.0000.0001", "true")}},
             {"a": {"30": _row("24576", "aaaa.0000.0001", True)}},
             {"a": {"1": {**_row(32769, "aaaa.0000.0001", True), "is_mst": True}},
              "b": {"1": _row(24577, "bbbb.0000.0002", True)}}]
    driver = tmp_path / "election.js"
    driver.write_text(fn + "\nprocess.stdout.write(JSON.stringify(" + json.dumps(cases)
                      + ".map(c=>stpRootElection(c))));\n", encoding="utf-8")
    out = subprocess.run([node, str(driver)], capture_output=True, text=True, encoding="utf-8", timeout=60)
    assert out.returncode == 0, out.stderr
    want = [{v: {"state": r["state"], "reason": r["reason"], "root": r["root"] or "", "claimants": r["claimants"],
                 "default_election": r["default_election"], "root_priority": r["root_priority"]}
             for v, r in _election(c)["pvst_vlan"].items()} for c in cases]
    assert json.loads(out.stdout) == want

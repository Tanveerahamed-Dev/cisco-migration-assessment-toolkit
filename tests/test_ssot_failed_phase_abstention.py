"""G14: a section whose analysis phase FAILED must never read as "collected, nothing found".

`_run_phase` is fail-soft: a phase that raises returns its `_default` (usually `[]` / `{}`) and the
failure is recorded in `assessment_integrity.failed_phases`. The SSOT abstention core
(`ssot.abstention_reason`) never read that record, so the fallback -- byte-identical to a clean
empty result -- was classified `collected_but_empty` ("collected, nothing of this kind found (not a
blind spot)"). The coverage census, the state-assertion pack (a `not_contains "Critical"` check over
a crashed punch-list PASSED) and the Law-8 eval all inherited that lie.

The honest token is `analysis_unavailable` -- the engine already uses it for a failed protocol
analysis -- attributed through ONE registry (`ssot.PHASE_SECTIONS`) that an AST ratchet below keeps
complete against `COLLECT_PARSE_V3_23_0.main()`: a phase added there without a classification fails
this file rather than silently tainting nothing.
"""
import ast
import json
import os
import pathlib
import shutil
import sys

import pytest

from cisco_toolkit import ssot

NODE_BIN = shutil.which("node")
ROOT = pathlib.Path(__file__).resolve().parent.parent
ENGINE = ROOT / "COLLECT_PARSE_V3_23_0.py"
GOLDEN = ROOT / "tests" / "golden" / "snapshot.json"
SAMPLE = ROOT / "webapp" / "sample_data" / "sample_fleet.snapshot.json"
AU = "analysis_unavailable"
CBE = "collected_but_empty"


def _failed(labels, **sections):
    return dict(sections, assessment_integrity={"failed_phases": labels})


# --------------------------------------------------------------------------------------------------
# abstention_reason
# --------------------------------------------------------------------------------------------------
def test_t1_failed_phase_section_is_unavailable_and_an_independent_empty_is_not():
    snap = _failed(["QoS audit"], qos_audit={}, vpc={})
    assert ssot.abstention_reason(snap, "qos_audit") == AU
    assert ssot.abstention_reason(snap, "vpc") == CBE
    assert ssot.ANALYSIS_UNAVAILABLE == AU


def test_t2_a_populated_error_default_is_unavailable_not_published():
    default = {"schema": "traffic_assurance_set/1", "state": "error", "results": [],
               "summary": {"n": 0, "proven": 0, "refuted": 0}}
    assert ssot.abstention_reason(_failed(["Traffic assurance"], traffic_assurance=default),
                                  "traffic_assurance") == AU


def test_t3_unavailable_sentinel_taints_the_section_and_its_subpaths():
    snap = {"executive_brief": {"_unavailable": True, "ssot": {"n_facts": 3}}}
    assert ssot.abstention_reason(snap, "executive_brief") == AU
    assert ssot.abstention_reason(snap, "executive_brief.ssot.n_facts") == AU


def test_t4_integrity_stamp_on_a_section_key_is_honoured_and_a_non_section_stamp_is_not():
    snap = {"assessment_integrity": {"architecture_review": "compute_failed"},
            "architecture_review": {"grade": "A"}}
    assert ssot.abstention_reason(snap, "architecture_review") == AU
    drift = {"assessment_integrity": {"ssot_reconciliation": "failed"}, "vpc": {}}
    assert ssot.abstention_reason(drift, "vpc") == CBE


def test_t5_an_absent_section_of_a_failed_phase_is_unavailable_not_a_blind_spot():
    assert ssot.abstention_reason(_failed(["State assertion pack"]), "state_assertions") == AU
    assert ssot.abstention_reason({}, "state_assertions") == "not_collected"


def test_t6_a_failed_workbook_sheet_writer_does_not_taint_the_section():
    assert ssot.abstention_reason(_failed(["QoS Audit sheet"], qos_audit={}), "qos_audit") == CBE


def test_t7_a_failed_intermediate_taints_every_empty_section_but_not_populated_ones():
    snap = _failed(["dependency map"], cross_layer=[], vpc={"sw1": {"role": "primary"}})
    assert ssot.abstention_reason(snap, "cross_layer") == AU
    assert ssot.abstention_reason(snap, "vpc") == "published"


def test_t8_an_unknown_label_fails_closed():
    snap = _failed(["X new phase"], qos_audit={}, vpc={"sw1": {"role": "primary"}})
    assert ssot.abstention_reason(snap, "qos_audit") == AU
    assert ssot.abstention_reason(snap, "vpc") == "published"


def test_t9_a_malformed_failure_record_fails_closed_and_garbage_does_not_crash():
    snap = {"assessment_integrity": {"failed_phases": "QoS audit"}, "qos_audit": {}}
    assert ssot.abstention_reason(snap, "qos_audit") == AU
    garbage = {"assessment_integrity": "garbage", "vpc": {}}
    assert ssot.abstention_reason(garbage, "vpc") == CBE
    assert ssot.failed_sections(garbage) == (frozenset(), False)
    assert ssot.failed_sections(None) == (frozenset(), False)


def test_t10_a_device_blind_spot_still_wins():
    snap = _failed(["QoS audit"], qos_audit={},
                   collection_completeness={"devices": [{"host": "sw9", "status": "not collected"}]})
    assert ssot.abstention_reason(snap, "qos_audit", device="sw9") == "not_collected"


def test_t11_a_non_string_subject_is_total():
    """The docstring promises a total function; `"." in 5` raised TypeError."""
    assert ssot.abstention_reason({"a": 1}, 5) == "not_collected"
    assert ssot.abstention_reason({"a": 1}, ["a"]) == "not_collected"
    assert ssot.abstention_reason({"a": 1}, None) == "not_collected"


# --------------------------------------------------------------------------------------------------
# one hop downstream: a canonical fact DERIVED from a failed phase's fallback
# --------------------------------------------------------------------------------------------------
def _crashed_health(**posture):
    return {"assessment_integrity": {"failed_phases": ["Health Scores", "Endpoint identity"]},
            "health_scores": [], "endpoint_identity": [],
            "executive_brief": {"scale": {"n_devices": 0, "n_endpoints": 0, "n_domains": 2},
                                "posture": dict({"avg_health": None, "n_critical": 0, "n_poor": 0,
                                                 "worst_band": None}, **posture)},
            "application_intelligence": {"summary": {"n_domains": 2}, "domains": [{}, {}]},
            "collection_completeness": {"summary": {"inventory": 3, "complete": 3}, "devices": []}}


def test_d1_lineage_marks_facts_derived_from_a_crashed_phase_unavailable():
    """With 'Health Scores' failed, lineage published n_devices / n_critical / n_poor = 0 as
    `collected_but_empty` -- 'looked, found nothing' about a count taken over the crash fallback."""
    states = {f["name"]: f["state"] for f in ssot.compute_fact_lineage(_crashed_health())["facts"]}
    for name in ("n_devices", "n_endpoints", "avg_health", "n_critical", "n_poor", "worst_band"):
        assert states[name] == AU, (name, states[name])
    # NON-VACUITY: facts whose basis did not fail keep their own state
    assert states["n_domains"] == "published"


def test_d2_an_assertion_over_a_derived_fact_of_a_crashed_phase_abstains():
    from cisco_toolkit import assertions
    spec = {"id": "no-critical-switches", "subject": "executive_brief.posture.n_critical",
            "all_of": [{"type": "comparison", "op": "==", "value": 0}]}
    r = assertions.evaluate_assertion(_crashed_health(), spec)
    assert r["status"] == assertions.NOT_OBSERVED and r["abstention"] == AU, r
    # NON-VACUITY: the same assertion over a clean, genuinely-zero posture still passes
    clean = {"health_scores": [{"switch": "a", "band": "Good", "score": 80}],
             "executive_brief": {"posture": {"n_critical": 0}}}
    assert assertions.evaluate_assertion(clean, spec)["status"] == "pass"


def test_d3_every_derived_canonical_fact_names_its_raw_basis():
    """Ratchet: a headline fact stored under `executive_brief` is computed from OTHER sections; one
    whose basis is not registered would read `collected_but_empty` over a crashed producer again."""
    import re
    known = {s for secs in ssot.PHASE_SECTIONS.values() for s in secs}
    for prefix, basis in ssot.DERIVED_FACT_BASIS.items():
        assert basis and set(basis) <= known, (prefix, basis)
    for name, (path, concept) in ssot.CANONICAL_FACTS.items():
        basis = ssot.fact_basis(path)
        assert basis[0] == path.split(".", 1)[0], (name, basis)
        if path.startswith("executive_brief."):
            assert len(basis) > 1, f"{name}: derived headline fact with no registered raw basis"
        for hint in re.findall(r"len\((?:analyze\.)?(\w+)\)", concept):
            if hint in known:
                assert hint in basis, (name, hint, basis)


def test_d4_committed_sample_lineage_is_unchanged():
    s = json.loads(SAMPLE.read_text(encoding="utf-8"))
    if "fact_lineage" not in s:
        pytest.skip("sample carries no embedded fact_lineage")
    assert ssot.compute_fact_lineage(s) == s["fact_lineage"]


def test_failed_sections_reports_attribution():
    direct, unattributed = ssot.failed_sections(
        _failed(["Golden-config drift", "QoS Audit sheet", "HTML Explorer"]))
    assert {"golden_drift", "feature_compliance"} <= direct and not unattributed
    _d, unattributed = ssot.failed_sections(_failed(["dependency map"]))
    assert unattributed


# --------------------------------------------------------------------------------------------------
# consumers: census / assertions / Law-8 eval
# --------------------------------------------------------------------------------------------------
def test_c1_census_counts_the_new_state_honestly():
    census = ssot.compute_schema_census(_failed(["QoS audit"], qos_audit={}, vpc={}))
    rows = {r["key"]: r for r in census["sections"]}
    assert rows["qos_audit"]["state"] == AU and rows["vpc"]["state"] == CBE
    note = rows["qos_audit"]["note"].lower()
    assert "failed" in note and not any(w in note.split() for w in ("ok", "healthy"))
    s = census["summary"]
    assert s["n_analysis_unavailable"] == 1
    assert (s["n_published"] + s["n_collected_but_empty"] + s["n_not_collected"]
            + s["n_analysis_unavailable"]) == s["n_sections"]
    clean = ssot.compute_schema_census({"vpc": {}, "qos_audit": {"a": 1}})["summary"]
    assert set(clean) == {"n_published", "n_collected_but_empty", "n_not_collected", "n_sections"}


def test_c2_coverage_sheet_accounts_for_the_new_state():
    """The Coverage Schema totals row omitted `analysis_unavailable`, so its counts silently stopped
    summing to n_sections; the crashed rows carried no fill at all."""
    import re
    from openpyxl import Workbook
    from cisco_toolkit.excel import COVERAGE_SCHEMA_SHEET_NAME, write_coverage_schema_sheet
    census = ssot.compute_schema_census(_failed(["QoS audit"], qos_audit={}, vpc={}, fhrp={"a": 1}))
    wb = Workbook()
    write_coverage_schema_sheet(wb, census)
    ws = wb[COVERAGE_SCHEMA_SHEET_NAME]
    totals = str(ws.cell(3, 5).value)
    assert "analysis FAILED 1" in totals, totals
    assert sum(int(x) for x in re.findall(r"\d+", totals)) == ws.cell(3, 4).value == census["summary"]["n_sections"]
    row = next(r for r in range(4, ws.max_row + 1) if ws.cell(r, 1).value == "qos_audit")
    assert ws.cell(row, 2).value == AU
    assert str(ws.cell(row, 2).fill.fgColor.rgb).upper().endswith("F4CCCC")
    # NON-VACUITY: a clean census keeps its historical totals text exactly
    clean_wb = Workbook()
    write_coverage_schema_sheet(clean_wb, ssot.compute_schema_census({"vpc": {}, "fhrp": {"a": 1}}))
    assert clean_wb[COVERAGE_SCHEMA_SHEET_NAME].cell(3, 5).value == (
        "published 1 · collected-but-empty 1 · NOT collected 0 (blind spots)")


def test_a1_assertion_over_a_crashed_section_abstains_instead_of_passing():
    from cisco_toolkit import assertions
    snap = {"assessment_integrity": {"failed_phases": ["Migration Punch-List"]}, "punchlist": []}
    spec = {"id": "no-crit", "subject": "punchlist",
            "all_of": [{"type": "not_contains", "value": "Critical"}]}
    r = assertions.evaluate_assertion(snap, spec)
    assert r["status"] == assertions.NOT_OBSERVED and r["abstention"] == AU, r
    pack = assertions.evaluate_pack(snap, {"assertions": [spec]})
    assert pack["summary"]["n_pass"] == 0 and pack["summary"]["n_assessed"] == 0, pack["summary"]


def test_a1b_for_each_over_a_crashed_section_abstains():
    from cisco_toolkit import assertions
    snap = {"assessment_integrity": {"architecture_review": "compute_failed"},
            "architecture_review": {"sw1": {"grade": "A"}}}
    spec = {"id": "fe", "for_each": "architecture_review",
            "field_rules": [{"field": "grade", "type": "contains", "value": "A"}]}
    r = assertions.evaluate_pack(snap, {"for_each": [spec]})["results"][0]
    assert r["status"] == assertions.NOT_OBSERVED and r["abstention"] == AU, r


def test_e1_law8_census_accounts_for_the_new_state():
    from cisco_toolkit import eval_harness
    snap = _failed(["QoS audit"], qos_audit={}, vpc={})
    assert ssot.compute_schema_census(snap)["summary"].get("n_analysis_unavailable") == 1
    checks = {c.check_id: c for c in eval_harness._check_not_observed(snap, None)}
    assert checks["law8.census"].status == eval_harness._PASS, checks["law8.census"]


def test_e1b_law8_rejects_a_healthy_note_on_an_unavailable_row(monkeypatch):
    from cisco_toolkit import eval_harness
    monkeypatch.setitem(ssot._CENSUS_NOTE, AU, "ok")
    checks = {c.check_id: c for c in eval_harness._check_not_observed(
        _failed(["QoS audit"], qos_audit={}), None)}
    assert checks["law8.census"].status == eval_harness._FAIL


# --------------------------------------------------------------------------------------------------
# the offline explorer's Ask-Atlas port of the abstention core (executed in the real embedded script)
# --------------------------------------------------------------------------------------------------
def _explorer_run(driver, tmp_path):
    if str(ROOT / "tests") not in sys.path:
        sys.path.insert(0, str(ROOT / "tests"))
    import test_explorer_render_safety as harness      # the repo's DOM-stubbed full-script runner
    return harness._run(driver, tmp_path)


def _with_census(snap):
    snap = dict(snap)
    snap["schema_census"] = ssot.compute_schema_census(snap)
    return snap


_JS_CASES = {
    # the engine's own projection (schema_census) is present: the explorer reads it, no second registry
    "census": _with_census(_failed(["QoS audit", "Syslog intelligence", "Platform health"],
                                   qos_audit={}, syslog_intelligence={}, platform_health={}, vpc={})),
    # an older snapshot with no census: only section-keyed signals attribute; the rest fails CLOSED
    "no_census": _failed(["QoS audit"], qos_audit={}, vpc={"sw1": {"role": "primary"}}),
    "sentinel": {"platform_health": {"_unavailable": True}, "vpc": {}},
    "stamp": {"assessment_integrity": {"qos_audit": "compute_failed"}, "qos_audit": {"sw1": {"x": 1}}},
    "clean": {"qos_audit": {}, "vpc": {"sw1": {"role": "primary"}}},
}
_JS_SUBJECTS = ("qos_audit", "syslog_intelligence", "platform_health", "vpc")


@pytest.mark.skipif(not NODE_BIN, reason="node is not installed -- executed explorer check skipped")
def test_x1_explorer_abstention_reads_failures_and_never_greens_a_crash(tmp_path):
    """abAbstention ignored assessment_integrity: a crashed QoS / syslog / platform phase got the green
    'NOT OBSERVED · collected, none found … a real negative, not a blind spot' pill."""
    payload = json.dumps(_JS_CASES)
    out = _explorer_run(r"""
      const CASES=""" + payload + r""";
      const res={};
      for(const [k,s] of Object.entries(CASES)){globalThis.__S=s;__EV("SNAP=globalThis.__S");
        res[k]={};
        for(const subj of """ + json.dumps(list(_JS_SUBJECTS)) + r""")res[k][subj]=__EV("abAbstention")(subj);
        const r=__EV("abH_notobserved")("what is my qos posture",{});
        res[k].qos_pill=(r.html.match(/pill (pl-\w+)/)||[])[1]||"";
        res[k].qos_text=r.html.replace(/<[^>]+>/g,"");}
      console.log(JSON.stringify(res));
    """, tmp_path)
    for case, snap in _JS_CASES.items():
        for subj in _JS_SUBJECTS:
            py = ssot.abstention_reason(snap, subj)
            assert out[case][subj] == py, (case, subj, out[case][subj], py)
    for case in ("census", "no_census", "stamp"):
        assert out[case]["qos_audit"] == AU, (case, out[case])
        assert out[case]["qos_pill"] != "pl-ok", (case, out[case])
        assert "analysis failed" in out[case]["qos_text"].lower(), (case, out[case]["qos_text"])
        assert "real negative" not in out[case]["qos_text"], (case, out[case]["qos_text"])
    assert out["sentinel"]["platform_health"] == AU
    # NON-VACUITY: a clean empty is still the honest 'collected, none found'
    assert out["clean"]["qos_audit"] == CBE and out["clean"]["qos_pill"] == "pl-ok"


# --------------------------------------------------------------------------------------------------
# the committed snapshots are unchanged (no failure record in either)
# --------------------------------------------------------------------------------------------------
@pytest.mark.parametrize("path", [GOLDEN, SAMPLE], ids=["golden", "sample"])
def test_g1_committed_snapshots_classify_exactly_as_embedded(path):
    s = json.loads(path.read_text(encoding="utf-8"))
    assert ssot.failed_sections(s) == (frozenset(), False)
    embedded = {r["key"]: r["state"] for r in s["schema_census"]["sections"]}
    # The golden strips its volatile sections (generated_at, executive_brief, ...) AFTER the census was
    # embedded; only the sections it still carries can be re-classified from its own bytes.
    present = [k for k in embedded if k in s]
    assert len(present) >= 100, len(present)
    for key in present:
        assert ssot.abstention_reason(s, key) == embedded[key], key


def test_g1_sample_census_recomputes_byte_equal():
    s = json.loads(SAMPLE.read_text(encoding="utf-8"))
    assert ssot.compute_schema_census(s) == s["schema_census"]


def test_r2_the_token_is_the_engine_vocabulary():
    from cisco_toolkit import analyze
    assert ssot.ANALYSIS_UNAVAILABLE in analyze.PROTOCOL_ASSESSABILITY_STATES
    assert ssot.ANALYSIS_UNAVAILABLE in ssot.ABSTENTION_STATES
    assert set(ssot._CENSUS_NOTE) == set(ssot.ABSTENTION_STATES)


# --------------------------------------------------------------------------------------------------
# R1: the registry is complete against the engine's own source (AST ratchet)
# --------------------------------------------------------------------------------------------------
_RECORDERS = ("_run_phase", "_emit_artifact", "_record_phase_failure", "_record_mandatory_failure",
              "_atomic_receipt_refresh")


def _engine():
    tree = ast.parse(ENGINE.read_text(encoding="utf-8"))
    main = next(n for n in tree.body if isinstance(n, ast.FunctionDef) and n.name == "main")
    return tree, main


def _label_node(call):
    if isinstance(call, ast.Call) and isinstance(call.func, ast.Name) and call.func.id in _RECORDERS \
            and call.args:
        return call.args[0]
    return None


def _all_labels(tree):
    """Every literal failure label the engine can record, plus the leading text of f-string labels."""
    literal, fstring_prefixes = set(), set()
    for n in ast.walk(tree):
        node = _label_node(n)
        if isinstance(node, ast.Constant) and isinstance(node.value, str):
            literal.add(node.value)
        elif isinstance(node, ast.JoinedStr):
            head = node.values[0] if node.values else None
            fstring_prefixes.add(head.value if isinstance(head, ast.Constant) else "")
        # a Name label is either a recorder's own parameter or the data-authority loop variable below
        if isinstance(n, ast.For) and isinstance(n.target, ast.Tuple) and isinstance(n.iter, ast.Tuple):
            names = [t.id for t in n.target.elts if isinstance(t, ast.Name)]
            if "_axis" in names:
                idx = names.index("_axis")
                for elt in n.iter.elts:
                    if isinstance(elt, ast.Tuple) and isinstance(elt.elts[idx], ast.Constant):
                        literal.add(elt.elts[idx].value)
    return literal, fstring_prefixes


def _direct_edges(main):
    """section -> {labels}: `snap_dict["k"] = <expr>` where <expr> is a `_run_phase("L")` call or
    names a variable assigned from one."""
    var_labels, edges = {}, {}

    def calls(expr):
        return [c for c in ast.walk(expr) if isinstance(c, ast.Call) and isinstance(c.func, ast.Name)
                and c.func.id == "_run_phase" and c.args and isinstance(c.args[0], ast.Constant)]

    for n in ast.walk(main):
        if not isinstance(n, ast.Assign):
            continue
        cs = calls(n.value)
        for t in n.targets:
            if isinstance(t, ast.Name) and cs:
                var_labels.setdefault(t.id, set()).update(c.args[0].value for c in cs)
    for n in ast.walk(main):
        if not isinstance(n, ast.Assign):
            continue
        for t in n.targets:
            if (isinstance(t, ast.Subscript) and isinstance(t.value, ast.Name) and t.value.id == "snap_dict"
                    and isinstance(t.slice, ast.Constant)):
                labels = {c.args[0].value for c in calls(n.value)}
                for name in {x.id for x in ast.walk(n.value) if isinstance(x, ast.Name)}:
                    labels |= var_labels.get(name, set())
                if labels:
                    edges.setdefault(t.slice.value, set()).update(labels)
    return edges


def test_r1_every_direct_phase_to_section_edge_is_registered():
    _tree, main = _engine()
    edges = _direct_edges(main)
    assert len(edges) >= 50, "the edge extractor went blind -- a ratchet over nothing proves nothing"
    missing = sorted((k, lab) for k, labs in edges.items() for lab in labs
                     if k not in ssot.PHASE_SECTIONS.get(lab, ()))
    assert not missing, f"phase->section edges absent from ssot.PHASE_SECTIONS: {missing}"


def test_r1_every_failure_label_is_classified():
    tree, _main = _engine()
    literal, prefixes = _all_labels(tree)
    assert len(literal) >= 150, "the label extractor went blind"
    unclassified = sorted(lab for lab in literal if ssot.phase_classification(lab) == "unknown")
    assert not unclassified, f"classify these in ssot (sections / intermediate / non-section): {unclassified}"
    assert prefixes == {"interface rows ("}, prefixes
    assert all(ssot.phase_classification(p + "sw1)") == "non_section" for p in prefixes)


def test_r1_registry_sets_are_disjoint_and_sheet_writers_own_no_section():
    sec, mid, non = set(ssot.PHASE_SECTIONS), set(ssot.INTERMEDIATE_PHASES), set(ssot.NON_SECTION_PHASES)
    assert not (sec & mid) and not (sec & non) and not (mid & non)
    assert not [lab for lab in sec | mid if lab.endswith(" sheet")]
    _tree, main = _engine()
    for n in ast.walk(main):
        if not isinstance(n, ast.Assign):
            continue
        for c in ast.walk(n.value):
            node = _label_node(c)
            if (getattr(c, "func", None) is not None and getattr(c.func, "id", "") == "_run_phase"
                    and isinstance(node, ast.Constant) and node.value.endswith(" sheet")):
                pytest.fail(f"sheet writer {node.value!r} returns a value that is used")


# --------------------------------------------------------------------------------------------------
# P1: the real engine
# --------------------------------------------------------------------------------------------------
def _run_engine(tmp_path, monkeypatch):
    sys.path.insert(0, str(ROOT / "tests"))
    import synthetic_fixtures as fx
    from openpyxl import Workbook
    if str(ROOT) not in sys.path:
        sys.path.insert(0, str(ROOT))
    import COLLECT_PARSE_V3_23_0 as cp
    collection = fx.write_collection(str(tmp_path / "collection"))
    devices = tmp_path / "devices.json"
    devices.write_text(json.dumps(fx.DEVICES), encoding="utf-8")
    template = tmp_path / "template.xlsx"
    wb = Workbook(); wb.active.title = "Interface Data"; wb.active.append(["Hostname", "Port", "Status"])
    wb.save(str(template))
    out = tmp_path / "out.xlsx"
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(sys, "argv", ["cisco-assess", "--no-collect", "--collection-dir", collection,
                                      "--devices-file", str(devices), "--template", str(template),
                                      "--output", str(out), "--workers", "1", "--no-html"])
    return cp, out


def _census_rows(out):
    snap = json.loads(pathlib.Path(str(out)[:-len(".xlsx")] + ".snapshot.json").read_text(encoding="utf-8"))
    return snap, {r["key"]: r["state"] for r in snap["schema_census"]["sections"]}


def test_p1_real_engine_reports_crashed_phases_as_unavailable(tmp_path, monkeypatch):
    cp, out = _run_engine(tmp_path, monkeypatch)

    def _boom(*a, **kw):
        raise RuntimeError("synthetic phase failure")

    for name in ("compute_cross_layer_correlations", "_punchlist", "compute_qos_audit"):
        monkeypatch.setattr(cp, name, _boom)
    cp.main()
    snap, rows = _census_rows(out)
    for key in ("cross_layer", "punchlist", "qos_audit"):
        assert rows[key] == AU, (key, rows[key], snap.get("assessment_integrity"))
    for key in ("vpc", "mroute", "trunk_native", "link_phy", "addressing_conflicts"):
        if key in rows and rows[key] != "published":
            assert rows[key] == CBE, (key, rows[key])
    assert snap["schema_census"]["summary"]["n_analysis_unavailable"] >= 3


def test_p1_clean_real_engine_run_has_no_unavailable_row(tmp_path, monkeypatch):
    cp, out = _run_engine(tmp_path, monkeypatch)
    cp.main()
    snap, rows = _census_rows(out)
    assert "assessment_integrity" not in snap, snap.get("assessment_integrity")
    assert AU not in rows.values()
    assert "n_analysis_unavailable" not in snap["schema_census"]["summary"]

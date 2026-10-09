"""W48: every consumer that presents a failure-impact value, rank or severity reads it through the engine owner.

``cisco_toolkit/impact_assessability.py`` (W33, with W32's blind-link rule) is the one owner of whether a stored
``failure_impact`` row is a measurement: ``published``, ``lower_bound`` (the worst band and each positive count are
floors; a band below the worst and a zero are not measurements), ``not_assessed`` (a hold: none of the row's values
is a measurement) or ``ambiguous``. W33 moved the projection and the workbook, design, deck, RES-4, dossier and
explorer onto it. W48 moves the rest of the class: the MOP's per-wave blast radius and rollback trigger, the
runbook's §10 Risk Register, the operations handbook's §2.1 keystones, the MCP ``failure_impact`` tool and the
AssessHub dossier recompute. After its independent review it also owns the WAVE rule (``wave_blast``): the MOP
applies it to the owner's verdicts and AssessHub's cutover plan to the projection's rows (built from those verdicts),
so the two classify a wave by one rule. They remain two readings of the rows, so their agreement is pinned per variant
(``webapp/tests/test_impact_surfaces.py``), not guaranteed by construction.

Three things are pinned here:

1. **A structural guard** (part 1). The class is defined by what a unit READS, never by a hand-kept list of
   consumers. A unit is every function and method; every nested function (``create_app.route``), method of a nested
   class (``A.B.render``) and def inside a module-level ``if``/``try``; every class body; every module-level
   assignment to a name (a dispatch table); and each module's remaining top-level code, under ``cisco_toolkit/``,
   ``webapp/backend/`` and ``COLLECT_PARSE_V3_23_0.py``. A unit READS the stored section when it passes the section
   name as a call argument (``snap.get``, ``getattr``, ``pop``, a path tuple that starts with it), uses it as a
   subscript or mapping-pattern key, loads ``.failure_impact`` or matches a ``failure_impact=`` class pattern, loads
   a ``failure_impact`` name that is no project unit (a parameter or local), or compares a key that iterates data
   (``for k, v in snap.items()``) against it. The section name is folded only through a literal, a plain name or an
   imported module's attribute bound to a module-level, explicitly imported or once-assigned (single-name) local key
   constant, ``+``, f-strings of unconverted foldable parts, ``sep.join([...])`` of a literal list, and a loop name
   iterating a MODULE-LEVEL literal table that holds it; every other way of computing the key is a known limit
   (below). A unit also reads when it is a nested function loading a name its
   enclosing function assigned from a read, or a helper that reads a row field (``r['stranded']``,
   ``r.get('severity')``) after another unit handed it the rows, as the callee or as a function passed beside them
   (a phase runner).

   Every reader must reach the owner. A route is a CALL to an owner function or class (a module-level table may hold
   one uncalled), directly or through any other project unit it references: calls, dispatch tables, callbacks,
   ``self.method``, ``Class.method``, a local alias such as ``P = _PURE``, and a class whose ``__init__``/``__new__``/
   ``__call__``/``__post_init__`` reaches it. Naming an owner CONSTANT, or an owner function without calling it, is
   not a route. The owner, the projection and AssessHub's surfaces pass by that property, not by name. The only
   named entries are a RATCHET of the consumers W48 stopped on because their output is persisted (see each entry): a
   new raw reader fails, and fixing a ratchet entry fails until the entry and its strict ``xfail`` below are deleted.
   W50 narrowed the ``protocol_assurance`` entry to the one frozen EVIDENCE binder (``_rehearsal_impact_evidence_v1``):
   an execution receipt is re-verified by recomputing it on every read, so it binds the rows raw and must never
   consult the evolving owner, while every presentation of those rows goes through the owner at display time
   (AssessHub's live ``impacts_view``). That entry stays a raw reader by design and is pinned to the binder alone; any
   other raw reader in the module still fails. The receipt-recompute closure check (W50) reads its own conservative
   resolver (:func:`_graph`), not this guard.

   The line check is the scanner's independent cross-check, line by line: every line where the tokenizer sees the
   section name (a NAME token, or a string literal that is exactly it; comments and docstrings are neither) must hold
   an AST occurrence, and every occurrence must be a read attributed to a unit or one of the closed non-read
   positions: a literal table that is not a path, a dict-literal key, a comparison, a key constant's definition, an
   f-string part, a store, a parameter or keyword name, a def name, an import, or a reference to a project unit.
   Anything else fails as unclassified. The review's evasions of the first guard are pinned in ``_EVASIONS``.

   **Known limits** (each pinned by ``test_each_documented_known_limit_is_still_a_limit``, so this list stays true):

   * dataflow is followed one assignment from a read. A copy of a copy (``again = list(rows)``), rows kept in an
     attribute, a container or a return value and consumed elsewhere, and rows a helper passes on to a second helper
     are not followed; a helper handed the rows counts as a reader only when it reads a row field by a literal or
     foldable key;
   * the key folding has holes. Not folded: ``%`` or ``str.format``; a tuple-unpacked or re-assigned local; a local
     literal table; a method call, slice, ``or``, conditional expression or ``!s``/``!r`` conversion over a key
     constant (``KEY.strip()``, ``KEY[:]``, ``KEY or ''``, ``KEY if x else ''``, ``f'{KEY!s}'``); ``globals()['KEY']``;
     a star import; ``importlib``; ``b'...'.decode()``; and a key read from data, a dict-literal value or a class
     attribute. Each is pinned in ``_KNOWN_LIMITS``. The line check does NOT close these: each places the literal
     section name in a closed position (a key constant's definition, a literal table) or not at all (bytes), so such
     a read passes both checks. The line check catches only a literal section name in an unlisted position;
   * calls through an instance (``ctx.impact.row()``), ``getattr(module, name)()``, ``importlib``, registries filled
     at run time and decorators that replace a function are not resolved. A reference to a project unit counts as a
     route, so a unit that only names a routed helper, without calling it, is admitted;
   * the granularity is the unit: a unit that reaches the owner for one value can still print another raw value.
     Part 2 pins what each moved consumer prints;
   * only Python under the three roots is scanned. Atlas Scope's ``atlas-scope/tools/lib/compile-model.mjs`` reads
     the raw rows (a Codex handoff, ``docs/w48-impact-consumers-validation-2026-10-09.md``); the SPA, the explorer's
     in-browser JS, ``tools/`` and ``portable/`` are outside.

2. **Behaviour on the committed sample** (part 2). On the regenerated sample the owner bounds core1 (one
   inter-switch link with no trunk/STP evidence, W32). Each moved consumer must present core1 as the lower bound it
   is, and a held row (core1 with its ``off_scan_gw_vlans`` marker removed: a row older than the marker) as not
   assessed. Every expectation is read from the owner (``table_value`` / ``ranked_value`` / ``ranking_floor`` /
   ``disclose`` / ``wave_blast``), never cached from the sample. Each test fails on the pre-W48 code, which printed
   core1's raw 45. The MOP also never sizes a wave or its rollback trigger on a zero that is only a lower bound, and a
   wave whose every device is published keeps its bare count and quantified trigger.

3. **The owner's wave rule** (part 3): exact, lower bound and not assessed, a device with no row, a row that names no
   readable host, the projection's fleet qualifier, and a zero lower bound.

No test here runs a pipeline. Written for the hosted runners (owner GitHub-only rule); not run locally.
"""
from __future__ import annotations

import ast
import copy
import functools
import io
import json
import os
import pathlib
import re
import tokenize
from typing import NamedTuple, Optional

import pytest

from cisco_toolkit import impact_assessability as ia
from cisco_toolkit import mcp_server

ROOT = pathlib.Path(__file__).resolve().parent.parent
SAMPLE = ROOT / "webapp" / "sample_data" / "sample_fleet.snapshot.json"

# ---------------------------------------------------------------------------------------------------------------
# part 1: the structural guard
# ---------------------------------------------------------------------------------------------------------------
SECTION = "failure_impact"
OWNER = "cisco_toolkit.impact_assessability"
#: (directory, dotted package) roots whose Python modules are scanned, plus the pipeline entry module.
_PACKAGE_ROOTS = (("cisco_toolkit", "cisco_toolkit"), (os.path.join("webapp", "backend"), "webapp.backend"))
_PIPELINE = "COLLECT_PARSE_V3_23_0.py"
#: A stored failure-impact row's fields: a helper handed the stored rows that reads one of these presents them.
_ROW_FIELDS = frozenset({"severity", "vlans_impacted", "stranded", "hard", "backup", "fhrp", "detail",
                         "off_scan_gw_vlans", "blind_links"})
_CALLABLE = ("function", "method", "class")
#: Calling a class runs these, so a class reaches the owner when one of them does.
_CLASS_ENTRY = ("__init__", "__new__", "__call__", "__post_init__")
_SCOPED = (ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)
_COMPOUND = tuple(t for t in (ast.If, ast.For, ast.AsyncFor, ast.While, ast.With, ast.AsyncWith, ast.Try,
                              getattr(ast, "TryStar", None), ast.Match) if t is not None)
_COLLECTION = (ast.Tuple, ast.List, ast.Set)

#: Raw readers W48 STOPPED on, each because changing what it presents changes persisted output; routed to the
#: supervisor for a hosted regeneration (docs/w48-impact-consumers-validation-2026-10-09.md). Keyed by
#: (module, unit), a unit named as the guard names it (``f``, ``C.m``, ``f.inner``). Each has a strict xfail in part 2.
#: Delete the entry when its consumer reads the owner.
#: The one exception is the protocol_assurance entry (W50): a frozen evidence binder that must stay raw, pinned to
#: that single function and checked by a behavioural test of the display path instead of an xfail.
_RAW_RATCHET = {
    ("cisco_toolkit.design_advisor", "_signals"): (
        "nobackup_high counts raw High rows with a raw zero backup, so a lower-bound or held row's withheld zero "
        "is counted as a measured no-backup device. Its text and count are stored in the snapshot's "
        "design_blueprint (decisions[topology-triangles-not-squares-rings].evidence.summary and "
        "tradeoff_scorecard[availability].evidence), so the fix needs a hosted sample regeneration."),
    # Allowlisted for the frozen evidence binder ONLY (W50): no other protocol_assurance function may read the rows
    # (test_only_the_frozen_binder_reads_the_rows_raw_and_the_display_path_reaches_the_owner).
    ("cisco_toolkit.protocol_assurance", "_rehearsal_impact_evidence_v1"): (
        "binds raw evidence for recompute-on-read receipts; presentation goes through the owner at display time "
        "(W50)"),
}


def _module_map(root):
    """dotted module name -> path relative to `root` (POSIX separators)."""
    out = {}
    for base, package in _PACKAGE_ROOTS:
        top = os.path.join(root, base)
        for dirpath, dirnames, filenames in os.walk(top):
            dirnames[:] = sorted(d for d in dirnames if d != "__pycache__")
            for name in sorted(filenames):
                if not name.endswith(".py"):
                    continue
                rel_to_top = os.path.relpath(os.path.join(dirpath, name), top).replace(os.sep, "/")
                dotted = package + "." + rel_to_top[:-3].replace("/", ".")
                if dotted.endswith(".__init__"):
                    dotted = dotted[: -len(".__init__")]
                out[dotted] = os.path.relpath(os.path.join(dirpath, name), root).replace(os.sep, "/")
    if os.path.exists(os.path.join(root, _PIPELINE)):
        out[_PIPELINE[:-3]] = _PIPELINE
    return out


def _resolve_from(module, rel, node):
    """The dotted module an ``ImportFrom`` names, relative imports resolved against `module`."""
    if not node.level:
        return node.module or ""
    package = module if rel.endswith("__init__.py") else module.rpartition(".")[0]
    parts = package.split(".")
    if node.level > 1:
        parts = parts[: len(parts) - (node.level - 1)]
    return ".".join(parts + ([node.module] if node.module else []))


def _aliases(module, rel, tree, modules):
    """Every local name an import binds to a project module or to a name inside one, wherever the import sits
    (lazy imports included): name -> ("module", dotted) | ("name", dotted, attribute)."""
    out = {}
    for node in ast.walk(tree):
        if isinstance(node, ast.Import):
            for alias in node.names:
                if alias.asname and alias.name in modules:
                    out[alias.asname] = ("module", alias.name)
                elif not alias.asname:                  # `import a.b` binds `a`; a.b.f resolves through it
                    out.setdefault(alias.name.split(".")[0], ("module", alias.name.split(".")[0]))
        elif isinstance(node, ast.ImportFrom):
            source = _resolve_from(module, rel, node)
            for alias in node.names:
                local = alias.asname or alias.name
                full = f"{source}.{alias.name}" if source else alias.name
                if full in modules:
                    out[local] = ("module", full)
                elif source in modules:
                    out[local] = ("name", source, alias.name)
    return out


class _Unit(NamedTuple):
    """One scanned unit: a function, method, class body, module-level assignment or the module's other code. A
    nested function, a method of a nested class and a def inside a module-level ``if``/``try`` are units of their own
    (``f.inner``, ``A.B.m``)."""
    name: str
    node: ast.AST
    kind: str                      # function | method | class | data | module
    scope: Optional[str]           # the enclosing unit (None: the module)
    cls: Optional[str]             # the class a method's self/cls names (a function nested in a method keeps it)
    roots: tuple                   # what the unit evaluates itself


def _blocks(stmt):
    """The statement lists a compound statement holds (never a def's or a class's body)."""
    for field in ("body", "orelse", "finalbody"):
        yield getattr(stmt, field, None) or []
    for handler in getattr(stmt, "handlers", None) or []:
        yield handler.body
    for case in getattr(stmt, "cases", None) or []:
        yield case.body


def _scoped_defs(stmts):
    """The defs and classes a block binds in its own scope, at any compound-statement depth."""
    for st in stmts:
        if isinstance(st, _SCOPED):
            yield st
        elif isinstance(st, _COMPOUND):
            for block in _blocks(st):
                yield from _scoped_defs(block)


def _module_assigns(stmts):
    """Module-scope assignments to plain names, at any compound-statement depth (a module-level table)."""
    for st in stmts:
        if isinstance(st, (ast.Assign, ast.AnnAssign)) and st.value is not None:
            if any(isinstance(t, ast.Name) for t in (st.targets if isinstance(st, ast.Assign) else [st.target])):
                yield st
        elif isinstance(st, _COMPOUND):
            for block in _blocks(st):
                yield from _module_assigns(block)


def _units(tree):
    """``(units, bindings)``: every unit by qualified name, and scope -> simple name -> the units it binds."""
    units, bindings, taken = {}, {}, set()

    def fresh(q, node):
        if q in taken:                                  # a name redefined in another branch keeps both units
            q = f"{q}@{node.lineno}"
        taken.add(q)
        return q

    def visit(stmts, prefix, scope, cls, in_class):
        for st in _scoped_defs(stmts):
            q = fresh(prefix + st.name, st)
            bindings.setdefault(scope, {}).setdefault(st.name, []).append(q)
            if isinstance(st, ast.ClassDef):
                units[q] = _Unit(q, st, "class", scope, None, tuple(st.body))
                visit(st.body, q + ".", q, q, True)
            else:
                a = st.args
                params = tuple(a.posonlyargs + a.args + a.kwonlyargs + [x for x in (a.vararg, a.kwarg) if x])
                units[q] = _Unit(q, st, "method" if in_class else "function", scope, cls, params + tuple(st.body))
                visit(st.body, q + ".", q, cls, False)

    visit(tree.body, "", None, None, False)
    assigns = list(_module_assigns(tree.body))
    for st in assigns:
        for target in (st.targets if isinstance(st, ast.Assign) else [st.target]):
            if isinstance(target, ast.Name):
                q = fresh(target.id, st)
                bindings.setdefault(None, {}).setdefault(target.id, []).append(q)
                units[q] = _Unit(q, st, "data", None, None, (st,))
    skip = frozenset(id(st) for st in assigns)
    units["<module>"] = _Unit("<module>", tree, "module", None, None, tuple(s for s in tree.body if id(s) not in skip))
    return units, bindings, skip


def _own(roots, skip=frozenset()):
    """Every node a unit evaluates itself: a nested def or class contributes only its decorators, defaults and
    bases (its body is its own unit); a module-level table is its own unit too."""
    stack = list(reversed(roots))
    while stack:
        node = stack.pop()
        if id(node) in skip:
            continue
        yield node
        if isinstance(node, ast.ClassDef):
            stack.extend(reversed(node.decorator_list + node.bases + [k.value for k in node.keywords]))
        elif isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)):
            a = node.args
            stack.extend(reversed(node.decorator_list + a.defaults + [d for d in a.kw_defaults if d is not None]))
        else:
            stack.extend(reversed(list(ast.iter_child_nodes(node))))


def _dotted(node):
    parts = []
    while isinstance(node, ast.Attribute):
        parts.append(node.attr)
        node = node.value
    if isinstance(node, ast.Name):
        return ".".join([node.id] + parts[::-1])
    return None


class _Module:
    """One parsed module and its name resolution."""

    def __init__(self, name, rel, text, modules):
        self.name, self.rel, self.text = name, rel, text
        self.tree = ast.parse(text, filename=rel)
        self.aliases = _aliases(name, rel, self.tree, modules)
        self.units, self.bindings, self.skip = _units(self.tree)
        self._parent = None
        self.consts = {}
        self.literal_tables = {}
        self.local_aliases = {}                         # function -> local name -> the units it was bound to

    @property
    def parent(self):
        """child id -> parent node, built on first use (only a module naming the section literally needs it)."""
        if self._parent is None:
            self._parent = {id(c): n for n in ast.walk(self.tree) for c in ast.iter_child_nodes(n)}
        return self._parent

    def lookup(self, unit, name):
        """The units of this module a bare name in `unit` binds, by Python's scope rule: the unit's own nested defs,
        then each enclosing function (never an enclosing class body), then the module."""
        scope, start = (unit.name if unit.kind in _CALLABLE else None), True
        while scope is not None:
            here = self.units[scope]
            if start or here.kind != "class":
                hits = self.bindings.get(scope, {}).get(name)
                if hits:
                    return [(self.name, h) for h in hits]
                if name in self.local_aliases.get(scope, {}):
                    return list(self.local_aliases[scope][name])
            scope, start = here.scope, False
        return [(self.name, h) for h in self.bindings.get(None, {}).get(name, [])]


def _names_bound(target):
    """The plain names an assignment target binds (never the base of a subscript or attribute it stores into)."""
    if isinstance(target, ast.Name):
        return [target.id]
    if isinstance(target, (ast.Tuple, ast.List)):
        return [name for e in target.elts for name in _names_bound(e)]
    if isinstance(target, ast.Starred):
        return _names_bound(target.value)
    return []


class _Scan:
    """The whole-tree scan: units, their references and calls, the section reads, and the owner routes."""

    def __init__(self, root):
        self.modules, self._own = {}, {}
        names = _module_map(root)
        for module, rel in names.items():
            with open(os.path.join(root, rel), encoding="utf-8") as fh:
                self.modules[module] = _Module(module, rel, fh.read(), names)
        for _ in range(2):                              # module-level key constants, imported ones included
            for mod in self.modules.values():
                for st in _module_assigns(mod.tree.body):
                    targets = st.targets if isinstance(st, ast.Assign) else [st.target]
                    if len(targets) == 1 and isinstance(targets[0], ast.Name):
                        value = self.fold(mod, None, st.value, {})
                        if value is not None:
                            mod.consts[targets[0].id] = value
        for mod in self.modules.values():              # module-level literal tables a loop may iterate
            for st in _module_assigns(mod.tree.body):
                targets = st.targets if isinstance(st, ast.Assign) else [st.target]
                if len(targets) == 1 and isinstance(targets[0], ast.Name) and isinstance(st.value, _COLLECTION):
                    mod.literal_tables[targets[0].id] = st.value.elts
        for mod in self.modules.values():              # a local bound once to a project unit (P = _PURE)
            for unit in mod.units.values():
                if unit.kind not in ("function", "method"):
                    continue
                assigns = [n for n in self.own(mod, unit) if isinstance(n, ast.Assign)]
                names = [t.id for n in assigns for t in n.targets if isinstance(t, ast.Name)]
                for n in assigns:
                    if (len(n.targets) == 1 and isinstance(n.targets[0], ast.Name)
                            and names.count(n.targets[0].id) == 1 and isinstance(n.value, (ast.Name, ast.Attribute))):
                        hits = self.resolve(mod, unit, n.value)
                        if hits:
                            mod.local_aliases.setdefault(unit.name, {})[n.targets[0].id] = hits

    # --- resolution ------------------------------------------------------------------------------------------
    def own(self, mod, unit):
        """The nodes `unit` evaluates itself (:func:`_own`), walked once."""
        key = (mod.name, unit.name)
        if key not in self._own:
            self._own[key] = list(_own(unit.roots, mod.skip if unit.kind == "module" else frozenset()))
        return self._own[key]

    def fields(self, key):
        """Whether a unit reads a stored row field by a literal or foldable key (``r.get('stranded')``, ``r['hard']``)."""
        mod, unit = self.modules[key[0]], self.unit(key)
        return any(
            (isinstance(n, ast.Call) and isinstance(n.func, ast.Attribute) and n.func.attr == "get" and n.args
             and self.fold(mod, unit, n.args[0], {}) in _ROW_FIELDS)
            or (isinstance(n, ast.Subscript) and isinstance(n.ctx, ast.Load)
                and self.fold(mod, unit, n.slice, {}) in _ROW_FIELDS)
            for n in self.own(mod, unit))

    def unit(self, key):
        mod = self.modules.get(key[0])
        return mod.units.get(key[1]) if mod is not None else None

    def resolve(self, mod, unit, node):
        """The project units a Name / Attribute load names (empty when it names none)."""
        if isinstance(node, ast.Name):
            hits = mod.lookup(unit, node.id) if unit is not None else []
            if hits:
                return hits
            bound = mod.aliases.get(node.id)
            if bound and bound[0] == "name" and self.unit((bound[1], bound[2])) is not None:
                return [(bound[1], bound[2])]
            return []
        if not isinstance(node, ast.Attribute):
            return []
        base = node.value
        if isinstance(base, ast.Name) and base.id in ("self", "cls") and unit is not None and unit.cls:
            q = f"{unit.cls}.{node.attr}"
            return [(mod.name, q)] if q in mod.units else []
        dotted = _dotted(base)
        if dotted is None:
            return []
        head, _, rest = dotted.partition(".")
        bound = mod.aliases.get(head)
        if bound is not None:
            full = (bound[1] if bound[0] == "module" else f"{bound[1]}.{bound[2]}") + ("." + rest if rest else "")
            if full in self.modules and node.attr in self.modules[full].units:
                return [(full, node.attr)]
            return []
        if not rest and unit is not None:               # Class.method
            return [(mod.name, f"{q}.{node.attr}") for m, q in mod.lookup(unit, head)
                    if f"{q}.{node.attr}" in mod.units]
        return []

    def fold(self, mod, unit, node, local):
        """The string an expression folds to: a literal, a key constant (local, module-level or imported), ``+``
        concatenation, an f-string of foldable parts, or ``sep.join([...])`` of them. None otherwise."""
        if isinstance(node, ast.Constant):
            return node.value if isinstance(node.value, str) else None
        if isinstance(node, ast.Name):
            if node.id in local:
                return local[node.id]
            if node.id in mod.consts:
                return mod.consts[node.id]
            bound = mod.aliases.get(node.id)
            if bound and bound[0] == "name" and bound[1] in self.modules:
                return self.modules[bound[1]].consts.get(bound[2])
            return None
        if isinstance(node, ast.Attribute) and isinstance(node.value, ast.Name):
            bound = mod.aliases.get(node.value.id)
            if bound and bound[0] == "module" and bound[1] in self.modules:
                return self.modules[bound[1]].consts.get(node.attr)
            return None
        if isinstance(node, ast.BinOp) and isinstance(node.op, ast.Add):
            left, right = self.fold(mod, unit, node.left, local), self.fold(mod, unit, node.right, local)
            return left + right if left is not None and right is not None else None
        if isinstance(node, ast.JoinedStr):
            parts = []
            for value in node.values:
                if isinstance(value, ast.Constant) and isinstance(value.value, str):
                    parts.append(value.value)
                elif (isinstance(value, ast.FormattedValue) and value.conversion == -1
                      and value.format_spec is None):
                    part = self.fold(mod, unit, value.value, local)
                    if part is None:
                        return None
                    parts.append(part)
                else:
                    return None
            return "".join(parts)
        if (isinstance(node, ast.Call) and isinstance(node.func, ast.Attribute) and node.func.attr == "join"
                and len(node.args) == 1 and not node.keywords and isinstance(node.args[0], _COLLECTION)):
            sep = self.fold(mod, unit, node.func.value, local)
            parts = [self.fold(mod, unit, e, local) for e in node.args[0].elts]
            return sep.join(parts) if sep is not None and None not in parts else None
        return None

    # --- one unit ----------------------------------------------------------------------------------------------
    def _has_section(self, mod, unit, node, local):
        if isinstance(node, _COLLECTION):
            return any(self._has_section(mod, unit, e, local) for e in node.elts)
        return self.fold(mod, unit, node, local) == SECTION

    def analyse(self, mod, unit):
        """One unit's facts: ``occ`` (line, kind) for every occurrence of the section name it evaluates, ``reads``
        (the read lines), ``refs`` (target, called), ``sinks`` (callee, line) handed the stored rows, ``fields``
        (whether it reads a row field), ``tainted`` (its names holding the rows) and ``free`` (name, line) loads it
        does not bind."""
        nodes = self.own(mod, unit)
        once, local = {}, {}
        if unit.kind in ("function", "method"):
            for n in nodes:
                if isinstance(n, ast.Assign) and len(n.targets) == 1 and isinstance(n.targets[0], ast.Name):
                    once[n.targets[0].id] = once.get(n.targets[0].id, 0) + 1
            for n in nodes:
                if (isinstance(n, ast.Assign) and len(n.targets) == 1 and isinstance(n.targets[0], ast.Name)
                        and once[n.targets[0].id] == 1):
                    value = self.fold(mod, unit, n.value, {})
                    if value is not None:
                        local[n.targets[0].id] = value
        callees = {id(n.func) for n in nodes if isinstance(n, ast.Call)}
        # loop keys: a name iterating data (a mapping's items or keys) that is compared against the section name;
        # a name iterating a literal table that holds the section name may BE it
        loop_keys, may_keys, loop_values = set(), set(), set()
        for n in nodes:
            if isinstance(n, (ast.For, ast.AsyncFor)):
                pairs = [(n.target, n.iter)]
            elif isinstance(n, (ast.ListComp, ast.SetComp, ast.DictComp, ast.GeneratorExp)):
                pairs = [(g.target, g.iter) for g in n.generators]
            else:
                continue
            for target, it in pairs:
                table = (it.elts if isinstance(it, _COLLECTION)
                         else mod.literal_tables.get(it.id) if isinstance(it, ast.Name) else None)
                if table is not None:
                    if isinstance(target, ast.Name) and any(self.fold(mod, unit, e, local) == SECTION for e in table):
                        may_keys.add(target.id)
                    continue
                if isinstance(target, ast.Name):
                    loop_keys.add(target.id)
                elif isinstance(target, ast.Tuple) and target.elts and isinstance(target.elts[0], ast.Name):
                    loop_keys.add(target.elts[0].id)
                    loop_values.update(e.id for e in target.elts[1:] if isinstance(e, ast.Name))

        def key_expr(node):
            return self.fold(mod, unit, node, local) == SECTION or (isinstance(node, ast.Name) and node.id in may_keys)

        def loop_compare(cmp):
            sides = [cmp.left] + list(cmp.comparators)
            if not any(isinstance(op, (ast.Eq, ast.In, ast.Is)) for op in cmp.ops):
                return False
            return (any(isinstance(s, ast.Name) and s.id in loop_keys for s in sides)
                    and any(self._has_section(mod, unit, s, local) for s in sides))

        occ, reads, refs, sinks, read_nodes, binders = [], [], [], [], set(), set()

        def hit(line, kind, node=None):
            occ.append((line, kind))
            if kind == "read":
                reads.append(line)
                if node is not None:
                    read_nodes.add(id(node))

        for n in nodes:
            if isinstance(n, (ast.Name, ast.Attribute)) and isinstance(n.ctx, ast.Load):
                for target in self.resolve(mod, unit, n):
                    refs.append((target, id(n) in callees))
            if isinstance(n, ast.Constant) and n.value == SECTION:
                kind, where = self.classify(mod, n, loop_compare)
                hit(n.lineno, kind, where)
            elif isinstance(n, ast.Call):
                for arg in list(n.args) + [k.value for k in n.keywords]:
                    arg = arg.value if isinstance(arg, ast.Starred) else arg
                    for e in (arg.elts[:1] if isinstance(arg, _COLLECTION) else [arg]):
                        if not isinstance(e, ast.Constant) and key_expr(e):
                            hit(e.lineno, "read", n)
            elif isinstance(n, ast.Subscript) and not isinstance(n.slice, ast.Constant) and key_expr(n.slice):
                hit(n.lineno, "read" if isinstance(n.ctx, ast.Load) else "store", n)
            elif isinstance(n, ast.Attribute) and n.attr == SECTION:
                if not isinstance(n.ctx, ast.Load):
                    hit(n.end_lineno, "store")
                else:
                    hit(n.end_lineno, "reference" if self.resolve(mod, unit, n) else "read", n)
            elif isinstance(n, ast.Name) and n.id == SECTION:
                if not isinstance(n.ctx, ast.Load):
                    hit(n.lineno, "store")
                else:
                    hit(n.lineno, "reference" if self.resolve(mod, unit, n) else "read", n)
            elif isinstance(n, ast.arg) and n.arg == SECTION:
                hit(n.lineno, "parameter")
            elif isinstance(n, ast.keyword) and n.arg == SECTION:
                hit(n.lineno, "keyword-name")
            elif isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)) and n.name == SECTION:
                hit(n.lineno, "definition")
            elif isinstance(n, ast.alias) and SECTION in (n.name.rpartition(".")[2], n.asname):
                hit(n.lineno, "import")
            elif isinstance(n, (ast.Global, ast.Nonlocal)) and SECTION in n.names:
                hit(n.lineno, "store")
            elif isinstance(n, ast.MatchMapping):
                for key, pattern in zip(n.keys, n.patterns):
                    if key_expr(key):
                        if not isinstance(key, ast.Constant):
                            hit(key.lineno, "read", n)
                        binders.update(p.name for p in ast.walk(pattern) if isinstance(p, ast.MatchAs) and p.name)
            elif isinstance(n, ast.MatchClass) and SECTION in n.kwd_attrs:
                hit(n.lineno, "read", n)
                for attr, pattern in zip(n.kwd_attrs, n.kwd_patterns):
                    if attr == SECTION:
                        binders.update(p.name for p in ast.walk(pattern) if isinstance(p, ast.MatchAs) and p.name)
            elif isinstance(n, ast.Compare) and loop_compare(n):
                if not any(isinstance(s, ast.Constant) and s.value == SECTION
                           for s in [n.left] + list(n.comparators)):
                    hit(n.lineno, "read", n)
                binders.update(loop_values)
        if any(isinstance(n, ast.Compare) and loop_compare(n) for n in nodes):
            binders.update(loop_values)
        # the names that hold the stored rows: a mapping pattern's capture, the value of a loop whose key is compared
        # against the section name, or a name assigned from an expression that reads the section (one hop: a name
        # assigned from such a name is not followed)
        def reads_in(node):
            stack = [node]
            while stack:
                s = stack.pop()
                if id(s) in read_nodes:
                    return True
                if not isinstance(s, ast.Dict):             # a snapshot-shaped literal is read by its key
                    stack.extend(ast.iter_child_nodes(s))
            return False

        tainted = set(binders)
        for n in (nodes if read_nodes else ()):
            if isinstance(n, ast.Assign):
                pairs = [(t, n.value) for t in n.targets]
            elif isinstance(n, (ast.AnnAssign, ast.AugAssign, ast.NamedExpr)) and n.value is not None:
                pairs = [(n.target, n.value)]
            elif isinstance(n, (ast.For, ast.AsyncFor, ast.comprehension)):
                pairs = [(n.target, n.iter)]
            elif isinstance(n, ast.withitem) and n.optional_vars is not None:
                pairs = [(n.optional_vars, n.context_expr)]
            else:
                continue
            for target, value in pairs:
                if reads_in(value):
                    tainted.update(_names_bound(target))

        def carries(node):
            return reads_in(node) or any(isinstance(s, ast.Name) and isinstance(s.ctx, ast.Load) and s.id in tainted
                                         for s in ast.walk(node))

        # a call handed the rows hands them to its callee, and to every project function passed beside them (a phase
        # runner: _run_phase("sheet", write_sheet, wb, rows))
        for n in (nodes if read_nodes or tainted else ()):
            args = list(n.args) + [k.value for k in n.keywords] if isinstance(n, ast.Call) else []
            if any(carries(a) for a in args):
                for fn in [n.func] + [a for a in args if isinstance(a, (ast.Name, ast.Attribute))]:
                    for target in self.resolve(mod, unit, fn):
                        if target[0] != OWNER:
                            sinks.append((target, n.lineno))
        bound = {n.arg for n in nodes if isinstance(n, ast.arg)}
        bound |= {t.id for t in nodes if isinstance(t, ast.Name) and not isinstance(t.ctx, ast.Load)}
        if unit.kind in _CALLABLE:
            bound |= set(mod.bindings.get(unit.name, {}))
        free = [(n.id, n.lineno) for n in nodes if isinstance(n, ast.Name) and isinstance(n.ctx, ast.Load)
                and n.id not in bound]
        return {"occ": occ, "reads": reads, "refs": refs, "sinks": sinks, "tainted": tainted, "bound": bound,
                "free": free}

    def classify(self, mod, node, loop_compare):
        """Where a literal section name sits: ``(kind, read node)``. A read is a call argument (a path tuple of one
        included), a subscript key, a mapping-pattern key, or a comparison with a key iterating data. Everything
        else must be one of the closed non-read positions -- a literal table, a dict-literal key, a comparison, a
        key constant's definition, an f-string part -- or it is unclassified (kind None)."""
        child, q, in_table, first = node, mod.parent.get(id(node)), False, True
        while isinstance(q, _COLLECTION):
            first = first and q.elts[0] is child
            child, q, in_table = q, mod.parent.get(id(q)), True
        if isinstance(q, ast.Starred):
            child, q = q, mod.parent.get(id(q))
        if isinstance(q, ast.Compare):
            return ("read", q) if loop_compare(q) else ("compare", None)
        path = not in_table or first                # a path into the snapshot starts at the section name
        if path and isinstance(q, ast.Call) and any(a is child for a in q.args):
            return "read", q
        if path and isinstance(q, ast.keyword):
            return "read", mod.parent.get(id(q))
        if path and isinstance(q, ast.Subscript) and q.slice is child:
            return ("read" if isinstance(q.ctx, ast.Load) else "store"), q
        if isinstance(q, ast.MatchMapping) and any(k is child for k in q.keys):
            return "read", q
        if in_table:
            return "table", None
        if isinstance(q, ast.Dict) and any(k is child for k in q.keys):
            return "dict-key", None
        if isinstance(q, (ast.Assign, ast.AnnAssign)) and q.value is child:
            return "const", None
        if isinstance(q, ast.MatchValue):
            return "compare", None
        if isinstance(q, (ast.JoinedStr, ast.FormattedValue)):
            return "fstring", None
        return None, None


@functools.lru_cache(maxsize=None)
def _scan(root):
    """``(readers, routed, occurrences, why)``. `readers`: unit -> the lines where it reads the stored section (a
    unit is ``(module, qualified name)``); `routed`: every unit with a route to the owner; `occurrences`: module ->
    ``(line, kind)`` for every occurrence of the section name the AST sees; `why`: unit -> how it reads. Cached per
    root: the callers only read the result."""
    scan = _Scan(root)
    facts = {}
    for mod in scan.modules.values():
        for unit in mod.units.values():
            facts[(mod.name, unit.name)] = scan.analyse(mod, unit)
    readers, why, occurrences = {}, {}, {}

    def read(key, line, how):
        readers.setdefault(key, set()).add(line)
        why.setdefault(key, set()).add(how)

    for key, f in facts.items():
        occurrences.setdefault(key[0], []).extend(f["occ"])
        for line in f["reads"]:
            read(key, line, "reads the section")
    # a nested unit loading a name its enclosing function bound to the stored rows reads them too
    for (module, name), f in facts.items():
        mod = scan.modules[module]
        for var, line in f["free"]:
            scope = mod.units[name].scope
            while scope is not None:
                outer = facts[(module, scope)]
                if mod.units[scope].kind in ("function", "method"):
                    if var in outer["tainted"]:
                        read((module, name), line, f"loads {var!r}, which {scope} bound to the rows")
                        break
                    if var in outer["bound"]:
                        break
                scope = mod.units[scope].scope
    # a helper handed the stored rows that reads a row field presents them, whatever its parameter is called
    for key, f in facts.items():
        for target, line in f["sinks"]:
            if target in facts and scan.fields(target):
                read(target, scan.unit(target).node.lineno, f"is handed the rows by {key[0]}:{key[1]} (line {line})")
    # routes: a CALL to an owner function or class (a module-level table may hold one uncalled); a reference to any
    # other project unit is an edge, so dispatch tables and callbacks route through it
    direct, edges = set(), {}
    for key, f in facts.items():
        unit = scan.unit(key)
        out = edges.setdefault(key, set())
        if key[0] == OWNER and unit.kind in _CALLABLE:
            direct.add(key)
        for target, called in f["refs"]:
            tu = scan.unit(target)
            if tu is None or target == key:
                continue
            if target[0] == OWNER:
                if tu.kind in _CALLABLE and (called or unit.kind == "data"):
                    direct.add(key)
                continue
            out.add(target)
        if unit.kind == "class":
            members = scan.modules[key[0]].units
            out.update((key[0], f"{key[1]}.{m}") for m in _CLASS_ENTRY if f"{key[1]}.{m}" in members)
    routed = set(direct)
    changed = True
    while changed:
        changed = False
        for key, targets in edges.items():
            if key not in routed and targets & routed:
                routed.add(key)
                changed = True
    return ({k: sorted(v) for k, v in readers.items()}, routed, occurrences, {k: sorted(v) for k, v in why.items()})


# --- the receipt-recompute closure resolver (W50; read only by tests/test_operator_evidence_contract.py) ------------
# The W48 guard above (:func:`_scan`) admits a reader only on a CALL route, so an extra route there would admit a raw
# reader. The W50 closure check needs the opposite bias: a missed route would hide an owner dependency of a receipt
# recomputation. So it keeps its own conservative resolver -- the first W48 guard's (``e7c00e12``), with W50's
# constructor routes and owner references -- over its own top-level unit split (:func:`_graph_units`), and shares only
# the module map, import resolution and alias table with the guard.
class _Graph(NamedTuple):
    """The resolved unit graph of one tree (see :func:`_graph`)."""
    units: dict          # (module, name) -> (node, kind)
    edges: dict          # unit -> the units it names, resolved (calls, references, dispatch tables, self.method)
    direct: frozenset    # units that are the owner, or name it through a resolved import alias
    constructs: dict     # unit -> every method of each project class it names (a constructor runs its methods)
    owner_refs: frozenset  # units naming the owner by an unaliased dotted path or importing it by any spelling


def _graph_units(tree):
    """(unit name, node, kind) for each top-level function, each method of a top-level class, and each module-level
    assignment to a plain name (a dispatch table routes through it). A nested function belongs to its enclosing
    unit, so every edge it has is that unit's (conservative for a closure)."""
    for node in tree.body:
        if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)):
            yield node.name, node, "function"
        elif isinstance(node, ast.ClassDef):
            for sub in node.body:
                if isinstance(sub, (ast.FunctionDef, ast.AsyncFunctionDef)):
                    yield f"{node.name}.{sub.name}", sub, "method"
        elif isinstance(node, (ast.Assign, ast.AnnAssign)) and node.value is not None:
            targets = node.targets if isinstance(node, ast.Assign) else [node.target]
            for target in targets:
                if isinstance(target, ast.Name):
                    yield target.id, node.value, "data"


@functools.lru_cache(maxsize=None)
def _graph(root):
    """The resolved unit graph of the scanned tree under `root`, built once per root (callers only read it).

    ``edges`` and ``direct`` (a unit that is the owner or names it through a resolved import alias) plus the extra,
    conservative ``constructs`` and ``owner_refs`` routes are what the W50 receipt-closure check reads
    (``tests/test_operator_evidence_contract.py``); the W48 guard does not read this graph (see the note above)."""
    modules = _module_map(root)
    units, edges, direct = {}, {}, set()
    constructs, owner_refs = {}, set()
    classes = {}
    parsed = {}
    for module, rel in modules.items():
        with open(os.path.join(root, rel), encoding="utf-8") as fh:
            tree = ast.parse(fh.read(), filename=rel)
        parsed[module] = (rel, tree, _aliases(module, rel, tree, modules))
        for name, node, kind in _graph_units(tree):
            units[(module, name)] = (node, kind)
            if kind == "method":
                classes.setdefault((module, name.split(".")[0]), set()).add((module, name))
    for module, (rel, tree, aliases) in parsed.items():
        own = {name for (m, name) in units if m == module}
        for (m, name), (node, _kind) in units.items():
            if m != module:
                continue
            unit = (module, name)
            if module == OWNER:
                direct.add(unit)
            cls = name.split(".")[0] if "." in name else None
            out, built = set(), set()
            for sub in ast.walk(node):
                if isinstance(sub, ast.Name) and isinstance(sub.ctx, ast.Load):
                    if sub.id in own:
                        out.add((module, sub.id))
                    built |= classes.get((module, sub.id), set())
                    bound = aliases.get(sub.id)
                    if bound and bound[0] == "name":
                        out.add((bound[1], bound[2]))
                        built |= classes.get((bound[1], bound[2]), set())
                        if bound[1] == OWNER:
                            direct.add(unit)
                    elif bound and bound[1] == OWNER:
                        direct.add(unit)
                elif isinstance(sub, ast.Attribute) and isinstance(sub.value, ast.Name):
                    bound = aliases.get(sub.value.id)
                    if bound and bound[0] == "module":
                        out.add((bound[1], sub.attr))
                        built |= classes.get((bound[1], sub.attr), set())
                        if bound[1] == OWNER:
                            direct.add(unit)
                    elif sub.value.id in ("self", "cls") and cls:
                        out.add((module, f"{cls}.{sub.attr}"))
                elif isinstance(sub, ast.ImportFrom):
                    source = _resolve_from(module, rel, sub)
                    if source == OWNER or any(f"{source}.{a.name}" == OWNER for a in sub.names):
                        direct.add(unit)
                if isinstance(sub, ast.Attribute) and sub.attr == OWNER.rpartition(".")[2]:
                    owner_refs.add(unit)
                elif isinstance(sub, ast.Import) and any(a.name == OWNER for a in sub.names):
                    owner_refs.add(unit)
            edges[unit] = {e for e in out if e in units and e != unit}
            constructs[unit] = {e for e in built if e in units and e != unit}
    return _Graph(units, edges, frozenset(direct), constructs, frozenset(owner_refs))


def _closure(graph, roots):
    """Every unit reachable from `roots` along resolved edges and constructor routes: unit -> the unit it was reached
    from (``None`` for a root), breadth-first so each recorded path is a shortest one."""
    parent = {root: None for root in roots}
    queue = list(roots)
    while queue:
        unit = queue.pop(0)
        for target in sorted(graph.edges.get(unit, set()) | graph.constructs.get(unit, set())):
            if target not in parent:
                parent[target] = unit
                queue.append(target)
    return parent


def _path(parent, unit):
    """The recorded route to `unit`, root first, as ``module:name`` strings."""
    path = []
    while unit is not None:
        path.append(f"{unit[0]}:{unit[1]}")
        unit = parent[unit]
    return " -> ".join(reversed(path))


def _raw(root):
    readers, routed, _occurrences, _why = _scan(root)
    return {unit: lines for unit, lines in readers.items() if unit not in routed}


def _textual(text):
    """Every line where the tokenizer sees the section name: a NAME token, or a string literal whose value is exactly
    the section name. Comments, docstrings and every other string are not occurrences."""
    lines = set()
    for tok in tokenize.generate_tokens(io.StringIO(text).readline):
        if tok.type == tokenize.NAME and tok.string == SECTION:
            lines.add(tok.start[0])
        elif tok.type == tokenize.STRING:
            try:
                value = ast.literal_eval(tok.string)
            except (ValueError, SyntaxError):
                continue
            if value == SECTION:
                lines.add(tok.start[0])
    return lines


def _line_check(root):
    """``(rel, line, why)`` for every occurrence of the section name the line check cannot account for: an AST
    occurrence in no closed position (unclassified), or a line where the tokenizer sees the section name and the AST
    scan records no occurrence at all."""
    _readers, _routed, occurrences, _why = _scan(root)
    out = []
    for module, rel in _module_map(root).items():
        with open(os.path.join(root, rel), encoding="utf-8") as fh:
            text = fh.read()
        seen = occurrences.get(module, [])
        out += [(rel, line, "unclassified position") for line, kind in seen if kind is None]
        out += [(rel, line, "no AST occurrence") for line in sorted(_textual(text) - {ln for ln, _k in seen})]
    return sorted(set(out))


def test_every_reader_of_the_stored_failure_impact_rows_reaches_the_owner():
    readers, routed, _occ, why = _scan(str(ROOT))
    raw = {unit: lines for unit, lines in readers.items() if unit not in routed}
    new = {f"{m}:{n}": (lines, why[(m, n)]) for (m, n), lines in sorted(raw.items()) if (m, n) not in _RAW_RATCHET}
    assert not new, (
        "these units read the stored failure_impact rows with no CALL route to the engine owner "
        "(cisco_toolkit.impact_assessability). Read the rows through it (rows_with_verdicts / table_value / "
        "ranked_value / ranks / ranking_floor / disclose / wave_blast, or the projection) so a held row reads as not "
        f"assessed and a lower bound as a lower bound, never as an exact measurement: {new}")
    fixed = sorted(f"{m}:{n}" for (m, n) in _RAW_RATCHET if (m, n) not in raw)
    assert not fixed, ("these ratchet entries now reach the owner or no longer read the section: delete them from "
                       f"_RAW_RATCHET and remove their strict xfail below: {fixed}")


def test_every_line_naming_the_section_is_a_detected_read_or_a_closed_non_read():
    """Line-level non-vacuity on the real tree: every line where the tokenizer sees the section name (a NAME token, or a
    string literal that is exactly it; comments and docstrings are not tokens of either kind) holds an AST occurrence,
    and every occurrence is a read attributed to a unit (judged above) or one of the closed non-read positions. A
    scanner that stopped seeing a read shape would leave its line unaccounted for and fail here."""
    assert not _line_check(str(ROOT)), _line_check(str(ROOT))
    readers, _routed, occurrences, _why = _scan(str(ROOT))
    # the scan spans all three roots: the pipeline, the engine and AssessHub each read the section
    assert any(m == _PIPELINE[:-3] for m, _n in readers)
    assert any(m.startswith("cisco_toolkit.") for m, _n in readers)
    assert any(m.startswith("webapp.backend.") for m, _n in readers)
    # and it is not vacuous: every read line is attributed to a reader unit of its module
    for module, seen in occurrences.items():
        attributed = {line for (m, _n), lines in readers.items() if m == module for line in lines}
        assert {line for line, kind in seen if kind == "read"} <= attributed, module


def test_the_line_check_reports_a_position_it_cannot_classify(tmp_path):
    """The closed position list is not decoration: a section name in a position that is neither a read nor a listed
    non-read (a parameter default, a dict-literal value) is reported, and a comment or docstring naming it is not."""
    _write(tmp_path, "cisco_toolkit/__init__.py", "")
    _write(tmp_path, "cisco_toolkit/impact_assessability.py", "def rows_with_verdicts(snap):\n    return []\n")
    _write(tmp_path, "cisco_toolkit/odd.py",
           '"""Reads failure_impact; snap.get("failure_impact") in a docstring is not a read."""\n'
           "LABELS = {'impact': 'failure_impact'}\n"
           "def render(snap, section='failure_impact'):   # snap['failure_impact'] in a comment\n"
           "    return snap.get(section)\n")
    assert _line_check(str(tmp_path)) == [("cisco_toolkit/odd.py", 2, "unclassified position"),
                                          ("cisco_toolkit/odd.py", 3, "unclassified position")]


def _write(root, rel, text):
    path = root / rel
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")


_OWNER_STUB = ("NOT_ASSESSED_CELL = 'not assessed'\n"
               "def rows_with_verdicts(snap):\n    return [(r, None) for r in snap.get('failure_impact') or []]\n")


def test_the_guard_flags_a_raw_reader_and_admits_each_route_to_the_owner(tmp_path):
    """The guard is not decoration: on a synthetic tree it flags each raw read shape and admits each route."""
    _write(tmp_path, "cisco_toolkit/__init__.py", "")
    _write(tmp_path, "cisco_toolkit/impact_assessability.py", _OWNER_STUB)
    _write(tmp_path, "cisco_toolkit/raw_get.py",
           "def render(snap):\n    return [r.get('stranded') for r in snap.get('failure_impact') or []]\n")
    _write(tmp_path, "cisco_toolkit/raw_key.py",
           "def render(snap):\n    return snap['failure_impact'][0]['stranded']\n")
    _write(tmp_path, "cisco_toolkit/raw_param.py",
           "def render(failure_impact):\n    return max(r['stranded'] for r in failure_impact)\n")
    _write(tmp_path, "cisco_toolkit/raw_attr.py",
           "class Writer:\n    def render(self, ctx):\n        return ctx.failure_impact\n")
    _write(tmp_path, "cisco_toolkit/routed_direct.py",
           "from cisco_toolkit import impact_assessability\n"
           "def render(snap):\n    rows = snap.get('failure_impact')\n"
           "    return impact_assessability.rows_with_verdicts(snap), rows\n")
    _write(tmp_path, "cisco_toolkit/routed_lazy.py",
           "def render(snap):\n    from cisco_toolkit.impact_assessability import rows_with_verdicts\n"
           "    return rows_with_verdicts(snap), snap['failure_impact']\n")
    _write(tmp_path, "cisco_toolkit/routed_dotted.py",
           "import cisco_toolkit.impact_assessability\n"
           "def render(snap):\n    return cisco_toolkit.impact_assessability.rows_with_verdicts(snap), "
           "snap['failure_impact']\n")
    _write(tmp_path, "cisco_toolkit/routed_transitive.py",
           "from . import routed_direct\n"
           "def _helper(snap):\n    return routed_direct.render(snap)\n"
           "def render(snap):\n    rows = snap.get('failure_impact')\n    return _helper(snap), rows\n")
    _write(tmp_path, "cisco_toolkit/routed_method.py",
           "from cisco_toolkit import impact_assessability as ia\n"
           "class Writer:\n    def _owner(self, snap):\n        return ia.rows_with_verdicts(snap)\n"
           "    def render(self, snap):\n        return self._owner(snap), snap['failure_impact']\n")
    _write(tmp_path, "cisco_toolkit/routed_table.py",
           "from cisco_toolkit import impact_assessability\n"
           "def failure_impact(snap):\n    return impact_assessability.rows_with_verdicts(snap)\n"
           "_PURE = {'failure_impact': failure_impact}\n"
           "def serve(snap):\n    return _PURE['failure_impact'](snap)\n"
           "def build(snap):\n    P = _PURE\n    def tool():\n        return P['failure_impact'](snap)\n    return tool\n")
    _write(tmp_path, "cisco_toolkit/routed_owner_table.py",
           "from cisco_toolkit import impact_assessability as ia\n"
           "_T = {'rows': ia.rows_with_verdicts}\n"
           "def serve(snap):\n    return _T['rows'](snap), snap.get('failure_impact')\n")
    _write(tmp_path, "cisco_toolkit/routed_nested.py",
           "from cisco_toolkit import impact_assessability as ia\n"
           "def create_app(snap):\n    def route():\n        return ia.rows_with_verdicts(snap), snap['failure_impact']\n"
           "    return route\n")
    _write(tmp_path, "cisco_toolkit/not_a_read.py",
           "SECTIONS = ('devices', 'failure_impact')\n"
           "def name_only(name):\n    return name == 'failure_impact'\n"
           "def literal():\n    return {'failure_impact': []}\n"
           "def store(snap, rows):\n    snap['failure_impact'] = rows\n"
           "def schema(build):\n    return build('Page', ('host', 'failure_impact'))\n"
           "def skip(snap):\n    return {k: v for k, v in snap.items() if k != 'failure_impact'}\n")
    _write(tmp_path, "webapp/backend/__init__.py", "")
    _write(tmp_path, "webapp/backend/engine.py",
           "from cisco_toolkit import routed_direct as _rd\n"
           "def projection(snap):\n    return _rd.render(snap)\n")
    _write(tmp_path, "webapp/backend/view.py",
           "from . import engine\n"
           "def table(snap):\n    return engine.projection(snap), snap.get('failure_impact')\n"
           "def raw_table(snap):\n    return snap.get('failure_impact')\n")
    _write(tmp_path, _PIPELINE,
           "from cisco_toolkit import raw_param\n"
           "def main(failure_impact):\n    return raw_param.render(failure_impact)\n")
    readers, routed, _occ, _why = _scan(str(tmp_path))
    raw = {f"{m}:{n}" for (m, n) in readers if (m, n) not in routed}
    assert raw == {"cisco_toolkit.raw_get:render", "cisco_toolkit.raw_key:render", "cisco_toolkit.raw_param:render",
                   "cisco_toolkit.raw_attr:Writer.render", "webapp.backend.view:raw_table",
                   "COLLECT_PARSE_V3_23_0:main"}, raw
    admitted = {f"{m}:{n}" for (m, n) in readers if (m, n) in routed}
    assert {"cisco_toolkit.routed_direct:render", "cisco_toolkit.routed_lazy:render",
            "cisco_toolkit.routed_dotted:render", "cisco_toolkit.routed_transitive:render",
            "cisco_toolkit.routed_method:Writer.render", "cisco_toolkit.routed_table:serve",
            "cisco_toolkit.routed_table:build.tool", "cisco_toolkit.routed_owner_table:serve",
            "cisco_toolkit.routed_nested:create_app.route", "webapp.backend.view:table",
            "cisco_toolkit.impact_assessability:rows_with_verdicts"} <= admitted, admitted
    assert not any(m == "cisco_toolkit.not_a_read" for m, _n in readers), readers
    assert not _line_check(str(tmp_path)), _line_check(str(tmp_path))


#: The independent review's evasions of the first guard (W48 refutation, P2-3), plus the read shapes and routes the
#: redesign closes. Each module must be flagged as a raw reader, under exactly the unit named. A case that the guard
#: cannot close statically belongs in _KNOWN_LIMITS below, never here.
_EVASIONS = {
    # a route nested in a function that reaches the owner: the nested function is its own unit
    "nested_route": ("from cisco_toolkit import impact_assessability as ia\n"
                     "def create_app(snap):\n    ia.rows_with_verdicts(snap)\n"
                     "    def route():\n        return snap.get('failure_impact')[:5]\n    return route\n",
                     "create_app.route"),
    # the key held in a module-level constant, an imported one, a local one, a concatenation, an f-string, a join
    "key_constant": ("KEY = 'failure_impact'\ndef render(snap):\n    return snap.get(KEY)\n", "render"),
    "key_imported": ("from cisco_toolkit.keys import SECTION_KEY\ndef render(snap):\n    return snap[SECTION_KEY]\n",
                     "render"),
    "key_module_attr": ("from cisco_toolkit import keys\ndef render(snap):\n    return snap.get(keys.SECTION_KEY)\n",
                        "render"),
    "key_local": ("def render(snap):\n    key = 'failure_' + 'impact'\n    return snap[key]\n", "render"),
    "key_concat": ("def render(snap):\n    return snap.get('failure_' + 'impact')\n", "render"),
    "key_fstring": ("PART = 'impact'\ndef render(snap):\n    return snap.get(f'failure_{PART}')\n", "render"),
    "key_join": ("def render(snap):\n    return snap.get('_'.join(['failure', 'impact']))\n", "render"),
    "key_table_loop": ("def render(snap):\n    for key in ('devices', 'failure_impact'):\n"
                       "        yield snap.get(key)\n", "render"),
    # defs the first guard folded away or never saw
    "def_in_try": ("try:\n    def render(snap):\n        return snap.get('failure_impact')\nexcept Exception:\n"
                   "    pass\n", "render"),
    "def_in_if": ("import sys\nif sys:\n    def render(snap):\n        return snap['failure_impact']\n", "render"),
    "nested_class_method": ("class A:\n    class B:\n        def render(self, snap):\n"
                            "            return snap['failure_impact']\n", "A.B.render"),
    "decorated": ("import functools\n@functools.lru_cache\ndef render(snap):\n    return snap.get('failure_impact')\n",
                  "render"),
    "lambda_table": ("RENDER = lambda snap: snap.get('failure_impact')\n", "RENDER"),
    "module_level": ("import json\nROWS = json.loads('{}').get('failure_impact')\n", "ROWS"),
    # an owner CONSTANT is not a route; neither is an owner function named but never called
    "owner_constant": ("from cisco_toolkit import impact_assessability as ia\n"
                       "def render(snap):\n"
                       "    return [r['stranded'] for r in snap.get('failure_impact')] or ia.NOT_ASSESSED_CELL\n",
                       "render"),
    "owner_uncalled": ("from cisco_toolkit import impact_assessability as ia\n"
                       "def render(snap):\n    keep = ia.rows_with_verdicts\n    return snap.get('failure_impact'), keep\n",
                       "render"),
    # read shapes: a mapping pattern, a class pattern, a .items() loop (== and in), getattr
    "match_mapping": ("def render(snap):\n    match snap:\n        case {'failure_impact': rows}:\n"
                      "            return rows\n", "render"),
    "match_class": ("def render(ctx):\n    match ctx:\n        case object(failure_impact=rows):\n"
                    "            return rows\n", "render"),
    "items_loop": ("def render(snap):\n    for k, v in snap.items():\n        if k == 'failure_impact':\n"
                   "            return v\n", "render"),
    "items_loop_in": ("def render(snap):\n    return [v for k, v in snap.items() if k in ('failure_impact',)]\n",
                      "render"),
    "getattr": ("def render(ctx):\n    return getattr(ctx, 'failure_impact')\n", "render"),
    # the rows handed to a helper under another parameter name, directly, through a local, or beside a phase runner;
    # a nested function reading the rows its enclosing function bound
    "param_direct": ("from cisco_toolkit import impact_assessability as ia\n"
                     "def render(fi_rows):\n    return max(r['stranded'] for r in fi_rows)\n"
                     "def caller(snap):\n    ia.rows_with_verdicts(snap)\n    return render(snap.get('failure_impact'))\n",
                     "render"),
    "param_local": ("from cisco_toolkit import impact_assessability as ia\n"
                    "def render(rows):\n    return [r.get('severity') for r in rows]\n"
                    "def caller(snap):\n    ia.rows_with_verdicts(snap)\n    rows = snap['failure_impact']\n"
                    "    return render(rows)\n", "render"),
    "param_runner": ("from cisco_toolkit import impact_assessability as ia\n"
                     "def write_sheet(wb, rows):\n    return [r['hard'] for r in rows]\n"
                     "def _run(name, fn, *args):\n    return fn(*args)\n"
                     "def main(snap):\n    ia.rows_with_verdicts(snap)\n"
                     "    return _run('sheet', write_sheet, None, snap.get('failure_impact'))\n", "write_sheet"),
    "closure": ("from cisco_toolkit import impact_assessability as ia\n"
                "def create_app(snap):\n    ia.rows_with_verdicts(snap)\n    rows = snap.get('failure_impact')\n"
                "    def route():\n        return [r['stranded'] for r in rows]\n    return route\n", "create_app.route"),
}


@pytest.mark.parametrize("case", sorted(_EVASIONS))
def test_each_known_evasion_of_the_guard_is_flagged(tmp_path, case):
    text, unit = _EVASIONS[case]
    _write(tmp_path, "cisco_toolkit/__init__.py", "")
    _write(tmp_path, "cisco_toolkit/impact_assessability.py", _OWNER_STUB)
    _write(tmp_path, "cisco_toolkit/keys.py", "SECTION_KEY = 'failure' + '_impact'\n")
    _write(tmp_path, f"cisco_toolkit/{case}.py", text)
    raw = _raw(str(tmp_path))
    assert set(raw) == {(f"cisco_toolkit.{case}", unit)}, raw
    assert not _line_check(str(tmp_path)), _line_check(str(tmp_path))


#: What the guard does NOT see, pinned so the docstring's known limits stay true: each module reads the stored rows
#: raw and passes (the key-folding holes pass the line check too). When the guard learns to see one, its test fails:
#: move the case to _EVASIONS and update the docstring.
_KNOWN_LIMITS = {
    # dataflow is followed one assignment from a read: a copy of a copy reaches the helper unseen
    "two_hop": ("from cisco_toolkit import impact_assessability as ia\n"
                "def render(fi):\n    return [r['stranded'] for r in fi]\n"
                "def caller(snap):\n    ia.rows_with_verdicts(snap)\n    rows = snap.get('failure_impact')\n"
                "    again = list(rows)\n    return render(again)\n"),
    # a key computed with % or str.format, or read from data, is not folded
    "key_formatted": ("def render(snap):\n    return snap.get('%s_%s' % ('failure', 'impact'))\n"),
    # a call through an object or a runtime registry is not resolved: a unit that names a routed helper is admitted
    "reference_only": ("from cisco_toolkit import impact_assessability as ia\n"
                       "def _helper(snap):\n    return ia.rows_with_verdicts(snap)\n"
                       "def render(snap):\n    keep = _helper\n    return snap.get('failure_impact'), keep\n"),
    # a unit that reaches the owner for one value can still print another raw value (granularity: the unit)
    "same_unit": ("from cisco_toolkit import impact_assessability as ia\n"
                  "def render(snap):\n    ia.rows_with_verdicts(snap)\n"
                  "    return [r['stranded'] for r in snap.get('failure_impact')]\n"),
    # W51 (the W48 re-verification's P3): the key folding's holes. Folding reads a literal, a module-level, imported
    # or once-assigned local key constant through a plain name, ``+``, an f-string of unconverted parts,
    # ``sep.join([...])`` and a loop over a module-level literal table; every other way of computing the key is
    # unfolded, so the read is not seen
    "key_tuple_unpack": ("def render(snap):\n    key, other = 'failure_impact', 'devices'\n"
                         "    return snap.get(key), other\n"),
    "key_reassigned": ("def render(snap):\n    key = 'devices'\n    key = 'failure_impact'\n    return snap.get(key)\n"),
    "key_local_table": ("def render(snap):\n    keys = ('devices', 'failure_impact')\n    for key in keys:\n"
                        "        yield snap.get(key)\n"),
    "key_method": ("KEY = 'failure_impact'\ndef render(snap):\n    return snap.get(KEY.strip())\n"),
    "key_slice": ("KEY = 'failure_impact'\ndef render(snap):\n    return snap.get(KEY[:])\n"),
    "key_boolop": ("KEY = 'failure_impact'\ndef render(snap):\n    return snap.get(KEY or '')\n"),
    "key_ifexp": ("KEY = 'failure_impact'\ndef render(snap):\n    return snap.get(KEY if snap else '')\n"),
    "key_globals": ("KEY = 'failure_impact'\ndef render(snap):\n    return snap.get(globals()['KEY'])\n"),
    "key_fstring_conversion": ("KEY = 'failure_impact'\ndef render(snap):\n    return snap.get(f'{KEY!s}')\n"),
    "key_star_import": ("from cisco_toolkit.keys import *\ndef render(snap):\n    return snap.get(SECTION_KEY)\n"),
    "key_importlib": ("import importlib\ndef render(snap):\n"
                      "    return snap.get(importlib.import_module('cisco_toolkit.keys').SECTION_KEY)\n"),
    "key_bytes_decode": ("def render(snap):\n    return snap.get(b'failure_impact'.decode())\n"),
}


@pytest.mark.parametrize("case", sorted(_KNOWN_LIMITS))
def test_each_documented_known_limit_is_still_a_limit(tmp_path, case):
    _write(tmp_path, "cisco_toolkit/__init__.py", "")
    _write(tmp_path, "cisco_toolkit/impact_assessability.py", _OWNER_STUB)
    _write(tmp_path, "cisco_toolkit/keys.py", "SECTION_KEY = 'failure' + '_impact'\n")     # as the evasions have it
    _write(tmp_path, f"cisco_toolkit/{case}.py", _KNOWN_LIMITS[case])
    assert not _raw(str(tmp_path)), _raw(str(tmp_path))


def test_every_ratchet_entry_is_still_a_raw_reader_with_a_stated_reason():
    raw = _raw(str(ROOT))
    for unit, reason in _RAW_RATCHET.items():
        assert unit in raw, unit
        assert len(reason) > 80, unit


# ---------------------------------------------------------------------------------------------------------------
# part 2: behaviour on the committed sample
# ---------------------------------------------------------------------------------------------------------------
_RISK_FIELDS = ("severity", "vlans_impacted", "stranded", "hard", "backup", "fhrp")   # runbook §10 column order
_KEYSTONE_FIELDS = ("severity", "stranded", "vlans_impacted")                         # ops §2.1 column order


@pytest.fixture(scope="module")
def sample():
    if not SAMPLE.exists():
        pytest.skip("sample_fleet.snapshot.json not present")
    return json.loads(SAMPLE.read_text(encoding="utf-8"))


def _core1(snap):
    """core1's stored row and the owner's verdict on it."""
    pairs = [(row, verdict) for row, verdict in ia.rows_with_verdicts(snap) if row.get("host") == "core1"]
    assert len(pairs) == 1, pairs
    return pairs[0]


@pytest.fixture(scope="module")
def bounded(sample):
    """The committed sample. Precondition, read from the owner: core1 is a lower bound with a positive floor."""
    row, verdict = _core1(sample)
    assert verdict.assessable == ia.LOWER_BOUND, verdict.as_dict()
    assert ia.ranking_floor(verdict) == row["stranded"] > 0, verdict.as_dict()
    return sample


@pytest.fixture(scope="module")
def held(sample):
    """The sample with core1's off_scan_gw_vlans marker removed: a row older than the marker, which the owner holds.
    Precondition, read from the owner: core1 is not assessed."""
    snap = copy.deepcopy(sample)
    stored = next(r for r in snap["failure_impact"] if isinstance(r, dict) and r.get("host") == "core1")
    del stored["off_scan_gw_vlans"]
    _row, verdict = _core1(snap)
    assert verdict.assessable == ia.NOT_ASSESSED and "legacy_row" in verdict.codes, verdict.as_dict()
    return snap


def _docx_lines(path):
    """Every paragraph, then every table row with its cells joined by ' | '."""
    from docx import Document
    doc = Document(path)
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(c.text.replace("\n", " ") for c in row.cells))
    return lines


def _raw_row(row, fields):
    """The row as the pre-W48 writers printed it: the stored values verbatim."""
    return " | ".join([str(row["host"])] + [str(row[f]) for f in fields])


# --- MOP: the per-wave max blast radius and its quantified rollback trigger ---------------------------------------
def _mop_lines(tmp_path, snap, name):
    pytest.importorskip("docx")
    from cisco_toolkit.mop import write_mop_docx
    out = str(tmp_path / name)
    write_mop_docx(out, snap, "W48 Fleet")
    return _docx_lines(out)


def _mop_scope(lines, device):
    """The index of the §x.1 'Devices in scope' row of the wave whose scope names `device`."""
    return next(i for i, ln in enumerate(lines) if ln.startswith("Devices in scope | ")
                and device in ln.split(" | ", 1)[1].split(", "))


def _mop_wave(lines, device):
    """The §x.1 max-blast-radius cell and the blast-radius rollback trigger row of the wave naming `device`."""
    start = _mop_scope(lines, device)
    blast = next(ln for ln in lines[start:] if ln.startswith("Max blast radius"))
    trigger = next(ln for ln in lines[start:] if ln.startswith("Blast-radius / outage overrun"))
    return blast.split(" | ", 1)[1], trigger


def _mop_overview(lines):
    """§1's overview table: wave name -> its 'Max blast' column."""
    return {cells[0]: cells[5] for cells in (ln.split(" | ") for ln in lines
                                             if ln.startswith(("Wave ", "Group ")) and ln.count(" | ") == 6)}


def _mop_core1_wave(tmp_path, snap, name):
    lines = _mop_lines(tmp_path, snap, name)
    cell, trigger = _mop_wave(lines, "core1")
    return lines, cell, trigger


def _with_waves(snap, *waves):
    """`snap` with its candidate wave plan replaced by `waves` (the MOP writes one section per candidate wave)."""
    snap["design_blueprint"]["target_state"]["wave_plan"]["waves"] = [
        {"wave": i, "kind": "pilot", "switches": list(switches)} for i, switches in enumerate(waves, 1)]
    return snap


def _published(snap):
    """`snap` with every stored row's evidence a measurement: no inter-switch link without trunk/STP evidence, and every
    uncollected cable-map peer that could carry endpoints shown as collected. Precondition, read from the owner: every
    row is published."""
    for row in snap["failure_impact"]:
        row["blind_links"] = 0
    for node in snap["cable_map"]["nodes"]:
        if node.get("collected") is False and node.get("kind") not in ia.IMPACT_EDGE_KINDS:
            node["collected"] = True
    assert {v.assessable for v in ia.assess_failure_impact(snap)} == {ia.PUBLISHED}
    return snap


def test_mop_sizes_the_wave_holding_core1_as_a_lower_bound(tmp_path, bounded):
    row, verdict = _core1(bounded)
    lines, cell, trigger = _mop_core1_wave(tmp_path, bounded, "mop_bounded.docx")
    match = re.match(r"^≥ (\d+) " + re.escape(ia.LOWER_BOUND_MARK) + " — ", cell)
    assert match, cell                                   # pre-W48: the bare count '45'
    assert int(match.group(1)) >= ia.ranking_floor(verdict)
    assert f"core1 (strands at least {row['stranded']} endpoint(s); {verdict.summary})" in cell, cell
    # §1's overview column carries the short form of the same figure
    overview = [ln.split(" | ") for ln in lines if ln.startswith(("Wave ", "Group ")) and ln.count(" | ") == 6]
    assert any(cells[5] == f"≥ {match.group(1)} {ia.LOWER_BOUND_MARK}" for cells in overview), overview
    assert "more endpoints than the" in trigger and "only a lower bound" in trigger, trigger


def test_mop_names_a_held_core1_as_not_assessed_and_never_sizes_on_its_45(tmp_path, held):
    row, verdict = _core1(held)
    _lines, cell, _trigger = _mop_core1_wave(tmp_path, held, "mop_held.docx")
    assert f"core1 ({verdict.summary})" in cell, cell     # pre-W48: the bare count '45'
    assert cell.strip() != str(row["stranded"]), cell
    assert not cell.startswith(f"≥ {row['stranded']} "), cell


def test_mop_never_sizes_a_wave_or_its_trigger_on_a_zero_lower_bound(tmp_path, sample):
    """Refutation P2-1: a wave whose only published counts are zeros, beside a device whose own zero the owner
    withholds (a lower-bound row with no positive floor: it carries an inter-switch link with no trunk/STP evidence).
    The largest count is then a lower bound of 0, which is not a measurement of none. Before, the cell read '≥ 0 (lower
    bound)' and the trigger 'more endpoints than the §x.1 figure', i.e. more than ≥ 0, which holds at the start of
    every window. Now the figure is not assessed and the endpoint clause is withdrawn."""
    snap = copy.deepcopy(sample)
    pairs = ia.rows_with_verdicts(snap)
    zero = [row["host"] for row, v in pairs if v.published and row.get("stranded") == 0]
    floorless = [row["host"] for row, v in pairs if v.assessable == ia.LOWER_BOUND and ia.ranking_floor(v) is None]
    assert zero and floorless, (zero, floorless)         # precondition from the owner (the sample's dist2 and dist1)
    _with_waves(snap, [zero[0], floorless[0]])
    wave = ia.wave_blast([zero[0], floorless[0]], ia.wave_rows(snap), **_FLEET_CLEAR)
    assert wave.assessable == ia.NOT_ASSESSED and wave.zero_bound and wave.value == 0, wave
    lines = _mop_lines(tmp_path, snap, "mop_zero.docx")
    cell, trigger = _mop_wave(lines, floorless[0])
    assert cell.startswith(f"{ia.NOT_ASSESSED_CELL} — {ia.R_WAVE_ZERO}: "), cell
    assert f"{floorless[0]} (" in cell and cell.endswith("Do NOT read this as zero") and "≥ 0" not in cell, cell
    assert "more endpoints than the" not in trigger and "≥ 0" not in trigger, trigger
    assert "figure is not assessed" in trigger and "NOT a threshold of zero" in trigger, trigger
    assert _mop_overview(lines)["Wave 1 (pilot)"] == ia.NOT_ASSESSED_CELL


def test_mop_writes_a_fully_published_wave_as_a_bare_count_with_its_quantified_trigger(tmp_path, sample):
    """The positive control for the lower-bound and not-assessed wording (refutation P3): when the owner publishes
    every row of a wave, the cell is the bare count, with no lower-bound or not-assessed wording, and the trigger keeps
    its quantified endpoint clause with no lower-bound caveat."""
    snap = _published(copy.deepcopy(sample))
    lines = _mop_lines(tmp_path, snap, "mop_published.docx")
    cell, trigger = _mop_wave(lines, "core1")
    scope = lines[_mop_scope(lines, "core1")].split(" | ", 1)[1].split(", ")
    want = max(ia.table_value(v, "stranded") for row, v in ia.rows_with_verdicts(snap) if row["host"] in scope)
    assert cell == str(want), cell
    assert "lower bound" not in cell and "≥" not in cell and ia.NOT_ASSESSED_CELL not in cell, cell
    assert "more endpoints than the" in trigger and "only a lower bound" not in trigger, trigger
    assert str(want) in _mop_overview(lines).values(), _mop_overview(lines)


@pytest.mark.parametrize("variant", ["fleet_caveat", "hostless"])
def test_mop_bounds_a_published_wave_by_the_fleet_and_by_a_row_naming_no_host(tmp_path, sample, variant):
    """Refutation P2-2: the wave rule's two fleet-wide bounds, which AssessHub's cutover plan applied and the MOP did
    not. A device the collection reached only in part, or never, was simulated without its evidence, and a row that
    names no readable host could be any device of the wave: either makes even a fully published wave a lower bound."""
    snap = _published(copy.deepcopy(sample))
    if variant == "fleet_caveat":
        snap["collection_completeness"]["devices"] = [
            {"host": "ghost1", "status": "not collected", "data_quality": 0, "missing": ["version/inventory"]}]
        # the producer counts ghost1 (outside the devices map) in its inventory; W51 reconciles that count with the
        # roster, so an uncounted one would add a record the engine cannot read beside the blind device
        snap["collection_completeness"]["summary"]["inventory"] += 1
        said = ia.R_WAVE_FLEET_BLIND.format(n=1)
    else:
        snap["failure_impact"].append(dict(snap["failure_impact"][0], host=None))
        said = ("name no readable switch, so each could describe any device in this wave: "
                f"/failure_impact/{len(snap['failure_impact']) - 1} (")
    lines = _mop_lines(tmp_path, snap, f"mop_{variant}.docx")
    cell, trigger = _mop_wave(lines, "core1")
    scope = lines[_mop_scope(lines, "core1")].split(" | ", 1)[1].split(", ")
    want = max(ia.table_value(v, "stranded") for row, v in ia.rows_with_verdicts(snap)
               if row["host"] in scope and v.published)
    assert cell.startswith(f"≥ {want} {ia.LOWER_BOUND_MARK} — the worst case may be larger: "), cell
    assert said in cell, cell
    assert "more endpoints than the" in trigger and "only a lower bound" in trigger, trigger


# --- runbook §10 Risk Register ------------------------------------------------------------------------------------
def _runbook_lines(tmp_path, snap, name):
    pytest.importorskip("docx")
    from cisco_toolkit.runbook import write_runbook_docx
    out = str(tmp_path / name)
    write_runbook_docx(out, snap, "W48 Fleet")
    return _docx_lines(out)


def test_runbook_risk_register_writes_core1_as_a_lower_bound(tmp_path, bounded):
    row, verdict = _core1(bounded)
    lines = _runbook_lines(tmp_path, bounded, "rb_bounded.docx")
    expected = " | ".join(["core1"] + [str(ia.ranked_value(verdict, f)) for f in _RISK_FIELDS])
    assert f"≥ {row['stranded']} {ia.LOWER_BOUND_MARK}" in expected      # the expectation itself is a floor
    assert expected in lines, [ln for ln in lines if ln.startswith("core1 | ")]
    assert _raw_row(row, _RISK_FIELDS) not in lines                       # pre-W48: 'core1 | High | 3 | 45 | ...'
    assert any("publish their counts only as lower bounds" in ln
               and f"core1 (strands at least {row['stranded']} endpoint(s); {verdict.summary})" in ln
               for ln in lines)


def test_runbook_risk_register_names_a_held_core1_as_not_assessed(tmp_path, held):
    row, verdict = _core1(held)
    lines = _runbook_lines(tmp_path, held, "rb_held.docx")
    assert _raw_row(row, _RISK_FIELDS) not in lines                       # pre-W48: ranked first, raw
    assert any("are not ranked above" in ln and f"core1 ({verdict.summary})" in ln for ln in lines), \
        [ln for ln in lines if "not ranked" in ln]


def test_runbook_risk_register_ranks_a_lower_bound_by_its_floor_never_by_a_withheld_band(tmp_path, bounded):
    """Refutation P3: §10 sorted a lower-bound row by its raw severity, which the owner withholds when it is below the
    worst band (a partial row may understate it). core1 with a Medium band and its 45-endpoint floor must rank by that
    floor, above every row that strands fewer, as the workbook, design, deck, archreview and ops rankings order it."""
    snap = copy.deepcopy(bounded)
    next(r for r in snap["failure_impact"] if isinstance(r, dict) and r.get("host") == "core1")["severity"] = "Medium"
    row, verdict = _core1(snap)
    assert verdict.withholds("severity") and ia.ranking_floor(verdict) == row["stranded"] > 0, verdict.as_dict()
    pairs = {r["host"]: v for r, v in ia.rows_with_verdicts(snap)}

    def line(host):
        return " | ".join([host] + [str(ia.ranked_value(pairs[host], f)) for f in _RISK_FIELDS])

    lines = _runbook_lines(tmp_path, snap, "rb_order.docx")
    assert line("core1") in lines, [ln for ln in lines if ln.startswith("core1 | ")]
    # every published High row that strands fewer, and is shown, follows core1 (pre-fix: every one of them led it)
    fewer = [r["host"] for r, v in ia.rows_with_verdicts(snap) if v.published and r["severity"] == ia.IMPACT_WORST
             and 0 < r["stranded"] < row["stranded"] and line(r["host"]) in lines]
    assert fewer, "precondition: a published High row that strands fewer than core1 is shown"
    assert all(lines.index(line(host)) > lines.index(line("core1")) for host in fewer), fewer


# --- operations handbook §2.1 keystones ---------------------------------------------------------------------------
def _ops_lines(tmp_path, snap, name):
    pytest.importorskip("docx")
    from cisco_toolkit.ops import write_ops_handbook_docx
    out = str(tmp_path / name)
    write_ops_handbook_docx(out, snap, "W48 Fleet")
    return _docx_lines(out)


def test_ops_keystones_write_core1_as_a_lower_bound(tmp_path, bounded):
    row, verdict = _core1(bounded)
    lines = _ops_lines(tmp_path, bounded, "ops_bounded.docx")
    expected = " | ".join(["core1"] + [str(ia.ranked_value(verdict, f)) for f in _KEYSTONE_FIELDS])
    assert expected in lines, [ln for ln in lines if ln.startswith("core1 | ")]
    assert _raw_row(row, _KEYSTONE_FIELDS) not in lines                   # pre-W48: 'core1 | High | 45 | 3'
    assert any("publish their counts only as lower bounds" in ln and f"core1 (strands at least {row['stranded']}"
               in ln for ln in lines)


def test_ops_keystones_name_a_held_core1_as_not_assessed(tmp_path, held):
    row, verdict = _core1(held)
    lines = _ops_lines(tmp_path, held, "ops_held.docx")
    assert _raw_row(row, _KEYSTONE_FIELDS) not in lines                   # pre-W48: the top keystone, raw
    assert any("not ranked as keystones" in ln and f"core1 ({verdict.summary})" in ln for ln in lines), \
        [ln for ln in lines if "keystone" in ln.lower()]


# --- MCP failure_impact tool --------------------------------------------------------------------------------------
def _mcp_core1(snap):
    items = [item for item in mcp_server.failure_impact(snap, limit=10_000) if item.get("host") == "core1"]
    assert len(items) == 1, items
    return items[0]


def test_mcp_failure_impact_hands_the_assistant_core1_as_a_lower_bound(bounded):
    row, verdict = _core1(bounded)
    item = _mcp_core1(bounded)
    assert item.get("assessable") == ia.LOWER_BOUND and item.get("why") == verdict.why, item
    assert item["stranded"] == f"≥ {row['stranded']}", item                # pre-W48: the raw 45
    for field in ia.IMPACT_MEASURES:
        assert item[field] == ia.table_value(verdict, field), (field, item)
    assert item["severity"] != row["severity"] and item["backup"] == ia.NOT_ASSESSED_CELL, item


def test_mcp_failure_impact_hands_the_assistant_a_held_core1_as_not_assessed(held):
    row, verdict = _core1(held)
    item = _mcp_core1(held)
    assert item.get("assessable") == ia.NOT_ASSESSED, item
    assert all(item[field] == ia.NOT_ASSESSED_CELL for field in ia.IMPACT_MEASURES), item
    assert row["detail"] not in item["detail"] and item["detail"] == ia.table_detail(verdict), item


# ---------------------------------------------------------------------------------------------------------------
# part 3: the owner's wave rule (impact_assessability.wave_blast), which the MOP and the cutover plan both read
# ---------------------------------------------------------------------------------------------------------------
_NO_SUCH = "no-such-switch"
#: The wave rule's fleet counts for a collection that reached every device (wave_blast has no default for them).
_FLEET_CLEAR = {"blind": 0, "blind_unread": 0}


def test_the_wave_rule_publishes_only_a_wave_whose_every_device_is_a_measurement(sample):
    snap = _published(copy.deepcopy(sample))
    rows = ia.wave_rows(snap)
    by_host = {r.key: r for r in rows}
    exact = ia.wave_blast(["core1", "access1"], rows, **_FLEET_CLEAR)
    assert exact.assessable == ia.PUBLISHED and exact.complete and exact.n_not_ranked == 0, exact
    assert exact.value == max(by_host["core1"].stranded, by_host["access1"].stranded) and ia.wave_why(exact) == ""
    # a device with no row: the largest count is only a floor
    floor = ia.wave_blast(["core1", _NO_SUCH], rows, **_FLEET_CLEAR)
    assert floor.assessable == ia.LOWER_BOUND and floor.value == by_host["core1"].stranded and not floor.complete
    assert ia.wave_why(floor) == ia.R_WAVE_NO_ROW.format(n=1, names=_NO_SUCH)
    # beside only published zeros that floor is 0: not assessed, never a threshold of zero
    zeros = [r.key for r in rows if r.ranked and r.stranded == 0]
    assert zeros, "precondition: the sample has switches that strand nobody"
    zero = ia.wave_blast([zeros[0], _NO_SUCH], rows, **_FLEET_CLEAR)
    assert zero.assessable == ia.NOT_ASSESSED and zero.zero_bound and zero.value == 0 and zero.observed, zero
    # an exact zero stays a measurement; a wave none of whose devices has a row is not observed at all
    assert ia.wave_blast(zeros, rows, **_FLEET_CLEAR)[:2] == (ia.PUBLISHED, 0)
    none = ia.wave_blast([_NO_SUCH], rows, **_FLEET_CLEAR)
    assert none.assessable == ia.NOT_ASSESSED and none.value is None and not none.observed and not none.zero_bound


def test_the_wave_rule_never_lets_a_withheld_zero_or_a_held_row_size_a_wave(sample):
    snap = copy.deepcopy(sample)
    rows = ia.wave_rows(snap)
    pairs = ia.rows_with_verdicts(snap)
    floorless = [row["host"] for row, v in pairs if v.assessable == ia.LOWER_BOUND and ia.ranking_floor(v) is None]
    assert floorless, "precondition: a lower-bound row whose zero the owner withholds"
    alone = ia.wave_blast(floorless[:1], rows, **_FLEET_CLEAR)
    assert alone.assessable == ia.NOT_ASSESSED and alone.value is None and alone.observed and not alone.ranked
    assert ia.wave_why(alone).startswith(f"{floorless[0]} (lower bound — "), ia.wave_why(alone)


def test_the_wave_rule_bounds_every_wave_by_a_row_naming_no_host_and_by_the_fleet(sample):
    snap = _published(copy.deepcopy(sample))
    snap["failure_impact"].append(dict(snap["failure_impact"][0], host=None))
    rows = ia.wave_rows(snap)
    wave = ia.wave_blast(["core1"], rows, **_FLEET_CLEAR)
    assert wave.assessable == ia.LOWER_BOUND and len(wave.hostless) == 1 and wave.n_not_ranked == 1, wave
    assert f"/failure_impact/{len(rows) - 1} (" in ia.wave_why(wave), ia.wave_why(wave)
    clean = ia.wave_rows(_published(copy.deepcopy(sample)))
    assert ia.wave_blast(["core1"], clean, **_FLEET_CLEAR).assessable == ia.PUBLISHED
    for blind, said in ((2, ia.R_WAVE_FLEET_BLIND.format(n=2)), (None, ia.R_WAVE_FLEET_UNREAD),
                        (True, ia.R_WAVE_FLEET_UNREAD), (-1, ia.R_WAVE_FLEET_UNREAD)):
        bounded = ia.wave_blast(["core1"], clean, blind=blind, blind_unread=0)
        assert bounded.assessable == ia.LOWER_BOUND and not bounded.complete, (blind, bounded)
        assert ia.wave_why(bounded) == said, (blind, ia.wave_why(bounded))
    # W51: a collection record the projection cannot read as a blind device bounds the wave too, worded as what it
    # is, after any blind devices and in the order the cutover plan's blind-spot note uses
    for unread, said in ((1, ia.R_WAVE_FLEET_BLIND_UNREAD.format(n=1)), (None, ia.R_WAVE_FLEET_UNREAD)):
        bounded = ia.wave_blast(["core1"], clean, blind=0, blind_unread=unread)
        assert bounded.assessable == ia.LOWER_BOUND and not bounded.complete, (unread, bounded)
        assert ia.wave_why(bounded) == said, (unread, ia.wave_why(bounded))
    both = ia.wave_blast(["core1"], clean, blind=2, blind_unread=1)
    assert ia.wave_why(both) == "; ".join([ia.R_WAVE_FLEET_BLIND.format(n=2),
                                           ia.R_WAVE_FLEET_BLIND_UNREAD.format(n=1)]), ia.wave_why(both)


def test_the_wave_rule_takes_no_default_fleet_and_never_reads_a_hostless_row_as_unobserved(sample):
    """W51 (the W48 re-verification's P3s): ``wave_blast`` has no default for either fleet count, so a caller that
    forgets the fleet fails loudly instead of reading as exact, and ``None`` is unknown, never exact. A row that names
    no readable host could describe any device of a wave, so a wave none of whose devices has a row of its own is
    still observed (not assessed, with that row named), never [NOT OBSERVED] in the MOP while the cutover plan names
    it."""
    from cisco_toolkit import mop
    snap = _published(copy.deepcopy(sample))
    rows = ia.wave_rows(snap)
    for kwargs in ({}, {"blind": 0}, {"blind_unread": 0}):
        with pytest.raises(TypeError):
            ia.wave_blast(["core1"], rows, **kwargs)
    unknown = ia.wave_blast(["core1"], rows, blind=None, blind_unread=None)
    assert unknown.assessable == ia.LOWER_BOUND and not unknown.complete, unknown
    assert ia.wave_why(unknown) == ia.R_WAVE_FLEET_UNREAD
    assert mop._blast_for([_NO_SUCH], rows, (0, 0)) is None          # control: no row could describe the wave
    snap["failure_impact"].append(dict(snap["failure_impact"][0], host=None))
    rows = ia.wave_rows(snap)
    lone = ia.wave_blast([_NO_SUCH], rows, **_FLEET_CLEAR)
    assert lone.observed and lone.assessable == ia.NOT_ASSESSED and lone.value is None and len(lone.hostless) == 1
    assert "name no readable switch" in ia.wave_why(lone), ia.wave_why(lone)
    assert mop._blast_for([_NO_SUCH], rows, (0, 0)) == lone


def test_fleet_blind_reads_the_projections_fleet_qualifier(sample):
    """``(blind devices, records it cannot read as one)``, told apart by the projection's own classifier (W51)."""
    from cisco_toolkit import ui_projection

    def counts(snap):
        listing = ui_projection.project_topology(snap)["failure_impact"]
        return listing, ia.fleet_blind(listing, ui_projection.fleet_blind_spot_rows(snap))

    assert counts(copy.deepcopy(sample))[1] == (0, 0)
    snap = copy.deepcopy(sample)
    snap["collection_completeness"]["devices"] = [{"host": "ghost1", "status": "not collected"},
                                                  {"host": "access2", "status": "partial"}]
    snap["collection_completeness"]["summary"]["inventory"] += 1          # ghost1 is outside the devices map
    listing, blind = counts(snap)
    assert ia.FLEET_BLIND_CAVEAT in listing["caveats"] and blind == (2, 0)
    # a row the classifier cannot read as a blind spot is a record, never a blind device
    snap["collection_completeness"]["devices"].append(None)
    assert counts(snap)[1] == (2, 1)
    # the same witnesses with no classifier rows: every one is unread (never a blind device by default)
    assert ia.fleet_blind(counts(snap)[0], None) == (0, 3)
    # a record the snapshot does not carry: the qualifier with no witness under it is one record, never a device
    absent = copy.deepcopy(sample)
    del absent["collection_completeness"]
    listing, blind = counts(absent)
    assert ia.FLEET_BLIND_CAVEAT in listing["caveats"] and blind == (0, 1)
    assert ia.fleet_blind(None, []) is None and ia.fleet_blind({"items": "not a list"}, []) is None


# --- the consumer W48 STOPPED on (persisted output; see _RAW_RATCHET) ---------------------------------------------
_STOPPED = pytest.mark.xfail(
    strict=True, raises=AssertionError,
    reason="W48 STOP: persisted output, routed to the supervisor for a hosted regeneration (see _RAW_RATCHET). "
           "When the consumer reads the owner this XPASSes: delete this marker and the ratchet entry.")


@_STOPPED
@pytest.mark.parametrize("variant", ["bounded", "held"])
def test_design_advisor_never_counts_a_withheld_zero_as_a_measured_no_backup_device(variant, request):
    from cisco_toolkit.design_advisor import _signals
    snap = request.getfixturevalue(variant)
    measured = sum(1 for row, verdict in ia.rows_with_verdicts(snap)
                   if not verdict.withholds("severity") and row.get("severity") == "High"
                   and not verdict.withholds("backup") and row.get("backup") == 0)
    _row, verdict = _core1(snap)
    assert verdict.withholds("backup")                  # core1's zero is not a measurement in either variant
    assert _signals(snap)["nobackup_high"] == measured


# --- protocol_assurance: the frozen evidence binder; presentation through the owner at display time (W50) --------
# W48 stopped here because rehearsal.impacts is persisted in every AssessHub execution receipt and re-verified by
# recomputation on every read. W50: the receipt binds the stored rows as raw EVIDENCE through one frozen binder that
# never consults the owner (its bytes are frozen in tests/test_operator_evidence_contract.py), and every presentation
# of those rows is the owner's live reading at display time (webapp.backend.engine.rehearsal_impacts_view, the API's
# display-only impacts_view). The former strict xfail is now this check of the display path.
@pytest.mark.parametrize("variant", ["bounded", "held"])
def test_receipt_impact_rows_are_bound_raw_and_presented_only_through_the_owner(variant, request):
    from cisco_toolkit.protocol_assurance import cutover_operator_evidence
    from webapp.backend import engine as web_engine
    snap = request.getfixturevalue(variant)
    row, verdict = _core1(snap)
    # the receipt binds core1's stored row as evidence, raw ...
    assert [r for r in cutover_operator_evidence(snap)["rehearsal"]["impacts"] if r.get("host") == "core1"] == [row]
    # ... and the display reads it through the owner, live
    view = web_engine.rehearsal_impacts_view(snap, source_sha256="sha256:" + "0" * 64)
    assert view["display_only"] is True and view["available"] is True and view["owner"] == ia.SCHEMA
    items = [r for r in view["rows"] if r["host"] == "core1"]
    assert len(items) == 1, items
    item = items[0]
    assert item["assessable"] == verdict.assessable and item["state"] == verdict.state, item
    assert item["reasons"] == [{"code": code, "n": n} for code, n in verdict.code_counts], item
    for field in ia.IMPACT_MEASURES:
        cell = item["cells"][field]
        assert cell == ia.cell_reading(verdict, field)._asdict(), (field, cell)     # the owner decides; W50 r4
        if verdict.withholds(field):
            assert cell == {"kind": "withheld", "text": None, "state": verdict.withheld_state(field)}, (field, cell)
        else:
            assert cell["text"] == str(ia.table_value(verdict, field)), (field, cell)
    assert item["cells"]["backup"]["kind"] == "withheld"      # core1's stored 0 is never presented as a measured 0
    disclosed = [entry for entry in view["unranked"] if entry["index"] == verdict.index]
    if variant == "bounded":
        assert item["cells"]["stranded"] == {"kind": "floor", "text": f"≥ {row['stranded']}", "state": None}, item
        assert item["ranked"] is True and view["rows"][0]["host"] == "core1"     # its floor of 45 leads
        assert disclosed == []
    else:
        assert item["assessable"] == ia.NOT_ASSESSED and item["ranked"] is False, item
        assert item["state"] == ia.NOT_COLLECTED and item["reasons"] == [{"code": "legacy_row", "n": 0}], item
        # a held row is named in the owner's unranked disclosure, whatever cap a display puts on the ranked rows
        assert disclosed == [{key: item[key] for key in ("index", "host", "assessable", "state", "reasons")}]


_FROZEN_ROWS = ROOT / "tests" / "fixtures" / "operator_evidence_v1" / "after-rows.json"


@pytest.mark.parametrize("variant", ["bounded", "held", "frozen_rows"])
def test_the_owners_display_accessors_follow_its_own_rules(variant, request):
    """W50 round 4 (P3-3): the display path only formats; every reading is an owner accessor, and each agrees with
    the owner's own rules on the sample, a held row and the frozen corpus's odd-typed and non-object rows:
    ``cell_reading`` (withheld exactly when ``withholds``; a floor only on a lower bound, worded by ``table_value``;
    an unreadable value never a zero), ``RowVerdict.readable`` (the row is an object), ``unranked`` (exactly the
    rows ``ranks`` refuses, in stored order) and ``ranking_order`` (every ranked row first, largest floor or count
    first)."""
    snap = json.loads(_FROZEN_ROWS.read_text(encoding="utf-8")) if variant == "frozen_rows" else \
        request.getfixturevalue(variant)
    verdicts = ia.assess_failure_impact(snap)
    assert verdicts
    for verdict in verdicts:
        assert verdict.readable is isinstance(verdict.raw, dict)
        for field in ia.CELL_FIELDS:
            reading = ia.cell_reading(verdict, field)
            assert reading.kind in ia.CELL_KINDS, reading
            if verdict.withholds(field):
                assert reading == (ia.CELL_WITHHELD, None, verdict.withheld_state(field)), (field, reading)
            elif reading.kind == ia.CELL_FLOOR:
                assert verdict.assessable == ia.LOWER_BOUND and field in ia.IMPACT_MEASURES, reading
                assert reading.text == ia.table_value(verdict, field), reading
            elif reading.kind == ia.CELL_PUBLISHED:
                assert isinstance(reading.text, str) and reading.text.strip() and reading.state is None, reading
                raw = verdict.raw[field]
                if field == "severity":
                    assert reading.text == raw and raw in ia.IMPACT_SEVERITIES, reading
                elif field != "detail":
                    assert reading.text == str(ia.count_value(raw)), reading     # the owner's count, never "None"
            else:
                assert reading == (ia.CELL_UNREADABLE, None, None), reading
    with pytest.raises(ValueError):
        ia.cell_reading(verdicts[0], "host")
    assert [v.index for v in ia.unranked(verdicts)] == [v.index for v in verdicts if not ia.ranks(v)]
    ordered = sorted(verdicts, key=ia.ranking_order)
    flags = [ia.ranks(v) for v in ordered]
    assert flags == sorted(flags, reverse=True)
    counts = [ia.ranking_order(v)[1] for v in ordered if ia.ranking_order(v)[0] == 0]
    assert counts == sorted(counts), counts                 # largest floor or measured count first
    if variant == "frozen_rows":
        assert {v.index for v in verdicts if not v.readable} == {2, 4, 5, 6, 7}
        assert all(not ia.ranks(v) for v in verdicts if not v.readable)


def test_only_the_frozen_binder_reads_the_rows_raw_and_the_display_path_reaches_the_owner():
    """The protocol_assurance ratchet entry excuses the frozen evidence binder and nothing else: it is the module's
    only reader of the stored section, every other unit there stays subject to the guard, and the display path that
    presents the bound rows is routed to the owner."""
    readers, routed, _occ, _why = _scan(str(ROOT))
    assert {name for (module, name) in readers if module == "cisco_toolkit.protocol_assurance"} == {
        "_rehearsal_impact_evidence_v1"}
    assert [unit for unit in _RAW_RATCHET if unit[0] == "cisco_toolkit.protocol_assurance"] == [
        ("cisco_toolkit.protocol_assurance", "_rehearsal_impact_evidence_v1")]
    for unit in ("rehearsal_impacts_view", "receipt_impacts_view", "_trend_comparison_receipts"):
        assert ("webapp.backend.engine", unit) in routed, unit


def test_the_guard_still_flags_a_second_raw_presenter_beside_the_binder(tmp_path):
    """Non-vacuity for the narrowed entry: on a synthetic tree, a second protocol_assurance function that presents the
    rows raw is flagged even though the binder beside it is excused by the ratchet."""
    _write(tmp_path, "cisco_toolkit/__init__.py", "")
    _write(tmp_path, "cisco_toolkit/impact_assessability.py", "def rows_with_verdicts(snap):\n    return []\n")
    _write(tmp_path, "cisco_toolkit/protocol_assurance.py",
           "def _rehearsal_impact_evidence_v1(snap):\n    return list(snap.get('failure_impact') or [])\n"
           "def presenter(snap):\n    return [r.get('stranded') for r in snap['failure_impact']]\n")
    readers, routed, _occ, _why = _scan(str(tmp_path))
    flagged = {unit for unit in readers if unit not in routed and unit not in _RAW_RATCHET}
    assert flagged == {("cisco_toolkit.protocol_assurance", "presenter")}, flagged

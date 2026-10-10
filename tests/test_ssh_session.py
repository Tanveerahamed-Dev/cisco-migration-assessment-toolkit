"""W59 PR-1: SSH session disclosure -- the owner (`cisco_toolkit.ssh_session`) and its collector seams.

Design: docs/w59-legacy-ssh-design-2026-10-09.md (sections 4.2-4.4, 5, 6 and the test plan in section 8). Every
test names the regression it catches. Nothing here opens a network connection to anything but a socketpair;
the loopback interop tests against an independent SSH server live in tests/test_ssh_session_interop.py.

Versions are NEVER asserted against the floating test environment (design section 8, critique P2-7): exact
versions are read only from the lock file's text. In the floating environment every seam is tested by
behaviour.
"""
from __future__ import annotations

import ast
import json
import os
import re
import socket
import threading
import warnings
from pathlib import Path

import pytest

import COLLECT_PARSE_V3_23_0 as C
from cisco_toolkit import ssh_session as S

ROOT = Path(__file__).resolve().parents[1]
ENGINE = ROOT / "COLLECT_PARSE_V3_23_0.py"


# ------------------------------------------------------------------------------------------- helpers ---
def _lib(permits=False, paramiko="5.0.0", netmiko="4.8.0"):
    return S.library_block(paramiko_version=paramiko, netmiko_version=netmiko, transport_class="ObservedTransport",
                           default_permits_sha1=permits)


_DEFAULT_CONSENT = S.consent_for({"hostname": "sw1", "ip": "192.0.2.10"})
_LEGACY_CONSENT = {"device_profile": S.LEGACY_SHA1_PROFILE, "run_flag_profile": S.LEGACY_SHA1_PROFILE,
                   "named_on_run_flag": True, "effective_profile": S.LEGACY_SHA1_PROFILE}
_G14_SHA1, _GEX_SHA1 = S.LEGACY_SHA1_TIER_KEX
_RSA_SHA1 = S.LEGACY_SHA1_TIER_HOST_KEYS[0]
_HMAC_SHA1 = sorted(n for n in S.SHA1_MACS if "etm" not in n)[0]


def _observation(kex, hash_bytes, bits, host_key, mac="hmac-sha2-256", cipher="aes128-ctr", *,
                 engine_agrees=True, group_agrees=None, newkeys=True):
    server = {"kex": [kex], "host_key": [host_key], "cipher_c2s": [cipher], "cipher_s2c": [cipher],
              "mac_c2s": [mac], "mac_s2c": [mac]}
    client = {"kex": [kex, "ecdh-sha2-nistp256"], "host_key": [host_key, "rsa-sha2-512"], "cipher": [cipher],
              "mac": [mac, "hmac-sha2-512"]}
    return {"kexinit": True, "newkeys": newkeys, "server": server, "client": client,
            "negotiated": {"kex": kex, "kex_hash_bytes": hash_bytes, "dh_group_bits": bits,
                           "host_key_algorithm": host_key, "cipher_c2s": cipher, "cipher_s2c": cipher,
                           "mac_c2s": mac, "mac_s2c": mac, "strict_kex": False,
                           "server_software": "SSH-2.0-Cisco-1.25"} if newkeys else None,
            "engine_name_agrees": engine_agrees, "group_size_agrees": group_agrees, "dropped": 0}


def _row(record, *, live=True, eligible=()):
    return S.derive_row("sw1", S.render_record(record), evidence="sw1/_ssh_session.json", live=live,
                        eligible_hosts=eligible)


class _IncompatiblePeer(Exception):
    """paramiko's class NAME (the classifier matches by name and never imports paramiko)."""


_IncompatiblePeer.__name__ = "IncompatiblePeer"


class _WeakGroupRefused(_IncompatiblePeer):
    """W59 PR-2's GEX floor refusal, by NAME."""

    def __init__(self, offered_bits):
        super().__init__("weak group")
        self.offered_bits = offered_bits


_WeakGroupRefused.__name__ = "WeakGroupRefused"


class NetmikoTimeoutException(Exception):
    pass


def _wrapped(inner):
    """What netmiko 4.8 actually raises: its own timeout, with paramiko's exception as __context__."""
    try:
        try:
            raise inner
        except Exception:
            raise NetmikoTimeoutException("A paramiko SSHException occurred during connection creation")
    except NetmikoTimeoutException as outer:
        return outer


def _obs_with_lists(server, client):
    obs = S.SessionObservation()
    obs.server, obs.client, obs.kexinit_seen = server, client, True
    return obs


_P5_CLIENT = {"kex": ["curve25519-sha256@libssh.org", "ecdh-sha2-nistp256", "diffie-hellman-group14-sha256"],
              "host_key": ["ssh-ed25519", "rsa-sha2-512", "rsa-sha2-256"], "cipher": ["aes128-ctr"],
              "mac": ["hmac-sha2-256", _HMAC_SHA1]}


# =================================================================== T1: pristine tables, pure observer ===
_TABLE_ATTRS = ("_preferred_kex", "_kex_info", "_preferred_keys", "_key_info", "_preferred_pubkeys",
                "_preferred_ciphers", "_preferred_macs")


def _stock_snapshot():
    import paramiko
    from paramiko import kex_gex, kex_group14

    snap = {a: getattr(paramiko.Transport, a) for a in _TABLE_ATTRS}
    snap["RSAKey.HASHES"] = paramiko.RSAKey.HASHES
    snap["KexGexSHA256.min_bits"] = kex_gex.KexGexSHA256.min_bits
    snap["KexGroup14SHA256.hash_algo"] = kex_group14.KexGroup14SHA256.hash_algo
    return {k: (v, (dict(v) if isinstance(v, dict) else v)) for k, v in snap.items()}


def test_t1_observed_transport_leaves_every_stock_table_pristine():
    """T1 (observer half). Catches: a table written into an inherited class dict
    (`Transport._kex_info[...] = ...`), which would change negotiation for EVERY device and thread."""
    before = _stock_snapshot()
    a, b = socket.socketpair()
    try:
        t = C.ObservedTransport(a)
        assert isinstance(t, C._paramiko.Transport)
        t.close()
    finally:
        a.close()
        b.close()
    after = _stock_snapshot()
    for key, (obj, value) in before.items():
        assert after[key][0] is obj, f"{key} was replaced"
        assert after[key][1] == value, f"{key} was mutated"


def test_t1_the_mixin_defines_no_table_and_overrides_exactly_two_methods():
    """T1. Catches: an algorithm table (or any other override) added to the observer, which would make the
    default path's negotiation differ from stock paramiko."""
    for cls in (S.ObservingTransportMixin, C.ObservedTransport):
        keys = set(vars(cls))
        assert not {k for k in keys if k.startswith("_preferred_") or k.endswith("_info") or k == "HASHES"}, cls
        methods = {k for k, v in vars(cls).items() if callable(v) and not k.startswith("__")}
        assert methods == ({"_parse_kex_init", "_parse_newkeys"} if cls is S.ObservingTransportMixin else set()), \
            (cls, methods)
    assert C.ObservedTransport.__mro__[1:3] == (S.ObservingTransportMixin, C._paramiko.Transport)


def test_t1_the_observer_module_imports_no_network_library():
    """Catches: `import paramiko` (or socket) creeping into the owner, which would drop it out of the
    no-egress walk the attestation publishes."""
    from cisco_toolkit.attestation import NETWORK_IMPORTS, imported_names

    tree = ast.parse((ROOT / "cisco_toolkit" / "ssh_session.py").read_text(encoding="utf-8"))
    names = imported_names(tree)
    assert not {n for n in names if n in NETWORK_IMPORTS or n.split(".")[0] in NETWORK_IMPORTS}, names


# =================================================================== T7: the netmiko / paramiko seam ===
def test_t7_handler_table_binds_the_observing_parse_newkeys():
    """T7. Catches: paramiko no longer dispatching NEWKEYS through `self._parse_newkeys` (simulated upstream
    change) -- the record would silently stop being captured."""
    from paramiko.common import MSG_NEWKEYS

    a, b = socket.socketpair()
    try:
        t = C.ObservedTransport(a)
        handler = t._handler_table[MSG_NEWKEYS]
        assert handler.__func__ is S.ObservingTransportMixin._parse_newkeys
        t.close()
    finally:
        a.close()
        b.close()


def test_t7_netmiko_builds_its_client_through_the_overridden_hook():
    """T7. Catches: netmiko's `_build_ssh_client` no longer calling `_get_ssh_client_instance` (the hook renamed
    upstream) -- the observed client, and with it the whole record, would silently disappear."""
    driver = C._observed_driver_for("ios", C.ObservedTransport)
    base = C._NETMIKO_CLASS_MAPPER["cisco_ios"]
    assert driver.__mro__[1:3] == (C._ObservedClientMixin, base)
    inst = object.__new__(driver)
    inst.system_host_keys = False
    inst.alt_host_keys = False
    inst.alt_key_file = ""
    inst.key_policy = C._paramiko.AutoAddPolicy()
    recorder = S.SessionRecorder(None, consent=_DEFAULT_CONSENT, library=_lib())
    C._SESSION_HANDOFF.value = (recorder, C.ObservedTransport)
    try:
        client = inst._build_ssh_client()
    finally:
        C._SESSION_HANDOFF.value = None
    assert isinstance(client, C._ObservedSSHClient)
    assert client._ssh_recorder is recorder and client._ssh_transport_cls is C.ObservedTransport


def test_t7_observed_driver_refuses_construction_outside_the_factory():
    """Catches: an observed driver built without the hand-off (an unrecorded session would open)."""
    driver = C._observed_driver_for("nxos", C.ObservedTransport)
    inst = object.__new__(driver)
    C._SESSION_HANDOFF.value = None
    with pytest.raises(RuntimeError, match="outside _open_connection"):
        inst._get_ssh_client_instance()


def test_t7_observed_client_forces_its_transport_factory_and_binds_the_sink(monkeypatch):
    """T7. Catches: `connect()` honouring a caller's transport_factory (or none), which would build an
    unobserved stock transport; and a sink bound to the class or a thread-local instead of the instance."""
    captured = {}

    def fake_connect(self, *args, **kwargs):
        captured.update(kwargs)

    monkeypatch.setattr(C._paramiko.SSHClient, "connect", fake_connect)
    recorder = S.SessionRecorder(None, consent=_DEFAULT_CONSENT, library=_lib())
    obs = recorder.begin_attempt()
    client = C._ObservedSSHClient(recorder, C.ObservedTransport)
    client.connect("127.0.0.1", transport_factory=lambda *a, **k: None)
    factory = captured["transport_factory"]
    a, b = socket.socketpair()
    try:
        t = factory(a, disabled_algorithms=None)
        assert type(t) is C.ObservedTransport
        assert getattr(t, S.SINK_ATTRIBUTE) is obs
        assert not hasattr(C.ObservedTransport, S.SINK_ATTRIBUTE)
        t.close()
    finally:
        a.close()
        b.close()
    assert recorder.record and recorder.record["outcome"] == "established"   # written before connect() returned


#: The hosted test job's explicit flag (.github/workflows/ci.yml, "Run the complete default suite"): set there, and
#: only there, on the provisioned floating environment this test is about.
_HOSTED_FLOATING_ENV_FLAG = "ATLAS_SSH_FIXTURE_REQUIRED"


def test_t7_floating_environment_stock_tables_contain_no_sha1():
    """T7 / design section 3: the floating test environment resolves paramiko 5, whose stock tables cannot
    negotiate SHA-1 in the exchange hash or the host-key signature. Behavioural: the probe reads the live
    classes, never a version string. Catches: the environment silently regressing to a SHA-1-permitting
    paramiko (e.g. a netmiko[par4] cap), which would make every 'default path' claim here false.

    W59 PR-1 review (P3-k): the claim is about the HOSTED floating environment. On the hosted test job (the flag
    below is set) it is asserted exactly as before and can never skip; anywhere else (a workstation that holds the
    bundle's paramiko 4.0.0) it skips VISIBLY, naming the flag and the observed answer, instead of failing a suite
    that never claimed to be the floating environment."""
    import paramiko

    permits = S.permits_sha1(paramiko.Transport, paramiko.RSAKey)
    if os.environ.get(_HOSTED_FLOATING_ENV_FLAG) != "1":
        pytest.skip(f"not the hosted floating environment ({_HOSTED_FLOATING_ENV_FLAG} != 1); this paramiko "
                    f"{getattr(paramiko, '__version__', '?')} stock tables permit SHA-1: {permits}")
    assert permits is False


def test_t7_the_floating_environment_gate_is_set_on_the_hosted_test_job():
    """Non-vacuity of the P3-k gate: the hosted test job sets the flag on the suite step, so T7 above asserts there
    and cannot skip. Catches: the flag renamed or dropped from ci.yml, which would turn T7 into a silent skip."""
    import yaml

    ci = yaml.safe_load((ROOT / ".github" / "workflows" / "ci.yml").read_text(encoding="utf-8"))
    suite = [s for s in ci["jobs"]["test"]["steps"] if s.get("name") == "Run the complete default suite"]
    assert suite and suite[0]["env"].get(_HOSTED_FLOATING_ENV_FLAG) == "1"


def test_t7_the_probe_reads_tables_not_versions():
    """Catches: `permits_sha1` reduced to a constant or a version comparison."""
    class Old:
        _preferred_kex = ("ecdh-sha2-nistp256", _G14_SHA1)
        _kex_info = {}

    class New:
        _preferred_kex = ("ecdh-sha2-nistp256",)
        _preferred_keys = ("rsa-sha2-512",)

    class RSAOld:
        HASHES = {_RSA_SHA1: object}

    assert S.permits_sha1(Old) is True
    assert S.permits_sha1(New) is False
    assert S.permits_sha1(New, RSAOld) is True


def test_t7_exact_versions_are_read_only_from_the_lock_text():
    """Design section 8: the only exact-version assertion, against the shipped lock's TEXT (PR-3 re-locks)."""
    lock = (ROOT / "portable" / "windows-x64-requirements.lock").read_text(encoding="utf-8")
    assert re.search(r"(?m)^paramiko==4\.0\.0 \\$", lock)
    assert re.search(r"(?m)^netmiko==4\.7\.0 \\$", lock)


# ================================================================== T8: one patchable connection factory ===
def _engine_tree():
    return ast.parse(ENGINE.read_text(encoding="utf-8"))


# W59 PR-2 review (P3-e): the structural guard lives in ONE shared helper module, used by this file's T8 and by
# tests/test_legacy_ssh_consent.py's T8 (W59 PR-1 review P3-f, closed in round 2, moved there verbatim from here).
# W59 PR-2 review round 2 (P2): the scan itself is now owned by cisco_toolkit.attestation, whose published
# legacy_ssh_confined claim runs it too; the helper module keeps the collector's site maps and re-exports the scan.
from ssh_structural_support import (  # noqa: E402  (tests/ is on sys.path: root conftest.py)
    CONNECTION_ARGUMENT_SITES as _CONNECTION_ARGUMENT_SITES,
    CONNECTION_CONSTRUCTOR_SITES as _CONNECTION_CONSTRUCTOR_SITES,
    SSH_LIBRARIES as _SSH_LIBRARIES,
    callee_name as _callee_name,
    connection_constructor_calls as _connection_constructor_calls,
    connection_roots as _connection_roots,
    qualified_owners as _qualified_owners,
)


def test_t8_open_connection_is_the_only_constructor_of_a_netmiko_connection():
    """T8 (default half), structural (W59 PR-1 review, P3-f, closed in round 2). Every name the netmiko / paramiko
    imports bind is a root, and so is every dynamic route to a name (``importlib.import_module`` / ``__import__`` of
    either library or of a computed name, ``sys.modules[...]``, ``globals()`` / ``vars()`` / ``locals()`` /
    ``eval``). Taint flows through every binding and expression form: assignment and augmented assignment, tuple
    unpacking (element-wise), conditional expressions and ``and`` / ``or``, tuple / list / set / dict displays and
    comprehensions, ``for`` and comprehension targets, ``with ... as``, the walrus, attribute stores (keyed by the
    receiver, ``self`` by its class) and aliases of their receiver, subscript stores into a container, function
    parameters (bound from every resolvable call and from defaults), returns and ``yield`` (class factories),
    lambdas, ``functools.partial``, subclasses and ``type()`` compositions. Names resolve as Python resolves them
    (innermost enclosing function, then the module), so an unrelated local of the same name is never tainted.
    No connection-capable callee may be CALLED outside the four named sites, each constructs exactly through its
    declared callee, no connection-capable value may be HANDED to code outside the three named argument sites, and
    ``exec`` / ``eval`` / ``compile`` appear nowhere. Catches ``_netmiko.ConnectHandler(...)``,
    ``CLASS_MAPPER[...](...)``, ``getattr(_netmiko, ...)(...)``, an aliased class map, ``_paramiko.SSHClient()``, a
    second ``_observed_driver_for(...)(...)`` and every shape in the bypass list below -- any route that bypasses the
    observer and the live-safety patch point.

    What it does not establish: it is static and by name. An INSTANCE's own methods are not followed (constructing one
    is what it guards), a callee resolved only at run time through an object it cannot see is not followed except
    through the argument rule, and source text executed from a string is refused outright rather than analysed."""
    tree = _engine_tree()
    found, passes, tainted, factories = _connection_constructor_calls(tree)
    assert {"_netmiko", "_paramiko", "_NETMIKO_CLASS_MAPPER", "SSHDetect"} <= _connection_roots(tree)
    assert {"_observed_driver_for", "_transport_for_profile"} <= factories, factories     # non-vacuity
    assert ("_open_connection", "transport_cls") in tainted and ("_observed_driver_for", "cls") in tainted
    assert found == _CONNECTION_CONSTRUCTOR_SITES, found
    assert passes == _CONNECTION_ARGUMENT_SITES, passes
    assert not [n.lineno for n in ast.walk(tree) if isinstance(n, ast.Name) and n.id == "ConnectHandler"]
    owners = _qualified_owners(tree)
    for callee, callers in (("_open_connection", ["connect_device"]), ("autodetect_platform", ["connect_device"])):
        got = [owners.get(id(n)) for n in ast.walk(tree) if isinstance(n, ast.Call)
               and isinstance(n.func, ast.Name) and n.func.id == callee]
        assert got == callers, (callee, got)
    assert not hasattr(C, "ConnectHandler"), "a stale `C.ConnectHandler` patch must fail loudly, not pass silently"


def test_t8_an_unrelated_local_of_a_tainted_name_is_not_tainted():
    """The scan resolves names as Python does, so the guard is exact without false positives: a function elsewhere
    whose locals are also called ``base`` / ``cls`` / ``driver`` (names the factory's own locals use) constructs
    nothing connection-capable, and a method's ``self`` belongs to its own class."""
    extra = ("def _elsewhere(rows):\n    base = dict(rows)\n    cls = type(base)\n    driver = cls()\n"
             "    return driver\n\nclass _Other:\n    def __init__(self):\n        self._ssh_transport_cls = int\n\n"
             "    def run(self):\n        return self._ssh_transport_cls(3)\n")
    found, passes, _tainted, _factories = _connection_constructor_calls(
        ast.parse(ENGINE.read_text(encoding="utf-8") + "\n\n" + extra))
    assert found == _CONNECTION_CONSTRUCTOR_SITES and passes == _CONNECTION_ARGUMENT_SITES, (found, passes)


@pytest.mark.parametrize("bypass", [
    "def _bypass():\n    return _netmiko.ConnectHandler(host='x')\n",
    "def _bypass():\n    return _NETMIKO_CLASS_MAPPER['cisco_ios'](host='x')\n",
    "def _bypass():\n    mapper = _NETMIKO_CLASS_MAPPER\n    return mapper['cisco_ios'](host='x')\n",
    "def _bypass():\n    return getattr(_netmiko, 'ConnectHandler')(host='x')\n",
    "def _bypass():\n    return _observed_driver_for('ios', ObservedTransport)(host='x')\n",
    "def _bypass():\n    drv = _observed_driver_for('ios', ObservedTransport)\n    return drv(host='x')\n",
    "def _bypass():\n    client = _paramiko.SSHClient()\n    return client\n",
    "class _Sneaky(_paramiko.SSHClient):\n    pass\n\ndef _bypass():\n    return _Sneaky()\n",
    "def _bypass():\n    return SSHDetect(device_type='autodetect', host='x')\n",
    # W59 PR-1 review (P3-f, round 2): the nine shapes the first structural guard could not follow
    "def _bypass(flag):\n    cls = _netmiko.ConnectHandler if flag else None\n    return cls(host='x')\n",
    "def _bypass():\n    classes = [_netmiko.ConnectHandler]\n    return classes[0](host='x')\n",
    "def _bypass():\n    for cls in (_netmiko.ConnectHandler,):\n        return cls(host='x')\n",
    "def _bypass():\n    if (cls := _netmiko.ConnectHandler):\n        return cls(host='x')\n",
    "class _Holder:\n    def __init__(self):\n        self.cls = _netmiko.ConnectHandler\n\n"
    "    def _bypass(self):\n        return self.cls(host='x')\n",
    "def _open(factory):\n    return factory(host='x')\n\ndef _bypass():\n    return _open(_netmiko.ConnectHandler)\n",
    "import functools\n\ndef _bypass():\n    return functools.partial(_netmiko.ConnectHandler, host='x')()\n",
    "import importlib\n\ndef _bypass():\n    return importlib.import_module('netmiko').ConnectHandler(host='x')\n",
    "def _bypass():\n    return __import__('paramiko').SSHClient()\n",
    # and the further forms the closure covers
    "def _bypass():\n    return [cls(host='x') for cls in (_netmiko.ConnectHandler,)]\n",
    "def _bypass(flag):\n    cls = flag and _netmiko.ConnectHandler\n    return cls(host='x')\n",
    "def _bypass(cls=_netmiko.ConnectHandler):\n    return cls(host='x')\n",
    "def _bypass(*, cls=_paramiko.SSHClient):\n    return cls()\n",
    "def _bypass():\n    return sys.modules['netmiko'].ConnectHandler(host='x')\n",
    "def _bypass():\n    return globals()['_netmiko'].ConnectHandler(host='x')\n",
    "def _bypass():\n    make = lambda: _netmiko.ConnectHandler\n    return make()(host='x')\n",
    "def _bypass():\n    table = {'ios': _netmiko.ConnectHandler}\n    return table.get('ios')(host='x')\n",
    "def _bypass():\n    box = []\n    box.append(_netmiko.ConnectHandler)\n    return box\n",
    "def _bypass():\n    exec('import netmiko')\n",
    "def _gen():\n    yield _netmiko.ConnectHandler\n\ndef _bypass():\n    for cls in _gen():\n        return cls(host='x')\n",
    "def _bypass():\n    alias = _SESSION_HANDOFF\n    return alias.value[1](object())\n",
    "def _bypass():\n    return setattr(_SESSION_HANDOFF, 'x', _netmiko.ConnectHandler)\n",
    "def _bypass():\n    return _OBSERVED_DRIVERS.get(('cisco_ios', ObservedTransport))(host='x')\n",
])
def test_t8_the_structural_guard_catches_every_bypass_shape(bypass):
    """Non-vacuity of the guard above: each shape, appended to the real engine source, is found constructing through,
    or handing over, a connection-capable value in a site that is not one of the named ones."""
    tree = ast.parse(ENGINE.read_text(encoding="utf-8") + "\n\n" + bypass)
    found, passes, _tainted, _factories = _connection_constructor_calls(tree)
    owners = [o for o in list(found) + list(passes) if o.split(".")[-1] == "_bypass"]
    assert owners and (found != _CONNECTION_CONSTRUCTOR_SITES or passes != _CONNECTION_ARGUMENT_SITES), (
        found, passes)


def _dynamic_ssh_imports(tree):
    """Calls importing netmiko or paramiko by a string constant (``importlib.import_module`` / ``__import__``)."""
    out = []
    for node in ast.walk(tree):
        if isinstance(node, ast.Call) and _callee_name(node.func) in ("import_module", "__import__") \
                and node.args and isinstance(node.args[0], ast.Constant) and isinstance(node.args[0].value, str) \
                and node.args[0].value.split(".")[0] in _SSH_LIBRARIES:
            out.append(node.lineno)
    return out


def test_t8_no_other_shipped_module_imports_an_ssh_library():
    """The class the T8 guard scans, closed: only the collector and the W59 PR-2 legacy tier import netmiko or
    paramiko among the shipped sources (the wheel and bundle denominators of T10), and the legacy tier constructs
    no connection of any kind (the same structural scan finds no constructor call in it), so no other module can
    open an unobserved session."""
    importers, constructing = set(), {}
    for rel in _shipped_python_files():
        try:
            tree = _parse(ROOT / rel)
        except (SyntaxError, UnicodeDecodeError):
            continue
        if _connection_roots(tree) or _dynamic_ssh_imports(tree):
            importers.add(rel)
            if rel != "COLLECT_PARSE_V3_23_0.py":
                constructing[rel] = _connection_constructor_calls(tree)[0]
    assert importers == {"COLLECT_PARSE_V3_23_0.py", "cisco_toolkit/legacy_ssh.py"}, importers
    assert constructing == {"cisco_toolkit/legacy_ssh.py": {}}, constructing


def test_t8_patching_the_factory_intercepts_the_default_path(monkeypatch):
    """T8 (default half). Catches: `connect_device` reaching the network by any route the patch misses."""
    seen = []

    class _Dev:
        def send_command(self, cmd, read_timeout=None):
            return "ok\n"

    def fake_open(kwargs, platform, profile, recorder):
        seen.append((kwargs["host"], kwargs.get("port"), platform, profile, recorder.effective_profile))
        return _Dev()

    monkeypatch.setattr(C, "_open_connection", fake_open)
    dev, resolved = C.connect_device("192.0.2.10", "SW1", "u", "p", "ios", port=2222)
    assert isinstance(dev, _Dev) and resolved == "ios"
    assert seen == [("192.0.2.10", 2222, "ios", S.DEFAULT_PROFILE, S.DEFAULT_PROFILE)]


def test_t8_a_non_default_profile_never_falls_back_to_the_default_transport(monkeypatch):
    """W59 PR-1 shipped no legacy transport and refused the profile; W59 PR-2 resolves it to the legacy tier's
    frozen class. Either way a non-default profile must never silently fall back to the default transport: an
    unknown profile, and the legacy profile when the tier cannot honour consent here (stock paramiko already permits
    SHA-1), refuse with SshProfileRefused; otherwise the legacy profile gets the tier's own class."""
    with pytest.raises(C.SshProfileRefused):
        C._transport_for_profile("no-such-profile")
    from cisco_toolkit import legacy_ssh

    monkeypatch.setattr(legacy_ssh, "default_permits_sha1", lambda: True)
    with pytest.raises(C.SshProfileRefused):
        C._transport_for_profile(S.LEGACY_SHA1_PROFILE)
    monkeypatch.setattr(legacy_ssh, "default_permits_sha1", lambda: False)
    cls = C._transport_for_profile(S.LEGACY_SHA1_PROFILE)
    assert cls is legacy_ssh.LegacySHA1Transport and cls is not C.ObservedTransport
    assert C._transport_for_profile(S.DEFAULT_PROFILE) is C.ObservedTransport


# ======================================================= section 4.4: refusals are classified, never retried ===
def test_refusal_from_lists_is_classified_whatever_exception_surfaced():
    """Catches: classification by exception type only -- netmiko turns IncompatiblePeer into a timeout, and a
    KEXINIT followed by DISCONNECT surfaces as EOFError."""
    server = {"kex": [_G14_SHA1, "ext-info-s"], "host_key": [_RSA_SHA1], "cipher_c2s": ["aes128-ctr"],
              "cipher_s2c": ["aes128-ctr"], "mac_c2s": [_HMAC_SHA1], "mac_s2c": [_HMAC_SHA1]}
    obs = _obs_with_lists(server, _P5_CLIENT)
    for exc in (_wrapped(_IncompatiblePeer("no acceptable kex algorithm")), EOFError(), OSError("reset")):
        ref = S.classify_failure(exc, obs)
        assert ref["classification"] == "refused_legacy_only" and ref["category"] == "kex", (exc, ref)
        assert ref["names"] == [_G14_SHA1]                     # the pseudo-algorithm is not graded


@pytest.mark.parametrize("server_kex, expected, bits", [
    (["diffie-hellman-group1-sha1"], "refused_weak_dh", 1024),
    (["curve25519-sha256", "mlkem768x25519-sha256"], "refused_unsupported_modern", None),
    (["diffie-hellman-group1-sha1", _G14_SHA1], "refused_legacy_only", None),
    (["some-vendor-kex@example.org"], "unclassified", None),
    ([_G14_SHA1, "curve25519-sha256"], "refused_unsupported_modern", None),
])
def test_kex_refusal_grades(server_kex, expected, bits):
    """Catches: a weakness inferred from absence (an ungraded name read as legacy), and a modern-but-
    unimplemented offer graded as a device weakness."""
    server = {"kex": server_kex, "host_key": ["rsa-sha2-512"], "cipher_c2s": ["aes128-ctr"],
              "cipher_s2c": ["aes128-ctr"], "mac_c2s": ["hmac-sha2-256"], "mac_s2c": ["hmac-sha2-256"]}
    ref = S.classify_failure(EOFError(), _obs_with_lists(server, _P5_CLIENT))
    assert ref["classification"] == expected and ref["offered_group_bits"] == bits, ref


def test_host_key_and_cipher_mac_refusals():
    base = {"kex": ["ecdh-sha2-nistp256"], "host_key": ["rsa-sha2-512"], "cipher_c2s": ["aes128-ctr"],
            "cipher_s2c": ["aes128-ctr"], "mac_c2s": ["hmac-sha2-256"], "mac_s2c": ["hmac-sha2-256"]}
    hk = S.classify_failure(EOFError(), _obs_with_lists(dict(base, host_key=[_RSA_SHA1]), _P5_CLIENT))
    assert (hk["category"], hk["classification"]) == ("host_key", "refused_legacy_only")
    dss = S.classify_failure(EOFError(), _obs_with_lists(dict(base, host_key=["ssh-dss"]), _P5_CLIENT))
    assert dss["classification"] == "refused_legacy_only" and not S._within_legacy_tier(dss)
    mac = S.classify_failure(EOFError(), _obs_with_lists(dict(base, mac_s2c=["hmac-md5"]), _P5_CLIENT))
    assert (mac["category"], mac["classification"], mac["names"]) == ("mac", "refused_cipher_mac", ["hmac-md5"])


def test_exception_chain_evidence_when_no_lists_were_recorded():
    """Catches: a weak-group refusal raised as a plain SSHException being retried as `unknown` (critique P2-3),
    and a disconnect before any KEXINIT being treated as a refusal (T17's no-evidence case)."""
    weak = S.classify_failure(_wrapped(_WeakGroupRefused(1024)), S.SessionObservation())
    assert (weak["classification"], weak["offered_group_bits"]) == ("refused_weak_dh", 1024)
    inc = S.classify_failure(_wrapped(_IncompatiblePeer("Incompatible version (1.5 instead of 2.0)")), None)
    assert (inc["classification"], inc["category"]) == ("unclassified", "version")
    assert S.classify_failure(EOFError(), S.SessionObservation()) is None
    assert S.classify_failure(_wrapped(OSError("timed out")), None) is None


def test_connect_device_never_retries_a_refusal(monkeypatch):
    """Catches: a refusal retried CONNECT_MAX_ATTEMPTS times (three failed logins' worth of noise on a device
    that will refuse identically), and a refusal recorded as a plain connection failure."""
    calls = []
    server = {"kex": [_G14_SHA1], "host_key": [_RSA_SHA1], "cipher_c2s": ["aes128-ctr"],
              "cipher_s2c": ["aes128-ctr"], "mac_c2s": [_HMAC_SHA1], "mac_s2c": [_HMAC_SHA1]}

    def refusing_open(kwargs, platform, profile, recorder):
        calls.append(1)
        recorder.observation.server = server
        recorder.observation.client = _P5_CLIENT
        recorder.observation.kexinit_seen = True
        raise _wrapped(_IncompatiblePeer("Incompatible ssh peer (no acceptable kex algorithm)"))

    monkeypatch.setattr(C, "_open_connection", refusing_open)
    monkeypatch.setattr(C, "CONNECT_BACKOFF_BASE", 0)
    recorder = S.SessionRecorder(None, consent=_DEFAULT_CONSENT, library=_lib())
    dev, _ = C.connect_device("192.0.2.10", "SW1", "u", "p", "ios", session=recorder)
    assert dev is None and len(calls) == 1 and recorder.attempts == 1
    assert recorder.record["outcome"] == "negotiation_refused"
    assert recorder.record["refusal"]["classification"] == "refused_legacy_only"


def test_connect_device_keeps_the_same_profile_retry_for_an_unevidenced_failure(monkeypatch):
    """Catches: every disconnect classified as a refusal (no retry at all for a transient drop)."""
    calls = []

    def flaky_open(kwargs, platform, profile, recorder):
        calls.append(profile)
        raise _wrapped(EOFError())

    monkeypatch.setattr(C, "_open_connection", flaky_open)
    monkeypatch.setattr(C, "CONNECT_BACKOFF_BASE", 0)
    recorder = S.SessionRecorder(None, consent=_DEFAULT_CONSENT, library=_lib())
    dev, _ = C.connect_device("192.0.2.10", "SW1", "u", "p", "ios", session=recorder)
    assert dev is None and calls == [S.DEFAULT_PROFILE] * C.CONNECT_MAX_ATTEMPTS
    assert recorder.record["outcome"] == "connect_failed" and recorder.record["attempts"] == C.CONNECT_MAX_ATTEMPTS


def test_session_record_error_is_neither_retried_nor_an_sshexception(monkeypatch):
    """Catches: a post-authentication record failure converted by netmiko into a retried timeout (it must not
    subclass SSHException or OSError)."""
    import paramiko

    assert not issubclass(S.SessionRecordError, (paramiko.SSHException, OSError))
    calls = []

    def failing_open(kwargs, platform, profile, recorder):
        calls.append(1)
        raise S.SessionRecordError("session record could not be written (OSError)")

    monkeypatch.setattr(C, "_open_connection", failing_open)
    dev, _ = C.connect_device("192.0.2.10", "SW1", "u", "p", "ios",
                              session=S.SessionRecorder(None, consent=_DEFAULT_CONSENT, library=_lib()))
    assert dev is None and calls == [1]


# ========================================================== T12: the offline owner (sidecar -> row) ===
def test_t12_every_established_status_and_its_label():
    """T12. Catches: one label for both consent states (critique P2-1) and weak-DH losing to SHA-1."""
    cases = [
        (_observation(_G14_SHA1, 20, 2048, _RSA_SHA1, mac=_HMAC_SHA1), _DEFAULT_CONSENT, _lib(True, "4.0.0"),
         "legacy_sha1", "Medium", "negotiated SHA-1 SSH without opt-in (the collector permitted SHA-1: paramiko 4.0.0)"),
        (_observation(_G14_SHA1, 20, 2048, _RSA_SHA1), _LEGACY_CONSENT, _lib(), "legacy_sha1", "Medium",
         "collected over SHA-1 SSH by explicit opt-in (device row and run flag)"),
        (_observation("diffie-hellman-group1-sha1", 20, 1024, _RSA_SHA1), _DEFAULT_CONSENT, _lib(True, "4.0.0"),
         "weak_dh", "High", "negotiated a 1024-bit Diffie-Hellman group, below 2048 (without opt-in)"),
        (_observation("diffie-hellman-group-exchange-sha256", 32, 1536, "rsa-sha2-512"), _DEFAULT_CONSENT, _lib(),
         "weak_dh", "High", "1536-bit"),
        (_observation("curve25519-sha256@libssh.org", 32, None, "ssh-ed25519"), _DEFAULT_CONSENT, _lib(),
         "modern", None, None),
        (_observation("diffie-hellman-group14-sha256", 32, 2048, "rsa-sha2-256"), _DEFAULT_CONSENT, _lib(),
         "modern", None, None),
    ]
    for obs, consent, lib, status, severity, label_part in cases:
        rec = S.build_record(outcome="established", consent=consent, library=lib, attempts=1, observation=obs)
        assert S.validate_record(rec) == []
        row = _row(rec)
        assert (row["status"], row["severity"]) == (status, severity), (obs["negotiated"], row)
        if label_part is None:
            assert row["label"] is None and row["finding"] == "closed"
        else:
            assert label_part in row["label"] and "host key not verified" in row["label"]
        assert row["host_key_verified"] is False


def test_t12_auth_failure_keeps_its_negotiated_posture():
    """The password crossed the negotiated session either way (design section 6.2)."""
    rec = S.build_record(outcome="auth_failed", consent=_DEFAULT_CONSENT, library=_lib(True, "4.0.0"), attempts=1,
                         observation=_observation(_G14_SHA1, 20, 2048, _RSA_SHA1), failure_class="NetmikoAuthenticationException")
    assert _row(rec)["status"] == "legacy_sha1"
    no_kex = S.build_record(outcome="auth_failed", consent=_DEFAULT_CONSENT, library=_lib(), attempts=1)
    assert _row(no_kex)["status"] == "unknown"


def test_t12_inconsistent_observation_is_unknown_never_modern():
    """T12. Catches: a broken observer reporting `modern` (engine name SHA-256, exchange hash 20 bytes)."""
    obs = _observation("diffie-hellman-group14-sha256", 20, 2048, "rsa-sha2-512")
    row = _row(S.build_record(outcome="established", consent=_DEFAULT_CONSENT, library=_lib(), attempts=1,
                              observation=obs))
    assert (row["status"], row["reason"], row["finding"]) == ("unknown", "inconsistent observation", "verify")
    disagreeing = _observation("ecdh-sha2-nistp256", 32, None, "rsa-sha2-512", engine_agrees=False)
    row = _row(S.build_record(outcome="established", consent=_DEFAULT_CONSENT, library=_lib(), attempts=1,
                              observation=disagreeing))
    assert row["status"] == "unknown"


def test_t12_pending_records_on_each_profile():
    """T12 / T16 offline half. Catches: a `pending` record rendered `not_recorded` (absence as health) and a
    legacy-authorized device with no completed record falling back to `verify` (critique P2-6)."""
    default = _row(S.build_record(outcome="pending", consent=_DEFAULT_CONSENT, library=_lib()), live=False)
    assert (default["status"], default["finding"]) == ("unknown", "verify")
    legacy = _row(S.build_record(outcome="pending", consent=_LEGACY_CONSENT, library=_lib()), live=False)
    assert (legacy["status"], legacy["finding"], legacy["severity"]) == ("legacy_unrecorded", "exposed", "Medium")
    assert "never completed" in legacy["label"]


def test_t12_absent_sidecar_live_and_offline():
    """Catches: an absent record rendered modern; an eligible live device with no record not exposed."""
    off = S.derive_row("sw1", None, evidence="sw1/_ssh_session.json", live=False)
    assert (off["status"], off["finding"], off["evidence"]) == ("not_recorded", "verify", None)
    live_eligible = S.derive_row("sw1", None, evidence="x", live=True, eligible_hosts={"sw1"})
    assert (live_eligible["status"], live_eligible["finding"]) == ("legacy_unrecorded", "exposed")
    live_other = S.derive_row("sw1", None, evidence="x", live=True, eligible_hosts={"sw2"})
    assert live_other["status"] == "not_recorded"


@pytest.mark.parametrize("cls, status, finding", [
    ("refused_legacy_only", "refused_legacy_only", "exposed"),
    ("refused_weak_dh", "refused_weak_dh", "exposed"),
    ("refused_unsupported_modern", "refused_unsupported_modern", "verify"),
    ("refused_cipher_mac", "refused_cipher_mac", "verify"),
    ("unclassified", "unknown", "verify"),
])
def test_t12_each_refusal_class_maps_to_its_status(cls, status, finding):
    ref = {"category": "kex", "classification": cls, "detail": "no_common_kex",
           "offered_group_bits": 1024 if cls == "refused_weak_dh" else None, "names": [_G14_SHA1]}
    obs = _observation(_G14_SHA1, 20, 2048, _RSA_SHA1, newkeys=False)
    rec = S.build_record(outcome="negotiation_refused", consent=_DEFAULT_CONSENT, library=_lib(), attempts=1,
                         observation=obs, refusal=ref, failure_class="NetmikoTimeoutException")
    assert S.validate_record(rec) == []
    row = _row(rec)
    assert (row["status"], row["finding"]) == (status, finding), row


def test_t12_refusal_wording_follows_what_the_opt_in_can_reach():
    """Catches: 'requires the legacy-sha1 opt-in' printed for a device (DSA-only) that no profile can collect."""
    rsa = {"category": "host_key", "classification": "refused_legacy_only", "names": [_RSA_SHA1]}
    dsa = {"category": "host_key", "classification": "refused_legacy_only", "names": ["ssh-dss"]}
    assert "requires the legacy-sha1 opt-in" in S.disclosure_sentence({"status": "refused_legacy_only", "refusal": rsa})
    assert "no profile of this collector implements" in S.disclosure_sentence(
        {"status": "refused_legacy_only", "refusal": dsa})


def test_t12_malformed_sidecars_are_unknown_with_their_reason():
    for data in (b"not json", b'{"schema": "ssh_session/1"}', b'{"a": 1, "a": 2}', b"[NaN]"):
        row = S.derive_row("sw1", data, evidence="e", live=False)
        assert row["status"] == "unknown" and row["reason"], data
    row = S.derive_row("sw1", None, evidence="e", live=False, read_error="session record failed its custody read")
    assert row["status"] == "unknown" and "custody" in row["reason"]


def test_t12_consent_fields_come_only_from_the_sidecar():
    """Catches: consent re-derived from the current devices.json / command line at analysis time."""
    rec = S.build_record(outcome="established", consent=_LEGACY_CONSENT, library=_lib(), attempts=1,
                         observation=_observation(_G14_SHA1, 20, 2048, _RSA_SHA1))
    block = S.compute_ssh_sessions(["sw1"], lambda h: (S.render_record(rec), None), live=False,
                                   consent=S.offline_consent_block(), evidence_path=lambda h: f"{h}/x")
    row = block["rows"][0]
    assert (row["device_profile"], row["named_on_run_flag"], row["effective_profile"]) == \
        (S.LEGACY_SHA1_PROFILE, True, S.LEGACY_SHA1_PROFILE)
    assert block["consent"] == S.offline_consent_block()
    assert all(block["consent"][k] is None for k in block["consent"] if k != "mode")


def test_t12_closed_schema_rejects_identifying_content():
    """T12 / critique P2-5. Catches: a hostname, an address, a fingerprint or an out-of-grammar name entering
    the record (which the redacted-collection path would then leak)."""
    good = S.build_record(outcome="established", consent=_DEFAULT_CONSENT, library=_lib(), attempts=1,
                          observation=_observation("ecdh-sha2-nistp256", 32, None, "rsa-sha2-512"))
    assert S.validate_record(good) == []
    for mutate in (
        lambda r: r.update(hostname="core1"),
        lambda r: r["negotiated"].update(server_software="10.1.2.3"),
        lambda r: r["negotiated"].update(server_software="SSH-2.0-x 10.1.2.3"),
        lambda r: r["host_key"].update(fingerprint="SHA256:abc"),
        lambda r: r["server_offered"]["kex"].append("has space"),
        lambda r: r["server_offered"]["kex"].append("a,b"),
        lambda r: r["server_offered"]["kex"].append("x" * 65),
        lambda r: r["library"].update(paramiko="10.0.0.1"),
        lambda r: r["negotiated"].update(kex_hash_bytes=True),
        lambda r: r.update(refusal={"category": "kex"}),
    ):
        r = json.loads(json.dumps(good))
        mutate(r)
        assert S.validate_record(r), mutate


def test_t12_banner_comment_is_stripped_and_names_are_screened():
    assert S.server_software_from_banner("SSH-2.0-Cisco-1.25 core1.example 10.1.2.3") == ("SSH-2.0-Cisco-1.25", 0)
    assert S.server_software_from_banner("SSH-2.0-0011.2233.4455") == (None, 1)
    names, dropped = S._sanitize_names(["aes128-ctr", "10.0.0.1", "a b", "ok@example.org"])
    assert names == ["aes128-ctr", "ok@example.org"] and dropped == 2


def test_t12_offline_block_shape_and_software_risk_projection():
    rec = S.build_record(outcome="established", consent=_DEFAULT_CONSENT, library=_lib(True, "4.0.0"), attempts=1,
                         observation=_observation(_G14_SHA1, 20, 2048, _RSA_SHA1))
    store = {"sw1": S.render_record(rec)}
    block = S.compute_ssh_sessions(["sw1", "sw2"], lambda h: (store.get(h), None), live=False, consent=None,
                                   evidence_path=lambda h: f"{h}/{S.SIDECAR_FILENAME}")
    assert block["schema"] == S.SET_SCHEMA and block["owner"] == S.OWNER
    assert [r["status"] for r in block["rows"]] == ["legacy_sha1", "not_recorded"]
    assert block["summary"]["by_status"]["legacy_sha1"] == 1 and block["summary"]["mode"] == "offline"
    proj = S.software_risk_projection(block)
    assert proj["sw1"]["finding"] == "exposed" and proj["sw1"]["severity"] == "Medium"
    assert proj["sw2"]["finding"] == "verify"


def test_software_risk_surface_is_a_projection_of_the_owner():
    """Catches: the surface computed from config text (it is session evidence), or an absent row read closed."""
    from cisco_toolkit.analyze import compute_software_risk

    rec = S.build_record(outcome="established", consent=_DEFAULT_CONSENT, library=_lib(True, "4.0.0"), attempts=1,
                         observation=_observation("diffie-hellman-group1-sha1", 20, 1024, _RSA_SHA1))
    block = S.compute_ssh_sessions(["sw1", "sw2"], lambda h: ((S.render_record(rec) if h == "sw1" else None), None),
                                   live=False, consent=None, evidence_path=lambda h: h)
    sr = compute_software_risk({}, {}, {}, ["sw1", "sw2", "sw3"], ssh_sessions=block)
    surf = {d["host"]: d["surfaces"].get(S.SURFACE_KIND) for d in sr["per_device"]}
    assert surf == {"sw1": "exposed", "sw2": "verify", "sw3": "verify"}
    [finding] = [f for f in sr["findings"] if f["kind"] == S.SURFACE_KIND]
    assert finding["host"] == "sw1" and finding["severity"] == "High" and finding["evidence_verbatim"] is False
    assert "1024-bit" in finding["detail"]
    # direct callers that pass no block see no surface at all (backward compatible)
    assert S.SURFACE_KIND not in str(compute_software_risk({}, {}, {}, ["sw1"]))
    # a failed block ({}) reads verify for every host, never closed
    failed = compute_software_risk({}, {}, {}, ["sw1"], ssh_sessions={})
    assert failed["per_device"][0]["surfaces"][S.SURFACE_KIND] == "verify"


# ================================================================ recorder write timing (section 6.1) ===
def test_recorder_write_failures_are_recorded_and_raised(tmp_path):
    def broken(path, record):
        raise PermissionError("denied")

    rec = S.SessionRecorder(str(tmp_path / S.SIDECAR_FILENAME), consent=_DEFAULT_CONSENT, library=_lib(),
                            writer=broken)
    with pytest.raises(PermissionError):
        rec.write_pending()
    rec.begin_attempt()
    with pytest.raises(S.SessionRecordError):
        rec.complete_established()
    assert rec.finish_failure("connect_failed", EOFError()) is False
    assert [f["stage"] for f in rec.write_failures] == ["pending", "established", "connect_failed"]


def test_recorder_writes_canonical_lf_bytes(tmp_path):
    path = tmp_path / S.SIDECAR_FILENAME
    rec = S.SessionRecorder(str(path), consent=_DEFAULT_CONSENT, library=_lib())
    rec.write_pending()
    data = path.read_bytes()
    assert b"\r" not in data and data.endswith(b"\n")
    assert json.loads(data)["outcome"] == "pending"
    assert not [p for p in tmp_path.iterdir() if p.name != S.SIDECAR_FILENAME], "a temp file was left behind"


def test_t16_pending_write_failure_leaves_the_device_unconnected(monkeypatch, tmp_path):
    """T16 (default half). Catches: a device connected although its record could not be written first."""
    opened = []
    monkeypatch.setattr(C, "_open_connection", lambda *a, **k: opened.append(a) or None)
    original = C._session_recorder_for

    def failing_recorder(devinfo, dev_dir):
        r = original(devinfo, dev_dir)
        r._writer = lambda path, record: (_ for _ in ()).throw(PermissionError("denied"))
        return r

    monkeypatch.setattr(C, "_session_recorder_for", failing_recorder)
    failures = []
    platform, cmd = C._collect_live_device(
        {"hostname": "SW1", "ip": "192.0.2.10", "username": "u", "password": "p", "platform": "ios"},
        str(tmp_path / "SW1"), claimed=set(), lock=threading.Lock(),
        on_record_failure=lambda *a: failures.append(a))
    assert cmd is None and opened == []
    assert failures == [("SW1", "pending", "PermissionError", 0)]


def test_duplicate_device_folders_cannot_share_one_record(monkeypatch, tmp_path):
    """Catches: two devices.json rows resolving to one folder silently overwriting each other's record."""
    monkeypatch.setattr(C, "_open_connection", lambda *a, **k: None)
    monkeypatch.setattr(C, "CONNECT_MAX_ATTEMPTS", 1)
    claimed, failures = set(), []
    dev = {"hostname": "SW1", "ip": "192.0.2.10", "username": "u", "password": "p", "platform": "ios"}
    C._collect_live_device(dev, str(tmp_path / "SW1"), claimed=claimed, lock=threading.Lock(),
                           on_record_failure=lambda *a: failures.append(a))
    C._collect_live_device(dict(dev), str(tmp_path / "SW1"), claimed=claimed, lock=threading.Lock(),
                           on_record_failure=lambda *a: failures.append(a))
    assert failures == [("SW1", "pending", "FileExistsError", 0)]


@pytest.mark.parametrize("port", ["22", 0, 70000, True, 2.5])
def test_an_invalid_port_is_refused_not_replaced_by_22(port):
    with pytest.raises(ValueError, match="port"):
        C._ssh_port({"hostname": "SW1", "port": port})
    assert C._ssh_port({"hostname": "SW1"}) is None and C._ssh_port({"port": 2222}) == 2222


# ======================================================================= consent block (section 5) ===
def test_pr1_consent_is_recorded_and_never_effective():
    """PR-1 has no run flag: a row that requests legacy-sha1 is recorded as requested-not-named and collected
    on the default path. Catches: a row request alone enabling a legacy profile (a stale standing grant)."""
    rows = [{"hostname": "a", "ip": "192.0.2.1", "ssh_profile": S.LEGACY_SHA1_PROFILE},
            {"hostname": "b", "ip": "192.0.2.2", "ssh_profile": True},
            {"hostname": "c", "ip": "192.0.2.3"}]
    block = S.live_consent_block(rows, None)
    assert block["devices_requesting_legacy"] == ["a"] and block["requested_not_named"] == ["a"]
    assert block["devices_eligible"] == [] and block["run_flag"] == {"profile": None, "hosts_named": []}
    assert S.consent_for(rows[0])["effective_profile"] == S.DEFAULT_PROFILE
    assert S.consent_for(rows[1])["device_profile"] == S.DEFAULT_PROFILE     # no truthiness coercion
    flagged = S.consent_for(rows[0], {"profile": S.LEGACY_SHA1_PROFILE, "hosts_named": ["a"]})
    assert flagged["effective_profile"] == S.LEGACY_SHA1_PROFILE


# ===================================================================== T10: SHA-1 identifier scope ===
# W59 PR-2 review (P3-e/P3-f): the shipped-source denominator is shared with tests/test_legacy_ssh.py's T10.
from ssh_structural_support import (  # noqa: E402  (tests/ is on sys.path: root conftest.py)
    parse_source as _parse,
    shipped_python_files as _shipped_python_files,
)

def test_t10_ssh_sha1_algorithm_literals_appear_only_in_the_vocabulary_owner():
    """T10 (vocabulary half). Catches: an SSH SHA-1 algorithm name (e.g. the RSA SHA-1 host-key name) added to
    any other shipped module -- a forgotten allowlist re-creating the advisory class with no scanner signal."""
    files = _shipped_python_files()
    assert "cisco_toolkit/ssh_session.py" in files and "COLLECT_PARSE_V3_23_0.py" in files
    pattern = re.compile(r"(?<![A-Za-z0-9@.\-])(?:" + "|".join(
        re.escape(n) for n in sorted(S.SSH_SHA1_ALGORITHM_NAMES, key=len, reverse=True)) + r")(?![A-Za-z0-9@.\-])")
    offenders = []
    for rel in files:
        if rel == "cisco_toolkit/ssh_session.py":
            continue
        try:
            tree = _parse(ROOT / rel)
        except (SyntaxError, UnicodeDecodeError):
            continue
        for node in ast.walk(tree):
            if isinstance(node, ast.Constant) and isinstance(node.value, str) and pattern.search(node.value):
                offenders.append((rel, node.lineno))
    assert not offenders, offenders


def test_t10_sha1_hash_primitives_stay_at_their_declared_sites():
    """T10. `hashlib.sha1` only at the declared git-blob identity sites and the W59 PR-2 legacy tier (its SHA-1
    exchange hash); `hashes.SHA1` only in the legacy tier (its SHA-1 RSA signature). Catches: a SHA-1 primitive
    added to a default-path module."""
    allowed_hashlib = {"portable/release_contract.py", "webapp/backend/observe_ui_projection_contract.py",
                       "cisco_toolkit/legacy_ssh.py"}
    found_hashlib, found_crypto = set(), set()
    for rel in _shipped_python_files():
        try:
            tree = _parse(ROOT / rel)
        except (SyntaxError, UnicodeDecodeError):
            continue
        for node in ast.walk(tree):
            if isinstance(node, ast.Attribute) and isinstance(node.value, ast.Name):
                if node.value.id == "hashlib" and node.attr == "sha1":
                    found_hashlib.add(rel)
                if node.value.id == "hashes" and node.attr == "SHA1":
                    found_crypto.add(rel)
    assert found_hashlib <= allowed_hashlib, found_hashlib
    assert "cisco_toolkit/legacy_ssh.py" in found_hashlib, "the scan no longer sees the tier's own SHA-1 site"
    assert found_crypto == {"cisco_toolkit/legacy_ssh.py"}, found_crypto


# ======================================================= sidecar basename owner and its restatements ===
def test_collection_sidecar_owner_and_restatements_agree():
    """Design section 6.1 custody. Catches: a new collection sidecar left out of the wheel auditor (it would
    ship client evidence) or the verifier's statement."""
    from cisco_toolkit.capture_integrity import COLLECTION_SIDECAR_BASENAMES
    from tools.audit_wheel import _COLLECTION_SIDECAR_BASENAMES
    from webapp.backend.redaction_verify import SSH_SESSION_RECORD_BASENAME

    assert S.SIDECAR_FILENAME in COLLECTION_SIDECAR_BASENAMES
    assert set(_COLLECTION_SIDECAR_BASENAMES) == set(COLLECTION_SIDECAR_BASENAMES)
    assert SSH_SESSION_RECORD_BASENAME == S.SIDECAR_FILENAME


def test_evidence_records_seal_a_refused_devices_record(tmp_path):
    """Design section 6.1. Catches: the record of a REFUSED device (absent from the command map) left unsealed,
    so the evidence its finding rests on could change unnoticed."""
    root = tmp_path / "collection"
    (root / "edge1").mkdir(parents=True)
    rec = S.build_record(outcome="pending", consent=_DEFAULT_CONSENT, library=_lib())
    (root / "edge1" / S.SIDECAR_FILENAME).write_bytes(S.render_record(rec))
    sealed = C._evidence_records({}, str(root), session_hosts=["edge1"])
    assert [(f["path"], f["commands"]) for f in sealed["files"]] == \
        [(f"edge1/{S.SIDECAR_FILENAME}", [S.CUSTODY_ROLE])]
    assert C._evidence_records({}, str(root))["files"] == []        # not named -> not sealed (explicit denominator)


# ======================================================== asyncssh stays a test-only, isolated peer ===
def test_asyncssh_is_absent_from_every_shipped_or_installed_declaration():
    """Design section 8. Catches: the EPL/GPL test peer leaking into the bundle lock or any pyproject extra."""
    lock = (ROOT / "portable" / "windows-x64-requirements.lock").read_text(encoding="utf-8").lower()
    assert "asyncssh" not in lock
    pyproject = (ROOT / "pyproject.toml").read_text(encoding="utf-8").lower()
    assert "asyncssh" not in pyproject
    for name in ("requirements.txt", "requirements-dev.txt"):
        p = ROOT / name
        if p.exists():
            assert "asyncssh" not in p.read_text(encoding="utf-8").lower(), name
    fixture = (ROOT / "tools" / "requirements-ssh-fixture-test.txt").read_text(encoding="utf-8")
    pins = re.findall(r"(?m)^([A-Za-z0-9_.\-]+)==([^\s\\]+) \\$", fixture)
    assert ("asyncssh", "2.24.1") in pins
    for name, _version in pins:
        block = fixture.split(f"{name}==", 1)[1].split("\n# ", 1)[0]
        assert "--hash=sha256:" in block, name


def test_no_test_module_imports_asyncssh_or_the_fixture_server():
    """Design section 8. Catches: collection depending on asyncssh (the test interpreter never has it)."""
    offenders = []
    for path in sorted((ROOT / "tests").rglob("*.py")) + sorted((ROOT / "webapp" / "tests").rglob("*.py")):
        if path == ROOT / "tests" / "ssh_fixture" / "server.py":
            continue
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", SyntaxWarning)      # an unrelated file's escape-sequence warning
            tree = ast.parse(path.read_text(encoding="utf-8"))
        for node in ast.walk(tree):
            names = []
            if isinstance(node, ast.Import):
                names = [a.name for a in node.names]
            elif isinstance(node, ast.ImportFrom):
                names = [node.module or ""] + [f"{node.module}.{a.name}" for a in node.names]
            if any(n == "asyncssh" or n.startswith("asyncssh.") or n.endswith("ssh_fixture.server")
                   or n == "ssh_fixture.server" for n in names):
                offenders.append((str(path.relative_to(ROOT)), node.lineno))
    assert not offenders, offenders


def test_ci_provisions_the_fixture_and_fails_closed_on_the_required_legs():
    """Design section 8. Catches: the interop gate passing vacuously (fixture never provisioned, or its absence
    allowed to skip on the legs that are supposed to run it), and the fixture file left unaudited."""
    import yaml

    ci = yaml.safe_load((ROOT / ".github" / "workflows" / "ci.yml").read_text(encoding="utf-8"))
    steps = ci["jobs"]["test"]["steps"]
    provision = [s for s in steps if "requirements-ssh-fixture-test.txt" in str(s.get("run", ""))]
    assert len(provision) == 1 and "--require-hashes" in provision[0]["run"]
    assert "ATLAS_SSH_FIXTURE_PYTHON" in provision[0]["run"]
    suite = [s for s in steps if s.get("name") == "Run the complete default suite"]
    assert suite and suite[0]["env"].get("ATLAS_SSH_FIXTURE_REQUIRED") == "1"
    assert steps.index(provision[0]) < steps.index(suite[0])
    audit_runs = [s.get("run", "") for s in ci["jobs"]["dependency-audit"]["steps"]]
    assert any("--disable-pip -r tools/requirements-ssh-fixture-test.txt" in r for r in audit_runs)


def test_assesshub_never_carries_the_legacy_ssh_opt_in():
    """Design section 5: AssessHub never collects live, so the run-level opt-in may never appear in its
    sources (it arrives in the engine with W59 PR-2). Catches: a webapp path that could pass it."""
    hits = [str(p.relative_to(ROOT)) for p in (ROOT / "webapp" / "backend").rglob("*.py")
            if "--allow-legacy-ssh" in p.read_text(encoding="utf-8")]
    assert hits == []


# ====================================================================================================================
# W59 PR-1 review fixes. Each test names the review finding it closes; all are written for the hosted suite.
# ====================================================================================================================
def _modern_record_bytes():
    rec = S.build_record(outcome="established", consent=_DEFAULT_CONSENT, library=_lib(), attempts=1,
                         observation=_observation("curve25519-sha256@libssh.org", 32, None, "ssh-ed25519"))
    assert S.validate_record(rec) == []
    return S.render_record(rec)


def _live_block(store, written):
    return S.compute_ssh_sessions(["sw1"], lambda h: (store.get(h), None), live=True,
                                  consent=S.live_consent_block([{"hostname": "sw1", "ip": "192.0.2.10"}], None),
                                  evidence_path=lambda h: f"{h}/{S.SIDECAR_FILENAME}", run_written=written)


def test_p2a_a_live_run_reads_only_the_records_it_wrote():
    """P2-a (owner half). Catches: a MODERN record an earlier run left in the folder read as this run's posture --
    absence of this run's record rendered as health."""
    store = {"sw1": _modern_record_bytes()}
    stale = _live_block(store, [])["rows"][0]
    assert (stale["status"], stale["finding"], stale["reason"]) == \
        ("unknown", "verify", S.REASON_NOT_WRITTEN_BY_THIS_RUN)
    assert _live_block(store, None)["rows"][0]["status"] == "unknown"        # a live run that wrote nothing
    assert _live_block(store, ["sw1"])["rows"][0]["status"] == "modern"      # non-vacuity: this run's own record
    assert _live_block({}, [])["rows"][0]["status"] == "not_recorded"        # no record at all stays not_recorded
    offline = S.compute_ssh_sessions(["sw1"], lambda h: (store.get(h), None), live=False, consent=None,
                                     evidence_path=lambda h: h)
    assert offline["rows"][0]["status"] == "modern"                          # a re-analysis reads the sealed set


def _dev(**extra):
    return {"hostname": "SW1", "ip": "192.0.2.10", "username": "u", "password": "p", "platform": "ios", **extra}


def _row_after(dev_dir, written):
    """The row the analysis derives for SW1 from whatever sidecar the folder now holds, given what this run wrote."""
    sidecar = dev_dir / S.SIDECAR_FILENAME
    data = sidecar.read_bytes() if sidecar.exists() else None
    run_written = sorted(h for h, state in written.items() if state == C.RUN_RECORD_WRITTEN)
    return S.compute_ssh_sessions(["SW1"], lambda h: (data, None), live=True, consent=None,
                                  evidence_path=lambda h: h, run_written=run_written)["rows"][0]


def _stale_modern_folder(tmp_path):
    dev_dir = tmp_path / "SW1"
    dev_dir.mkdir()
    (dev_dir / S.SIDECAR_FILENAME).write_bytes(_modern_record_bytes())
    return dev_dir


def test_p2a_a_stale_modern_record_and_a_failed_pending_write_never_read_modern(monkeypatch, tmp_path):
    """P2-a. Catches: the pending write failing while an earlier run's MODERN record still sits in the folder,
    which the analysis then read as this device's posture."""
    opened = []
    monkeypatch.setattr(C, "_open_connection", lambda *a, **k: opened.append(a))
    dev_dir = _stale_modern_folder(tmp_path)
    original = C._session_recorder_for

    def failing_recorder(devinfo, d):
        r = original(devinfo, d)
        r._writer = lambda path, record: (_ for _ in ()).throw(PermissionError("denied"))
        return r

    monkeypatch.setattr(C, "_session_recorder_for", failing_recorder)
    written, failures = {}, []
    _platform, cmd = C._collect_live_device(_dev(), str(dev_dir), claimed=set(), lock=threading.Lock(),
                                            on_record_failure=lambda *a: failures.append(a), written=written)
    assert cmd is None and opened == []
    assert failures == [("SW1", "pending", "PermissionError", 0)]
    assert written == {"SW1": C.RUN_RECORD_FAILED}
    assert not (dev_dir / S.SIDECAR_FILENAME).exists()          # the stale record went before the write was tried
    row = _row_after(dev_dir, written)
    assert (row["status"], row["finding"]) == ("not_recorded", "verify")


def test_p2a_a_stale_modern_record_and_an_invalid_port_never_read_modern(monkeypatch, tmp_path):
    """P2-a. Catches: a device refused before connecting (here an invalid port) leaving the stale MODERN record to
    stand as its posture."""
    monkeypatch.setattr(C, "_open_connection", lambda *a, **k: pytest.fail("an invalid port must never connect"))
    dev_dir = _stale_modern_folder(tmp_path)
    written, failures = {}, []
    _platform, cmd = C._collect_live_device(_dev(port="22"), str(dev_dir), claimed=set(), lock=threading.Lock(),
                                            on_record_failure=lambda *a: failures.append(a), written=written)
    assert cmd is None and failures == [("SW1", "pending", "ValueError", 0)]
    row = _row_after(dev_dir, written)
    assert (row["status"], row["finding"], row["reason"]) == ("unknown", "verify", S.REASON_NOT_WRITTEN_BY_THIS_RUN)


def test_p2a_a_stale_record_that_cannot_be_removed_refuses_the_connection(monkeypatch, tmp_path):
    """P2-a. Catches: a session opened beside a stale record the collector could not remove."""
    def cannot_remove(path):
        raise PermissionError("held by an on-access scan")

    monkeypatch.setattr(C, "_unlink_stale_session_record", cannot_remove)
    monkeypatch.setattr(C, "_open_connection", lambda *a, **k: pytest.fail("must not connect"))
    dev_dir = _stale_modern_folder(tmp_path)
    written, failures = {}, []
    _platform, cmd = C._collect_live_device(_dev(), str(dev_dir), claimed=set(), lock=threading.Lock(),
                                            on_record_failure=lambda *a: failures.append(a), written=written)
    # a stale-record removal failure carries its own attempt count as the fifth field (0 here: the patched remover
    # reports none); the atomic-replace count of a removal is 0, never another stage's count
    assert cmd is None and failures == [("SW1", "stale_record_unlink", "PermissionError", 0, 0)]
    assert _row_after(dev_dir, written)["status"] == "unknown"


def test_p2a_a_second_row_claiming_the_folder_makes_the_record_ambiguous(monkeypatch, tmp_path):
    """P2-a. Catches: the record the first devices.json row wrote read as the posture of BOTH rows that resolve to
    one folder."""
    monkeypatch.setattr(C, "_open_connection", lambda *a, **k: None)
    monkeypatch.setattr(C, "CONNECT_MAX_ATTEMPTS", 1)
    claimed, written, failures = set(), {}, []
    for _ in range(2):
        C._collect_live_device(_dev(), str(tmp_path / "SW1"), claimed=claimed, lock=threading.Lock(),
                               on_record_failure=lambda *a: failures.append(a), written=written)
    assert written == {"SW1": C.RUN_RECORD_CONFLICT}
    assert ("SW1", "pending", "FileExistsError", 0) in failures
    assert _row_after(tmp_path / "SW1", written)["status"] == "unknown"


def test_p2a_the_engine_derives_live_rows_only_from_the_records_it_wrote():
    """P2-a, the wiring: main() passes the run's own record set to the owner on a live run (None offline)."""
    text = ENGINE.read_text(encoding="utf-8")
    call = text[text.index('"SSH session disclosure", ssh_session.compute_ssh_sessions'):]
    call = call[:call.index("_default={})")]
    assert "run_written=" in call and "RUN_RECORD_WRITTEN" in call and "args.no_collect" in call
    assert "written=_ssh_run_records" in text


# ------------------------------------------------------------------------------------- P2-b: refused devices ---
def _dsa_only_host_key():
    names = [n for n in sorted(S.SHA1_HOST_KEY_NAMES)
             if n not in S.LEGACY_SHA1_TIER_HOST_KEYS and "cert" not in n and not n.startswith("x509")]
    assert len(names) == 1, names
    return names


def _session_block():
    """Three devices from the REAL producer path: edge1 refused (DSA-only host key: refused_legacy_only, Medium),
    edge2 refused (a 1024-bit group: refused_weak_dh, High), access4 negotiated SHA-1 without opt-in (Medium)."""
    server = {"kex": [_G14_SHA1], "host_key": _dsa_only_host_key(), "cipher_c2s": ["aes128-ctr"],
              "cipher_s2c": ["aes128-ctr"], "mac_c2s": ["hmac-sha2-256"], "mac_s2c": ["hmac-sha2-256"]}
    obs = _obs_with_lists(server, dict(_P5_CLIENT, kex=_P5_CLIENT["kex"] + [_G14_SHA1]))
    refusal = S.classify_failure(EOFError(), obs)
    assert (refusal["category"], refusal["classification"]) == ("host_key", "refused_legacy_only")
    weak = S.classify_failure(_wrapped(_WeakGroupRefused(1024)), S.SessionObservation())
    records = {
        "edge1": S.build_record(outcome="negotiation_refused", consent=_DEFAULT_CONSENT, library=_lib(), attempts=1,
                                observation=obs.snapshot(), refusal=refusal, failure_class="NetmikoTimeoutException"),
        "edge2": S.build_record(outcome="negotiation_refused", consent=_DEFAULT_CONSENT, library=_lib(), attempts=1,
                                refusal=weak, failure_class="NetmikoTimeoutException"),
        "access4": S.build_record(outcome="established", consent=_DEFAULT_CONSENT, library=_lib(), attempts=1,
                                  observation=_observation(_G14_SHA1, 20, 2048, _RSA_SHA1)),
    }
    return S.compute_ssh_sessions(sorted(records), lambda h: (S.render_record(records[h]), None), live=False,
                                  consent=None, evidence_path=lambda h: f"{h}/{S.SIDECAR_FILENAME}")


def test_p2b_a_refused_device_keeps_its_exposed_finding_in_the_dossier():
    """P2-b. Catches: the dossier's not-assessable gate (no config, no version) firing BEFORE the session-evidenced
    surface, so a device refused at collection lost its exposed ssh-legacy-transport finding."""
    from cisco_toolkit import analyze

    block = _session_block()
    sr = analyze.compute_software_risk({}, {}, {}, ["edge1", "edge2"], ssh_sessions=block)
    rows = {r["host"]: {e["axis"]: e for e in r["exposures"]}
            for r in analyze.compute_device_dossiers(software_risk=sr, ssh_sessions=block)["per_device"]}
    assert (rows["edge1"]["Software risk"]["state"], rows["edge2"]["Software risk"]["state"]) == ("watch", "risk")
    for host in ("edge1", "edge2"):
        axis = rows[host]["Software risk"]
        assert axis["input_state"] == "published", axis
        assert S.SURFACE_KIND in axis["label"] and "not screened" in axis["label"]
        assert "PSIRT" not in axis["label"]                  # P3-h: never called an advisory surface to validate
    # non-vacuity: with no session evidence the same device is still not assessable, never clean
    plain = analyze.compute_software_risk({}, {}, {}, ["edge1"])
    axis = {e["axis"]: e for e in analyze.compute_device_dossiers(software_risk=plain)["per_device"][0]["exposures"]}
    assert axis["Software risk"]["state"] == "na"


def test_p2b_the_dossier_names_the_session_surface_beside_a_screened_config():
    """P3-h. Catches: a session-evidenced finding on a config-screened device labelled as an advisory surface to
    validate with the PSIRT checker, and a clean config axis read as a clean SSH transport."""
    from cisco_toolkit import analyze

    block = _session_block()
    configs = {"access4": "hostname access4\nno ip http server\nno ip http secure-server\nno vstack\n",
               "clean1": "hostname clean1\nno ip http server\nno ip http secure-server\nno vstack\n"}
    devices = {h: {"model": "C9300", "sw_version": "17.9.4"} for h in configs}
    sr = analyze.compute_software_risk(configs, devices, {}, sorted(configs), ssh_sessions=block)
    rows = {r["host"]: {e["axis"]: e for e in r["exposures"]}
            for r in analyze.compute_device_dossiers(software_risk=sr, ssh_sessions=block)["per_device"]}
    assert rows["access4"]["Software risk"]["state"] == "watch"
    assert S.SURFACE_KIND in rows["access4"]["Software risk"]["label"]
    assert "PSIRT" not in rows["access4"]["Software risk"]["label"]
    clean = rows["clean1"]["Software risk"]
    assert clean["state"] == "ok" and "SSH transport not recorded" in clean["label"]


def test_p2b_the_sample_pins_a_refused_legacy_only_record_only_device():
    """P2-b. The sample's record-only device classifies refused_legacy_only from paramiko 4.0.0's own client tables
    (DSA removed), so the regenerated sample carries a refused device's exposed finding on every surface."""
    import importlib.util

    spec = importlib.util.spec_from_file_location("build_sample_under_test",
                                                  ROOT / "webapp" / "sample_data" / "build_sample.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    records = module._ssh_session_records()
    row = S.derive_row("edge1", records["edge1"], evidence="edge1/x", live=False)
    assert (row["status"], row["finding"], row["severity"]) == ("refused_legacy_only", "exposed", "Medium")
    assert "no profile of this collector implements" in row["label"]
    statuses = {h: S.derive_row(h, data, evidence="x", live=False)["status"] for h, data in records.items()}
    assert statuses == {"access4": "legacy_sha1", "core1": "modern", "edge1": "refused_legacy_only"}


def _disclosure_snapshot():
    snap = json.loads((ROOT / "tests" / "golden" / "snapshot.json").read_text(encoding="utf-8"))
    block = _session_block()
    snap["ssh_sessions"] = block
    labels = sorted({r["label"] for r in block["rows"] if r["status"] != "modern"})
    assert len(labels) == 3
    return snap, labels


def _docx_text(path):
    from docx import Document

    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def test_p2b_the_collection_integrity_deliverables_carry_the_per_device_disclosure(tmp_path):
    """P2-b / design section 6.3. Catches: the runbook, operations handbook and executive deck silent on what the
    collection sessions negotiated -- every device that is not modern must carry its owned disclosure label."""
    pytest.importorskip("docx")
    pptx = pytest.importorskip("pptx")
    from cisco_toolkit.deck import write_executive_deck_pptx
    from cisco_toolkit.ops import write_ops_handbook_docx
    from cisco_toolkit.runbook import write_runbook_docx

    snap, labels = _disclosure_snapshot()
    runbook = str(tmp_path / "rb.docx")
    write_runbook_docx(runbook, snap, "W59 disclosure")
    ops = str(tmp_path / "ops.docx")
    write_ops_handbook_docx(ops, snap, "W59 disclosure")
    deck = str(tmp_path / "deck.pptx")
    write_executive_deck_pptx(deck, snap, "W59 disclosure")
    deck_text = "\n".join(sh.text_frame.text for sl in pptx.Presentation(deck).slides for sh in sl.shapes
                          if sh.has_text_frame)
    for name, text in (("runbook", _docx_text(runbook)), ("ops", _docx_text(ops)), ("deck", deck_text)):
        for label in labels:
            assert label in text, (name, label)
        for host in ("edge1", "edge2", "access4"):
            assert host in text, (name, host)
    # non-vacuity: an older snapshot without the block adds nothing, and never claims a modern session
    old = json.loads((ROOT / "tests" / "golden" / "snapshot.json").read_text(encoding="utf-8"))
    old.pop("ssh_sessions", None)
    old_rb = str(tmp_path / "old.docx")
    write_runbook_docx(old_rb, old, "W59 disclosure")
    assert "Collection transport (SSH session disclosure)" not in _docx_text(old_rb)


def test_p2b_disclosure_groups_are_total_and_worst_first():
    block = _session_block()
    groups = S.disclosure_groups(block)
    assert [(g["finding"], g["severity"]) for g in groups] == [("exposed", "High"), ("exposed", "Medium"),
                                                             ("exposed", "Medium")]
    assert sorted(h for g in groups for h in g["hosts"]) == ["access4", "edge1", "edge2"]
    for hostile in (None, [], {"rows": "x"}, {"rows": [None, 5, {"host": ["x"]}, {"host": "h", "status": {}}]}):
        S.disclosure_groups(hostile)                         # total: never raises
    assert S.disclosure_groups({"rows": [{"host": "h", "status": "modern"}]}) == []


# ---------------------------------------------------------------------------- P3-c: per-severity punch fold ---
def test_p3c_the_punch_list_folds_per_severity_and_keeps_each_devices_detail():
    """P3-c. Catches: ssh-legacy-transport findings grouped by kind under the FIRST finding's severity and detail
    (a High weak-DH device folded into a Medium row, and one device's disclosure standing in for another's)."""
    from cisco_toolkit import analyze

    block = _session_block()
    sr = analyze.compute_software_risk({}, {}, {}, sorted(r["host"] for r in block["rows"]), ssh_sessions=block)
    items = [i for i in analyze.compute_migration_punchlist([], {}, {}, [], [], [], {}, [], [], software_risk=sr,
                                                            ssh_sessions=block) if i["title"] == S.SURFACE_LABEL]
    by_sev = {i["severity"]: i for i in items}
    assert len(items) == 2 and set(by_sev) == {"High", "Medium"}
    assert by_sev["High"]["devices"] == ["edge2"] and by_sev["Medium"]["devices"] == ["access4", "edge1"]
    medium = by_sev["Medium"]["detail"]
    assert medium.startswith("2 device(s). ") and "[access4]" in medium and "[edge1]" in medium
    index = {r["host"]: i for i, r in enumerate(block["rows"])}
    refs = {r["ref"] for r in by_sev["Medium"]["evidence_refs"]}
    assert {f"/ssh_sessions/rows/{index['access4']}", f"/ssh_sessions/rows/{index['edge1']}"} <= refs
    assert all(analyze.punch_row_session_evidenced(i) for i in items)
    assert not analyze.punch_row_session_evidenced({"evidence_refs": [{"ref": "/software_risk/findings/0"}]})
    # the generic fold: one kind, two severities -> two rows; a detail shared by two devices reads once
    si = {"detections": [
        {"host": "sw1", "kind": "k", "label": "L", "severity": "High", "detail": "a", "recommendation": "r"},
        {"host": "sw2", "kind": "k", "label": "L", "severity": "Medium", "detail": "b", "recommendation": "r"},
        {"host": "sw3", "kind": "k", "label": "L", "severity": "Medium", "detail": "b", "recommendation": "r"}]}
    rows = {i["severity"]: i for i in analyze.compute_migration_punchlist([], {}, {}, [], [], [], {}, [], [],
                                                                         syslog_intelligence=si)}
    assert rows["High"]["devices"] == ["sw1"] and rows["Medium"]["devices"] == ["sw2", "sw3"]
    assert rows["Medium"]["detail"] == "2 device(s). b"


# ------------------------------------------------------------------- P3-d: the failed-phase ratchet reaches ---
def test_p3d_the_consent_block_states_no_negotiation_before_any_session():
    """P3-d. Catches: ``devices_negotiated_sha1: []`` written before any connection -- read as 'no device
    negotiated SHA-1' by any manifest whose disclosure phase then failed."""
    block = S.live_consent_block([{"hostname": "sw1", "ip": "192.0.2.10"}], None)
    assert block["devices_negotiated_sha1"] is None
    rec = S.build_record(outcome="established", consent=_DEFAULT_CONSENT, library=_lib(True, "4.0.0"), attempts=1,
                         observation=_observation(_G14_SHA1, 20, 2048, _RSA_SHA1))
    out = S.compute_ssh_sessions(["sw1"], lambda h: (S.render_record(rec), None), live=True, consent=block,
                                 evidence_path=lambda h: h, run_written=["sw1"])
    assert out["consent"]["devices_negotiated_sha1"] == ["sw1"]
    assert block["devices_negotiated_sha1"] is None                 # the pre-connection copy is not rewritten


def test_p3d_a_failed_session_phase_reaches_every_reader():
    """P3-d. Catches: a failed 'SSH session disclosure' phase leaving software_risk, the dossier's Software risk axis
    and the punch-list rollup published over its fallback."""
    from cisco_toolkit import analyze, ssot
    from cisco_toolkit import ui_projection as uip

    assert "ssh_sessions" in analyze.DOSSIER_AXIS_INPUTS["Software risk"]
    assert "ssh_sessions" in uip.PUNCHLIST_INPUTS
    assert "ssh_sessions" in ssot.fact_basis("software_risk.findings")
    label = next(lab for lab, secs in ssot.PHASE_SECTIONS.items() if "ssh_sessions" in secs)
    software = {"per_device": [{"host": "sw1", "config_assessable": True}], "findings": []}
    snap = {"assessment_integrity": {"failed_phases": [label]}, "ssh_sessions": {}, "software_risk": software}
    assert ssot.abstention_reason(snap, "software_risk") == ssot.ANALYSIS_UNAVAILABLE
    assert ssot.abstention_reason({"software_risk": software}, "software_risk") == "published"   # non-vacuity
    sr = analyze.compute_software_risk({"sw1": "hostname sw1\nno ip http server\n"}, all_hosts=["sw1"],
                                       ssh_sessions={})
    axes = {e["axis"]: e for e in analyze.compute_device_dossiers(
        software_risk=sr, ssh_sessions={}, input_failures=ssot.failed_sections(snap))["per_device"][0]["exposures"]}
    assert (axes["Software risk"]["state"], axes["Software risk"]["input_state"]) == ("na", "analysis_unavailable")


# ------------------------------------------------------------------------- P3-e: no device-controlled text ---
def test_p3e_the_record_stores_no_device_controlled_free_text():
    """P3-e. Catches: a crafted server-offered name or a custom banner (an organisation or host name) stored in a
    record the redaction verifiers then report as covered by schema."""
    obs = _observation("ecdh-sha2-nistp256", 32, None, "rsa-sha2-512")
    obs["server"]["kex"] = ["ecdh-sha2-nistp256", "acme-corp-hq-kex@acme.example"]
    obs["negotiated"]["server_software"] = "SSH-2.0-AcmeCorp_HQ"
    rec = S.build_record(outcome="established", consent=_DEFAULT_CONSENT, library=_lib(), attempts=1, observation=obs)
    assert rec["server_offered"]["kex"] == ["ecdh-sha2-nistp256"] and rec["negotiated"]["server_software"] is None
    assert rec["dropped_names"] == 2 and S.validate_record(rec) == []
    assert "acme" not in S.render_record(rec).decode("ascii").lower()
    for mutate in (lambda r: r["server_offered"]["kex"].append("acme-corp-kex"),
                   lambda r: r["negotiated"].update(server_software="SSH-2.0-AcmeCorp_1.0"),
                   lambda r: r["client_offered"]["mac"].append("acme-mac")):
        r = json.loads(json.dumps(rec))
        mutate(r)
        assert S.validate_record(r), mutate
    assert S.server_software_from_banner("SSH-2.0-AcmeCorp_HQ") == (None, 1)
    assert S.server_software_from_banner("SSH-2.0-OpenSSH_9.6p1 Ubuntu-3") == ("SSH-2.0-OpenSSH_9.6p1", 0)
    assert S.server_software_from_banner("SSH-1.99-Cisco-1.25") == ("SSH-1.99-Cisco-1.25", 0)


def test_p3e_a_refusal_is_graded_on_every_name_but_stores_only_vocabulary_names():
    server = {"kex": ["acme-vendor-kex"], "host_key": ["rsa-sha2-512"], "cipher_c2s": ["aes128-ctr"],
              "cipher_s2c": ["aes128-ctr"], "mac_c2s": ["hmac-sha2-256"], "mac_s2c": ["hmac-sha2-256"]}
    obs = _obs_with_lists(server, _P5_CLIENT)
    refusal = S.classify_failure(EOFError(), obs)
    assert refusal["classification"] == "unclassified"         # an ungraded name is never read as legacy
    rec = S.build_record(outcome="negotiation_refused", consent=_DEFAULT_CONSENT, library=_lib(), attempts=1,
                         observation=obs.snapshot(), refusal=refusal, failure_class="NetmikoTimeoutException")
    assert rec["refusal"]["names"] == [] and rec["server_offered"]["kex"] == [] and rec["dropped_names"] == 2
    assert S.validate_record(rec) == []


# ------------------------------------------------------------------- P3-g: list evidence must be complete ---
class _SSHException(Exception):
    """paramiko's SSHException, by NAME (the classifier never imports paramiko)."""


_SSHException.__name__ = "SSHException"


def test_p3g_list_evidence_needs_a_complete_observation():
    """P3-g. Catches: a refusal classified from lists the observer did not record in full -- an observation error,
    an empty client list (read as 'no common algorithm') or a dropped server name (possibly the common one)."""
    server = {"kex": [_G14_SHA1], "host_key": [_RSA_SHA1], "cipher_c2s": ["aes128-ctr"],
              "cipher_s2c": ["aes128-ctr"], "mac_c2s": [_HMAC_SHA1], "mac_s2c": [_HMAC_SHA1]}
    assert S.classify_failure(EOFError(), _obs_with_lists(server, _P5_CLIENT))["classification"] == \
        "refused_legacy_only"                                          # non-vacuity
    errored = _obs_with_lists(server, _P5_CLIENT)
    errored.errors = 1
    assert S.classify_failure(EOFError(), errored) is None
    modern = dict(server, kex=["ecdh-sha2-nistp256"], host_key=["rsa-sha2-512"])
    assert S.classify_failure(EOFError(), _obs_with_lists(modern, dict(_P5_CLIENT, mac=[]))) is None
    dropped = _obs_with_lists(server, _P5_CLIENT)
    dropped.dropped = 1
    assert S.classify_failure(EOFError(), dropped) is None
    incompatible = _wrapped(_IncompatiblePeer("Incompatible ssh peer (no acceptable kex algorithm)"))
    assert S.classify_failure(incompatible, errored)["classification"] == "unclassified"   # the chain still decides


def test_p3g_paramikos_gex_out_of_range_is_a_deterministic_refusal():
    """P3-g. Catches: paramiko's deterministic group-exchange refusal (the server's prime outside 1024-8192 bits)
    retried as an ordinary failure, and a too-small group not graded weak."""
    small = S.classify_failure(_wrapped(_SSHException(
        "Server-generated gex p (don't ask) is out of range (1024 bits)")), S.SessionObservation())
    assert (small["classification"], small["offered_group_bits"], small["detail"]) == \
        ("refused_weak_dh", 1024, "weak_group_refused")
    big = S.classify_failure(_wrapped(_SSHException(
        "Server-generated gex p (don't ask) is out of range (16384 bits)")), None)
    assert (big["classification"], big["offered_group_bits"]) == ("unclassified", 16384)
    assert S.classify_failure(_wrapped(_SSHException("Negotiation failed.")), None) is None


def test_p3g_a_gex_out_of_range_refusal_is_never_retried(monkeypatch):
    calls = []

    def gex_open(kwargs, platform, profile, recorder):
        calls.append(1)
        raise _wrapped(_SSHException("Server-generated gex p (don't ask) is out of range (1024 bits)"))

    monkeypatch.setattr(C, "_open_connection", gex_open)
    monkeypatch.setattr(C, "CONNECT_BACKOFF_BASE", 0)
    recorder = S.SessionRecorder(None, consent=_DEFAULT_CONSENT, library=_lib())
    dev, _ = C.connect_device("192.0.2.10", "SW1", "u", "p", "ios", session=recorder)
    assert dev is None and calls == [1]
    assert recorder.record["outcome"] == "negotiation_refused"
    row = S.derive_row("SW1", S.render_record(recorder.record), evidence="e", live=True)
    assert (row["status"], row["severity"]) == ("refused_weak_dh", "High")


# ------------------------------------------------------------------------ P3-i: bounded exponential backoff ---
def test_p3i_the_atomic_replace_backoff_is_exponential_bounded_and_counted(tmp_path):
    """P3-i. Catches: four replace attempts 0.1 s apart (shorter than an ordinary on-access scan) and a record
    failure that does not say how hard the writer tried."""
    schedule = S.REPLACE_BACKOFF_S
    assert all(later == 2 * earlier for earlier, later in zip(schedule, schedule[1:]))
    assert 2.0 <= sum(schedule) <= 5.0
    rec = S.build_record(outcome="pending", consent=_DEFAULT_CONSENT, library=_lib())
    path = tmp_path / S.SIDECAR_FILENAME
    sleeps, calls = [], []

    def flaky(src, dst):
        calls.append(1)
        if len(calls) <= 3:
            raise PermissionError("held")
        os.replace(src, dst)

    S.write_record_atomic(str(path), rec, replace=flaky, sleep=sleeps.append)
    assert sleeps == list(schedule[:3]) and json.loads(path.read_bytes())["outcome"] == "pending"
    sleeps.clear()

    def held(src, dst):
        raise PermissionError("held")

    with pytest.raises(PermissionError) as info:
        S.write_record_atomic(str(path), rec, replace=held, sleep=sleeps.append)
    assert getattr(info.value, S.REPLACE_ATTEMPTS_ATTRIBUTE) == len(schedule) + 1
    assert sleeps == list(schedule)
    assert [p.name for p in tmp_path.iterdir()] == [S.SIDECAR_FILENAME]          # no temp file left behind


def test_p3i_the_record_failure_discloses_the_replace_attempts(monkeypatch, tmp_path):
    def held_writer(path, record):
        S.write_record_atomic(path, record, replace=lambda s, d: (_ for _ in ()).throw(PermissionError("held")),
                              sleep=lambda _seconds: None)

    recorder = S.SessionRecorder(str(tmp_path / S.SIDECAR_FILENAME), consent=_DEFAULT_CONSENT, library=_lib(),
                                 writer=held_writer)
    with pytest.raises(PermissionError):
        recorder.write_pending()
    attempts = len(S.REPLACE_BACKOFF_S) + 1
    assert recorder.write_failures == [{"stage": "pending", "error_class": "PermissionError",
                                        "replace_attempts": attempts}]
    original = C._session_recorder_for

    def held_recorder(devinfo, d):
        r = original(devinfo, d)
        r._writer = held_writer
        return r

    monkeypatch.setattr(C, "_session_recorder_for", held_recorder)
    monkeypatch.setattr(C, "_open_connection", lambda *a, **k: pytest.fail("must not connect"))
    failures = []
    C._collect_live_device(_dev(), str(tmp_path / "SW1"), claimed=set(), lock=threading.Lock(),
                           on_record_failure=lambda *a: failures.append(a))
    assert failures == [("SW1", "pending", "PermissionError", attempts)]


# --------------------------------------------------------------------------- P3-j: the port at load time ---
@pytest.mark.parametrize("allow_prompt", [True, False])
def test_p3j_an_invalid_port_fails_at_load_time_before_any_connection(tmp_path, allow_prompt):
    """P3-j. Catches: devices.json `port` validated only at connect time -- a --no-collect run (allow_prompt False)
    accepted a bad port silently, and a live run connected other devices first."""
    row = {"hostname": "SW1", "ip": "192.0.2.10", "username": "u", "password": "p"}
    bad = tmp_path / "bad.json"
    bad.write_text(json.dumps([dict(row, port="22")]), encoding="utf-8")
    with pytest.raises(ValueError, match="'port' must be an integer"):
        C.load_devices(str(bad), allow_prompt=allow_prompt)
    good = tmp_path / "good.json"
    good.write_text(json.dumps([dict(row, port=2222), dict(row, hostname="SW2")]), encoding="utf-8")
    loaded = C.load_devices(str(good), allow_prompt=allow_prompt)
    assert loaded[0]["port"] == 2222 and "port" not in loaded[1]


# =========================================================================== PR-1 review, round 2 ===
class _ProbeDev:
    def send_command(self, cmd, read_timeout=None):
        return "ok\n"


def test_r2_the_autodetect_probe_reaches_the_rows_own_port(monkeypatch):
    """Round 2 (P2). Catches: a devices.json ``port`` honoured by the observed session but not by the ``platform:
    auto`` probe, which then sent the device's credentials over an unobserved session to whatever answers on 22."""
    probes, opened = [], []

    class _Detect:
        def __init__(self, **kwargs):
            probes.append(kwargs)
            self.connection = None

        def autodetect(self):
            return "cisco_ios"

    def fake_open(kwargs, platform, profile, recorder):
        opened.append(kwargs.get("port"))
        return _ProbeDev()

    monkeypatch.setattr(C, "SSHDetect", _Detect)
    monkeypatch.setattr(C, "_open_connection", fake_open)
    _dev_obj, resolved = C.connect_device("192.0.2.10", "SW1", "u", "p", "auto", port=2201)
    assert resolved == "ios" and len(probes) == 1 and probes[0]["port"] == 2201 and opened == [2201]
    probes.clear()
    opened.clear()
    C.connect_device("192.0.2.10", "SW1", "u", "p", "auto")
    assert len(probes) == 1 and "port" not in probes[0] and opened == [None]   # no row port: netmiko's default
    probes.clear()
    C.connect_device("192.0.2.10", "SW1", "u", "p", "ios", port=2201)
    assert probes == []                                   # an explicit platform never probes


def test_r2_a_stale_record_removal_is_retried_over_the_bounded_backoff(tmp_path):
    """Round 2 (P3). Catches: one ``os.unlink`` attempt on the record an earlier run left, so a scanner's transient
    handle left the device uncollected for the whole run."""
    path = tmp_path / S.SIDECAR_FILENAME
    assert C._unlink_stale_session_record(str(path)) == 0                # nothing to remove: no attempt
    path.write_bytes(b"{}")
    sleeps, calls = [], []

    def flaky(p):
        calls.append(p)
        if len(calls) <= 3:
            raise PermissionError("held by an on-access scan")
        os.unlink(p)

    assert C._unlink_stale_session_record(str(path), unlink=flaky, sleep=sleeps.append) == 4
    assert not path.exists() and sleeps == list(S.REPLACE_BACKOFF_S[:3])
    path.write_bytes(b"{}")
    sleeps.clear()

    def held(p):
        raise PermissionError("held")

    with pytest.raises(PermissionError) as info:
        C._unlink_stale_session_record(str(path), unlink=held, sleep=sleeps.append)
    assert getattr(info.value, C.UNLINK_ATTEMPTS_ATTRIBUTE) == len(S.REPLACE_BACKOFF_S) + 1
    assert sleeps == list(S.REPLACE_BACKOFF_S) and path.exists()


def test_r2_a_failed_removal_discloses_its_attempts_and_never_connects(monkeypatch, tmp_path):
    original = C._unlink_stale_session_record

    def held_remover(p):
        return original(p, unlink=lambda q: (_ for _ in ()).throw(PermissionError("held")), sleep=lambda _s: None)

    monkeypatch.setattr(C, "_unlink_stale_session_record", held_remover)
    monkeypatch.setattr(C, "_open_connection", lambda *a, **k: pytest.fail("must not connect"))
    failures = []
    _platform, cmd = C._collect_live_device(_dev(), str(_stale_modern_folder(tmp_path)), claimed=set(),
                                            lock=threading.Lock(), on_record_failure=lambda *a: failures.append(a))
    assert cmd is None
    assert failures == [("SW1", S.RECORD_STAGE_STALE_UNLINK, "PermissionError", 0, len(S.REPLACE_BACKOFF_S) + 1)]


def test_r2_every_manifest_record_failure_has_one_fixed_shape():
    """The run manifest's ``record_failures`` entries (main()'s ``_ssh_record_failure``) always carry both attempt
    counts, so a reader never has to guess which operation a count belongs to."""
    tree = _engine_tree()
    fn = next(n for n in ast.walk(tree) if isinstance(n, ast.FunctionDef) and n.name == "_ssh_record_failure")
    dicts = [n for n in ast.walk(fn) if isinstance(n, ast.Dict)]
    assert len(dicts) == 1
    assert {k.value for k in dicts[0].keys} == {"host", "stage", "error_class", "replace_attempts", "unlink_attempts"}


def _eligible_rows(failures, *, live=True):
    block = (dict(S.live_consent_block([], None), devices_eligible=["sw1"], record_failures=failures)
             if live else S.offline_consent_block())
    return S.compute_ssh_sessions(["sw1"], lambda h: (None, None), live=live, consent=block,
                                  evidence_path=lambda h: h, run_written=[] if live else None)["rows"][0]


def test_r2_a_pre_connect_record_failure_is_not_recorded_live_and_offline_alike():
    """Design re-check P3 (section 6.1 step 1 against section 6.2). Catches: a device left UNCONNECTED by a failed
    pre-connect record write derived ``legacy_unrecorded`` (exposed, Medium) on the live run -- nothing was negotiated
    -- while an offline re-analysis of the same folder says ``not_recorded``."""
    for stage in sorted(S.PRE_CONNECT_RECORD_STAGES):
        row = _eligible_rows([{"host": "sw1", "stage": stage, "error_class": "PermissionError",
                               "replace_attempts": 0, "unlink_attempts": 7}])
        assert (row["status"], row["finding"], row["reason"]) == ("not_recorded", "verify", S.REASON_NOT_CONNECTED)
    assert _eligible_rows(None, live=False)["status"] == "not_recorded"
    # non-vacuity: an authorized device with no record and no pre-connect failure stays exposed, and neither a
    # post-connect failure nor another host's failure changes that
    assert _eligible_rows([])["status"] == "legacy_unrecorded"
    assert _eligible_rows([{"host": "sw1", "stage": "established", "error_class": "OSError",
                            "replace_attempts": 7, "unlink_attempts": 0}])["status"] == "legacy_unrecorded"
    assert _eligible_rows([{"host": "sw2", "stage": S.RECORD_STAGE_PENDING, "error_class": "OSError",
                            "replace_attempts": 7, "unlink_attempts": 0}])["status"] == "legacy_unrecorded"
    assert S.PRE_CONNECT_RECORD_STAGES == {S.RECORD_STAGE_PENDING, S.RECORD_STAGE_STALE_UNLINK}
    text = ENGINE.read_text(encoding="utf-8")
    assert "stage = ssh_session.RECORD_STAGE_STALE_UNLINK" in text and "stage = ssh_session.RECORD_STAGE_PENDING" in text


@pytest.mark.parametrize("server_change, opt_in_collects", [
    ({}, True),                                                          # SHA-1 kex and host key: the tier closes both
    ({"host_key": ["rsa-sha2-512"]}, True),                              # only the kex is refused
    ({"host_key": ["ssh-dss"]}, False),                                  # SHA-1 kex, DSA-only host key
    ({"cipher_c2s": ["blowfish-cbc"], "cipher_s2c": ["blowfish-cbc"]}, False),   # no common cipher either
])
def test_r2_the_opt_in_advice_needs_the_tier_to_close_every_empty_category(server_change, opt_in_collects):
    """Design re-check P3 (section 4.4). Catches: "requires the legacy-sha1 opt-in" for a device the opt-in still
    cannot collect -- the classifier names the FIRST refused category, but the advice must hold for every category with
    no common algorithm. The collector's log line (full sink lists) and the row's label (the sealed record's lists)
    agree."""
    server = dict({"kex": [_G14_SHA1, "ext-info-s"], "host_key": [_RSA_SHA1], "cipher_c2s": ["aes128-ctr"],
                   "cipher_s2c": ["aes128-ctr"], "mac_c2s": ["hmac-sha2-256"], "mac_s2c": ["hmac-sha2-256"]},
                  **server_change)
    obs = _obs_with_lists(server, _P5_CLIENT)
    refusal = S.classify_failure(EOFError(), obs)
    assert (refusal["category"], refusal["classification"]) == ("kex", "refused_legacy_only")
    assert S.legacy_tier_closes(server, _P5_CLIENT) is opt_in_collects
    rec = S.build_record(outcome="negotiation_refused", consent=_DEFAULT_CONSENT, library=_lib(), attempts=1,
                         observation=obs.snapshot(), refusal=refusal, failure_class="NetmikoTimeoutException")
    label = S.derive_row("sw1", S.render_record(rec), evidence="e", live=False)["label"]
    message = S.refusal_message(refusal, _DEFAULT_CONSENT, obs.snapshot())
    for text in (label, message):
        assert ("requires the legacy-sha1 opt-in" in text) is opt_in_collects, text
        assert ("no profile of this collector implements" in text) is not opt_in_collects, text


def test_r2_a_session_only_high_finding_on_a_high_impact_device_carries_no_psirt_step():
    """Round 2 (P3). Catches: the session-evidenced ssh-legacy-transport finding counted and worded as a configuration
    advisory surface to validate with the PSIRT checker -- compound pattern CR-04 on a high-impact device and the
    executive brief's Software risk axis."""
    from cisco_toolkit import analyze

    block = _session_block()             # edge2: a 1024-bit group, refused (High); edge1 / access4: SHA-1 (Medium)
    hosts = sorted(r["host"] for r in block["rows"])
    sr = analyze.compute_software_risk({}, {}, {}, hosts, ssh_sessions=block)
    impact = [{"host": "edge2", "severity": "High", "stranded": 0, "vlans_impacted": 3}]
    rows = {r["host"]: r for r in analyze.compute_device_dossiers(
        software_risk=sr, failure_impact=impact, ssh_sessions=block)["per_device"]}
    cr04 = [c for c in rows["edge2"]["compound"] if c["code"] == "CR-04"]
    assert len(cr04) == 1 and cr04[0]["title"] == "Legacy SSH transport on a high-impact asset", cr04
    assert S.SURFACE_KIND in cr04[0]["basis"] and "session-evidenced" in cr04[0]["basis"]
    # the session-only basis disclaims the PSIRT step in words; it never carries the configuration trigger's
    # PSIRT-checker instruction (first hosted run: a bare "PSIRT" substring test matched the disclaimer itself)
    assert "PSIRT Software Checker" not in cr04[0]["basis"]
    assert cr04[0]["basis"].endswith("no PSIRT step applies.") and cr04[0]["basis"].count("PSIRT") == 1
    axis = next(a for a in analyze.compute_executive_brief(software_risk=sr)["axes"] if a["axis"] == "Software risk")
    assert axis["severity"] == "High" and f"{len(sr['findings'])} {S.SURFACE_COUNT_NOUN}" in axis["headline"]
    assert "configuration surfaces not assessable" in axis["headline"] and S.SURFACE_NO_PSIRT in axis["detail"]
    # a configuration trigger on the same device keeps its PSIRT wording and names the session finding beside it
    config = {"edge2": "hostname edge2\nip http server\n"}
    sr_cfg = analyze.compute_software_risk(config, {"edge2": {"model": "C9300", "sw_version": "17.9.4"}}, {},
                                           ["edge2"], ssh_sessions=block)
    cfg_cr = [c for r in analyze.compute_device_dossiers(software_risk=sr_cfg, failure_impact=impact,
                                                         ssh_sessions=block)["per_device"]
              for c in r["compound"] if c["code"] == "CR-04"]
    assert len(cfg_cr) == 1 and cfg_cr[0]["title"] == "Open advisory surface on a high-impact asset"
    assert "PSIRT Software Checker" in cfg_cr[0]["basis"] and S.SURFACE_KIND in cfg_cr[0]["basis"]
    axis = next(a for a in analyze.compute_executive_brief(software_risk=sr_cfg)["axes"]
                if a["axis"] == "Software risk")
    n_session = sum(1 for f in sr_cfg["findings"] if f["kind"] == S.SURFACE_KIND)
    assert n_session == len(hosts)                         # every host the session block names keeps its finding
    assert axis["headline"].startswith(f"1 exposed advisory surface(s) · {n_session} {S.SURFACE_COUNT_NOUN}"), axis
    # non-vacuity: without a session block both read exactly as before W59
    sr_old = analyze.compute_software_risk(config, {"edge2": {"model": "C9300", "sw_version": "17.9.4"}}, {},
                                           ["edge2"])
    old_cr = [c for r in analyze.compute_device_dossiers(software_risk=sr_old, failure_impact=impact)["per_device"]
              for c in r["compound"] if c["code"] == "CR-04"]
    assert old_cr[0]["basis"].endswith("validate with the Cisco PSIRT Software Checker before the window.")
    axis = next(a for a in analyze.compute_executive_brief(software_risk=sr_old)["axes"]
                if a["axis"] == "Software risk")
    assert S.SURFACE_KIND not in axis["headline"] and axis["detail"] == (
        "Screening, not a scan — validate releases with the Cisco PSIRT Software Checker.")


def test_r2_the_runbook_and_handbook_count_the_session_findings_apart(tmp_path):
    """Round 2 (P3). Catches: runbook section 6.12 and the operations handbook (known issues, section 5) counting
    the session-evidenced findings as exposed advisory surfaces with a PSIRT step."""
    pytest.importorskip("docx")
    from cisco_toolkit import analyze
    from cisco_toolkit.ops import write_ops_handbook_docx
    from cisco_toolkit.runbook import write_runbook_docx

    snap, _labels = _disclosure_snapshot()
    block = snap["ssh_sessions"]
    session = analyze.compute_software_risk({}, {}, {}, sorted(r["host"] for r in block["rows"]),
                                            ssh_sessions=block)["findings"]
    config = list(snap["software_risk"]["findings"])
    assert session and config and all(f["kind"] != S.SURFACE_KIND for f in config)
    snap["software_risk"]["findings"] = config + session
    snap["software_risk"]["summary"]["n_findings"] = len(config) + len(session)
    runbook = str(tmp_path / "rb.docx")
    write_runbook_docx(runbook, snap, "W59 round 2")
    ops = str(tmp_path / "ops.docx")
    write_ops_handbook_docx(ops, snap, "W59 round 2")
    rb_text, ops_text = _docx_text(runbook), _docx_text(ops)
    noun = f"{len(session)} {S.SURFACE_COUNT_NOUN}"
    assert f"Separately, {noun}" in rb_text and S.SURFACE_NO_PSIRT in rb_text
    assert f"{len(config)} exposed advisory / hardening surface(s) open at assessment" in ops_text
    assert "Software Risk (SSH transport)" in ops_text and S.SURFACE_NO_PSIRT in ops_text
    assert f"{len(config)} exposed advisory surface(s) were open" in ops_text and f"Separately, {noun}" in ops_text
    assert f"{len(config) + len(session)} exposed advisory" not in rb_text + ops_text

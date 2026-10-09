"""W27 (follow-up F2): every AssessHub failure-impact surface honours the W23 projection holds.

``analyze.compute_failure_impact`` writes Info, zero counts and a clean-bill detail for a switch it could not
simulate, and a partial simulation or an uncollected neighbour makes a low band or a zero only a lower bound. The
engine-owned projection (``ui_projection``'s one shared failure-impact row builder) withholds exactly those cells,
each with its state and reason. These tests hold the AssessHub surfaces that used to read the raw stored rows to it:

* the dashboard keystones (``summary._keystones`` -> ``summarize``) and the cutover plan's per-wave keystone tags;
* the cutover plan's worst-case blast radius (``cutover._worst_blast_radius``), on screen and in the .docx;
* the snapshot "Failure impact" tab (``GET /section/failure_impact``);
* the /graph keystone badge and a summary cached by the old raw-row ranking.

A held, duplicated, INDETERMINATE or uncollected-neighbour row never yields a published "no impact", a zero or an
Info-ranked worst; a clean, fully assessed row renders exactly as before. Rows come from the REAL producers
(``analyze.compute_failure_impact`` and ``analyze.compute_cable_map`` over constructed interfaces) and from the
shipped sample snapshot. The dossier recompute is deliberately left on the stored rows; its pin is the last test.

The refutation follow-up adds the PUBLISHED lower bound. On a row simulated only in part (a positive
off_scan_gw_vlans) or cabled to an uncollected neighbour, the projection still publishes High and each positive
count, but as lower bounds, citing a witness ref on each such measure. Every surface flags them ("at least N",
``lower_bound``, ``≥ N`` in the tab), a wave whose ranked switches include one is never ``complete``, and an exact row
is never flagged. It also pins that a stored ``executive_brief.keystones`` list is never read, that a section that is
not a list shows the projection's own disclosure, and that a row with an empty host is disclosed like one with none.

W45 (the W36 + W33 + W32 integration train, after this file merged): the projection's row rules now live in the engine
owner ``cisco_toolkit/impact_assessability.py`` (W33), which also bounds a row by W32's per-row ``blind_links`` count
(inter-switch links with no trunk/STP evidence) or by its absence (a row older than that count). So the sample rows
each surface flags are derived from each stored row's own evidence (:func:`_sample_bounds`), never from a host list:
core2's uncollected router as before, plus every row whose stored count is positive or absent.
"""
from __future__ import annotations

import copy
import json
import sys
from dataclasses import asdict
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))  # make `backend` importable

from backend import cutover, summary  # noqa: E402  (also bootstraps sys.path for cisco_toolkit)
from backend.app import create_app  # noqa: E402
from cisco_toolkit import analyze  # noqa: E402
from cisco_toolkit import impact_assessability as ia  # noqa: E402
from cisco_toolkit import ui_projection as ui  # noqa: E402
from cisco_toolkit.model import InterfaceData  # noqa: E402

SAMPLE = Path(__file__).resolve().parents[1] / "sample_data" / "sample_fleet.snapshot.json"
#: The failure-impact cells that measure a simulated blast radius (host, off_scan_gw_vlans and detail do not).
MEASURES = ("severity", "vlans_impacted", "stranded", "hard", "backup", "fhrp")
#: Two blind spots, neither of them a sample switch with a row: one never reached, one reached in part.
BLIND = [{"host": "ghost1", "status": "not collected", "data_quality": 0, "missing": ["version/inventory"]},
         {"host": "access2", "status": "partial", "data_quality": 75, "missing": ["interface status"]}]
CLEAN_BILL = "No reachability impact from removing this switch (within the scan)."


@pytest.fixture(scope="module")
def sample():
    return json.loads(SAMPLE.read_text(encoding="utf-8"))


@pytest.fixture()
def client(tmp_path):
    app = create_app(db_path=str(tmp_path / "test.db"))
    # base_url=localhost so the default Host passes the no-token DNS-rebinding guard (app.py
    # _request_host_allowed) -- the dev server this emulates is reached over loopback.
    with TestClient(app, base_url="http://localhost") as c:
        yield c


def _upload(client, snap) -> int:
    cid = client.post("/api/campaigns", json={"name": "impact"}).json()["id"]
    r = client.post(f"/api/campaigns/{cid}/snapshots",
                    files={"file": ("s.json", json.dumps(snap).encode(), "application/json")},
                    data={"label": "s"})
    assert r.status_code == 201, r.text
    return r.json()["id"]


# --------------------------------------------------------------------------------------------------
# REAL producer fixtures (the W23 shapes): every interface carries run_config_observed, as build.py marks each
# interface its scoped interface running-config capture parsed -- the only source of an SVI's gateway address.
# --------------------------------------------------------------------------------------------------
def _trunk(port, peer, peer_port, vlans="", blocked=""):
    return InterfaceData(port=port, status="connected", switchport_mode="Trunk", cdp_neighbor=peer,
                         neighbor_port=peer_port, endpoint_type="Switch", trunk_allowed_vlans=vlans,
                         stp_fwd_vlans="" if blocked else vlans, stp_blk_vlans=blocked, run_config_observed=True)


def _access(port, vid, mac):
    return InterfaceData(port=port, status="connected", switchport_mode="Access", vlan=str(vid), end_host_mac=mac,
                         run_config_observed=True)


def _svi(vid, address, fhrp=""):
    return InterfaceData(port=f"Vlan{vid}", svi_ip=address, hsrp_behavior=fhrp, run_config_observed=True)


def _edge(port, peer, endpoint_type, platform=""):
    """A CDP-seen trunk to a peer the collection never reached."""
    return InterfaceData(port=port, status="connected", switchport_mode="Trunk", cdp_neighbor=peer,
                         neighbor_port="Gi0/1", endpoint_type=endpoint_type, neighbor_platform=platform,
                         trunk_allowed_vlans="10", stp_fwd_vlans="10", run_config_observed=True)


def _unsimulatable_fleet():
    """`acc` carries VLAN 20, whose gateway was never scanned, and nothing it can simulate; `gw` gateways VLAN 10 for
    `acc` (a measured hard partition) and also carries VLAN 20; `x1` and `x2` share a trunk with no VLAN-carriage
    evidence on either end. So acc, x1 and x2 get the producer's INDETERMINATE Info rows with zero counts."""
    return {"gw": {"Gi1": _trunk("Gi1", "acc", "Gi1", "10,20"), "Vlan10": _svi(10, "10.10.0.1/24")},
            "acc": {"Gi1": _trunk("Gi1", "gw", "Gi1", "10,20"), "Gi10": _access("Gi10", 10, "0000.0000.000a"),
                    "Gi20": _access("Gi20", 20, "0000.0000.0014")},
            "x1": {"Gi1": _trunk("Gi1", "x2", "Gi1")}, "x2": {"Gi1": _trunk("Gi1", "x1", "Gi1")}}


def _downstream_fleet():
    """`gw` gateways VLAN 10 for `acc` (removing it is a measured hard partition, High) and trunks up to the router
    `wan`; `acc` trunks down to the switch `dsw`. The collection reached neither, so the producer writes its clean
    bill for `acc` whatever hangs behind `dsw`."""
    return {"gw": {"Gi1": _trunk("Gi1", "acc", "Gi1", "10"), "Gi47": _edge("Gi47", "wan", "Router", "cisco ISR4331/K9"),
                   "Vlan10": _svi(10, "10.10.0.1/24")},
            "acc": {"Gi1": _trunk("Gi1", "gw", "Gi1", "10"), "Gi10": _access("Gi10", 10, "0000.0000.000a"),
                    "Gi48": _edge("Gi48", "dsw", "Switch")}}


def _snapshot(interfaces, waves=()):
    """The REAL producers' rows over `interfaces`, as a stored snapshot carries them, with one hard-cutover wave per
    ``(group, hosts)``. JSON round-tripped: what the store hands every read route."""
    snap = {"schema": "collect_parse_snapshot/1", "devices": {host: {"hostname": host} for host in interfaces},
            "interfaces": {host: {port: asdict(row) for port, row in ports.items()}
                           for host, ports in interfaces.items()},
            "failure_impact": analyze.compute_failure_impact(interfaces),
            "cable_map": analyze.compute_cable_map(interfaces), "routes": {}}
    if waves:
        snap["wave_sequencing"] = [{"group": group, "make_before_break": [], "hard_cutover": list(hosts)}
                                   for group, hosts in waves]
    return json.loads(json.dumps(snap))


def _by_host(snap):
    return {row["host"]: row for row in snap["failure_impact"]}


def _node(snap, host):
    same = [n for n in snap["cable_map"]["nodes"] if n["host"] == host]
    assert len(same) == 1, (host, same)
    return same[0]


def _cable(snap, a, b):
    """The pointer of the one stored cable_map.cables row joining `a` and `b` (either orientation)."""
    hits = [j for j, c in enumerate(snap["cable_map"]["cables"]) if {c["a"], c["b"]} == {a, b}]
    assert len(hits) == 1, (a, b, hits)
    return f"/cable_map/cables/{hits[0]}"


def _row_pointer(snap, host):
    hits = [i for i, row in enumerate(snap["failure_impact"]) if row["host"] == host]
    assert len(hits) == 1, (host, hits)
    return f"/failure_impact/{hits[0]}"


def _old_entry(row):
    """The pre-W27 keystone / blast-radius shape of a stored row, for a row the engine publishes in full."""
    return {field: row[field] for field in ("host", "severity", "stranded", "vlans_impacted", "detail")}


def _exact(row):
    """A fully assessed row's keystone / blast-radius entry: the pre-W27 values, flagged as no lower bound."""
    return {**_old_entry(row), "lower_bound": False}


def _bound_entry(row, pointers, reasons):
    """A ranked row the engine publishes only as a lower bound: its measured values kept, flagged with why, and its
    detail opened by the "at least" disclosure."""
    return {**_old_entry(row), "lower_bound": True, "lower_bound_reasons": reasons, "lower_bound_pointers": pointers,
            "detail": f"LOWER BOUND, at least {row['stranded']} endpoint(s) stranded: {'; '.join(reasons)}. "
                      f"{row['detail']}"}


#: The sample's one uncollected neighbour that is not edge gear (W23 pins it as the sample's only uncollected-neighbour
#: bound: tests/test_ui_projection_device_impact.py::test_a_simulated_device_selects_exactly_its_fleet_rows).
WAN = "wan-edge-rtr1.lab"


def _sample_index(sample, host):
    hits = [i for i, row in enumerate(sample["failure_impact"]) if row["host"] == host]
    assert len(hits) == 1, (host, hits)
    return hits[0]


def _sample_bounds(sample, i):
    """Independent of the projection and the owner: how stored sample row `i`'s own evidence bounds it, as
    ``(pointers, reasons)`` in the projection's cite order. First its blind-link count (W32): a positive count cites
    the count, and a row without the field predates it and cites the row itself. Then core2's cable to the
    uncollected router. ``([], [])``: the row is a measurement."""
    row = sample["failure_impact"][i]
    pointer = f"/failure_impact/{i}"
    pointers, reasons = [], []
    if "blind_links" not in row:
        pointers.append(pointer)
        reasons.append(summary._R_BOUND_BLIND_LEGACY)
    elif row["blind_links"]:
        pointers.append(pointer + "/blind_links")
        reasons.append(summary._R_BOUND_BLIND)
    if row["host"] == "core2":
        pointers.append(_cable(sample, "core2", WAN))
        reasons.append(summary._R_BOUND_PEERS.format(k=1))
    return pointers, reasons


def _sample_bounded(sample, host):
    return bool(_sample_bounds(sample, _sample_index(sample, host))[0])


def _sample_ranks(sample, i):
    """Whether the projection publishes stored sample row `i`'s host, severity and stranded count: always for a
    measurement; on a bounded row only a High band and a positive count (a bound withholds a lower band and a 0)."""
    row = sample["failure_impact"][i]
    return not _sample_bounds(sample, i)[0] or (row["severity"] == "High" and row["stranded"] != 0)


def _sample_ranked(sample):
    """The stored sample rows a ranking may place, in stored order."""
    return [row for i, row in enumerate(sample["failure_impact"]) if _sample_ranks(sample, i)]


def _sample_unranked(sample):
    """The hosts of the stored sample rows no ranking places (always disclosed instead), in stored order."""
    return [row["host"] for i, row in enumerate(sample["failure_impact"]) if not _sample_ranks(sample, i)]


def _sample_entry(sample):
    """The entry each ranked sample row ranks as: a flagged lower bound when its own evidence bounds it (core2's
    uncollected router, a positive or absent blind_links count), else exact."""
    def entry(row):
        pointers, reasons = _sample_bounds(sample, _sample_index(sample, row["host"]))
        return _bound_entry(row, pointers, reasons) if pointers else _exact(row)
    return entry


def _old_keystones(rows, top=8, entry=_exact):
    """The pre-W27 ranking (severity, then stranded; stable over the stored order) of fully published rows."""
    ranked = sorted(rows, key=lambda r: (summary._SEV_RANK.get(r["severity"], 99), -r["stranded"]))
    return [entry(row) for row in ranked[:top]]


def _old_worst(rows, switches):
    """The pre-W27 worst case of a wave over fully published rows (first minimal in the stored order)."""
    return min((r for r in rows if r["host"] in switches),
               key=lambda r: (cutover._SEV_RANK.get(r["severity"], 99), -r["stranded"]))


def _held_core1(sample, mode):
    """The sample with core1's row withheld by the projection, the mode naming the hold."""
    snap = copy.deepcopy(sample)
    rows = snap["failure_impact"]
    k = next(i for i, row in enumerate(rows) if row["host"] == "core1")
    assert rows[k]["severity"] == "High" and rows[k]["stranded"] > 0       # the sample's top keystone, as stored
    if mode == "no_run_config":
        for port in snap["interfaces"]["core1"].values():
            port.pop("run_config_observed", None)            # what html.sparsify_interfaces writes for false
        return snap, "run_config_observed"
    if mode == "legacy":
        del rows[k]["off_scan_gw_vlans"]
        return snap, "predates the producer's assessability marker"
    rows.append(dict(rows[k], severity="Info", vlans_impacted=0, stranded=0, hard=0, backup=0, fhrp=0,
                     detail=CLEAN_BILL))                                   # a contradicting second row for core1
    return snap, "name this exact host"


# --------------------------------------------------------------------------------------------------
# the hand-written constants are the projection's own
# --------------------------------------------------------------------------------------------------
def test_the_impact_fields_and_tokens_are_the_projections_own(sample):
    item = ui.project_topology(sample)["failure_impact"]["items"][0]
    assert [k for k in item if k not in ("index", "pointer", "node_refs", "style")] == list(summary.IMPACT_FIELDS)
    assert summary._PUBLISHED == ui._PUB and summary._COLLECTED_BUT_EMPTY == ui._CBE
    assert summary._UNVERIFIED == ui._UV
    assert summary._IMPACT_BLIND_CAVEAT in {lim["id"] for lim in ui.LIMITATIONS}
    assert summary.IMPACT_NOT_ASSESSED == cutover.IMPACT_NOT_ASSESSED == cutover.GATE_NOT_ASSESSED
    # the lower-bound signal: the projection's witness role on its measure cells
    assert summary.IMPACT_MEASURES == ui._IMPACT_MEASURES == MEASURES
    assert summary._WITNESS_ROLE in ui.REF_ROLES
    # a cable-row bound's kind is read from the projection's own words, in each reason its bound writes: the
    # neighbours it counts and how many fail closed (unreadable or ambiguous cable evidence). Since W33 those words
    # are the engine owner's (impact_assessability.R_PEERS / R_PEERS_CLOSED / make_bound), which the projection carries.
    for n, k in ((1, 0), (2, 1), (3, 3)):
        clause = ia.R_PEERS.format(n=n, closed=ia.R_PEERS_CLOSED.format(k=k) if k else "")
        for reason in ia.make_bound(ia.NOT_COLLECTED, clause, [], "uncollected_neighbours", n)[1:4]:
            assert summary._impact_peers_said({"fhrp": (False, None, reason)}) == (n, k), reason
    assert summary._impact_peers_said({"fhrp": (True, 0, "")}) is None
    # the blind-link bound (W32, through the owner) is worded by its witness, which is exactly what the owner cites: the
    # row's blind_links count, or the row itself when it predates that count
    assert ia.blind_bound({"blind_links": 2}, ("failure_impact", 7)).witnesses == [
        ("witness", ("failure_impact", 7, "blind_links"))]
    assert ia.blind_bound({}, ("failure_impact", 7)).witnesses == [("witness", ("failure_impact", 7))]
    assert ia.blind_bound({"blind_links": 0}, ("failure_impact", 7)) is None


# --------------------------------------------------------------------------------------------------
# keystones: the dashboard panel, the cutover plan's per-wave tags
# --------------------------------------------------------------------------------------------------
def test_keystones_never_rank_a_row_the_engine_could_not_simulate():
    snap = _snapshot(_unsimulatable_fleet())
    raw = _by_host(snap)
    held = ("acc", "x1", "x2")
    for host in held:                       # the producer's own rows: Info and zeros that look like measurements
        assert raw[host]["severity"] == "Info" and raw[host]["stranded"] == 0, raw[host]
        assert raw[host]["detail"].startswith(ui.IMPACT_INDETERMINATE_PREFIX), raw[host]
    gw = raw["gw"]
    assert gw["severity"] == "High" and gw["stranded"] > 0 and gw["off_scan_gw_vlans"] == 1, gw
    keystones = summary.summarize(snap)["keystones"]
    # before W27 the panel listed all four rows, three of them as Info with 0 stranded. High and gw's positive counts
    # stay published, as lower bounds that cite its off-scan count, so gw ranks and is flagged "at least"
    off_scan = _row_pointer(snap, "gw") + "/off_scan_gw_vlans"
    assert keystones[:-1] == [_bound_entry(gw, [off_scan], [summary._R_BOUND_OFF_SCAN.format(n=1)])]
    note = keystones[-1]
    assert note["host"] == "" and note["severity"] == summary.IMPACT_NOT_ASSESSED
    assert note["stranded"] is None and note["vlans_impacted"] is None and note["n_not_ranked"] == 3
    for host in held:
        assert f"{host}, " in note["detail"] or f"{host} — " in note["detail"], (host, note["detail"])
    assert "could not simulate" in note["detail"]
    assert not any(k["severity"] == "Info" or k["stranded"] == 0 for k in keystones), keystones
    assert cutover._keystone_hosts(snap) == {"gw"}      # the disclosure names no host, so it tags no wave


def test_keystones_never_rank_a_clean_bill_an_uncollected_neighbour_bounds():
    snap = _snapshot(_downstream_fleet())
    raw = _by_host(snap)
    acc, gw = raw["acc"], raw["gw"]
    assert acc["severity"] == "Info" and acc["stranded"] == 0 and acc["off_scan_gw_vlans"] == 0, acc
    assert acc["detail"] == CLEAN_BILL
    assert gw["severity"] == "High" and gw["stranded"] == 1, gw
    dsw, wan = _node(snap, "dsw"), _node(snap, "wan")
    assert dsw["collected"] is False and dsw["kind"] == "switch", dsw        # the REAL cable-map producer's nodes
    assert wan["collected"] is False and wan["kind"] not in ia.IMPACT_EDGE_KINDS, wan
    # gw is cabled to the uncollected router wan: High and its positive count are published lower bounds
    gw_bound = _bound_entry(gw, [_cable(snap, "gw", "wan")], [summary._R_BOUND_PEERS.format(k=1)])
    keystones = summary.summarize(snap)["keystones"]
    assert keystones[:-1] == [gw_bound]
    note = keystones[-1]
    assert note["severity"] == summary.IMPACT_NOT_ASSESSED and note["n_not_ranked"] == 1
    assert "acc — " in note["detail"] and "uncollected neighbour" in note["detail"], note["detail"]
    assert CLEAN_BILL not in note["detail"]
    # the control: the same stored row once the cable map shows that peer collected is a published clean bill,
    # ranked and rendered exactly as before, with no disclosure; gw stays a flagged lower bound behind wan
    dsw["collected"] = True
    assert summary.summarize(snap)["keystones"] == [gw_bound, _exact(acc)]
    # and once wan reads collected too, nothing bounds gw: both rows are exact and render as before
    wan["collected"] = True
    assert summary.summarize(snap)["keystones"] == [_exact(gw), _exact(acc)]


def test_an_unreadable_cable_row_is_a_bound_worded_as_unreadable_never_as_an_uncollected_neighbour():
    """Codex's P3 on #620 (`d48d558a`): ui_projection._impact_peers also cites a cable row it cannot read, because
    that row could name the switch. The bound and its pointer stay, but its reason is the projection's own: cable
    evidence that cannot be read or is ambiguous, never "cables this switch to a neighbour it does not show as
    collected". The real producer's rows with both known peers collected, plus one appended null cable row."""
    snap = _snapshot(_downstream_fleet())
    _node(snap, "dsw")["collected"] = True
    _node(snap, "wan")["collected"] = True
    raw = _by_host(snap)
    gw, acc = raw["gw"], raw["acc"]
    assert gw["severity"] == "High" and gw["stranded"] == 1 and gw["backup"] == 0 and gw["fhrp"] == 0, gw
    # the control: with both known peers collected nothing bounds either row
    assert summary.summarize(snap)["keystones"] == [_exact(gw), _exact(acc)]
    snap["cable_map"]["cables"].append(None)
    unread = f"/cable_map/cables/{len(snap['cable_map']['cables']) - 1}"
    k = int(_row_pointer(snap, "gw").rsplit("/", 1)[1])
    # the projection's own signal: gw's published measures cite the null row, and its withheld zeros say it fails
    # closed
    item = ui.project_topology(snap)["failure_impact"]["items"][k]
    for field in ("severity", "vlans_impacted", "stranded", "hard"):
        cited = [ref["pointer"] for ref in item[field]["refs"] if ref["role"] == "witness"]
        assert item[field]["state"] == ui._PUB and cited == [unread], (field, item[field])
    for field in ("backup", "fhrp"):
        assert item[field]["state"] == ui._NC and "1 of them fail closed" in item[field]["reason"], item[field]
    why = summary._R_BOUND_PEERS_UNREAD.format(k=1)
    rows = {row["key"]: row for row in summary.impact_view(snap)["rows"]}
    assert rows["gw"]["ranked"] is True and rows["gw"]["lower_bound"] is True
    assert rows["gw"]["bound_fields"] == ("severity", "vlans_impacted", "stranded", "hard")
    assert rows["gw"]["bound_pointers"] == [unread] and rows["gw"]["bound_reasons"] == [why]
    assert summary._R_BOUND_PEERS.format(k=1) not in rows["gw"]["bound_reasons"]
    # every surface: the keystone keeps gw's measured place, flagged with the unreadable wording; acc is not ranked
    keystones = summary.summarize(snap)["keystones"]
    assert keystones[:-1] == [_bound_entry(gw, [unread], [why])]
    note = keystones[-1]
    assert note["severity"] == summary.IMPACT_NOT_ASSESSED and note["n_not_ranked"] == 1
    assert "acc — " in note["detail"] and "1 of them fail closed" in note["detail"], note["detail"]
    tab = summary.failure_impact_table(snap)
    assert tab[k]["stranded"] == f"≥ 1 — a lower bound, not an exact measurement: {why}", tab[k]["stranded"]


def test_keystones_on_the_sample_keep_their_order_flag_only_the_rows_their_own_evidence_bounds(sample):
    s = summary.summarize(copy.deepcopy(sample))
    assert s["keystone_contract"] == summary.KEYSTONE_CONTRACT_VERSION
    shown = [k for k in s["keystones"] if k["host"]]
    # the same ranking and values as before over the rows a ranking may place; each row its own stored evidence bounds
    # (core2's uncollected router, a positive or absent blind_links count) is flagged "at least", and only those
    assert shown == _old_keystones(_sample_ranked(sample), entry=_sample_entry(sample))
    assert [k["host"] for k in shown if k["lower_bound"]] == [k["host"] for k in shown
                                                              if _sample_bounded(sample, k["host"])]
    assert "core2" in {k["host"] for k in shown if k["lower_bound"]}
    # pinned on the sample (W45 refutation), never only derived from its stored counts: the flagged rows are exactly
    # core1 (blind_links 1) and core2 (its router), and the one row no ranking places is dist1 (a bounded Low band)
    assert [k["host"] for k in shown if k["lower_bound"]] == ["core1", "core2"], shown
    assert _sample_unranked(sample) == ["dist1"]
    # a bounded band below High or a bounded zero is never ranked and always disclosed, with any lower bound below
    # the cut, in the one NOT ASSESSED entry
    unranked = _sample_unranked(sample)
    order = _old_keystones(_sample_ranked(sample), top=len(sample["failure_impact"]))
    below = [k["host"] for k in order[len(shown):] if _sample_bounded(sample, k["host"])]
    notes = [k for k in s["keystones"] if not k["host"]]
    assert len(notes) == (1 if unranked or below else 0), notes
    if notes:
        assert notes[0]["severity"] == summary.IMPACT_NOT_ASSESSED and notes[0]["n_not_ranked"] == len(unranked)
        assert all(host in notes[0]["detail"] for host in unranked + below), notes[0]["detail"]
    assert cutover._keystone_hosts(copy.deepcopy(sample)) == {k["host"] for k in shown}


def test_the_samples_core2_is_a_lower_bound_on_every_surface(client, sample):
    """The refuter's counterexample: core2 is High with 42 stranded, but the stored cable map cables it to
    wan-edge-rtr1.lab (collected: false, kind router), so 42 is only a lower bound. Every surface says so."""
    cable = _cable(sample, "core2", WAN)
    why = summary._R_BOUND_PEERS.format(k=1)
    # the projection's own signal: each published measure cites exactly its own row's bounds (core2's cable; a row's
    # positive blind_links count, or the row itself when it predates that count), and a measurement's cite none
    items = ui.project_topology(sample)["failure_impact"]["items"]
    flagged = []
    for i, (item, src) in enumerate(zip(items, sample["failure_impact"])):
        for field in MEASURES:
            if item[field]["state"] == "published":
                cited = [ref["pointer"] for ref in item[field]["refs"] if ref["role"] == "witness"]
                assert cited == _sample_bounds(sample, i)[0], (src["host"], field, cited)
                if cited and src["host"] not in flagged:
                    flagged.append(src["host"])
    rows = {row["key"]: row for row in summary.impact_view(copy.deepcopy(sample))["rows"]}
    assert [host for host, row in rows.items() if row["lower_bound"]] == flagged and "core2" in flagged
    for host in flagged:
        pointers, reasons = _sample_bounds(sample, _sample_index(sample, host))
        assert rows[host]["bound_pointers"] == pointers and rows[host]["bound_reasons"] == reasons, rows[host]
    assert rows["core2"]["ranked"] is True and rows["core2"]["bound_pointers"] == [cable]
    assert rows["core2"]["bound_fields"] == ("severity", "vlans_impacted", "stranded", "hard")
    assert rows["core2"]["bound_reasons"] == [why]
    sid = client.post("/api/demo/seed").json()["snapshot"]["id"]
    # the dashboard's keystone panel: core2 keeps its measured place, flagged, its detail opening "at least 42"
    keystones = client.get(f"/api/snapshots/{sid}").json()["summary"]["keystones"]
    core2 = next(k for k in keystones if k["host"] == "core2")
    assert core2["stranded"] == 42 and core2["lower_bound"] is True, core2
    assert core2["lower_bound_pointers"] == [cable] and core2["lower_bound_reasons"] == [why]
    assert core2["detail"].startswith(f"LOWER BOUND, at least 42 endpoint(s) stranded: {why}. VLAN "), core2["detail"]
    assert all(k["lower_bound"] is (k["host"] in flagged) for k in keystones if k["host"]), keystones
    # the cutover planner: core2's wave is not complete (core2 can exceed its worst case, core1's 45, itself flagged
    # when core1's own stored evidence bounds it); core2 alone is a flagged lower bound, never complete
    plan = client.get(f"/api/snapshots/{sid}/cutover").json()
    wave = next(w for w in plan["waves"] if "core2" in w["switches"])
    br = wave["blast_radius"]
    assert br["host"] == "core1" and br["stranded"] == 45 and br["lower_bound"] is ("core1" in flagged)
    assert br["complete"] is False and "core2 — " + why in br["detail"], br
    alone = cutover._worst_blast_radius({"core2"}, summary.impact_view(copy.deepcopy(sample)))
    assert alone["host"] == "core2" and alone["stranded"] == 42 and alone["lower_bound"] is True
    assert alone["complete"] is False and alone["n_not_ranked"] == 0
    # the snapshot "Failure impact" tab: "≥ 42" with why
    table = client.get(f"/api/snapshots/{sid}/section/failure_impact").json()["data"]
    tab = next(row for row, src in zip(table, sample["failure_impact"]) if src["host"] == "core2")
    assert tab["stranded"] == f"≥ 42 — a lower bound, not an exact measurement: {why}", tab["stranded"]


@pytest.mark.parametrize("mode", ["no_run_config", "legacy", "duplicate"])
def test_a_withheld_sample_row_never_ranks_and_is_disclosed(sample, mode):
    snap, why = _held_core1(sample, mode)
    keystones = summary.summarize(snap)["keystones"]
    ranked, note = keystones[:-1], keystones[-1]
    others = [row for row in _sample_ranked(sample) if row["host"] != "core1"]
    # core1 is neither ranked first nor shown as Info
    assert ranked == _old_keystones(others, entry=_sample_entry(sample)), ranked
    assert note["host"] == "" and note["severity"] == summary.IMPACT_NOT_ASSESSED
    # core1's row(s), beside every other sample row its own evidence keeps from a ranking (always disclosed)
    also = [host for host in _sample_unranked(sample) if host != "core1"]
    assert note["n_not_ranked"] == (2 if mode == "duplicate" else 1) + len(also), note
    assert all(host in note["detail"] for host in also), note["detail"]
    assert "core1" in note["detail"] and why in note["detail"], note["detail"]
    assert "core1" not in cutover._keystone_hosts(snap)


def test_a_failed_phase_leaves_no_keystone_and_never_an_empty_clean_panel(sample):
    snap = copy.deepcopy(sample)
    snap["assessment_integrity"] = {"failure_impact": "failed"}
    keystones = summary.summarize(snap)["keystones"]
    assert len(keystones) == 1, keystones                          # never [] ("no keystone devices flagged")
    note = keystones[0]
    assert note["severity"] == summary.IMPACT_NOT_ASSESSED and note["n_not_ranked"] == len(sample["failure_impact"])
    assert note["stranded"] is None and note["vlans_impacted"] is None


def test_a_blind_spot_makes_the_keystone_ranking_a_disclosed_lower_bound(sample):
    snap = copy.deepcopy(sample)
    snap["collection_completeness"]["devices"] = copy.deepcopy(BLIND)
    keystones = summary.summarize(snap)["keystones"]
    # every row is still the engine's to rank (the blind spot names no sample row)
    assert keystones[:-1] == _old_keystones(_sample_ranked(sample), entry=_sample_entry(sample))
    note = keystones[-1]
    assert note["severity"] == summary.IMPACT_NOT_ASSESSED and note["n_not_ranked"] == len(_sample_unranked(sample))
    assert "lower bound" in note["detail"] and "collection_completeness lists 2 device(s)" in note["detail"]


def test_a_lower_bound_below_the_cut_is_disclosed_because_it_could_rank_among_the_devices_shown(sample):
    """A lower bound can be larger than it reads, so one ranked below the cut could belong above it."""
    view = summary.impact_view(copy.deepcopy(sample))
    ranked = sorted((row for row in view["rows"] if row["ranked"]), key=summary.impact_rank_key)
    place = [row["key"] for row in ranked].index("core2")
    assert place == 2                                       # core1 45, access1 42, core2 42 (stored order)
    keystones = summary._keystones(copy.deepcopy(sample), top=place)        # core2 just below the cut
    assert [k["host"] for k in keystones[:-1]] == ["core1", "access1"]
    # each shown row is flagged exactly when its own stored evidence bounds it (core1 by a positive blind_links count)
    assert [k["lower_bound"] for k in keystones[:-1]] == [_sample_bounded(sample, h) for h in ("core1", "access1")]
    note = keystones[-1]
    assert note["severity"] == summary.IMPACT_NOT_ASSESSED and note["n_not_ranked"] == len(_sample_unranked(sample))
    order = [row["key"] for row in ranked]
    below = [host for host in order[place:] if _sample_bounded(sample, host)]
    assert "core2" in below
    assert (f"{len(below)} ranked row(s) below the devices shown publish only lower bounds"
            in note["detail"]), note["detail"]
    assert "core2 — " + summary._R_BOUND_PEERS.format(k=1) in note["detail"]
    # the control: a cut that shows core2 flags it in place and no longer discloses it as below the cut. Its
    # precondition is asserted, never branched on, so a future sample cannot skip the control silently: the sample's
    # only ranked bounds are core1 (first) and core2 (third) -- dist1's bounded Low band never ranks -- so no bounded
    # ranked row lies below core2
    assert [host for host in order if _sample_bounded(sample, host)] == ["core1", "core2"], order
    assert not [host for host in order[place + 1:] if _sample_bounded(sample, host)], order
    keystones = summary._keystones(copy.deepcopy(sample), top=place + 1)
    shown = [k for k in keystones if k["host"]]
    assert [k["host"] for k in shown] == ["core1", "access1", "core2"] and shown[-1]["lower_bound"] is True
    assert not any("publish only lower bounds" in k["detail"] for k in keystones if not k["host"]), keystones


@pytest.mark.parametrize("mode", ["held", "failed", "published"])
def test_a_stored_executive_brief_keystone_list_is_never_read(client, sample, mode):
    """No engine producer writes executive_brief.keystones (compute_executive_brief takes no failure-impact input),
    so a stored one is uploaded or hand-made. Read first, it bypassed every projection hold on the dashboard, the
    /graph badge and the cutover plan's wave tags. Each surface now ranks only the projection, list or not."""
    if mode == "held":
        snap, _why = _held_core1(sample, "no_run_config")       # the engine withholds core1's row
    else:
        snap = copy.deepcopy(sample)
        if mode == "failed":
            snap["assessment_integrity"] = {"failure_impact": "failed"}
    without = copy.deepcopy(snap)
    # a crafted list naming the withheld core1 first and the sample's Info podacc1 as Critical
    snap["executive_brief"]["keystones"] = [
        {"host": "core1", "severity": "High", "stranded": 45, "vlans_impacted": 3, "detail": CLEAN_BILL},
        {"host": "podacc1", "severity": "Critical", "stranded": 999, "vlans_impacted": 9, "detail": "crafted"}]
    want = summary._keystones(without)
    s = summary.summarize(copy.deepcopy(snap))
    assert s["keystones"] == want                           # exactly the projection's ranking, as if no list existed
    assert all(k.get("stranded") != 999 and k.get("detail") != CLEAN_BILL for k in s["keystones"]), s["keystones"]
    sid = _upload(client, snap)
    assert client.get(f"/api/snapshots/{sid}").json()["summary"]["keystones"] == want
    nodes = {n["id"]: n for n in client.get(f"/api/snapshots/{sid}/graph").json()["nodes"]}
    tags = {tag for wave in cutover.build_plan(copy.deepcopy(snap))["waves"] for tag in wave["keystones"]}
    assert nodes["podacc1"]["keystone"] is False and "podacc1" not in tags
    if mode == "published":
        # the control: the published positive rows still surface, from the projection
        assert want[0]["host"] == "core1" and nodes["core1"]["keystone"] is True and "core1" in tags
    else:
        assert nodes["core1"]["keystone"] is False and "core1" not in tags
        assert "core1" not in {k["host"] for k in s["keystones"]}
    if mode == "failed":
        assert len(want) == 1 and want[0]["severity"] == summary.IMPACT_NOT_ASSESSED, want
        assert not tags and not any(n["keystone"] for n in nodes.values())


# --------------------------------------------------------------------------------------------------
# the cutover plan's worst-case blast radius
# --------------------------------------------------------------------------------------------------
def test_the_worst_case_never_ranks_a_row_the_engine_withholds():
    snap = _snapshot(_unsimulatable_fleet(), waves=[("held", ("acc", "x1")), ("mixed", ("gw", "acc")),
                                                    ("missing", ("gw", "ghost"))])
    raw = _by_host(snap)
    waves = {w["group"]: w for w in cutover.build_plan(snap)["waves"]}
    # before W27 this wave's worst case was acc's stored row: Info, 0 endpoints across 0 VLANs
    held = waves["held"]["blast_radius"]
    assert held["host"] == "" and held["severity"] == cutover.IMPACT_NOT_ASSESSED
    assert held["stranded"] is None and held["vlans_impacted"] is None
    assert held["complete"] is False and held["n_not_ranked"] == 2
    assert held["detail"].startswith("NOT ASSESSED") and "could not simulate" in held["detail"], held["detail"]
    assert "acc" in held["detail"] and "x1" in held["detail"]
    gw = raw["gw"]
    mixed = waves["mixed"]["blast_radius"]
    assert {k: mixed[k] for k in ("host", "severity", "stranded", "vlans_impacted")} == {
        k: gw[k] for k in ("host", "severity", "stranded", "vlans_impacted")}
    # gw's own counts are lower bounds (its off-scan VLAN), and acc is withheld: both are disclosed
    gw_bound = _bound_entry(gw, [_row_pointer(snap, "gw") + "/off_scan_gw_vlans"],
                            [summary._R_BOUND_OFF_SCAN.format(n=1)])
    assert mixed["lower_bound"] is True and mixed["lower_bound_pointers"] == gw_bound["lower_bound_pointers"]
    assert mixed["detail"].startswith(gw_bound["detail"] + " — LOWER BOUND"), mixed["detail"]
    assert "acc — " in mixed["detail"] and mixed["complete"] is False and mixed["n_not_ranked"] == 1
    missing = waves["missing"]["blast_radius"]
    assert missing["host"] == "gw" and missing["complete"] is False and missing["n_not_ranked"] == 1
    assert "ghost — not collected: no failure_impact row names this switch" in missing["detail"], missing["detail"]


def test_a_wave_behind_an_uncollected_neighbour_is_not_assessed_and_its_control_renders_as_before():
    snap = _snapshot(_downstream_fleet(), waves=[("edge", ("acc",)), ("both", ("gw", "acc"))])
    raw = _by_host(snap)
    waves = {w["group"]: w for w in cutover.build_plan(snap)["waves"]}
    edge = waves["edge"]["blast_radius"]
    assert edge["severity"] == cutover.IMPACT_NOT_ASSESSED and edge["stranded"] is None
    assert "uncollected neighbour" in edge["detail"] and CLEAN_BILL not in edge["detail"], edge["detail"]
    both = waves["both"]["blast_radius"]
    assert both["host"] == "gw" and both["severity"] == "High" and both["complete"] is False
    assert both["lower_bound"] is True                      # gw itself is cabled to the uncollected router wan
    # once dsw reads collected, acc's wave renders exactly its stored row, plus the new keys; gw's wave is still a
    # lower bound behind wan
    _node(snap, "dsw")["collected"] = True
    waves = {w["group"]: w for w in cutover.build_plan(snap)["waves"]}
    assert waves["edge"]["blast_radius"] == {**_exact(raw["acc"]), "complete": True, "n_not_ranked": 0}
    both = waves["both"]["blast_radius"]
    assert both["host"] == "gw" and both["complete"] is False and both["n_not_ranked"] == 0
    assert "1 ranked switch(es) in this wave publish their counts only as lower bounds" in both["detail"]
    # the control: once wan reads collected too, both waves render exactly the stored rows, plus the new keys
    _node(snap, "wan")["collected"] = True
    waves = {w["group"]: w for w in cutover.build_plan(snap)["waves"]}
    assert waves["edge"]["blast_radius"] == {**_exact(raw["acc"]), "complete": True, "n_not_ranked": 0}
    assert waves["both"]["blast_radius"] == {**_exact(raw["gw"]), "complete": True, "n_not_ranked": 0}


def test_a_lower_bound_worst_case_is_never_complete_and_an_exact_one_stays_complete():
    """Codex's case on #620: the REAL producer's partial row for gw (High, positive stranded, off_scan_gw_vlans 1)
    alone in a wave. Its positive values are legitimate and kept, but they are lower bounds, so the wave's worst case
    is not complete and says why. The unbounded positive control stays complete."""
    snap = _snapshot(_unsimulatable_fleet(), waves=[("alone", ("gw",))])
    gw = _by_host(snap)["gw"]
    assert gw["severity"] == "High" and gw["stranded"] > 0 and gw["off_scan_gw_vlans"] == 1, gw
    br = {w["group"]: w for w in cutover.build_plan(snap)["waves"]}["alone"]["blast_radius"]
    assert {k: br[k] for k in ("host", "severity", "stranded", "vlans_impacted")} == {
        k: gw[k] for k in ("host", "severity", "stranded", "vlans_impacted")}       # the measured values are kept
    assert br["complete"] is False and br["n_not_ranked"] == 0 and br["lower_bound"] is True
    off_scan = _row_pointer(snap, "gw") + "/off_scan_gw_vlans"
    assert br["lower_bound_pointers"] == [off_scan]
    assert br["lower_bound_reasons"] == [summary._R_BOUND_OFF_SCAN.format(n=1)]
    assert br["detail"].startswith(f"LOWER BOUND, at least {gw['stranded']} endpoint(s) stranded: 1 VLAN(s) on this "
                                   "switch have an off-scan gateway"), br["detail"]
    assert "the worst case may be larger" in br["detail"] and "gw — 1 VLAN(s)" in br["detail"], br["detail"]
    # the control: an unbounded positive worst case (gw behind a router the cable map now shows collected)
    ctrl = _snapshot(_downstream_fleet(), waves=[("alone", ("gw",))])
    _node(ctrl, "wan")["collected"] = True
    exact = {w["group"]: w for w in cutover.build_plan(ctrl)["waves"]}["alone"]["blast_radius"]
    assert exact == {**_exact(_by_host(ctrl)["gw"]), "complete": True, "n_not_ranked": 0}
    assert exact["severity"] == "High" and exact["stranded"] > 0


@pytest.mark.parametrize("host", ["", None])
def test_a_row_with_an_empty_host_is_disclosed_like_one_with_no_host(host):
    """An empty host names no readable host, exactly as an absent one: it could be any switch in the wave, so it is
    disclosed and the worst case is not complete (before, an empty host joined no switch and vanished silently)."""
    snap = _snapshot(_downstream_fleet(), waves=[("alone", ("gw",))])
    _node(snap, "wan")["collected"] = True                  # gw's row is then exact: a ranked, complete worst case
    clean = {w["group"]: w for w in cutover.build_plan(snap)["waves"]}["alone"]["blast_radius"]
    assert clean["host"] == "gw" and clean["complete"] is True and clean["n_not_ranked"] == 0
    snap["failure_impact"].append(dict(_by_host(snap)["gw"], host=host))
    k = len(snap["failure_impact"]) - 1
    br = {w["group"]: w for w in cutover.build_plan(snap)["waves"]}["alone"]["blast_radius"]
    assert br["host"] == "gw" and br["complete"] is False and br["n_not_ranked"] == 1, br
    assert f"/failure_impact/{k} (names no readable host, so it could be any switch here)" in br["detail"], br["detail"]


def test_the_sample_plan_keeps_its_worst_cases_and_a_duplicate_is_never_picked(sample):
    plan = cutover.build_plan(copy.deepcopy(sample))
    assert plan["waves"]
    entry = _sample_entry(sample)
    ranked_rows = _sample_ranked(sample)
    unranked = set(_sample_unranked(sample))
    with_core2 = [wave for wave in plan["waves"] if "core2" in wave["switches"]]
    assert with_core2 and len(with_core2) < len(plan["waves"])
    for wave in plan["waves"]:
        switches = set(wave["switches"])
        assert switches <= set(_by_host(sample)), wave["group"]          # every sample wave switch has its row
        br = wave["blast_radius"]
        not_ranked = sorted(switches & unranked)
        # a ranked switch whose counts its own evidence bounds (core2's router, a positive or absent blind_links
        # count) could exceed the wave's worst case, whether or not it is the worst
        bounded_in = [row["host"] for row in ranked_rows if row["host"] in switches
                      and _sample_bounded(sample, row["host"])]
        if not any(row["host"] in switches for row in ranked_rows):
            assert br["severity"] == cutover.IMPACT_NOT_ASSESSED and br["stranded"] is None, br
            assert br["complete"] is False and br["n_not_ranked"] == len(not_ranked), br
            continue
        want = _old_worst(ranked_rows, switches)
        if not not_ranked and not bounded_in:
            assert br == {**_exact(want), "complete": True, "n_not_ranked": 0}, wave["group"]
            continue
        assert {k: v for k, v in br.items() if k != "detail"} == {
            **{k: v for k, v in entry(want).items() if k != "detail"}, "complete": False,
            "n_not_ranked": len(not_ranked)}, wave["group"]
        assert br["detail"].startswith(entry(want)["detail"] + " — LOWER BOUND, the worst case may be larger")
        assert all(host in br["detail"] for host in not_ranked + bounded_in), (wave["group"], br["detail"])
        if "core2" in switches:
            assert "core2 — " + summary._R_BOUND_PEERS.format(k=1) in br["detail"], br["detail"]
    snap, why = _held_core1(sample, "duplicate")
    wave = next(w for w in cutover.build_plan(snap)["waves"] if "core1" in w["switches"])
    br = wave["blast_radius"]
    others = [row for row in ranked_rows if row["host"] != "core1"]
    want = _old_worst(others, set(wave["switches"]))
    assert br["host"] != "core1" and br["severity"] != "Info"
    assert {k: br[k] for k in ("host", "severity", "stranded", "vlans_impacted")} == {
        k: want[k] for k in ("host", "severity", "stranded", "vlans_impacted")}
    assert br["detail"].startswith(entry(want)["detail"] + " — LOWER BOUND") and why in br["detail"]
    also = sorted((set(wave["switches"]) & unranked) - {"core1"})
    assert br["complete"] is False and br["n_not_ranked"] == 1 + len(also)


def test_the_cutover_document_prints_not_assessed_never_a_zero(tmp_path):
    pytest.importorskip("docx")
    from docx import Document

    from backend.cutover_docx import write_cutover_docx

    snap = _snapshot(_unsimulatable_fleet(), waves=[("held", ("acc", "x1"))])
    out = tmp_path / "cutover.docx"
    write_cutover_docx(str(out), snap, "W27")
    lines = [p.text for p in Document(out).paragraphs if p.text.startswith("Worst-case blast radius:")]
    assert len(lines) == 1, lines
    assert "NOT ASSESSED" in lines[0] and "could not simulate" in lines[0], lines[0]
    assert "0 endpoint(s)" not in lines[0] and "(Info)" not in lines[0], lines[0]


def _blast_lines(snap, tmp_path):
    from docx import Document

    from backend.cutover_docx import write_cutover_docx

    out = tmp_path / "cutover.docx"
    write_cutover_docx(str(out), snap, "W27")
    return [p.text for p in Document(out).paragraphs if p.text.startswith("Worst-case blast radius:")]


def test_the_cutover_document_prints_a_lower_bound_as_at_least(tmp_path, sample):
    pytest.importorskip("docx")
    # the REAL producer's partial row for gw, and the sample's core2 behind its uncollected router, each alone
    snap = _snapshot(_unsimulatable_fleet(), waves=[("alone", ("gw",))])
    gw = _by_host(snap)["gw"]
    (line,) = _blast_lines(snap, tmp_path)
    assert (f"gw (High) — at least {gw['stranded']} endpoint(s) stranded across at least {gw['vlans_impacted']} "
            "VLAN(s). LOWER BOUND, at least") in line, line
    snap = copy.deepcopy(sample)
    snap["wave_sequencing"].append({"group": "Core2 alone", "make_before_break": [], "hard_cutover": ["core2"]})
    lines = _blast_lines(snap, tmp_path)
    assert any("core2 (High) — at least 42 endpoint(s) stranded across at least 3 VLAN(s). LOWER BOUND" in line
               for line in lines), lines
    # core1's wave prints core1's counts, exact unless core1's own stored evidence bounds them (then "at least"), with
    # core2's bound disclosed as the reason it may be larger
    core1 = _by_host(sample)["core1"]
    head = ("core1 (High) — at least {s} endpoint(s) stranded across at least {v} VLAN(s)."
            if _sample_bounded(sample, "core1") else "core1 (High) — {s} endpoint(s) stranded across {v} VLAN(s).")
    head = head.format(s=core1["stranded"], v=core1["vlans_impacted"])
    assert any(head in line and "core2 — " in line for line in lines), (head, lines)
    # the control: an exact worst case prints its counts bare
    ctrl = _snapshot(_downstream_fleet(), waves=[("alone", ("gw",))])
    _node(ctrl, "wan")["collected"] = True
    g = _by_host(ctrl)["gw"]
    (line,) = _blast_lines(ctrl, tmp_path)
    assert f"gw (High) — {g['stranded']} endpoint(s) stranded across {g['vlans_impacted']} VLAN(s)." in line, line
    assert "at least" not in line and "LOWER BOUND" not in line, line


# --------------------------------------------------------------------------------------------------
# the snapshot "Failure impact" tab
# --------------------------------------------------------------------------------------------------
def test_the_failure_impact_tab_shows_the_engine_reason_never_a_withheld_value(client):
    sid = _upload(client, _snapshot(_unsimulatable_fleet()))
    stored = json.loads(client.get(f"/api/snapshots/{sid}/raw").content)
    table = client.get(f"/api/snapshots/{sid}/section/failure_impact").json()["data"]
    items = ui.project_topology(stored)["failure_impact"]["items"]
    assert len(table) == len(items) == len(stored["failure_impact"])
    for row, item, src in zip(table, items, stored["failure_impact"]):
        assert list(row) == list(summary.IMPACT_FIELDS), row
        for field in summary.IMPACT_FIELDS:                  # each cell is the projection's: value, bound or reason
            fact = item[field]
            cited = [ref for ref in fact["refs"] if ref["role"] == "witness"]
            if fact["state"] == "published" and field in MEASURES and cited:
                # a published lower bound: "≥ value" and why, never the bare value
                assert isinstance(row[field], str), (src["host"], field, row[field])
                assert row[field].startswith(f"≥ {src[field]} — a lower bound, not an exact measurement: "), (
                    src["host"], field, row[field])
            elif fact["state"] == "published":
                assert row[field] == src[field], (src["host"], field)
            else:
                assert isinstance(row[field], str) and row[field] == fact["reason"], (src["host"], field)
    by_host = {src["host"]: row for row, src in zip(table, stored["failure_impact"])}
    for host in ("acc", "x1", "x2"):
        for field in MEASURES:                               # never the stored Info or 0
            cell = by_host[host][field]
            assert isinstance(cell, str) and cell.startswith("not collected:"), (host, field, cell)
            assert "could not simulate" in cell, (host, field, cell)
        # the producer's own INDETERMINATE disclosure stays published beside its held measures
        assert by_host[host]["detail"].startswith(ui.IMPACT_INDETERMINATE_PREFIX)
    gw, raw = by_host["gw"], _by_host(stored)["gw"]
    # simulated in part: High and the positive counts are lower bounds that name the off-scan VLAN
    why = summary._R_BOUND_OFF_SCAN.format(n=1)
    assert gw["severity"] == f"≥ High — a lower bound, not an exact measurement: {why}", gw["severity"]
    assert raw["stranded"] > 0, raw
    assert gw["stranded"] == f"≥ {raw['stranded']} — a lower bound, not an exact measurement: {why}", gw["stranded"]
    assert gw["host"] == "gw" and gw["off_scan_gw_vlans"] == 1 and gw["detail"] == raw["detail"]
    for field in ("backup", "fhrp"):                         # a zero under a partial simulation is a lower bound
        assert isinstance(gw[field], str) and "only a lower bound" in gw[field], (field, gw[field])


def _tab_against_bounds(table, snap):
    """The "Failure impact" tab of `snap`, checked cell by cell against each stored row's own bounds
    (:func:`_sample_bounds`). Returns which branch each bounded row's cells took, by host: ``withheld`` (a band below
    High or a zero), ``bounded`` (High or a positive count, published as a lower bound) and ``cleared`` (a clean-bill
    detail, which names no simulated VLAN, withheld as not a clean bill)."""
    assert len(table) == len(snap["failure_impact"])
    withheld, bounded, cleared = {}, {}, []
    for i, (row, src) in enumerate(zip(table, snap["failure_impact"])):
        assert list(row) == list(summary.IMPACT_FIELDS)
        pointers, reasons = _sample_bounds(snap, i)
        for field in summary.IMPACT_FIELDS:
            understates = field in MEASURES and (src[field] != "High" if field == "severity" else src[field] == 0)
            if pointers and understates:
                # a band below High and a zero: withheld, since the bound cannot vouch for them
                assert isinstance(row[field], str), (src["host"], field, row[field])
                assert ("may understate" if field == "severity" else "only a lower bound") in row[field], row[field]
                withheld.setdefault(src["host"], []).append(field)
            elif pointers and field in MEASURES:
                # High and each positive count: published, as the lower bounds they are
                assert row[field] == (f"≥ {src[field]} — a lower bound, not an exact measurement: "
                                      f"{'; '.join(reasons)}"), (src["host"], field, row[field])
                bounded.setdefault(src["host"], []).append(field)
            elif pointers and field == "detail" and not src["vlans_impacted"]:
                assert isinstance(row[field], str) and "not a clean bill" in row[field], row[field]
                assert src[field] not in row[field], (src["host"], row[field])   # never the stored clean bill
                cleared.append(src["host"])
            else:
                assert row[field] == src[field], (src["host"], field)      # every exact cell renders as before
    assert set(withheld) | set(bounded) == {src["host"] for i, src in enumerate(snap["failure_impact"])
                                            if _sample_bounds(snap, i)[0]}
    return ({host: sorted(fields) for host, fields in withheld.items()},
            {host: sorted(fields) for host, fields in bounded.items()}, cleared)


def test_the_sample_tab_changes_only_the_rows_their_own_evidence_bounds(client, sample):
    """core2's uncollected router bounds it, as before; since W32 (through the W33 owner) a row's positive or absent
    blind_links count bounds it the same way. Each bound is derived from the stored row (:func:`_sample_bounds`)."""
    sid = client.post("/api/demo/seed").json()["snapshot"]["id"]
    table = client.get(f"/api/snapshots/{sid}/section/failure_impact").json()["data"]
    withheld, bounded, cleared = _tab_against_bounds(table, sample)
    assert withheld["core2"] == ["backup", "fhrp"], withheld
    assert bounded["core2"] == ["hard", "severity", "stranded", "vlans_impacted"], bounded
    assert "uncollected neighbour" in table[_sample_index(sample, "core2")]["backup"]
    # pinned exactly on the sample (W45 refutation), never only derived from its stored counts: core2's router, and
    # the blind_links of core1 (High, 45 stranded) and dist1 (Low, FHRP-covered), 1 each
    assert withheld == {"core1": ["backup", "fhrp"], "core2": ["backup", "fhrp"],
                        "dist1": ["backup", "hard", "severity", "stranded"]}, withheld
    assert bounded == {"core1": ["hard", "severity", "stranded", "vlans_impacted"],
                       "core2": ["hard", "severity", "stranded", "vlans_impacted"],
                       "dist1": ["fhrp", "vlans_impacted"]}, bounded
    # the clean-bill detail branch, asserted rather than skipped: every bounded sample row simulated a VLAN, so its
    # per-VLAN detail lists what was simulated and renders as stored
    assert cleared == [], cleared
    assert {src["host"]: src["vlans_impacted"] for i, src in enumerate(sample["failure_impact"])
            if _sample_bounds(sample, i)[0]} == {"core1": 3, "core2": 3, "dist1": 2}
    # ... and that branch exercised, once per bound that reaches a detail: podacc1's stored clean bill made older than
    # the blind_links count (the row itself bounds it), and core2 storing a clean bill behind its uncollected router.
    # A POSITIVE count never takes this branch: with no VLAN simulated, the owner reads it as its blind_links_only
    # hold (impact_assessability.row_hold), never as a bound. Each tab row withholds every measure and the bill itself.
    snap = copy.deepcopy(sample)
    pod, core2 = _sample_index(snap, "podacc1"), _sample_index(snap, "core2")
    assert snap["failure_impact"][pod]["detail"] == CLEAN_BILL and snap["failure_impact"][pod]["blind_links"] == 0
    del snap["failure_impact"][pod]["blind_links"]
    snap["failure_impact"][core2].update(severity="Info", vlans_impacted=0, stranded=0, hard=0, backup=0, fhrp=0,
                                         detail=CLEAN_BILL)
    sid = _upload(client, snap)
    stored = json.loads(client.get(f"/api/snapshots/{sid}/raw").content)
    assert stored["failure_impact"] == snap["failure_impact"]
    table = client.get(f"/api/snapshots/{sid}/section/failure_impact").json()["data"]
    withheld, bounded, cleared = _tab_against_bounds(table, stored)
    assert cleared == ["core2", "podacc1"], cleared                         # stored order
    for host in cleared:
        assert withheld[host] == sorted(MEASURES) and host not in bounded, (host, withheld, bounded)
    for k, cause in ((core2, "uncollected neighbour"), (pod, "carries no blind_links")):
        assert table[k]["detail"].startswith("not collected: ") and cause in table[k]["detail"], table[k]["detail"]


def test_the_tab_withholds_unreadable_rows(client):
    sid = _upload(client, {"devices": {"sw1": {}},
                           "failure_impact": [1, "x", {"host": "sw1", "severity": "Info", "stranded": 0}]})
    table = client.get(f"/api/snapshots/{sid}/section/failure_impact").json()["data"]
    assert len(table) == 3
    for row in table:
        assert list(row) == list(summary.IMPACT_FIELDS)
        for field in MEASURES:
            assert isinstance(row[field], str) and row[field], (field, row[field])   # a reason, never Info or 0
    for row in table[:2]:                                    # a row that is not an object claims nothing at all
        assert all(isinstance(row[field], str) and row[field] for field in summary.IMPACT_FIELDS), row


@pytest.mark.parametrize("value", [
    {"severity": "Info", "stranded": 0, "detail": CLEAN_BILL},      # Codex's case on #620: a clean bill as an object
    5, "No reachability impact"])
def test_a_failure_impact_section_that_is_not_a_list_shows_the_projections_disclosure(client, value):
    """The tab never passes a non-list section through to the table: GenericTable rendered the stored object's own
    keys (severity Info, stranded 0, the clean-bill detail) as a result. It shows the projection's own disclosure of
    the list instead: its state and reason, a flat object of two strings (rendered as two labelled rows)."""
    sid = _upload(client, {"devices": {"sw1": {}}, "failure_impact": value})
    stored = json.loads(client.get(f"/api/snapshots/{sid}/raw").content)
    listing = ui.project_topology(stored)["failure_impact"]
    assert listing["state"] != "published" and listing["items"] == [], listing
    r = client.get(f"/api/snapshots/{sid}/section/failure_impact")
    assert r.status_code == 200, r.text
    data = r.json()["data"]
    assert data == {"state": listing["state"], "reason": listing["reason"]}, data
    assert all(isinstance(v, str) and v for v in data.values())
    assert CLEAN_BILL not in json.dumps(data) and "Info" not in data.values()
    # every other surface agrees: nothing is ranked, and the dashboard says why
    keystones = client.get(f"/api/snapshots/{sid}").json()["summary"]["keystones"]
    assert len(keystones) == 1 and keystones[0]["severity"] == summary.IMPACT_NOT_ASSESSED, keystones
    assert listing["reason"] in keystones[0]["detail"]


def test_a_valid_empty_failure_impact_list_stays_an_empty_table(client):
    """The control: a list the producer wrote empty, which the projection reads as collected but empty, is still the
    empty table (the tab says "Empty."); only a withheld list becomes a disclosure."""
    sid = _upload(client, {"devices": {"sw1": {}}, "failure_impact": []})
    stored = json.loads(client.get(f"/api/snapshots/{sid}/raw").content)
    assert ui.project_topology(stored)["failure_impact"]["state"] == "collected_but_empty"
    assert client.get(f"/api/snapshots/{sid}/section/failure_impact").json()["data"] == []
    # an empty list the projection withholds (its phase failed) is disclosed, never an empty clean table
    sid = _upload(client, {"devices": {"sw1": {}}, "failure_impact": [],
                           "assessment_integrity": {"failure_impact": "failed"}})
    stored = json.loads(client.get(f"/api/snapshots/{sid}/raw").content)
    listing = ui.project_topology(stored)["failure_impact"]
    assert listing["state"] != "collected_but_empty", listing
    data = client.get(f"/api/snapshots/{sid}/section/failure_impact").json()["data"]
    assert data == {"state": listing["state"], "reason": listing["reason"]}, data


# --------------------------------------------------------------------------------------------------
# the /graph keystone badge and a summary cached by the raw-row ranking
# --------------------------------------------------------------------------------------------------
def test_the_graph_badge_and_a_cached_summary_never_keep_a_withheld_keystone(client, sample):
    snap, _why = _held_core1(sample, "no_run_config")
    sid = _upload(client, snap)
    nodes = {n["id"]: n for n in client.get(f"/api/snapshots/{sid}/graph").json()["nodes"]}
    assert nodes["core1"]["keystone"] is False          # the raw ranking made it the first keystone (High, 45)
    assert nodes["access1"]["keystone"] is True
    # a summary cached by the raw-row ranking (no keystone contract) that still names core1
    store = client.app.state.store
    stale = dict(store.get_snapshot_meta(sid)["summary"])
    assert stale.pop("keystone_contract") == summary.KEYSTONE_CONTRACT_VERSION
    stale["keystones"] = [{"host": "core1", "severity": "High", "stranded": 45, "vlans_impacted": 3, "detail": ""}]
    store.update_summary(sid, stale)
    nodes = {n["id"]: n for n in client.get(f"/api/snapshots/{sid}/graph").json()["nodes"]}
    assert nodes["core1"]["keystone"] is False          # the badge reads the snapshot, never the stale cache
    fresh = client.get(f"/api/snapshots/{sid}").json()["summary"]
    assert fresh["keystone_contract"] == summary.KEYSTONE_CONTRACT_VERSION
    assert "core1" not in {k["host"] for k in fresh["keystones"]}
    assert store.get_snapshot_meta(sid)["summary"]["keystone_contract"] == summary.KEYSTONE_CONTRACT_VERSION


# --------------------------------------------------------------------------------------------------
# an owner fault withholds every surface; the dossier recompute keeps the engine's own input
# --------------------------------------------------------------------------------------------------
def test_a_projection_fault_withholds_every_surface(monkeypatch, sample):
    def fault(_snap):
        raise RuntimeError("owner fault")

    monkeypatch.setattr(summary.engine, "failure_impact_projection", fault)
    snap = copy.deepcopy(sample)
    keystones = summary.summarize(snap)["keystones"]
    assert len(keystones) == 1 and keystones[0]["severity"] == summary.IMPACT_NOT_ASSESSED, keystones
    assert "could not be built" in keystones[0]["detail"]
    for wave in cutover.build_plan(snap)["waves"]:
        br = wave["blast_radius"]
        assert br["severity"] == cutover.IMPACT_NOT_ASSESSED and br["stranded"] is None, br
    table = summary.failure_impact_table(snap)
    assert len(table) == len(sample["failure_impact"])
    assert all(row[field] == summary._R_IMPACT_FAULT for row in table for field in summary.IMPACT_FIELDS)


def test_the_dossier_recompute_forwards_the_stored_producer_rows_unchanged(client, monkeypatch):
    """Left on the stored rows on purpose: compute_device_dossiers is an engine function, and the pipeline hands it
    the same failure_impact list it stores (COLLECT_PARSE main). The webapp recompute must forward exactly that list,
    never a projection-derived substitute, so the engine keeps owning its own input."""
    seen = {}

    def capture(**kwargs):
        seen.update(kwargs)
        return {"summary": {"bands": {"Unassessed": 1}}, "per_device": []}

    monkeypatch.setattr(analyze, "compute_device_dossiers", capture)
    impact = [{"host": "good1", "severity": "Info", "vlans_impacted": 0, "stranded": 0, "hard": 0, "backup": 0,
               "fhrp": 0, "detail": CLEAN_BILL}]
    snap = {"devices": {"blind1": {}, "good1": {}},
            "health_scores": [{"switch": "blind1", "band": "Insufficient Data", "score": 90},
                              {"switch": "good1", "band": "Good", "score": 85}],
            "failure_impact": impact,
            "device_dossiers": {"summary": {"bands": {"Low": 1}}, "per_device": []}}
    sid = _upload(client, snap)
    r = client.get(f"/api/snapshots/{sid}/section/device_dossiers")
    assert r.status_code == 200, r.text
    assert seen["failure_impact"] == impact


def test_the_blind_spot_note_never_calls_a_record_the_engine_cannot_read_a_blind_device(sample):
    """W43 (F6): the projection's fleet qualifier also cites a collection_completeness record it cannot read as a
    blind spot (a row that is not an object or states no status of its owner's vocabulary, or a list or section of
    the wrong type). Each could be one, so it qualifies the ranking, but it is not a device the owner lists as partial
    or not collected: the note counts and words the two kinds apart, by the projection's own classifier."""
    from backend import engine
    snap = copy.deepcopy(sample)
    # every case below lists one device outside the devices map (ghost1, then x), which the producer counts in its
    # inventory; W51 reconciles that count with the roster, so an uncounted one would be a third kind of witness
    snap["collection_completeness"]["summary"]["inventory"] += 1
    snap["collection_completeness"]["devices"] = copy.deepcopy(BLIND)            # the control: two readable rows
    view = summary.impact_view(snap)
    assert (view["blind"], view["blind_unread"]) == (2, 0)
    assert summary.impact_blind_note(view) == summary._R_IMPACT_BLIND.format(n=2)

    snap["collection_completeness"]["devices"] = [{"host": "x", "status": "complete"}]
    assert engine.fleet_blind_spot_rows(snap) == [] == ui.fleet_blind_spot_rows(snap)
    listing = engine.failure_impact_projection(snap)
    assert "fleet_lists_exclude_blind_devices" in listing["caveats"]               # the engine qualifies the list
    view = summary.impact_view(snap)
    assert (view["blind"], view["blind_unread"]) == (0, 1)
    note = summary.impact_blind_note(view)
    assert note == summary._R_IMPACT_BLIND_UNREAD.format(n=1), note
    assert "as partial or not collected:" not in note and "lists 1 device(s)" not in note
    keystones = summary.summarize(snap)["keystones"]
    tail = keystones[-1]
    assert tail["host"] == "" and tail["severity"] == summary.IMPACT_NOT_ASSESSED, tail
    assert f"The ranking is a lower bound: {note}" in tail["detail"] and "lists 1 device(s)" not in tail["detail"]
    alone = cutover._worst_blast_radius({"core2"}, view)
    assert alone["complete"] is False and note in alone["detail"], alone

    snap["collection_completeness"]["devices"] = copy.deepcopy(BLIND) + [None]  # both kinds, each as what it is
    assert engine.fleet_blind_spot_rows(snap) == [0, 1]
    view = summary.impact_view(snap)
    assert (view["blind"], view["blind_unread"]) == (2, 1)
    assert summary.impact_blind_note(view) == "; ".join(
        [summary._R_IMPACT_BLIND.format(n=2), summary._R_IMPACT_BLIND_UNREAD.format(n=1)])

    absent = copy.deepcopy(sample)                  # W51: a record the snapshot does not carry is no listed device
    del absent["collection_completeness"]
    assert "fleet_lists_exclude_blind_devices" in engine.failure_impact_projection(absent)["caveats"]
    view = summary.impact_view(absent)
    assert (view["blind"], view["blind_unread"]) == (0, 1)
    assert summary.impact_blind_note(view) == summary._R_IMPACT_BLIND_UNREAD.format(n=1)
    assert "lists 1 device(s)" not in summary.impact_blind_note(view)

    snap["collection_completeness"]["devices"] = {"core1": {"status": "not collected"}}   # a list it cannot read
    view = summary.impact_view(snap)
    assert (view["blind"], view["blind_unread"]) == (0, 1)
    assert summary.impact_blind_note(view) == summary._R_IMPACT_BLIND_UNREAD.format(n=1)

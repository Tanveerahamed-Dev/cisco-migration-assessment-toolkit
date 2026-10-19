"""Persisted selected-plan output identity must not become assessment or approval authority."""

from __future__ import annotations

import copy
import json
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from backend import deliverables, engine  # noqa: E402
from backend.protocol_portfolio import build_protocol_single_snapshot_bundle  # noqa: E402
from backend.storage import saved_plan_digest  # noqa: E402
from cisco_toolkit.protocol_assurance import (  # noqa: E402
    bind_snapshot_json_bytes,
    bound_snapshot_source,
)


def test_catalogue_declares_a_closed_saved_plan_basis_for_every_writer():
    from backend.app import _plan_catalogue_row

    derived = {"design", "crd", "mop", "nrfu", "deck"}
    baseline = {"cutover", "engagement", "runbook", "archreview", "opshandbook"}
    assert derived.isdisjoint(baseline)
    assert set(deliverables.SPECS) == derived | baseline
    rows = {row["key"]: _plan_catalogue_row(row) for row in deliverables.catalogue()}
    assert set(rows) == derived | baseline
    for key in derived:
        assert rows[key]["saved_plan_basis"] == "selected-requirements"
    for key in baseline:
        assert rows[key]["saved_plan_basis"] == "assessment-baseline"
    assert deliverables.saved_plan_basis("future-unreviewed-writer") is None


def test_future_unclassified_writer_cannot_export_a_saved_plan(monkeypatch):
    monkeypatch.setitem(deliverables.SPECS, "future-unreviewed-writer", deliverables.SPECS["mop"])
    # Refusal precedes snapshot/plan parsing and any document writer invocation.
    with pytest.raises(deliverables.SavedPlanError, match="unsupported"):
        deliverables.generate_saved_plan("future-unreviewed-writer", {}, "EXAMPLE", {})


def _snapshot(*, count: int = 2):
    hosts = [f"EXAMPLE-SW-{index}" for index in range(count)]
    validation_item = {
        "device": hosts[0], "platform": "ios", "wave": "EXAMPLE-G1", "category": "Routing",
        "severity": "High", "check": "EXAMPLE observed routing baseline",
        "command": "show ip ospf neighbor", "expect": "PRE-CUTOVER DEGRADED",
        "why": "EXAMPLE observed fault remains a blocker",
        "evidence_state": "degraded", "projection_custody": "embedded_unverified",
        "source_key": "routing_neighbors.EXAMPLE-SW-0",
    }
    baseline = {
        "script_version": "3.23.0",
        "generated_at": "2026-10-15T00:00:00Z",
        "devices": {host: {"hostname": host, "platform": "ios"} for host in hosts},
        "interfaces": {},
        "move_groups": [{"group": "EXAMPLE-G1", "switches": hosts, "endpoints": 0}],
        "migration_readiness": [{
            "group": "EXAMPLE-G1", "switches": hosts, "readiness": "CAUTION",
            "n_fail": 0, "n_warn": 1, "checks": [],
        }],
        "wave_sequencing": [{
            "group": "EXAMPLE-G1", "make_before_break": [], "hard_cutover": hosts,
            "hard_cutover_endpoints": 0,
        }],
        "validation_plan": {
            "items": [validation_item],
            "by_wave": {"EXAMPLE-G1": [copy.deepcopy(validation_item)]},
            "summary": {"n_items": 1, "n_waves": 1, "n_high": 1,
                        "by_category": {"Routing": 1}},
        },
        "collection_completeness": {
            "summary": {"inventory": count, "complete": count, "not_collected": 0},
        },
        "executive_brief": {"scale": {"n_devices": count, "n_collected": count}},
        "design_blueprint": {"summary": {"n_decisions": 999}, "decisions": []},
        "design_nrfu": {"n_items": 999, "items": []},
    }
    return bind_snapshot_json_bytes(json.dumps(baseline, separators=(",", ":")).encode("utf-8"))


def _assessment(snapshot, **changes):
    source = bound_snapshot_source(snapshot)
    value = {
        "source": "persisted snapshots.snapshot_json blob",
        "sha256": source["sha256"], "bytes": source["bytes"],
        "snapshot_id": 11, "campaign_id": 7, "engagement_id": "EXAMPLE-ENGAGEMENT",
        "label": "EXAMPLE assessment", "script_version": snapshot["script_version"],
    }
    value.update(changes)
    return value


def _plan(snapshot, requirements=None, *, revision=1, selection_id=1):
    value = {
        "schema": "atlas.saved_plan/1", "revision": revision,
        "label": "EXAMPLE selected plan", "created_at": "2026-10-15T00:00:00Z",
        "assessment": _assessment(snapshot),
        "projection": deliverables.prepare_saved_plan(snapshot, requirements or {}),
        "status": "UNAPPROVED_DRAFT", "selection_id": selection_id,
    }
    value["sha256"] = saved_plan_digest(value)
    return value


def _office(path: str, extension: str = "docx"):
    with zipfile.ZipFile(path) as package:
        if extension == "docx":
            names = ["word/document.xml"]
            text_tag = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"
        else:
            names = sorted(name for name in package.namelist()
                           if name.startswith("ppt/slides/slide") and name.endswith(".xml"))
            text_tag = "{http://schemas.openxmlformats.org/drawingml/2006/main}t"
        texts = [node.text or "" for name in names
                 for node in ElementTree.fromstring(package.read(name)).iter(text_tag)]
        core = ElementTree.fromstring(package.read("docProps/core.xml"))
    properties = {node.tag.rsplit("}", 1)[-1]: node.text or "" for node in core}
    return "\n".join(texts), properties


def test_preparation_preserves_source_and_publishes_only_the_plan_owned_sections(monkeypatch):
    from cisco_toolkit import design_advisor

    calls = []
    original_late = design_advisor.derive_plan_sections

    def counted_late(snapshot, requirements, **kwargs):
        calls.append(kwargs["profile"])
        return original_late(snapshot, requirements, **kwargs)

    monkeypatch.setattr(design_advisor, "derive_plan_sections", counted_late)
    snapshot = _snapshot()
    original = copy.deepcopy(dict(snapshot))
    marker = bound_snapshot_source(snapshot)
    plan = _plan(snapshot, {
        "availability_tier": "Gold", "critical_apps": "voice, video",
        "convergence_budget_ms": "0", "fabric_operating_model": "ACI",
    })
    projection = plan["projection"]
    assert projection["requirements"] == {
        "availability_tier": "gold", "critical_apps": ["voice", "video"],
        "convergence_budget_ms": 0, "fabric_operating_model": "aci",
    }
    assert projection["derivation"]["implementation_identity"] is None
    assert projection["derivation"]["implementation_identity_status"] == "not_established"
    assert projection["derivation"]["authority"] == "none"
    assert projection["schema"] == "atlas.saved_plan_projection/2"
    assert projection["derivation"]["schema"] == "atlas.saved_plan_derivation/2"
    assert projection["derivation"]["profile"] == "saved_plan"
    assert set(projection["sections"]) == {"requirements_register", "design_blueprint", "design_nrfu"}
    rendered = deliverables.render_saved_plan(snapshot, plan)
    assert calls == ["saved_plan"]
    assert rendered["design_blueprint"]["summary"]["n_decisions"] != 999
    assert rendered["design_nrfu"]["n_items"] == len(rendered["design_nrfu"]["items"])
    recommended = {row["id"] for row in rendered["design_blueprint"]["decisions"]
                   if row["status"] == "recommended"}
    assert {row["decision_id"] for row in rendered["design_nrfu"]["items"]} == recommended
    assert {key: value for key, value in rendered.items() if key not in projection["sections"]} == \
        {key: value for key, value in original.items() if key not in projection["sections"]}
    for absent in ("architecture_coverage", "coverage_matrix", "unknown_evidence", "schema_census", "fact_lineage"):
        assert absent not in original and absent not in rendered
    assert rendered["validation_plan"] == original["validation_plan"]
    assert engine.compute_current_baseline_gate(rendered["validation_plan"]) == \
        engine.compute_current_baseline_gate(original["validation_plan"])
    assert dict(snapshot) == original and bound_snapshot_source(snapshot) == marker
    assert bound_snapshot_source(rendered)["source_bound"] is False
    rendered["devices"]["EXAMPLE-SW-0"]["hostname"] = "changed rendering copy"
    assert dict(snapshot) == original


def test_reopen_consumes_persisted_sections_without_current_planner_recomputation(monkeypatch):
    snapshot = _snapshot()
    plan = _plan(snapshot, {"constraints": "budget, security"})
    from cisco_toolkit import design_advisor

    def forbidden(*_args, **_kwargs):
        raise AssertionError("reopening must not rerun the current planner")

    monkeypatch.setattr(design_advisor, "compute_design_blueprint", forbidden)
    monkeypatch.setattr(design_advisor, "compute_design_nrfu", forbidden)
    monkeypatch.setattr(design_advisor, "derive_plan_sections", forbidden)
    reopened = deliverables.render_saved_plan(snapshot, json.loads(json.dumps(plan)))
    for key, value in plan["projection"]["sections"].items():
        assert reopened[key] == value


@pytest.mark.parametrize("kind", [
    "design", "crd", "mop", "nrfu", "deck", "engagement", "runbook", "archreview", "cutover", "opshandbook",
])
def test_every_plan_export_preserves_hostile_observed_baseline_and_source(kind, monkeypatch):
    from cisco_toolkit import coverage_matrix, design_advisor, ssot, unknown_evidence

    raw = copy.deepcopy(dict(_snapshot()))
    # These retained observations are deliberately inconsistent with a fresh derivation.
    # A renderer must retain them and disclose drift, never silently repair the assessment.
    raw["executive_brief"]["scale"]["n_devices"] = 701
    raw["executive_brief"]["ssot"] = {"verified": False, "n_violations": 701}
    raw.update({
        "architecture_coverage": {"classes": [], "summary": {"n_observed": 702}},
        "coverage_matrix": {"rows": [], "summary": {"n_assessed": 703}},
        "unknown_evidence": {"state": "unavailable", "events": [], "EXAMPLE_marker": 704},
        "schema_census": {"sections": [], "summary": {"n_sections": 705}},
        "fact_lineage": {"facts": [{"name": "n_design_decisions", "value": 706, "state": "unverified"}]},
        "assessment_integrity": {
            "failed_phases": ["Design blueprint", "requirements"],
            "phase_failures": [{"phase": "Design blueprint", "detail": "EXAMPLE original failure"}],
        },
    })
    snapshot = bind_snapshot_json_bytes(json.dumps(raw, separators=(",", ":")).encode("utf-8"))
    original, marker = copy.deepcopy(dict(snapshot)), bound_snapshot_source(snapshot)
    assert ssot.audit(snapshot), "the observation fixture must actually require a drift disclosure"

    def forbidden(*_args, **_kwargs):
        raise AssertionError("saved plans must not rederive observed assessment sections")

    for module, name in (
        (design_advisor, "compute_architecture_coverage"),
        (coverage_matrix, "compute_coverage_matrix"),
        (unknown_evidence, "compute_unknown_evidence"),
        (ssot, "summary"), (ssot, "compute_schema_census"), (ssot, "compute_fact_lineage"),
    ):
        monkeypatch.setattr(module, name, forbidden)
    plan = _plan(snapshot, {"availability_tier": "gold", "constraints": ["EXAMPLE selected constraint"]})
    original_plan = copy.deepcopy(plan)
    sections = plan["projection"]["sections"]
    assert set(sections) == {"requirements_register", "design_blueprint", "design_nrfu"}
    rendered = deliverables.render_saved_plan(snapshot, plan)
    observed_keys = set(original) - set(sections)
    assert {key: rendered[key] for key in observed_keys} == {key: original[key] for key in observed_keys}
    assert rendered["design_blueprint"] != original["design_blueprint"]
    derived = kind in {"design", "crd", "mop", "nrfu", "deck"}
    expected_drift = ssot.audit(rendered if derived else original)
    assert expected_drift
    seen = []

    def capture_writer(path, writer_snapshot, _label, **_kwargs):
        seen.append(copy.deepcopy(writer_snapshot))
        # Attack the actual input, not the spy's retained copy: neither the stored assessment
        # nor the persisted selected revision may alias a renderer-owned nested container.
        writer_snapshot["executive_brief"]["scale"]["n_devices"] = -1
        writer_snapshot["design_blueprint"]["summary"]["n_decisions"] = -2
        Path(path).write_bytes(b"EXAMPLE writer boundary control")

    # Keep the actual generate() dispatch and reconciliation guard. Existing Office tests
    # separately exercise each real writer and the three stamping operations.
    monkeypatch.setattr(deliverables, "resolve_writer", lambda _spec: capture_writer)
    monkeypatch.setattr(deliverables, "_stamp_ssot_integrity", lambda *_args: None)
    monkeypatch.setattr(deliverables, "_stamp_unapproved_draft", lambda *_args: None)
    monkeypatch.setattr(deliverables, "_stamp_saved_plan", lambda *_args, **_kwargs: None)
    portfolio = build_protocol_single_snapshot_bundle(snapshot, plan["assessment"])
    path = deliverables.generate_saved_plan(
        kind, snapshot, "EXAMPLE observed preservation", plan,
        protocol_assurance_bundle=portfolio, document_gate=None,
    )
    try:
        assert len(seen) == 1
        writer_snapshot = seen[0]
        for key in observed_keys - {"assessment_integrity"}:
            assert json.dumps(writer_snapshot[key], sort_keys=True) == json.dumps(original[key], sort_keys=True), key
        if derived:
            for key, value in sections.items():
                assert writer_snapshot[key] == value
        else:
            assert {key: value for key, value in writer_snapshot.items() if key != "assessment_integrity"} == \
                {key: value for key, value in original.items() if key != "assessment_integrity"}
        # Existing render-time audit disclosure is still active and retains original failures.
        assert writer_snapshot["assessment_integrity"] == {**original["assessment_integrity"], **expected_drift}
        assert writer_snapshot["assessment_integrity"]["failed_phases"] == ["Design blueprint", "requirements"]
        assert dict(snapshot) == original and bound_snapshot_source(snapshot) == marker
        assert plan == original_plan
    finally:
        Path(path).unlink()


@pytest.mark.parametrize("section", [
    "executive_brief", "architecture_coverage", "coverage_matrix", "unknown_evidence",
    "schema_census", "fact_lineage", "assessment_integrity",
])
def test_rehashed_observed_section_injection_is_refused_for_every_plan_export(section):
    snapshot = _snapshot()
    plan = _plan(snapshot)
    plan["projection"]["sections"][section] = {"EXAMPLE": "re-authored observed facts"}
    plan["sha256"] = saved_plan_digest(plan)
    for kind in ("design", "crd", "mop", "nrfu", "deck"):
        with pytest.raises(deliverables.SavedPlanError, match="derived sections"):
            deliverables.generate_saved_plan(kind, snapshot, "EXAMPLE refused", plan)


@pytest.mark.parametrize("legacy", ["projection", "derivation", "observed-sections", "old-helper-owner"])
def test_old_projection_profiles_refuse_instead_of_reinterpreting_saved_bytes(legacy):
    snapshot = _snapshot()
    plan = _plan(snapshot)
    if legacy == "projection":
        plan["projection"]["schema"] = "atlas.saved_plan_projection/1"
    elif legacy == "derivation":
        plan["projection"]["derivation"]["schema"] = "atlas.saved_plan_derivation/1"
    elif legacy == "old-helper-owner":
        plan["projection"]["derivation"]["owners"][1] = "cisco_toolkit.ssot.derive_late_snapshot_sections"
    else:
        for key in ("architecture_coverage", "coverage_matrix", "unknown_evidence", "schema_census", "fact_lineage"):
            plan["projection"]["sections"][key] = {}
        plan["projection"]["sections"]["executive_brief"] = copy.deepcopy(snapshot["executive_brief"])
    plan["sha256"] = saved_plan_digest(plan)
    before = copy.deepcopy(plan)
    with pytest.raises(deliverables.SavedPlanError):
        deliverables.render_saved_plan(snapshot, plan)
    assert plan == before


@pytest.mark.parametrize("part", ["projection", "derivation"])
def test_rehashed_unexpected_projection_metadata_is_not_silently_ignored(part):
    snapshot = _snapshot()
    plan = _plan(snapshot)
    target = plan["projection"] if part == "projection" else plan["projection"]["derivation"]
    target["EXAMPLE_unrecognized_profile_field"] = True
    plan["sha256"] = saved_plan_digest(plan)
    with pytest.raises(deliverables.SavedPlanError):
        deliverables.render_saved_plan(snapshot, plan)


@pytest.mark.parametrize("requirements", [
    {"unknown": "EXAMPLE"}, {"availability_tier": "platinum"},
    {"convergence_budget_ms": True}, {"convergence_budget_ms": -1},
    {"convergence_budget_ms": "fast"}, {"critical_apps": ["voice", {}]},
    {"convergence_budget_ms": {}}, {"constraints": {}}, {"availability_tier": []},
    {"fabric_operating_model": "unknown"}, {"address_space": "invalid"},
    {"address_space": "2001:db8::/32"}, {"vlan_zones": {"guest": [True]}},
    {"vlan_zones": {"guest": [10], "storage": [10]}},
    {"vlan_zones": {"guest": [10], "storage": "invalid"}},
    {"constraints": float("nan")}, {"growth_horizon": "x" * 65536},
])
def test_invalid_or_oversized_requirements_are_explicit(requirements):
    with pytest.raises(deliverables.SavedPlanError):
        deliverables.prepare_saved_plan(_snapshot(), requirements)


def test_blank_optional_values_and_vlan_keys_reopen_canonically():
    snapshot = _snapshot()
    plan = _plan(snapshot, {
        "critical_apps": " , ", "growth_horizon": " ",
        "convergence_budget_ms": 0, "vlan_zones": {"guest": [10, "20"]},
    })
    assert plan["projection"]["requirements"] == {
        "convergence_budget_ms": 0, "vlan_zones": {"10": "guest", "20": "guest"},
    }
    assert deliverables.render_saved_plan(snapshot, plan)["requirements_register"] == \
        plan["projection"]["requirements"]


@pytest.mark.parametrize("owner", [None, "", "   ", 3230, False])
def test_unknown_script_owner_refuses_before_a_revision_can_be_persisted(owner):
    snapshot = {"devices": {"EXAMPLE-SW": {}}}
    if owner is not None:
        snapshot["script_version"] = owner
    with pytest.raises(deliverables.SavedPlanError, match="explicit assessment script owner"):
        deliverables.prepare_saved_plan(snapshot, {})


@pytest.mark.parametrize("field,value", [
    ("selection_id", True), ("selection_id", "1"), ("selection_id", 0),
    ("selection_id", 1 << 63),
    ("revision", True), ("revision", "1"), ("revision", 0),
    ("revision", 1 << 63),
])
def test_revision_and_selection_are_strict_positive_integer_identities(field, value):
    snapshot = _snapshot()
    plan = _plan(snapshot)
    plan[field] = value
    with pytest.raises(deliverables.SavedPlanError):
        deliverables.render_saved_plan(snapshot, plan)


@pytest.mark.parametrize("field", ["snapshot_id", "campaign_id"])
@pytest.mark.parametrize("value", [False, 0, -1, 1 << 63])
def test_rehashed_assessment_identifiers_must_remain_strict_positive_sqlite_integers(field, value):
    snapshot = _snapshot()
    plan = _plan(snapshot)
    plan["assessment"][field] = value
    plan["sha256"] = saved_plan_digest(plan)
    with pytest.raises(deliverables.SavedPlanError, match="assessment binding is invalid"):
        deliverables.render_saved_plan(snapshot, plan)


@pytest.mark.parametrize("where", ["requirements_register", "requirements_model", "assessment_facts"])
@pytest.mark.parametrize("number_type", ["boolean", "float"])
def test_rehashed_boolean_integer_and_float_aliases_never_pass_typed_plan_consistency(where, number_type):
    snapshot = _snapshot(count=1)
    plan = _plan(snapshot, {"convergence_budget_ms": 0})
    sections = plan["projection"]["sections"]
    if where == "requirements_register":
        sections[where]["convergence_budget_ms"] = False if number_type == "boolean" else 0.0
    elif where == "requirements_model":
        field = next(row for row in sections["design_blueprint"]["requirements_model"]["fields"]
                     if row["key"] == "convergence_budget_ms")
        field["value"] = False if number_type == "boolean" else 0.0
    else:
        # Observed sections cannot enter the projection even after rehashing a typed alias.
        sections["executive_brief"] = copy.deepcopy(snapshot["executive_brief"])
        sections["executive_brief"]["scale"]["n_devices"] = True if number_type == "boolean" else 1.0
    plan["sha256"] = saved_plan_digest(plan)
    with pytest.raises(deliverables.SavedPlanError):
        deliverables.render_saved_plan(snapshot, plan)


def test_preparation_propagates_the_shared_projection_allocation_bound_as_a_limit_error(monkeypatch):
    from backend.storage import SavedPlanLimitError, SAVED_PLAN_MAX_BYTES
    from cisco_toolkit import design_advisor

    original_blueprint = design_advisor.compute_design_blueprint

    def oversized_blueprint(*args, **kwargs):
        result = original_blueprint(*args, **kwargs)
        result["EXAMPLE-oversized-plan-field"] = "x" * SAVED_PLAN_MAX_BYTES
        return result

    monkeypatch.setattr(design_advisor, "compute_design_blueprint", oversized_blueprint)
    with pytest.raises(SavedPlanLimitError):
        deliverables.prepare_saved_plan(_snapshot(), {})


def test_large_observed_brief_is_preserved_without_becoming_persisted_plan_data():
    from backend.storage import SAVED_PLAN_MAX_BYTES, saved_plan_json_bytes

    snapshot = copy.deepcopy(dict(_snapshot()))
    observed = "x" * SAVED_PLAN_MAX_BYTES
    snapshot["executive_brief"]["EXAMPLE-oversized-observation"] = observed
    original = copy.deepcopy(snapshot)
    projection = deliverables.prepare_saved_plan(snapshot, {})
    assert len(saved_plan_json_bytes(projection)) < SAVED_PLAN_MAX_BYTES
    assert "executive_brief" not in projection["sections"]
    assert snapshot == original and snapshot["executive_brief"]["EXAMPLE-oversized-observation"] == observed


def test_digest_cross_source_and_derived_section_tampering_refuse():
    snapshot = _snapshot()
    original = _plan(snapshot, {"fabric_operating_model": "nxos-evpn", "data_classification": ["restricted"]})
    changed = copy.deepcopy(original)
    changed["projection"]["requirements"]["growth_horizon"] = "different"
    with pytest.raises(deliverables.SavedPlanError, match="digest"):
        deliverables.render_saved_plan(snapshot, changed)
    with pytest.raises(deliverables.SavedPlanError, match="original bound assessment"):
        deliverables.render_saved_plan(_snapshot(count=3), original)
    changed = copy.deepcopy(original)
    changed["projection"]["sections"]["validation_plan"] = {"items": []}
    changed["sha256"] = saved_plan_digest(changed)
    with pytest.raises(deliverables.SavedPlanError, match="derived sections"):
        deliverables.render_saved_plan(snapshot, changed)
    changed = copy.deepcopy(original)
    changed["projection"]["sections"]["executive_brief"] = copy.deepcopy(snapshot["executive_brief"])
    changed["projection"]["sections"]["executive_brief"]["scale"]["n_devices"] = 999
    changed["sha256"] = saved_plan_digest(changed)
    with pytest.raises(deliverables.SavedPlanError, match="derived sections"):
        deliverables.render_saved_plan(snapshot, changed)
    changed = copy.deepcopy(original)
    changed["projection"]["sections"]["design_nrfu"]["items"] = []
    changed["projection"]["sections"]["design_nrfu"]["n_items"] = 0
    changed["sha256"] = saved_plan_digest(changed)
    # The positive premise is explicit: removing these tests must create a genuine coverage gap.
    assert any(row["status"] == "recommended"
               for row in changed["projection"]["sections"]["design_blueprint"]["decisions"])
    with pytest.raises(deliverables.SavedPlanError, match="acceptance items"):
        deliverables.render_saved_plan(snapshot, changed)


@pytest.mark.parametrize("kind", [
    "engagement", "crd", "runbook", "design", "archreview", "mop", "cutover", "nrfu",
    "opshandbook", "deck",
])
def test_actual_office_outputs_keep_visible_identity_and_original_protocol_custody(kind):
    snapshot = _snapshot()
    original = copy.deepcopy(dict(snapshot))
    plan = _plan(snapshot, {"fabric_operating_model": "nxos-evpn", "convergence_budget_ms": 0})
    portfolio = build_protocol_single_snapshot_bundle(snapshot, plan["assessment"])
    gate = {"status": "pending", "missing": ["lld_approved"], "revoked": ["assessment_approved"]}
    path = deliverables.generate_saved_plan(
        kind, snapshot, "EXAMPLE selected export", plan,
        protocol_assurance_bundle=portfolio, document_gate=gate,
    )
    try:
        text, core = _office(path, deliverables.SPECS[kind].ext)
        for token in ("UNAPPROVED DRAFT", "Revision 1", "Selection 1",
                      "Original assessment SHA-256: " + plan["assessment"]["sha256"],
                      "Selected plan SHA-256: " + plan["sha256"]):
            assert token in text
        assert "Original assessment protocol evidence" in text
        assert plan["assessment"]["sha256"][7:] in core["identifier"]
        assert plan["sha256"][7:] in core["identifier"]
        assert "UNAPPROVED DRAFT" in core["contentStatus"]
        if kind in {"design", "mop"}:
            assert core["keywords"] == (
                "atlas_gate_status=pending;missing=lld_approved;revoked=assessment_approved"
            )
        if kind in {"mop", "runbook", "nrfu"}:
            assert portfolio["receipt"]["receipt_sha256"] in text
            assert "This renderer did not hash the parsed snapshot or mint source custody" not in text
        if kind in {"cutover", "engagement", "runbook", "archreview", "opshandbook"}:
            assert "Original assessment baseline/reference associated" in text
            assert "do not approve this plan revision" in text
        else:
            assert "Plan-derived output" in text
    finally:
        Path(path).unlink()
    assert dict(snapshot) == original and bound_snapshot_source(snapshot)["source_bound"] is True


def test_mop_wave_and_nrfu_coverage_consume_the_same_persisted_selected_design():
    snapshot = _snapshot(count=45)
    baseline_gate = engine.compute_current_baseline_gate(snapshot["validation_plan"])
    assert baseline_gate["integrity"] == {"valid": True, "failures": []}
    assert baseline_gate["verdict"] == "BLOCKED"
    assert baseline_gate["summary"]["n_blockers"] == 1
    plan = _plan(snapshot, {"fabric_operating_model": "nxos-evpn", "data_classification": ["restricted"]})
    rendered = deliverables.render_saved_plan(snapshot, plan)
    plan_before = copy.deepcopy(plan)
    snapshot_before = copy.deepcopy(dict(snapshot))
    blueprint = plan["projection"]["sections"]["design_blueprint"]
    waves = blueprint["target_state"]["wave_plan"]["waves"]
    assert [len(wave["switches"]) for wave in waves] == [40, 5]
    assert blueprint["evpn_migration"]["applicable"] is True
    assert blueprint["requirements_model"]["provided"] is True
    recommended = [row for row in blueprint["decisions"] if row["status"] == "recommended"]
    assert recommended
    assert {row["decision_id"] for row in rendered["design_nrfu"]["items"]} == \
        {row["id"] for row in recommended}
    portfolio = build_protocol_single_snapshot_bundle(snapshot, plan["assessment"])
    for kind in ("mop", "nrfu"):
        path = deliverables.generate_saved_plan(
            kind, snapshot, "EXAMPLE coherent plan", plan,
            protocol_assurance_bundle=portfolio, document_gate=None,
        )
        try:
            text, _core = _office(path)
            if kind == "mop":
                assert "Wave 1 (coupled-subwave)" in text and "Wave 2 (coupled-subwave)" in text
                assert "customer-confirmed via the requirements register" in text
            else:
                assert recommended[0]["title"] in text
                assert "Tested — " in text
                assert "NO acceptance test derived" not in text
                from docx import Document

                persisted = plan["projection"]["sections"]["design_nrfu"]
                coverage_table = next(table for table in Document(path).tables
                                      if [cell.text for cell in table.rows[0].cells] ==
                                      ["Target-state design decision", "Priority", "NRFU coverage"])
                coverage = {row.cells[0].text: row.cells[2].text for row in coverage_table.rows[1:]}
                for decision in recommended:
                    case = next(item for item in persisted["items"] if item["decision_id"] == decision["id"])
                    complete = all(isinstance(case.get(key), str) and case[key].strip()
                                   for key in ("description", "setup", "pass_criteria"))
                    complete = complete and bool(case["devices"]) and all(
                        isinstance(device, str) and device.strip() for device in case["devices"])
                    if complete:
                        assert coverage[decision["title"]].startswith("Tested — ")
                    else:
                        assert "NOT COVERED" in coverage[decision["title"]]
                        assert "Tested" not in coverage[decision["title"]]
                assert persisted["items"] and persisted["evpn_acceptance"]
                # Assert real saved instructions, not merely titles/coverage labels.
                for item in persisted["items"]:
                    assert item["description"] and item["setup"] and item["pass_criteria"]
                    for field in ("description", "setup", "pass_criteria"):
                        assert item[field] in text
                    for device in item["devices"]:
                        assert device in text
                for item in persisted["evpn_acceptance"]:
                    assert item["description"] in text
                    assert item["pass_criteria"] in text
            assert "PRE-CUTOVER DEGRADED" in text
        finally:
            Path(path).unlink()
    assert plan == plan_before
    assert dict(snapshot) == snapshot_before


def test_malformed_original_baseline_stays_indeterminate_in_selected_plan_mop():
    # Reproduce the former inconsistent synthetic baseline: neither its missing row owner
    # nor its contradictory grouping/summary may become an apparent observed blocker.
    malformed = copy.deepcopy(dict(_snapshot()))
    validation = malformed["validation_plan"]
    validation["items"][0].pop("platform")
    validation["by_wave"] = {}
    validation["summary"].pop("by_category")
    snapshot = bind_snapshot_json_bytes(json.dumps(malformed, separators=(",", ":")).encode("utf-8"))
    baseline_gate = engine.compute_current_baseline_gate(validation)
    assert baseline_gate["verdict"] == "INDETERMINATE"
    assert baseline_gate["integrity"]["valid"] is False
    assert baseline_gate["blockers"] == []
    assert baseline_gate["summary"]["n_blockers"] == 0
    plan = _plan(snapshot, {"fabric_operating_model": "nxos-evpn"})
    rendered = deliverables.render_saved_plan(snapshot, plan)
    assert rendered["validation_plan"] == validation
    assert engine.compute_current_baseline_gate(rendered["validation_plan"]) == baseline_gate
    portfolio = build_protocol_single_snapshot_bundle(snapshot, plan["assessment"])
    path = deliverables.generate_saved_plan(
        "mop", snapshot, "EXAMPLE malformed original baseline", plan,
        protocol_assurance_bundle=portfolio, document_gate=None,
    )
    try:
        text, _core = _office(path)
        assert "HOLD — current baseline INDETERMINATE" in text
        assert "EXAMPLE observed routing baseline" not in text
    finally:
        Path(path).unlink()


@pytest.mark.parametrize("phase", ["Design blueprint", "requirements"])
def test_design_export_honestly_retains_inherited_assessment_failure_limits(phase):
    raw = copy.deepcopy(dict(_snapshot()))
    raw["assessment_integrity"] = {
        "failed_phases": [phase],
        "phase_failures": [{"phase": phase, "detail": "EXAMPLE original assessment failure"}],
    }
    snapshot = bind_snapshot_json_bytes(json.dumps(raw, separators=(",", ":")).encode("utf-8"))
    marker = bound_snapshot_source(snapshot)
    plan = _plan(snapshot, {"fabric_operating_model": "nxos-evpn"})
    assert any(row["status"] == "recommended"
               for row in plan["projection"]["sections"]["design_blueprint"]["decisions"])
    rendered = deliverables.render_saved_plan(snapshot, plan)
    assert rendered["assessment_integrity"] == raw["assessment_integrity"]
    path = deliverables.generate_saved_plan("design", snapshot, "EXAMPLE inherited limit", plan,
                                           document_gate=None)
    try:
        text, _core = _office(path)
        assert "UNVERIFIED - target-state design synthesis did not complete for this run" in text
        assert "4.2 Recommended target-state design decisions" not in text
        assert "Original assessment integrity failures remain in force" in text
        assert "affected sections may remain withheld or unverified" in text
        assert "UNAPPROVED DRAFT" in text
    finally:
        Path(path).unlink()
    assert dict(snapshot) == raw and bound_snapshot_source(snapshot) == marker


def test_tampered_and_valid_cross_campaign_sidecars_refuse_without_dropping_controls():
    from cisco_toolkit.docmeta import protocol_assurance_receipt_view

    snapshot = _snapshot()
    plan = _plan(snapshot)
    valid = build_protocol_single_snapshot_bundle(snapshot, plan["assessment"])
    tampered = copy.deepcopy(valid)
    tampered["receipt"]["source_binding"]["bytes"] += 1
    cross = build_protocol_single_snapshot_bundle(snapshot, _assessment(snapshot, campaign_id=99))
    assert protocol_assurance_receipt_view(cross)["valid"] is True
    for bundle in (None, tampered, cross):
        with pytest.raises(deliverables.SavedPlanError, match="protocol sidecar"):
            deliverables.generate_saved_plan(
                "mop", snapshot, "EXAMPLE refused plan", plan,
                protocol_assurance_bundle=bundle, document_gate=None,
            )


def test_baseline_generate_keeps_its_existing_bytes_and_never_stamps_a_plan(tmp_path, monkeypatch):
    from docx import Document

    fixture = tmp_path / "baseline.docx"
    document = Document()
    document.add_paragraph("EXAMPLE baseline output")
    document.core_properties.identifier = "EXAMPLE-original-identity"
    document.save(fixture)
    baseline_bytes = fixture.read_bytes()

    def write(path, _snapshot, _label):
        Path(path).write_bytes(baseline_bytes)

    def forbidden(*_args, **_kwargs):
        raise AssertionError("ordinary exports must not gain a selected plan")

    monkeypatch.setattr(deliverables, "resolve_writer", lambda _spec: write)
    monkeypatch.setattr(deliverables, "_reconcile_gate", lambda *_args: None)
    monkeypatch.setattr(deliverables, "_stamp_saved_plan", forbidden)
    path = deliverables.generate("crd", {"devices": {}}, "EXAMPLE baseline", document_gate=None)
    try:
        assert Path(path).read_bytes() == baseline_bytes
        text, core = _office(path)
        assert "Selected plan SHA-256" not in text
        assert core["identifier"] == "EXAMPLE-original-identity"
    finally:
        Path(path).unlink()

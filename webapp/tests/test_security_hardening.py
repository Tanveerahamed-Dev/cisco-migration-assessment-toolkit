"""Client-data confidentiality hardening (Plan A / Tier-1 #4).

The confirmed exposure this closes: with `allow_origins=['*']` and a zero-auth API, ANY
open browser tab could cross-origin read client topology / IPs / serials / parsed configs
and forge gate sign-offs against the default localhost bind; the unsandboxed same-origin
explorer iframe was the innerHTML -> parent -> API escape path.

Model under test:
- CORS: localhost-only origin regex by default (+ ASSESSHUB_CORS_ORIGINS extras).
- ASSESSHUB_TOKEN set   -> Bearer required on ALL /api (except /api/health liveness + OPTIONS).
- ASSESSHUB_TOKEN unset -> /api serves LOOPBACK clients only (dev UX unchanged; the ASGI
  test harness counts as loopback), 403 with an actionable message otherwise.
- ASSESSHUB_TOKEN unset -> the Host header must ALSO name a loopback target (or an
  ASSESSHUB_ALLOWED_HOSTS entry): a DNS-rebinding page reaches a loopback peer with a foreign
  Host, so loopback position alone can't authorize the (blind) write. Token mode is Host-agnostic.
- The explorer iframe is sandboxed WITHOUT allow-same-origin (the explorer feature-detects
  storage for opaque origins, so this is loss-free)."""
import hashlib
import sys
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))  # make `backend` importable

from backend import app as app_module  # noqa: E402
from backend import deliverables as deliverables_module  # noqa: E402
from backend.app import create_app  # noqa: E402
from frontend_fixture import write_frontend_dist  # noqa: E402

_REPO = Path(__file__).resolve().parents[2]


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.delenv("ASSESSHUB_TOKEN", raising=False)
    monkeypatch.delenv("ASSESSHUB_ALLOWED_HOSTS", raising=False)
    a = create_app(db_path=str(tmp_path / "test.db"))
    # base_url=localhost so the default Host passes the no-token DNS-rebinding guard; the tests
    # that model rebinding pass an explicit foreign Host header (the ASGI peer stays the loopback
    # transport, so request.client.host is still loopback — exactly the rebinding shape).
    with TestClient(a, base_url="http://localhost") as c:
        yield c


# ---------------------------------------------------------------- CORS
def test_internet_origin_gets_no_cors_allow(client):
    r = client.get("/api/health", headers={"Origin": "https://evil.example"})
    assert "access-control-allow-origin" not in {k.lower() for k in r.headers}


def test_localhost_origin_is_allowed(client):
    r = client.get("/api/health", headers={"Origin": "http://localhost:5173"})
    assert r.headers.get("access-control-allow-origin") == "http://localhost:5173"


def test_preflight_from_internet_origin_is_refused(client):
    r = client.options("/api/campaigns", headers={
        "Origin": "https://evil.example",
        "Access-Control-Request-Method": "POST",
    })
    assert "access-control-allow-origin" not in {k.lower() for k in r.headers}


# ------------------------------------------------- no-token mode (loopback-only)
def test_loopback_dev_flow_needs_no_token(client):
    """The zero-config localhost workflow must be untouched: seed + read + write."""
    assert client.get("/api/campaigns").status_code == 200
    assert client.post("/api/demo/seed").status_code == 200


def test_nonloopback_client_is_refused_without_token(client, monkeypatch):
    """Bound non-loopback with no token configured -> every /api data route refuses,
    with a message that says HOW to enable remote access."""
    monkeypatch.setattr(app_module, "_client_is_loopback", lambda request: False)
    r = client.get("/api/campaigns")
    assert r.status_code == 403
    assert "ASSESSHUB_TOKEN" in r.json()["detail"]
    assert client.post("/api/demo/seed").status_code == 403
    # liveness stays open for monitoring probes (no client data in it)
    assert client.get("/api/health").status_code == 200


# ------------------------------------------------- token mode
def test_token_gates_all_api_routes(client, monkeypatch):
    monkeypatch.setenv("ASSESSHUB_TOKEN", "s3cret-token")
    assert client.get("/api/campaigns").status_code == 401
    assert client.get("/api/campaigns",
                      headers={"Authorization": "Bearer wrong"}).status_code == 401
    ok = client.get("/api/campaigns", headers={"Authorization": "Bearer s3cret-token"})
    assert ok.status_code == 200
    assert client.post("/api/demo/seed",
                       headers={"Authorization": "Bearer s3cret-token"}).status_code == 200
    # health liveness stays open even in token mode
    assert client.get("/api/health").status_code == 200


def test_token_applies_to_loopback_too(client, monkeypatch):
    """One mental model: once a token is configured it is required everywhere — a local
    process is not implicitly trusted on a multi-user machine."""
    monkeypatch.setenv("ASSESSHUB_TOKEN", "s3cret-token")
    assert client.post("/api/demo/seed").status_code == 401


# ----------------------------------- cross-site write / blind-CSRF protection
# CORS hides RESPONSES from foreign origins, but `multipart/form-data` and empty-body POSTs
# are CORS "simple requests" that need no preflight and still EXECUTE. Without a request-side
# guard a page on evil.example could blind-fire snapshot upload / ZIP ingest (a heavy engine
# subprocess) / demo-seed against the victim's zero-token loopback bind — store pollution +
# resource-exhaustion DoS. Every state-changing method is checked against the CORS allowlist.
def test_cross_origin_simple_post_is_rejected(client):
    """The pinned regression: a foreign-origin simple-request POST to a write route must be
    refused BEFORE it executes (this exact call returned 200 before the fix)."""
    r = client.post("/api/demo/seed", headers={"Origin": "https://evil.example"})
    assert r.status_code == 403
    assert "cross-site" in r.json()["detail"].lower()
    # and it must not have run: no campaign was created by the refused request
    assert client.get("/api/campaigns").json() == []


def test_cross_origin_multipart_upload_is_rejected(client):
    """The multipart upload route (a CORS simple request) is refused at the guard, before the
    body is even parsed — a 403, not a 404/422 from the route."""
    r = client.post("/api/campaigns/1/snapshots",
                    headers={"Origin": "https://evil.example"},
                    files={"file": ("x.json", b'{"devices": []}', "application/json")})
    assert r.status_code == 403


def test_cross_site_write_rejected_via_fetch_metadata(client):
    """Fallback signal: a cross-site write that (unusually) carries no Origin is still caught by
    the browser-set, JS-unforgeable Sec-Fetch-Site header."""
    r = client.post("/api/demo/seed", headers={"Sec-Fetch-Site": "cross-site"})
    assert r.status_code == 403


def test_cross_site_delete_is_rejected(client):
    """DELETE is preflight-protected in real browsers, but the server-side guard refuses a
    foreign-origin DELETE too (defense in depth, one rule for all unsafe methods)."""
    r = client.delete("/api/snapshots/1", headers={"Origin": "https://evil.example"})
    assert r.status_code == 403


def test_same_origin_write_is_allowed(client):
    """The SPA's own writes carry a localhost Origin (or Sec-Fetch-Site: same-origin) and must
    pass untouched — the write-side guard mirrors the read-side CORS allowlist, nothing more."""
    assert client.post("/api/demo/seed",
                       headers={"Origin": "http://localhost:5173"}).status_code == 200
    assert client.post("/api/demo/seed",
                       headers={"Sec-Fetch-Site": "same-origin"}).status_code == 200
    # a non-browser client (curl / TestClient) sends neither header and is likewise allowed —
    # this is the zero-config loopback path pinned by test_loopback_dev_flow_needs_no_token.
    # KNOWN, DELIBERATE residual: a request with NEITHER header is trusted, so a pre-2019 browser
    # that omits both on a cross-site write is undefended. That trade-off is load-bearing — denying
    # it would break every non-browser client (curl / the ASGI harness) and the zero-config UX —
    # and modern browsers (Chrome 76+/FF 90+/Safari 16.4+) always send at least one of the two.


def test_origin_null_write_is_rejected(client):
    """`Origin: null` is the classic allowlist-bypass: a foreign page mints it via a sandboxed
    iframe + data: URL (and Chrome sends it on cross-site redirects). It is neither a localhost
    origin nor a configured extra, so it must be refused — never whitelisted 'for local dev'."""
    assert client.post("/api/demo/seed", headers={"Origin": "null"}).status_code == 403


def test_localhost_lookalike_origins_are_rejected(client):
    """The allowlist regex is fully anchored: an attacker domain that merely embeds a trusted host
    must NOT match. Covers the suffix, subdomain, userinfo-`@`, and fragment/query tricks."""
    for origin in ("http://localhost.evil.example",
                   "http://127.0.0.1.evil.example",
                   "https://evil.example/localhost",
                   "http://localhost:8000@evil.example",
                   "http://evil.example#http://localhost"):
        r = client.post("/api/demo/seed", headers={"Origin": origin})
        assert r.status_code == 403, f"lookalike origin was not refused: {origin}"


def test_cross_site_ingest_is_rejected_before_the_subprocess(client):
    """The heaviest route (ZIP ingest spawns a bounded-but-expensive engine subprocess) is the
    prime resource-exhaustion target — the guard must refuse a foreign-origin ingest at the door,
    before any body is read or subprocess is launched."""
    r = client.post("/api/campaigns/1/ingest",
                    headers={"Origin": "https://evil.example"},
                    files={"file": ("c.zip", b"PK\x03\x04not-a-real-zip", "application/zip")})
    assert r.status_code == 403


def test_cross_site_write_refused_even_in_token_mode(client, monkeypatch):
    """CSRF protection runs BEFORE the Bearer check, so a foreign-origin write is refused in token
    mode too — even carrying the correct token (defense in depth; one rule for all writes)."""
    monkeypatch.setenv("ASSESSHUB_TOKEN", "s3cret-token")
    r = client.post("/api/demo/seed",
                    headers={"Origin": "https://evil.example",
                             "Authorization": "Bearer s3cret-token"})
    assert r.status_code == 403


# --- production (non-localhost single-origin) deployment: same-origin writes must WORK -------
# The documented production mode (app.py module docstring: "the whole platform runs from one origin
# in production") is a non-loopback bind behind ASSESSHUB_TOKEN, served at e.g. https://assesshub.example.com.
# The SPA's own writes are SAME-ORIGIN there — they must be allowed even though the origin is not
# localhost. The guard must trust the browser's Sec-Fetch-Site oracle (primary), falling back to an
# Origin-equals-own-host comparison for pre-Fetch-Metadata browsers — NOT a hardcoded localhost list.
def test_prod_same_origin_write_allowed_via_fetch_metadata(client):
    """Modern browser on a non-localhost deployment: Sec-Fetch-Site: same-origin is authoritative."""
    r = client.post("/api/demo/seed", headers={"Origin": "https://assesshub.example.com",
                                               "Sec-Fetch-Site": "same-origin"})
    assert r.status_code == 200


def test_prod_same_origin_write_allowed_via_origin_host_match(client, monkeypatch):
    """Pre-Fetch-Metadata browser (no Sec-Fetch-Site): Origin equals the host the request was addressed
    to => same-origin => allowed. Under the stricter combined model, a non-localhost deployment must ALSO
    ALLOWLIST its hostname (ASSESSHUB_ALLOWED_HOSTS) to clear the no-token DNS-rebinding Host guard (#384):
    a bare-hostname no-token prod bind is rebinding-vulnerable and otherwise refused."""
    monkeypatch.setenv("ASSESSHUB_ALLOWED_HOSTS", "assesshub.example.com")
    r = client.post("/api/demo/seed", headers={"Origin": "http://assesshub.example.com:8000",
                                               "Host": "assesshub.example.com:8000"})
    assert r.status_code == 200


def test_prod_cross_site_write_still_refused(client):
    """The fix must not over-open: on that same non-localhost deployment a genuine foreign origin is
    still refused — via Sec-Fetch-Site: cross-site, and via the no-metadata Origin!=Host fallback."""
    assert client.post("/api/demo/seed",
                       headers={"Host": "assesshub.example.com",
                                "Origin": "https://evil.example",
                                "Sec-Fetch-Site": "cross-site"}).status_code == 403
    assert client.post("/api/demo/seed",
                       headers={"Host": "assesshub.example.com",
                                "Origin": "https://evil.example"}).status_code == 403


def test_allowlisted_cross_origin_frontend_can_write(client, monkeypatch):
    """Split-origin deployment: the UI is a DIFFERENT site than the API, explicitly trusted via
    ASSESSHUB_CORS_ORIGINS (CORS already grants it reads). Its writes legitimately carry
    Sec-Fetch-Site: cross-site — an admin-allowlisted origin must be honored on the WRITE side too,
    for read/write parity. This must hold on a MODERN browser (Sec-Fetch-Site present)."""
    monkeypatch.setenv("ASSESSHUB_CORS_ORIGINS", "https://frontend.example.com")
    r = client.post("/api/demo/seed", headers={"Origin": "https://frontend.example.com",
                                               "Sec-Fetch-Site": "cross-site"})
    assert r.status_code == 200


def test_allowlist_does_not_extend_to_localhost_bypass(client):
    """The allowlist override must be EXACT and not leak the blanket localhost trust: a localhost-origin
    page issuing a cross-site write to a non-localhost deployment is still refused (Sec-Fetch-Site wins;
    only an explicit ASSESSHUB_CORS_ORIGINS entry overrides it, never the localhost regex)."""
    r = client.post("/api/demo/seed", headers={"Origin": "http://localhost:9999",
                                               "Sec-Fetch-Site": "cross-site"})
    assert r.status_code == 403


# ---------------------------------------- DNS rebinding (Host-header allowlist)
# A zero-token instance trusts loopback network position. An attacker who rebinds a domain they
# control to 127.0.0.1 reaches it from a loopback peer (so the loopback check passes) while the
# victim's browser still sends the ATTACKER's name in the Host header — a BLIND cross-origin write
# (CORS still hides the response, but the store-polluting / ingest-DoS side effect lands). Each
# case below sends a foreign `Host` header over the in-process (loopback) transport: exactly that
# shape (request.client.host stays "testclient" == loopback; only the Host header is attacker-named).
def test_dns_rebinding_read_is_refused(client):
    """A rebound page trying to read client data is blocked by the Host guard, even though it
    signals a same-origin fetch and presents a loopback peer."""
    r = client.get("/api/campaigns", headers={"Host": "evil.example",
                                              "Origin": "http://evil.example",
                                              "Sec-Fetch-Site": "same-origin"})
    assert r.status_code == 403
    assert "Host" in r.json()["detail"]


def test_dns_rebinding_write_is_refused(client):
    """The impactful vector: the blind write (store pollution + heavy ingest subprocess = a
    resource-exhaustion DoS) must be refused before it reaches a route."""
    r = client.post("/api/demo/seed", headers={"Host": "evil.example",
                                               "Origin": "http://evil.example",
                                               "Sec-Fetch-Site": "same-origin"})
    assert r.status_code == 403
    assert "Host" in r.json()["detail"]  # refused BY the host guard, not incidentally


@pytest.mark.parametrize("host", ["localhost", "localhost:8000", "127.0.0.1:8000",
                                  "[::1]", "[::1]:8000", "LOCALHOST"])
def test_loopback_hosts_pass(client, host):
    """The genuine local dev flow — any loopback Host, bracketed IPv6, mixed case, with or without
    a port — still reads AND writes, so the guard is a scalpel, not a wall."""
    assert client.get("/api/campaigns", headers={"Host": host}).status_code == 200
    assert client.post("/api/demo/seed", headers={"Host": host}).status_code == 200


def test_allowed_hosts_env_trusts_configured_host(client, monkeypatch):
    """ASSESSHUB_ALLOWED_HOSTS is the escape hatch for a trusted same-host reverse-proxy vhost;
    a host NOT on the list is still refused (a port variant of a listed host passes)."""
    monkeypatch.setenv("ASSESSHUB_ALLOWED_HOSTS", "assesshub.internal")
    assert client.get("/api/campaigns",
                      headers={"Host": "assesshub.internal:9000"}).status_code == 200
    assert client.get("/api/campaigns",
                      headers={"Host": "not-listed.example"}).status_code == 403


def test_health_stays_open_under_foreign_host(client):
    """Liveness carries no client data and no side effect, so it stays reachable for monitoring
    probes even under an unrecognized Host (mirrors its token/loopback exemption)."""
    assert client.get("/api/health", headers={"Host": "evil.example"}).status_code == 200


def test_token_mode_is_host_agnostic(client, monkeypatch):
    """Token mode needs no Host check: a rebound page cannot forge the Bearer credential, so the
    token — not the Host — is the authority. Pinned so nobody adds a Host check here later and
    silently breaks a legitimate token deployment served on its own hostname."""
    monkeypatch.setenv("ASSESSHUB_TOKEN", "s3cret-token")
    # the rebinding write (no/forged Bearer) is refused by the token gate, not the Host
    assert client.post("/api/demo/seed",
                       headers={"Host": "assesshub.corp.example"}).status_code == 401
    # a genuine client with the Bearer is served regardless of the (non-loopback) Host
    assert client.post("/api/demo/seed",
                       headers={"Host": "assesshub.corp.example",
                                "Authorization": "Bearer s3cret-token"}).status_code == 200


# ------------------------ Host-allowlist parser: exhaustive bypass matrix (unit)
# Grounds the certification (see docs: PortSwigger "host-header", OWASP WSTG host-injection, Django
# host_validation_re, Oligo "0.0.0.0-day"). The load-bearing fact: during rebinding the browser puts
# the ATTACKER's domain in Host (RFC 9110 §7.2; Host is a forbidden header name JS can't override),
# so an exact-match allowlist is the correct AND complete defense. IP-encoding / 0.0.0.0 / IPv4-mapped
# -IPv6 forms are SSRF / fuzzy-matching concerns — an exact-match allowlist fails closed on all of them.
class _ReqWithHost:
    """Minimal stand-in — _request_host_allowed only reads request.headers.get('host')."""
    def __init__(self, host):
        self.headers = {"host": host}


_LOOPBACK_OK = ["localhost", "localhost:8000", "127.0.0.1", "127.0.0.1:8000", "[::1]", "[::1]:8000",
                "LOCALHOST", "LocalHost:3000", "localhost:0", "localhost:65535"]

_HOST_REJECT = [
    "evil.example", "evil.example:8000",                         # rebinding: attacker's own domain
    "0.0.0.0", "0.0.0.0:8000", "[::]", "[::]:8000",              # 0.0.0.0-day / unspecified addr
    "2130706433", "0x7f000001", "0177.0.0.1", "127.1",           # decimal/hex/octal/short IPv4 encodings
    "[::ffff:127.0.0.1]", "[::ffff:7f00:1]", "[0:0:0:0:0:0:0:1]",  # IPv4-mapped / expanded IPv6 of loopback
    "localhost.evil.example", "127.0.0.1.evil.example",         # loopback name as a subdomain label
    "localhost:8000@evil.example",                              # userinfo confusion (bad-parser trap)
    "evil.example:8000:9000", "[::1]:8000:9000",                # extra colons
    "localhost:notaport",                                        # non-numeric port
    "localhost,evil.example",                                    # comma-joined duplicate Host headers
    "127.0.0.1:8000/../", "127.0.0.1%00", "loc alhost",         # path / null / space injection
    "localhost:8000\r\nX-Evil: 1",                               # CRLF header injection
    "xn--e1afmkfd.example",                                      # punycode / IDN homograph
    "localhost\t", "localhost\n", "127.0.0.1 ", "  localhost  ",  # control chars / surrounding whitespace
    "localhost.", "::1", "fe80::1", "[::1].", "",               # trailing dot / unbracketed IPv6 / empty
]


@pytest.mark.parametrize("host", _LOOPBACK_OK)
def test_host_parser_accepts_loopback(host, monkeypatch):
    monkeypatch.delenv("ASSESSHUB_ALLOWED_HOSTS", raising=False)
    assert app_module._request_host_allowed(_ReqWithHost(host)) is True


@pytest.mark.parametrize("host", _HOST_REJECT)
def test_host_parser_rejects_bypass_and_malformed(host, monkeypatch):
    """Every rebinding-bypass class and every malformed Host fails closed."""
    monkeypatch.delenv("ASSESSHUB_ALLOWED_HOSTS", raising=False)
    assert app_module._request_host_allowed(_ReqWithHost(host)) is False


def test_host_parser_allowed_hosts_is_exact_not_wildcard(monkeypatch):
    """ASSESSHUB_ALLOWED_HOSTS matches EXACTLY (case-insensitive, port-stripped). No suffix/subdomain
    wildcarding — a leading-dot/suffix entry would re-open rebinding for any name under that suffix."""
    monkeypatch.setenv("ASSESSHUB_ALLOWED_HOSTS", "assesshub.internal, proxy.local")
    for h in ["assesshub.internal", "assesshub.internal:9000", "ASSESSHUB.INTERNAL", "proxy.local"]:
        assert app_module._request_host_allowed(_ReqWithHost(h)) is True, h
    for h in ["sub.assesshub.internal", "assesshub.internal.evil", "evil.example",
              "assesshub.internal@evil"]:
        assert app_module._request_host_allowed(_ReqWithHost(h)) is False, h


@pytest.mark.parametrize("host", ["localhost:8000@evil.example", "localhost,evil.example",
                                  "localhost:notaport"])
def test_dns_rebinding_malformed_host_refused_end_to_end(client, host):
    """The strict parser is wired into the middleware: malformed / bypass-shaped Hosts are refused
    for a real write through the whole stack (httpx forwards these values verbatim), not just at the
    parser — so the guard can't be sidestepped by a Host the exact-match set wouldn't catch."""
    r = client.post("/api/demo/seed", headers={"Host": host})
    assert r.status_code == 403 and "Host" in r.json()["detail"]


# ------------------------------------------------- static source pins
def test_no_wildcard_cors_in_source():
    src = (_REPO / "webapp" / "backend" / "app.py").read_text(encoding="utf-8")
    assert 'allow_origins=["*"]' not in src and "allow_origins=['*']" not in src, \
        "CORS wildcard must not return"


def test_explorer_iframe_is_sandboxed_without_same_origin():
    """The iframe must carry a sandbox WITHOUT allow-same-origin — same-origin was the
    escape path from explorer innerHTML to the parent app and its API."""
    tsx = (_REPO / "webapp" / "frontend" / "src" / "pages" / "Snapshot.tsx").read_text(encoding="utf-8")
    assert "sandbox=" in tsx, "explorer iframe lost its sandbox attribute"
    sandbox_line = next(ln for ln in tsx.splitlines() if "sandbox=" in ln)
    assert "allow-scripts" in sandbox_line
    assert "allow-same-origin" not in sandbox_line, \
        "allow-same-origin would hand the iframe the parent origin + zero-auth API"


def test_spa_catchall_refuses_path_traversal(tmp_path, monkeypatch):
    """The unguarded SPA fallback serves only its startup-indexed distribution files."""
    dist = tmp_path / "dist"
    write_frontend_dist(dist)
    (dist / "app.js").write_text("console.log('legit-asset')", encoding="utf-8")
    secret = tmp_path / "secret.txt"                         # a sibling OUTSIDE dist
    secret.write_text("TOP-SECRET-SNAPSHOT-DB", encoding="utf-8")
    monkeypatch.setattr(app_module, "FRONTEND_DIST", dist)
    with TestClient(create_app(db_path=str(tmp_path / "t.db"))) as c:
        assert c.get("/app.js").text == "console.log('legit-asset')"       # a real in-dist asset is served
        assert "SPA-SHELL" in c.get("/campaigns/5").text                   # a deep link -> the SPA shell
        r = c.get("/%2e%2e/secret.txt")                                    # traversal to the sibling secret
        assert "TOP-SECRET" not in r.text and "SPA-SHELL" in r.text        # contained -> index.html, not the file
        assert "TOP-SECRET" not in c.get("/%2e%2e%2f%2e%2e%2fsecret.txt").text  # deeper encoded traversal too


def test_create_app_lifespan_closes_the_store(tmp_path, monkeypatch):
    database = tmp_path / "lifespan.db"
    app = create_app(db_path=str(database), dist_dir=tmp_path / "missing-dist")
    store = app.state.store
    real_close = store.close
    closes = 0

    def traced_close():
        nonlocal closes
        closes += 1
        real_close()

    monkeypatch.setattr(store, "close", traced_close)
    with TestClient(app, base_url="http://localhost") as client:
        assert client.get("/api/health").status_code == 200
    assert closes == 1
    database.replace(tmp_path / "closed.db")


@pytest.mark.parametrize(
    "hostile_path",
    (
        "//198.51.100.7/share/secret.txt",
        "\\\\198.51.100.7\\share\\secret.txt",
        "C:/Windows/win.ini",
        "C:\\Windows\\win.ini",
        "../secret.txt",
        "safe/../../secret.txt",
        "safe\\nested.js",
        "safe/secret.txt\x00",
        "index.html:alternate-stream",
        "CON",
        "missing.js",
    ),
)
def test_spa_requests_perform_no_request_derived_filesystem_operations(
    tmp_path, monkeypatch, hostile_path, request,
):
    """No request byte reaches Path construction, resolution, metadata, or response selection."""
    dist = tmp_path / "dist"
    index_bytes, _module = write_frontend_dist(dist)
    nested = dist / "safe"
    nested.mkdir()
    asset = nested / "nested.js"
    asset_bytes = b"console.log('trusted')"
    asset.write_bytes(asset_bytes)
    app = create_app(db_path=str(tmp_path / "hostile.db"), dist_dir=dist)
    request.addfinalizer(app.state.store.close)
    spa = next(route.endpoint for route in app.routes if route.path == "/{full_path:path}")

    def request_time_filesystem_operation_forbidden(*_args, **_kwargs):
        raise AssertionError("SPA request performed a filesystem path operation")

    for name in ("resolve", "exists", "is_dir", "is_file", "is_relative_to"):
        monkeypatch.setattr(Path, name, request_time_filesystem_operation_forbidden)
    http_request = app_module.Request({"type": "http", "method": "GET", "headers": []})
    response = spa(http_request, hostile_path)
    assert response.body == index_bytes

    trusted = spa(http_request, "safe/nested.js")
    assert trusted.body == asset_bytes


def test_spa_unknown_asset_requests_never_reach_realpath(tmp_path, monkeypatch):
    dist = tmp_path / "dist"
    write_frontend_dist(dist, asset_bytes=b"trusted")
    app = create_app(db_path=str(tmp_path / "asset-path.db"), dist_dir=dist)

    with TestClient(app, base_url="http://localhost") as client:
        def request_time_realpath_forbidden(*_args, **_kwargs):
            raise AssertionError("asset request reached os.path.realpath")

        monkeypatch.setattr(app_module.os.path, "realpath", request_time_realpath_forbidden)
        for path in (
            "/assets/C:%5CWindows%5Cwin.ini",
            "/assets/%5C%5C198.51.100.7%5Cshare%5Csecret.txt",
        ):
            assert client.get(path).status_code == 404


def test_spa_file_index_is_immutable_after_app_construction(tmp_path, request):
    dist = tmp_path / "dist"
    index_bytes, original_asset = write_frontend_dist(
        dist, asset_bytes=b"original asset",
    )
    assets = dist / "assets"
    app = create_app(db_path=str(tmp_path / "immutable.db"), dist_dir=dist)
    request.addfinalizer(app.state.store.close)
    spa = next(route.endpoint for route in app.routes if route.path == "/{full_path:path}")
    request = app_module.Request({"type": "http", "method": "GET", "headers": []})

    late = dist / "late.js"
    late.write_text("console.log('late')", encoding="utf-8")
    assert spa(request, "late.js").body == index_bytes
    (assets / "late.js").write_text("console.log('late asset')", encoding="utf-8")
    original_asset.write_bytes(b"replacement asset")
    assert spa(request, "assets/app.js").body == b"original asset"
    original_asset.unlink()
    assert spa(request, "assets/app.js").body == b"original asset"
    with TestClient(app, base_url="http://localhost") as client:
        assert client.get("/assets/late.js").status_code == 404
        assert client.get("/assets/app.js").content == b"original asset"


def test_spa_file_index_link_rejection_is_not_platform_vacuous(tmp_path, monkeypatch):
    dist = tmp_path / "dist"
    write_frontend_dist(dist)
    classified_link = dist / "linked.txt"
    classified_link.write_text("must not be indexed", encoding="utf-8")
    real_link_check = app_module._is_filesystem_link

    monkeypatch.setattr(
        app_module,
        "_is_filesystem_link",
        lambda path: path == classified_link or real_link_check(path),
    )
    assert app_module._frontend_file_index(dist) is None


def test_spa_file_index_rejects_windows_reparse_points_on_older_python(
    tmp_path, monkeypatch,
):
    candidate = tmp_path / "junction"
    candidate.mkdir()
    reparse_flag = 0x400
    monkeypatch.setattr(
        app_module.stat,
        "FILE_ATTRIBUTE_REPARSE_POINT",
        reparse_flag,
        raising=False,
    )
    monkeypatch.setattr(
        app_module.os,
        "lstat",
        lambda _path: type("ReparseStat", (), {"st_file_attributes": reparse_flag})(),
    )
    monkeypatch.setattr(Path, "is_symlink", lambda _path: False)
    monkeypatch.setattr(Path, "is_junction", lambda _path: False, raising=False)
    assert app_module._is_filesystem_link(candidate) is True


def test_spa_file_index_direct_symlink_check_is_independently_pinned(
    tmp_path, monkeypatch,
):
    candidate = tmp_path / "synthetic-link"
    candidate.write_text("fixture", encoding="utf-8")
    monkeypatch.setattr(
        app_module.os,
        "lstat",
        lambda _path: type("PlainStat", (), {"st_file_attributes": 0})(),
    )
    monkeypatch.setattr(Path, "is_symlink", lambda path: path == candidate)
    monkeypatch.setattr(Path, "is_junction", lambda _path: False, raising=False)
    assert app_module._is_filesystem_link(candidate) is True


def test_spa_file_index_refuses_partial_output_on_link_metadata_error(tmp_path, monkeypatch):
    dist = tmp_path / "dist"
    write_frontend_dist(dist)
    unreadable = dist / "unreadable.js"
    unreadable.write_text("fixture", encoding="utf-8")
    real_lstat = app_module.os.lstat

    def metadata_refused(path):
        if Path(path) == unreadable:
            raise PermissionError("synthetic link metadata refusal")
        return real_lstat(path)

    monkeypatch.setattr(app_module.os, "lstat", metadata_refused)
    with pytest.raises(PermissionError):
        app_module._is_filesystem_link(unreadable)
    assert app_module._frontend_file_index(dist) is None


def test_spa_file_index_rejects_a_link_classified_root(tmp_path, monkeypatch):
    dist = tmp_path / "dist"
    write_frontend_dist(dist)
    real_link_check = app_module._is_filesystem_link
    monkeypatch.setattr(
        app_module,
        "_is_filesystem_link",
        lambda path: path == dist or real_link_check(path),
    )
    assert app_module._frontend_file_index(dist) is None


def test_spa_file_index_independently_rejects_resolved_escape(tmp_path, monkeypatch):
    dist = tmp_path / "dist"
    write_frontend_dist(dist)
    candidate = dist / "escape.txt"
    candidate.write_text("placeholder", encoding="utf-8")
    outside = tmp_path / "outside.txt"
    outside.write_text("must not be indexed", encoding="utf-8")
    real_resolve = Path.resolve

    monkeypatch.setattr(app_module, "_is_filesystem_link", lambda _path: False)

    def escaping_resolve(path, *args, **kwargs):
        if path == candidate:
            return outside
        return real_resolve(path, *args, **kwargs)

    monkeypatch.setattr(Path, "resolve", escaping_resolve)
    monkeypatch.setattr(
        app_module,
        "_read_frontend_file",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            AssertionError("resolved escape reached the file reader")
        ),
    )
    assert app_module._frontend_file_index(dist) is None


def test_spa_same_handle_reader_rejects_parent_retarget_after_open(tmp_path, monkeypatch):
    dist = tmp_path / "dist"
    assets = dist / "assets"
    assets.mkdir(parents=True)
    candidate = assets / "app.js"
    candidate.write_bytes(b"trusted")
    outside = tmp_path / "outside.js"
    outside.write_bytes(b"outside")
    resolved_dist = dist.resolve()
    resolved_candidate = candidate.resolve()
    real_resolve = Path.resolve

    def retargeted_resolve(path, *args, **kwargs):
        if path == resolved_candidate:
            return outside
        return real_resolve(path, *args, **kwargs)

    monkeypatch.setattr(Path, "resolve", retargeted_resolve)
    assert app_module._read_frontend_file(resolved_candidate, resolved_dist) is None


def test_spa_file_index_rejects_hard_linked_members(tmp_path):
    dist = tmp_path / "dist"
    write_frontend_dist(dist)
    outside = tmp_path / "outside.js"
    outside.write_bytes(b"outside")
    linked = dist / "linked.js"
    try:
        app_module.os.link(outside, linked)
    except OSError as exc:
        pytest.skip(f"hard links are unavailable: {exc}")
    assert app_module._frontend_file_index(dist) is None


def test_spa_file_index_enforces_file_total_and_count_ceilings(tmp_path, monkeypatch):
    dist = tmp_path / "dist"
    dist.mkdir()
    index = dist / "index.html"
    asset = dist / "app.js"
    index.write_bytes(b"1234")
    asset.write_bytes(b"5678")

    monkeypatch.setattr(app_module, "_FRONTEND_MAX_FILE_BYTES", 3)
    assert app_module._frontend_file_index(dist) is None

    monkeypatch.setattr(app_module, "_FRONTEND_MAX_FILE_BYTES", 8)
    monkeypatch.setattr(app_module, "_FRONTEND_MAX_TOTAL_BYTES", 7)
    assert app_module._frontend_file_index(dist) is None

    monkeypatch.setattr(app_module, "_FRONTEND_MAX_TOTAL_BYTES", 16)
    monkeypatch.setattr(app_module, "_FRONTEND_MAX_FILES", 1)
    assert app_module._frontend_file_index(dist) is None

    monkeypatch.setattr(app_module, "_FRONTEND_MAX_FILES", 4)
    monkeypatch.setattr(app_module, "_FRONTEND_MAX_ENTRIES", 1)
    assert app_module._frontend_file_index(dist) is None


@pytest.mark.parametrize("mutation", ("add-member", "rewrite-index"))
def test_spa_file_index_reconciles_the_whole_tree_after_member_reads(
    tmp_path, monkeypatch, mutation,
):
    dist = tmp_path / "dist"
    _index_bytes, module = write_frontend_dist(dist)
    index = dist / "index.html"
    real_reader = app_module._read_frontend_file
    mutated = False

    def mutate_between_members(path, root):
        nonlocal mutated
        content = real_reader(path, root)
        if path == module and not mutated:
            mutated = True
            if mutation == "add-member":
                (dist / "assets" / "late.js").write_bytes(b"late member")
            else:
                index.write_bytes(
                    index.read_bytes().replace(b"SPA-SHELL", b"NEW-CROSS-TIME-SHELL")
                )
        return content

    monkeypatch.setattr(app_module, "_read_frontend_file", mutate_between_members)
    assert app_module._frontend_file_index(dist) is None
    assert mutated is True


def test_spa_file_index_aborts_on_the_first_guarded_reader_refusal(tmp_path, monkeypatch):
    dist = tmp_path / "dist"
    _index_bytes, module = write_frontend_dist(dist)
    real_reader = app_module._read_frontend_file
    calls = 0

    def refuse_first_read(path, root):
        nonlocal calls
        if path == module:
            calls += 1
            if calls == 1:
                return None
        return real_reader(path, root)

    monkeypatch.setattr(app_module, "_read_frontend_file", refuse_first_read)
    assert app_module._frontend_file_index(dist) is None
    assert calls == 1


@pytest.mark.parametrize(
    "invalid_index",
    (
        b"",
        b"<!doctype html><html><head></head><body><div id='root'></div></body></html>",
        b"<!doctype html><html><head><script type='module' src='/assets/app.js'></script>"
        b"</head><body><div id='root'></div>",
    ),
)
def test_spa_file_index_rejects_empty_or_truncated_boot_shells(tmp_path, invalid_index):
    dist = tmp_path / "dist"
    write_frontend_dist(dist)
    (dist / "index.html").write_bytes(invalid_index)
    assert app_module._frontend_file_index(dist) is None


@pytest.mark.parametrize(
    "mutation",
    (
        "base", "integrity", "meta-csp", "meta-refresh", "duplicate-src", "template",
        "self-closing-script",
    ),
)
def test_spa_file_index_rejects_browser_parser_differentials(tmp_path, mutation):
    dist = tmp_path / "dist"
    index_bytes, _module = write_frontend_dist(dist)
    text = index_bytes.decode("utf-8")
    if mutation == "base":
        text = text.replace("<head>", '<head><base href="/elsewhere/">')
    elif mutation == "integrity":
        text = text.replace("<script type=", '<script integrity="sha256-wrong" type=')
    elif mutation == "meta-csp":
        text = text.replace(
            "<head>",
            '<head><meta http-equiv="Content-Security-Policy" content="script-src none">',
        )
    elif mutation == "meta-refresh":
        text = text.replace(
            "<head>",
            '<head><meta http-equiv="refresh" content="0;url=https://example.invalid/">',
        )
    elif mutation == "duplicate-src":
        text = text.replace(
            'src="/assets/app.js"',
            'src="https://example.invalid/x.js" src="/assets/app.js"',
        )
    elif mutation == "template":
        text = text.replace("<script ", "<template><script ").replace(
            "</script>", "</script></template>",
        )
    else:
        text = text.replace("></script>", "/>")
    (dist / "index.html").write_text(text, encoding="utf-8")
    assert app_module._frontend_file_index(dist) is None


def test_spa_file_index_rejects_encoded_dot_segments_and_blank_module(tmp_path):
    dist = tmp_path / "encoded"
    write_frontend_dist(dist)
    encoded_dir = dist / "assets" / "%2e%2e"
    encoded_dir.mkdir()
    (encoded_dir / "app.js").write_bytes(b"encoded alias")
    index = dist / "index.html"
    index.write_bytes(index.read_bytes().replace(b"/assets/app.js", b"/assets/%2e%2e/app.js"))
    assert app_module._frontend_file_index(dist) is None

    blank = tmp_path / "blank"
    write_frontend_dist(blank, asset_bytes=b" \r\n")
    assert app_module._frontend_file_index(blank) is None


def test_spa_file_index_rejects_encoded_dot_segment_reference(tmp_path):
    dist = tmp_path / "dist"
    write_frontend_dist(dist)
    encoded_dir = dist / "assets" / "%2e%2e"
    encoded_dir.mkdir()
    (encoded_dir / "app.js").write_bytes(b"encoded alias")
    index = dist / "index.html"
    index.write_bytes(index.read_bytes().replace(b"/assets/app.js", b"/assets/%2e%2e/app.js"))
    assert app_module._frontend_file_index(dist) is None


def test_spa_file_index_rejects_blank_boot_module(tmp_path):
    dist = tmp_path / "dist"
    write_frontend_dist(dist, asset_bytes=b" \r\n")
    assert app_module._frontend_file_index(dist) is None


def test_spa_file_index_etags_bind_each_exact_representation(tmp_path):
    dist = tmp_path / "dist"
    _index_bytes, module = write_frontend_dist(dist, asset_bytes=b"module-one")
    other = dist / "assets" / "other.js"
    other.write_bytes(b"module-two")
    result = app_module._frontend_file_index(dist)
    assert result is not None
    indexed = result[1]
    first = indexed[module.relative_to(dist).as_posix()]
    second = indexed[other.relative_to(dist).as_posix()]
    assert first.etag == f'"{hashlib.sha256(first.content).hexdigest()}"'
    assert second.etag == f'"{hashlib.sha256(second.content).hexdigest()}"'
    assert first.etag != second.etag


def test_spa_file_index_pins_boot_media_types_against_host_registry(tmp_path, monkeypatch):
    dist = tmp_path / "dist"
    _index_bytes, module = write_frontend_dist(dist)
    monkeypatch.setattr(
        app_module.mimetypes,
        "guess_type",
        lambda _path: (_ for _ in ()).throw(AssertionError("pinned type reached host registry")),
    )
    result = app_module._frontend_file_index(dist)
    assert result is not None
    indexed = result[1]
    assert indexed["index.html"].media_type == "text/html"
    assert indexed[module.relative_to(dist).as_posix()].media_type == "text/javascript"


def test_spa_reader_rejects_divergent_second_read_and_closes_handle(tmp_path, monkeypatch):
    dist = tmp_path / "dist"
    dist.mkdir()
    candidate = dist / "app.js"
    candidate.write_bytes(b"stable")
    resolved = candidate.resolve()
    real_open = app_module.os.open
    real_read = app_module.os.read
    real_lseek = app_module.os.lseek
    opened = []
    second_pass = False

    def traced_open(*args, **kwargs):
        descriptor = real_open(*args, **kwargs)
        opened.append(descriptor)
        return descriptor

    def traced_lseek(*args, **kwargs):
        nonlocal second_pass
        second_pass = True
        return real_lseek(*args, **kwargs)

    def divergent_read(*args, **kwargs):
        data = real_read(*args, **kwargs)
        if second_pass and data:
            return b"X" + data[1:]
        return data

    monkeypatch.setattr(app_module.os, "open", traced_open)
    monkeypatch.setattr(app_module.os, "lseek", traced_lseek)
    monkeypatch.setattr(app_module.os, "read", divergent_read)
    assert app_module._read_frontend_file(resolved, dist.resolve()) is None
    assert len(opened) == 1
    with pytest.raises(OSError):
        app_module.os.fstat(opened[0])


def test_spa_reader_rejects_handle_metadata_drift(tmp_path, monkeypatch):
    dist = tmp_path / "dist"
    dist.mkdir()
    candidate = dist / "app.js"
    candidate.write_bytes(b"stable")
    real_fstat = app_module.os.fstat
    calls = 0

    def drifting_fstat(descriptor):
        nonlocal calls
        calls += 1
        value = real_fstat(descriptor)
        if calls == 2:
            return type("DriftedStat", (), {
                "st_dev": value.st_dev,
                "st_ino": value.st_ino,
                "st_size": value.st_size,
                "st_mtime_ns": value.st_mtime_ns + 1,
                "st_ctime_ns": value.st_ctime_ns,
                "st_nlink": value.st_nlink,
            })()
        return value

    monkeypatch.setattr(app_module.os, "fstat", drifting_fstat)
    assert app_module._read_frontend_file(candidate.resolve(), dist.resolve()) is None
    assert calls == 2


def test_spa_reader_closes_handle_when_read_fails(tmp_path, monkeypatch):
    dist = tmp_path / "dist"
    dist.mkdir()
    candidate = dist / "app.js"
    candidate.write_bytes(b"stable")
    real_open = app_module.os.open
    opened = []

    def traced_open(*args, **kwargs):
        descriptor = real_open(*args, **kwargs)
        opened.append(descriptor)
        return descriptor

    monkeypatch.setattr(app_module.os, "open", traced_open)
    monkeypatch.setattr(
        app_module.os,
        "read",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(OSError("synthetic read failure")),
    )
    assert app_module._read_frontend_file(candidate.resolve(), dist.resolve()) is None
    assert len(opened) == 1
    with pytest.raises(OSError):
        app_module.os.fstat(opened[0])


def test_spa_single_range_parser_is_ascii_bounded_and_case_insensitive():
    assert app_module._single_byte_range("BYTES=0-1", 10) == (0, 1)
    assert app_module._single_byte_range("bytes=-3", 10) == (7, 9)
    assert app_module._single_byte_range("bytes=²-", 10) is None
    assert app_module._single_byte_range("bytes=0-1,4-5", 10) is None
    assert app_module._single_byte_range("items=0-1", 10) is None

    entry = app_module._FrontendFile(
        content=b"0123456789",
        media_type="text/plain",
        etag='"content-etag"',
    )
    non_ascii = app_module.Request({
        "type": "http",
        "method": "GET",
        "headers": [(b"range", b"bytes=\xb2-")],
    })
    non_ascii_response = app_module._frontend_response(entry, non_ascii)
    assert non_ascii_response.status_code == 416
    assert non_ascii_response.headers["content-range"] == "bytes */10"
    unsupported = app_module.Request({
        "type": "http",
        "method": "GET",
        "headers": [(b"range", b"items=0-1")],
    })
    unsupported_response = app_module._frontend_response(entry, unsupported)
    assert unsupported_response.status_code == 200
    assert unsupported_response.body == entry.content
    head = app_module.Request({
        "type": "http",
        "method": "HEAD",
        "headers": [(b"range", b"bytes=0-1")],
    })
    head_response = app_module._frontend_response(entry, head)
    assert head_response.status_code == 200 and head_response.body == b""
    assert head_response.headers["content-length"] == "10"
    oversized_precondition = app_module.Request({
        "type": "http",
        "method": "GET",
        "headers": [(b"if-none-match", b"x" * 4_097)],
    })
    assert app_module._frontend_response(entry, oversized_precondition).status_code == 431


def test_spa_file_index_rejects_linked_directories_before_walk_descends(
    tmp_path, monkeypatch,
):
    dist = tmp_path / "dist"
    linked = dist / "linked"
    safe = dist / "safe"
    write_frontend_dist(dist)
    linked.mkdir()
    safe.mkdir()
    (linked / "secret.js").write_text("must not be indexed", encoding="utf-8")
    (safe / "app.js").write_text("must be indexed", encoding="utf-8")

    real_link_check = app_module._is_filesystem_link
    monkeypatch.setattr(
        app_module,
        "_is_filesystem_link",
        lambda path: path == linked or real_link_check(path),
    )
    assert app_module._frontend_file_index(dist) is None


def test_spa_file_index_revalidates_child_before_scandir(tmp_path, monkeypatch):
    dist = tmp_path / "dist"
    child = dist / "child"
    outside = tmp_path / "outside"
    write_frontend_dist(dist)
    child.mkdir()
    outside.mkdir()
    (child / "app.js").write_text("trusted", encoding="utf-8")
    real_resolve = Path.resolve
    real_scandir = app_module.os.scandir

    def retargeted_resolve(path, *args, **kwargs):
        if path == child:
            return outside
        return real_resolve(path, *args, **kwargs)

    def guarded_scandir(path):
        if Path(path) == child:
            raise AssertionError("stale child reached scandir")
        return real_scandir(path)

    monkeypatch.setattr(Path, "resolve", retargeted_resolve)
    monkeypatch.setattr(app_module.os, "scandir", guarded_scandir)
    assert app_module._frontend_file_index(dist) is None


def test_spa_file_index_fails_closed_on_enumeration_error(tmp_path, monkeypatch):
    dist = tmp_path / "dist"
    write_frontend_dist(dist)

    def failing_scandir(_root):
        raise PermissionError("synthetic enumeration refusal")

    monkeypatch.setattr(app_module.os, "scandir", failing_scandir)
    assert app_module._frontend_file_index(dist) is None


def test_spa_file_index_refuses_out_of_root_symlinks(tmp_path):
    dist = tmp_path / "dist"
    write_frontend_dist(dist)
    secret = tmp_path / "secret.txt"
    secret.write_text("TOP-SECRET-SNAPSHOT-DB", encoding="utf-8")
    link = dist / "linked.txt"
    try:
        link.symlink_to(secret)
    except OSError as exc:
        pytest.skip(f"file symlinks are unavailable: {exc}")

    assert app_module._frontend_file_index(dist) is None


def test_spa_file_index_rejects_same_root_alias_retarget(tmp_path, monkeypatch):
    dist = tmp_path / "dist"
    write_frontend_dist(dist)
    candidate = dist / "app.js"
    candidate.write_bytes(b"candidate")
    other = dist / "other.js"
    other.write_bytes(b"other")
    real_resolve = Path.resolve

    def aliased_resolve(path, *args, **kwargs):
        if path == candidate:
            return other
        return real_resolve(path, *args, **kwargs)

    monkeypatch.setattr(Path, "resolve", aliased_resolve)
    assert app_module._frontend_file_index(dist) is None


@pytest.mark.parametrize(
    ("body", "expected_detail"),
    (
        (
            b'{"ok":1,"secret\\r\\nFORGED LOG RECORD":2,'
            b'"secret\\r\\nFORGED LOG RECORD":3}',
            "Invalid JSON request body: duplicate JSON object key",
        ),
        (b'{"secret-material":', "Invalid JSON request body"),
        (b'{"name":"\xff"}', "Invalid JSON request body"),
    ),
)
def test_invalid_json_response_never_echoes_parser_or_request_detail(
    client, body, expected_detail,
):
    response = client.post(
        "/api/compare",
        content=body,
        headers={"content-type": "application/json"},
    )

    assert response.status_code == 400
    assert response.json() == {"detail": expected_detail}
    assert "secret" not in response.text.lower()
    assert "forged" not in response.text.lower()


# ── whole-repo review, 2026-07-28 ───────────────────────────────────────────────────
def _raw_asgi_get(app, path, headers, client_peer):
    """Drive the ASGI app with a hand-built scope, so `client` can be ABSENT — the shape a
    Unix-domain-socket bind and several ASGI adapters/proxies produce. TestClient cannot express it:
    it always stamps ('testclient', 50000) into the scope."""
    import anyio

    scope = {"type": "http", "asgi": {"version": "3.0", "spec_version": "2.1"},
             "http_version": "1.1", "method": "GET", "path": path, "raw_path": path.encode(),
             "query_string": b"", "root_path": "", "scheme": "http", "server": ("testserver", 80),
             "headers": [(k.lower().encode(), v.encode()) for k, v in headers.items()]}
    if client_peer is not None:
        scope["client"] = client_peer
    seen = {}

    async def receive():
        return {"type": "http.request", "body": b"", "more_body": False}

    async def send(msg):
        if msg["type"] == "http.response.start":
            seen["status"] = msg["status"]

    anyio.run(lambda: app(scope, receive, send))
    return seen["status"]


def test_unknown_asgi_peer_is_not_loopback(tmp_path, monkeypatch):
    """[#59-adjacent / #58] `_client_is_loopback`'s docstring promised "deliberately conservative: a peer
    with a non-loopback IP is NOT loopback" — but the UNKNOWN case failed OPEN (`if host is None ...:
    return True`). `request.client` is None for a Unix-domain-socket bind and several ASGI adapters and
    proxies, so in no-token mode a request through such a deployment satisfied the loopback half of the
    access guard, leaving only the Host allowlist — a value the client picks for itself. Unknown position
    is not local position: it must fail CLOSED."""
    monkeypatch.delenv("ASSESSHUB_TOKEN", raising=False)
    monkeypatch.delenv("ASSESSHUB_ALLOWED_HOSTS", raising=False)
    a = create_app(db_path=str(tmp_path / "peer.db"))
    hdrs = {"host": "localhost"}                       # the Host guard is satisfied; only the peer differs
    assert _raw_asgi_get(a, "/api/campaigns", hdrs, ("127.0.0.1", 5000)) == 200   # genuine loopback: served
    assert _raw_asgi_get(a, "/api/campaigns", hdrs, ("203.0.113.9", 5000)) == 403  # remote: refused
    assert _raw_asgi_get(a, "/api/campaigns", hdrs, None) == 403                  # UNKNOWN: was 200


def test_asgi_test_harness_peer_is_not_a_production_bypass(tmp_path, monkeypatch):
    """[#58] The hardcoded `host == "testclient"` allowance shipped in production code. It is not
    unreachable-by-construction: uvicorn's proxy-headers middleware copies X-Forwarded-For into
    scope["client"] VERBATIM without checking it parses as an IP, so an operator who widens
    forwarded_allow_ips would let a remote caller name itself "testclient". It is now honoured only
    while the process is executing a pytest test."""
    monkeypatch.delenv("ASSESSHUB_TOKEN", raising=False)
    monkeypatch.delenv("ASSESSHUB_ALLOWED_HOSTS", raising=False)
    a = create_app(db_path=str(tmp_path / "harness.db"))
    hdrs = {"host": "localhost"}
    assert _raw_asgi_get(a, "/api/campaigns", hdrs, ("testclient", 50000)) == 200   # under pytest
    monkeypatch.delenv("PYTEST_CURRENT_TEST", raising=False)                        # simulate production
    assert app_module._under_pytest() is False
    assert _raw_asgi_get(a, "/api/campaigns", hdrs, ("testclient", 50000)) == 403


# --- unbounded stored strings / unbounded request bodies (#61) ----------------------
def test_every_write_model_caps_its_strings():
    """[#61] `GateIn` carried `max_length` caps with a comment naming the exact vector — "stored
    verbatim, echoed by every board fetch and rendered into a DOCX table cell". EVERY sibling write
    model has the identical properties, and none had a cap. Assert the CLASS (a new write model must
    inherit the rule), not a hand-listed subset."""
    from pydantic import BaseModel

    models = [m for m in vars(app_module).values()
              if isinstance(m, type) and issubclass(m, BaseModel) and m is not BaseModel]
    assert len(models) >= 9, [m.__name__ for m in models]
    for model in models:
        for name, field in model.model_fields.items():
            if field.annotation is not str:
                continue
            caps = [getattr(m, "max_length", None) for m in field.metadata]
            assert any(c for c in caps), f"{model.__name__}.{name} accepts an unbounded string"


def test_unbounded_write_field_is_rejected(client):
    """[#61] End-to-end: an over-long value on a sibling write model is refused (422) instead of being
    stored verbatim and echoed back by every later fetch."""
    over = "A" * 5000
    assert client.post("/api/campaigns", json={"name": over}).status_code == 422
    cid = client.post("/api/campaigns", json={"name": "ok"}).json()["id"]
    assert client.get("/api/campaigns").json()[0]["name"] == "ok"
    sid = client.post("/api/demo/seed").json()["snapshot"]["id"]
    ex = client.post(f"/api/snapshots/{sid}/executions", json={"label": over})
    assert ex.status_code == 422
    eid = client.post(f"/api/snapshots/{sid}/executions", json={"label": "run"}).json()["id"]
    assert client.post(f"/api/executions/{eid}/event",
                       json={"kind": "note", "text": over}).status_code == 422
    assert cid  # campaign creation itself still works


def test_oversized_json_body_is_refused_before_it_is_parsed(client):
    """[#61] The per-field caps only fire AFTER Starlette has buffered the whole body and json.loads has
    materialised it, so a single huge JSON body was still a memory spike on an unauthenticated loopback
    POST. The declared Content-Length is refused up front with 413. Multipart is exempt — the upload
    routes do their own chunked read against the (much larger) archive limit."""
    r = client.post("/api/campaigns", json={"name": "x", "description": "B" * (2 * 1024 * 1024)})
    assert r.status_code == 413, r.status_code
    assert "limit" in r.json()["detail"]
    # the exemption still holds: a multipart snapshot upload well over the JSON cap is accepted
    import json as _json
    cid = client.post("/api/campaigns", json={"name": "big-upload"}).json()["id"]
    blob = _json.dumps({"devices": {f"sw{i}": {"model": "C9300", "serial": "S" * 200}
                                    for i in range(5000)}}).encode()
    assert len(blob) > app_module._max_json_body_bytes()
    up = client.post(f"/api/campaigns/{cid}/snapshots",
                     files={"file": ("s.json", blob, "application/json")})
    assert up.status_code == 201, up.text


# --- generated deliverables must not linger in %TEMP% (#62) -------------------------
def test_send_file_deletes_the_temp_deliverable_before_returning(tmp_path):
    """[#62] `_send_file` used to delete the generated temp file only from a Starlette BackgroundTask,
    which runs after the body is fully sent — so a client disconnect mid-download, or a killed process
    (the normal way a USB-stick field app ends), left a fully-rendered UNREDACTED client deliverable
    (hostnames, IPs, serials, parsed configs) in the OS temp dir permanently. There must be no path
    through this function that returns while the file is still on disk."""
    import os as _os

    p = tmp_path / "assesshub_probe.docx"
    p.write_bytes(b"UNREDACTED CLIENT DELIVERABLE")
    resp = app_module._send_file(str(p), "application/octet-stream", "Meridian reference fleet/2026", "_mop.docx",
                                 headers={"X-Gate-Status": "pending:design"})
    assert not _os.path.exists(p), "the rendered deliverable survived the response construction"
    assert resp.body == b"UNREDACTED CLIENT DELIVERABLE"          # bytes are unchanged
    assert resp.headers["x-gate-status"] == "pending:design"      # out-of-band notes survive
    assert 'filename="Meridian_reference_fleet_2026_mop.docx"' in resp.headers["content-disposition"]


def test_downloaded_deliverable_leaves_no_temp_file(client, tmp_path, monkeypatch):
    """[#62] End-to-end over the real routes: after a deliverable and a PIR download, the temp dir holds
    no `assesshub_*` residue. Pinned with a private TMPDIR so the assertion is exact rather than a
    best-effort scan of a shared %TEMP%."""
    import tempfile as _tempfile

    private = tmp_path / "tmp"
    private.mkdir()
    for var in ("TMPDIR", "TEMP", "TMP"):
        monkeypatch.setenv(var, str(private))
    monkeypatch.setattr(_tempfile, "tempdir", None)
    sid = client.post("/api/demo/seed").json()["snapshot"]["id"]
    r = client.get(f"/api/snapshots/{sid}/deliverable/mop")
    if r.status_code == 503:
        pytest.skip("python-docx not installed on this runner")
    assert r.status_code == 200
    eid = client.post(f"/api/snapshots/{sid}/executions", json={}).json()["id"]
    assert client.get(f"/api/executions/{eid}/report").status_code == 200
    assert list(private.glob("assesshub_*")) == [], "a rendered client deliverable was left in %TEMP%"


@pytest.mark.parametrize("kind", sorted(deliverables_module.SPECS))
def test_deliverable_temp_path_is_independent_of_request_kind(
    tmp_path, monkeypatch, kind,
):
    """The closed registry key selects a writer; it never needs to become path material."""
    seen = []
    real_mkstemp = deliverables_module.tempfile.mkstemp

    def traced_mkstemp(*, suffix, prefix):
        seen.append((suffix, prefix))
        return real_mkstemp(suffix=suffix, prefix=prefix, dir=tmp_path)

    def writer(path, _snap, _label, **_kwargs):
        Path(path).write_bytes(b"generated")

    monkeypatch.setattr(deliverables_module.tempfile, "mkstemp", traced_mkstemp)
    monkeypatch.setattr(deliverables_module, "resolve_writer", lambda _spec: writer)
    monkeypatch.setattr(deliverables_module, "_reconcile_gate", lambda *_args: None)

    path = deliverables_module.generate(kind, {}, "label", document_gate=None)
    try:
        assert seen == [("." + deliverables_module.SPECS[kind].ext, "assesshub_")]
        assert Path(path).parent == tmp_path
    finally:
        Path(path).unlink(missing_ok=True)


def test_deliverable_logs_keep_only_aggregate_security_state(
    tmp_path, monkeypatch, caplog,
):
    """Uploaded snapshot and gate values cannot forge or disclose adjacent log records."""
    from cisco_toolkit import ssot

    violation = "client-secret-value\r\nFORGED SSOT RECORD"
    drift = {
        "ssot_reconciliation": "failed",
        "n_violations": 1,
        "violations": [violation],
    }
    monkeypatch.setattr(ssot, "audit", lambda _snap: drift)
    caplog.set_level("WARNING", logger=deliverables_module.__name__)
    assert deliverables_module._reconcile_gate({}) == drift
    assert "bounded reconciliation details are embedded in the artifact" in caplog.text
    assert "client-secret-value" not in caplog.text
    assert "FORGED SSOT RECORD" not in caplog.text

    caplog.clear()

    rendered = []

    def writer_for(spec):
        def capture_writer(path, snapshot, _label, **_kwargs):
            rendered.append(snapshot)
            if spec.ext == "docx":
                from docx import Document
                Document().save(path)
            else:
                from pptx import Presentation
                Presentation().save(path)
        return capture_writer

    monkeypatch.setattr(deliverables_module, "resolve_writer", writer_for)
    disclosure = {
        "status": "pending",
        "missing": ["assessment_approved"],
        "revoked": [],
        "generator": "client-secret-value\r\nFORGED GATE RECORD",
        "detail": "C:\\private\\client-secret-value",
    }
    source_snapshot = {}
    path = deliverables_module.generate("runbook", source_snapshot, "label", document_gate=disclosure)
    try:
        from docx import Document
        assert rendered[0]["assessment_integrity"] == drift
        assert "assessment_integrity" not in source_snapshot
        doc_text = "\n".join(paragraph.text for paragraph in Document(path).paragraphs)
        assert "ASSESSMENT INTEGRITY WARNING" in doc_text
        assert "client-secret-value FORGED SSOT RECORD" in doc_text
        assert "[GATE UNSATISFIED]" in caplog.text
        assert "exact status is disclosed in the response and document" in caplog.text
        assert "client-secret-value" not in caplog.text
        assert "FORGED GATE RECORD" not in caplog.text
        assert "C:\\private" not in caplog.text
    finally:
        Path(path).unlink(missing_ok=True)

    deck_path = deliverables_module.generate(
        "deck", source_snapshot, "label", document_gate=None,
    )
    try:
        from pptx import Presentation
        deck_text = "\n".join(
            shape.text
            for slide in Presentation(deck_path).slides
            for shape in slide.shapes
            if hasattr(shape, "text")
        )
        assert "ASSESSMENT INTEGRITY WARNING" in deck_text
        assert "client-secret-value FORGED SSOT RECORD" in deck_text
    finally:
        Path(deck_path).unlink(missing_ok=True)


# ------------------------------------- OpenAPI / docs routes (guard COMPLETENESS)
# The access guard's reach was written as the path PREFIX "/api/", but FastAPI also generates
# /openapi.json, /docs, /docs/oauth2-redirect and /redoc — routes that describe THIS API and do not
# live under /api/. They were therefore outside the guard entirely. Measured before the fix: with
# ASSESSHUB_TOKEN set, `GET /openapi.json` returned 200 and the complete route + request-model schema
# to a caller sending no Bearer; from a NON-loopback peer on a zero-token instance it also returned
# 200 while /api/campaigns returned 403. Same class as every other finding in this file — a guard
# stated as a named subset rather than as the structural property it was meant to express.
def _doc_paths(app):
    """From the app's own attributes, exactly like the guard: a renamed docs_url stays covered."""
    return [p for p in (app.openapi_url, app.docs_url, app.redoc_url,
                        app.swagger_ui_oauth2_redirect_url) if p]


def test_openapi_and_docs_require_the_token(client, monkeypatch):
    monkeypatch.setenv("ASSESSHUB_TOKEN", "s3cret-token")
    paths = _doc_paths(client.app)
    assert len(paths) == 4, paths          # the enumeration must not go silently empty
    for p in paths:
        assert client.get(p).status_code == 401, p
        assert client.get(
            p, headers={"Authorization": "Bearer s3cret-token"}).status_code == 200, p


def test_openapi_and_docs_are_loopback_only_without_a_token(client, monkeypatch):
    """No token -> the API schema is client-adjacent metadata and follows /api's network rule."""
    monkeypatch.setattr(app_module, "_client_is_loopback", lambda request: False)
    for p in _doc_paths(client.app):
        assert client.get(p).status_code == 403, p
    assert client.get("/api/health").status_code == 200      # liveness stays open


def test_zero_token_loopback_dev_flow_still_serves_the_docs(client):
    """The default developer posture is untouched — the guard adds a network/auth condition, it
    does not switch the documentation off."""
    for p in _doc_paths(client.app):
        assert client.get(p).status_code == 200, p
